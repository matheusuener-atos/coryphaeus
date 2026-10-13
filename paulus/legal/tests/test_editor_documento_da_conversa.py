"""
O documento da conversa (docs/PLANO-EDITOR.md, E1), pelo servidor.

  - o trabalho redigido vira rascunho DA CONVERSA: fora da lista do Editor,
    sem as "Notas para o advogado" no corpo, listado na conversa;
  - o .docx do Acervo entra com a estrutura (título, negrito, quadro no lugar,
    a numeração do Word) e abrir de novo na mesma conversa volta ao mesmo
    rascunho - pelo conteúdo do arquivo; outra conversa tem o seu;
  - o PDF continua PDF: sem o pedido explícito, 409 (e a conversa oferece a
    versão editável num cartão, sem criar nada); com o pedido, o texto entra
    limpo (sem "[pagina N]", palavras partidas reunidas);
  - gravar com a revisão de que a tela partiu: a gravação velha não passa por
    cima (409 com o que está gravado); o autosave não cria versão a cada pausa;
  - sugestão pendente impede PDF, DOCX e Acervo (409); aceitar todas libera;
  - o rascunho da conversa não vai ao Acervo nem ao Word sem concluir;
  - concluir: o .docx do Acervo volta remendado do original, ao lado dele, com
    outro nome (o original não muda), cabeçalho e rodapé intactos; nome
    repetido pergunta; o concluído não aceita gravação e pode ser reaberto;
  - a leitura para PDF ignora o texto que uma sugestão tira e põe o número
    automático do Word;
  - apagar a conversa leva o rascunho para a lixeira; o concluído fica.

    PYTHONIOENCODING=utf-8 venv/Scripts/python.exe tests/test_editor_documento_da_conversa.py
"""

from __future__ import annotations

import io
import json
import os
import sys
import tempfile
from pathlib import Path
from types import SimpleNamespace

RAIZ = Path(__file__).parent.parent
TMP = Path(tempfile.mkdtemp(prefix="paulus-doc-conversa-"))
os.environ["PAULUS_DADOS"] = str(TMP / "dados")
os.environ["PAULUS_SEM_AVISOS"] = "1"
sys.path.insert(0, str(RAIZ / "src"))
sys.path.insert(0, str(RAIZ))

_falhas: list[str] = []


def checar(cond, nome: str, detalhe=None) -> None:
    print(("  ok   " if cond else "  FALHOU  ") + nome + ("" if cond or detalhe is None else f"  -> {detalhe!r}"))
    if not cond:
        _falhas.append(nome)


def eventos(texto: str) -> list[tuple[str, dict]]:
    saida = []
    for bloco in texto.split("\n\n"):
        tipo = next((l[7:] for l in bloco.splitlines() if l.startswith("event: ")), "")
        dado = next((l[6:] for l in bloco.splitlines() if l.startswith("data: ")), "")
        if tipo:
            saida.append((tipo, json.loads(dado) if dado else {}))
    return saida


def docx_de_teste() -> bytes:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    d = Document()
    d.sections[0].header.paragraphs[0].text = "SILVA ADVOGADOS ASSOCIADOS"
    d.sections[0].footer.paragraphs[0].text = "Rua da Bahia, 1.148 - Belo Horizonte"
    d.add_heading("CONTRATO DE PRESTAÇÃO DE SERVIÇOS", level=1)
    p = d.add_paragraph("Contratante: ")
    p.add_run("EMPRESA ALFA LTDA.").bold = True
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    d.add_paragraph("Os serviços são:")
    d.add_paragraph("consultoria tributária", style="List Number")
    d.add_paragraph("planejamento sucessório", style="List Number")
    t = d.add_table(rows=2, cols=2)
    for (i, j), v in {(0, 0): "Parcela", (0, 1): "Valor", (1, 0): "1", (1, 1): "R$ 5.000,00"}.items():
        t.cell(i, j).text = v
    d.add_paragraph("CLÁUSULA 1ª – DO PREÇO").runs[0].bold = True
    d.add_paragraph("O preço total é de R$ 5.000,00 (cinco mil reais).")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def pdf_de_teste() -> bytes:
    import documento as D

    temas = ["o prazo de pagamento", "a multa contratual", "os juros de mora", "a correção monetária", "o protesto",
             "a negativação", "o foro", "as custas", "os honorários", "a mora", "o vencimento", "a cobrança judicial"]
    blocos = D.ler_html("<h1>NOTIFICAÇÃO EXTRAJUDICIAL</h1>" + "".join(
        f"<p>Parágrafo {i} da notificação, que trata de {temas[i % len(temas)]} com texto bastante longo para ocupar "
        f"mais de uma linha na folha do PDF e obrigar a extração a partir a frase em linhas visuais diferentes, como "
        f"acontece de verdade em contratos, item {i * 7}.</p>"
        for i in range(1, 40)))
    return D.para_pdf(blocos, "Notificação", "rodapé do escritório")


def main() -> int:
    print("=" * 60)
    print("  O documento da conversa (E1)")
    print("=" * 60)

    import requests

    import api
    import blocos as B
    import documento as D
    import llama_client
    import nuvem
    from extract import file_sha1
    from fastapi.testclient import TestClient

    llama_client.requests = SimpleNamespace(post=lambda *a, **k: (_ for _ in ()).throw(RuntimeError("sem modelo")),
                                           get=lambda *a, **k: SimpleNamespace(status_code=200, json=lambda: {"models": []},
                                                                              raise_for_status=lambda: None),
                                           exceptions=requests.exceptions)
    api.estado.saber.ligada = False
    api._juiz = lambda: None
    api.check_ollama = lambda *a, **k: (True, "")
    api.plano_mod.liberada = lambda estado: True
    nuvem.chave = lambda estado, provedor: ""
    local = TestClient(api.app, headers=api.cabecalho_local())

    pasta = Path(api.estado.pasta)
    (pasta / "Clientes").mkdir(parents=True, exist_ok=True)
    original_docx = pasta / "Clientes" / "Contrato Alfa.docx"
    original_docx.write_bytes(docx_de_teste())
    sha_original = file_sha1(original_docx)
    original_pdf = pasta / "Notificação Beta.pdf"
    original_pdf.write_bytes(pdf_de_teste())
    api.estado.recarregar()

    def nova(titulo: str = "p") -> str:
        return local.post("/api/trabalhos", json={"pedido": titulo}).json()["id"]

    def lista_do_editor() -> set:
        return {d["id"] for d in local.get("/api/documentos").json()["documentos"]}

    # -------------------------------------------------------- redigido
    print("\no trabalho redigido vira documento da conversa")
    tid = nova("Contrato Gama")
    t = api.estado.trabalhos.obter(tid)
    redigido = ("Segue a minuta do contrato, conforme pedido.\n\nCONTRATO DE LOCAÇÃO\n\nCLÁUSULA 1ª – DO OBJETO\n\n"
                "O imóvel da Rua A.\n\n---\n\nNotas para o advogado\n\nConfira o endereço do imóvel.")
    t.dizer("paulus", redigido, proposta={"tipo": "levar_ao_editor", "titulo": "Contrato de locação", "campos": {}})
    api.estado.trabalhos.salvar(t)
    r = local.post(f"/api/trabalhos/{tid}/fazer", json={"tipo": "levar_ao_editor", "campos": {}})
    did = r.json()["proposta"]["campos"]["id"]
    doc = local.get(f"/api/documentos/{did}").json()
    checar(doc["situacao"] == "conversa" and doc["trabalho_id"] == tid and doc["da_conversa"], "é da conversa", doc["situacao"])
    checar("Notas para o advogado" not in doc["corpo"] and "Confira o endereço" not in doc["corpo"],
           "as notas para o advogado ficam fora do corpo")
    checar("Segue a minuta" not in doc["corpo"] and "CLÁUSULA 1ª" in doc["corpo"], "a frase de abertura também")
    checar(doc["origem"].get("tipo") == "redigido" and "Confira o endereço" in doc["origem"].get("notas", ""),
           "as notas guardadas na origem")
    checar('data-b="' in doc["corpo"], "o corpo gravado já vem com os blocos nomeados")
    checar(did not in lista_do_editor(), "fora da lista do Editor")
    lista = local.get(f"/api/trabalhos/{tid}/documentos").json()
    checar([d["id"] for d in lista["documentos"]] == [did], "listado na conversa")

    # ------------------------------------------------------------ .docx
    print("\no .docx do Acervo, com a estrutura")
    r = local.post("/api/documentos/importar", json={"nome": "Contrato Alfa.docx", "trabalho_id": tid})
    checar(r.status_code == 200, "importa", r.text[:200])
    alfa = r.json()["id"]
    doc = local.get(f"/api/documentos/{alfa}").json()
    blocos = B.ler(doc["corpo"])
    tipos = [b.tipo for b in blocos]
    checar(tipos[0] == "h1" and "table" in tipos and tipos.index("table") == 5, "título e quadro no lugar", tipos)
    checar(any(b.num == "1." for b in blocos) and any(b.num == "2." for b in blocos), "a numeração do Word como rótulo",
           [b.num for b in blocos])
    checar("<b>EMPRESA ALFA LTDA.</b>" in doc["corpo"] and 'text-align: justify' in doc["corpo"], "negrito e alinhamento")
    checar(doc["origem"]["formato"] == "docx" and doc["origem"]["cabecalho"] == ["SILVA ADVOGADOS ASSOCIADOS"],
           "a origem guarda o formato e o cabeçalho", doc["origem"])
    r2 = local.post("/api/documentos/importar", json={"nome": "Contrato Alfa.docx", "trabalho_id": tid})
    checar(r2.json()["id"] == alfa and r2.json()["reaberto"], "de novo na mesma conversa: o mesmo rascunho")
    tid_b = nova("Outra conversa")
    r3 = local.post("/api/documentos/importar", json={"nome": "Contrato Alfa.docx", "trabalho_id": tid_b})
    checar(r3.json()["id"] != alfa, "outra conversa: outro rascunho")

    # -------------------------------------------------------------- PDF
    print("\no PDF continua PDF")
    antes = api.estado.base.um("SELECT COUNT(*) AS n FROM documentos")["n"]
    r = local.post("/api/documentos/importar", json={"nome": "Notificação Beta.pdf", "trabalho_id": tid})
    checar(r.status_code == 409 and r.json().get("pdf") and "versão editável" in r.json()["detail"],
           "sem o pedido: 409, e nada criado", r.text[:200])
    checar(api.estado.base.um("SELECT COUNT(*) AS n FROM documentos")["n"] == antes, "nenhum documento novo")
    resposta = local.post(f"/api/trabalhos/{tid}/perguntar",
                          json={"pergunta": "abra a Notificação Beta no editor", "apenas": ["Notificação Beta.pdf"]})
    ev = eventos(resposta.text)
    proposta = next((d for k, d in ev if k == "proposta"), {})
    checar(proposta.get("tipo") == "editavel_de_pdf" and proposta["campos"]["nome"] == "Notificação Beta.pdf",
           "pela conversa: o cartão que oferece a versão editável", proposta)
    checar(api.estado.base.um("SELECT COUNT(*) AS n FROM documentos")["n"] == antes, "e o cartão não cria nada")
    r = local.post("/api/documentos/importar", json={"nome": "Notificação Beta.pdf", "trabalho_id": tid, "de_pdf": True})
    checar(r.status_code == 200, "com o pedido: a versão editável", r.text[:200])
    beta = local.get(f"/api/documentos/{r.json()['id']}").json()
    checar("[pagina" not in beta["corpo"] and "Parágrafo 7 da notificação" in beta["corpo"], "texto limpo, sem marca de página")
    paragrafos = [b.texto for b in B.ler(beta["corpo"])]
    checar(any(p.startswith("Parágrafo 7 da notificação") and p.endswith("item 49.") for p in paragrafos),
           "as linhas do PDF juntadas no parágrafo", [p[:60] for p in paragrafos[:3]])
    checar(len(paragrafos) == 40 and all("rodapé do escritório" not in p and "página" not in p for p in paragrafos),
           "um bloco por parágrafo, sem o rodapé nem o número da página", (len(paragrafos), paragrafos[-1][-60:]))
    checar(beta["origem"]["tipo"] == "pdf" and beta["situacao"] == "conversa", "origem PDF, da conversa")

    # -------------------------------------------------------- gravar
    print("\ngravar com a revisão")
    doc = local.get(f"/api/documentos/{did}").json()
    rev = doc["revisao"]
    versoes = len(local.get(f"/api/documentos/{did}/versoes").json()["versoes"])
    corpo2 = doc["corpo"].replace("O imóvel da Rua A.", "O imóvel da Rua B.")
    r = local.post(f"/api/documentos/{did}", json={"corpo": corpo2, "revisao": rev})
    checar(r.status_code == 200 and r.json()["revisao"] == rev + 1, "a gravação da revisão atual entra", r.text[:200])
    checar(len(local.get(f"/api/documentos/{did}/versoes").json()["versoes"]) == versoes,
           "o autosave não cria versão a cada pausa")
    r = local.post(f"/api/documentos/{did}", json={"corpo": corpo2.replace("Rua B", "Rua C"), "revisao": rev})
    checar(r.status_code == 409 and r.json().get("conflito") and "Rua B" in r.json()["corpo"],
           "a gravação de uma tela velha não passa por cima (409 com o gravado)", r.text[:200])
    r = local.post(f"/api/documentos/{did}", json={"corpo": corpo2.replace("Rua B", "Rua D"), "nota": "rua corrigida"})
    checar(r.json()["nova_versao"], "com nota, versão nova")

    # ----------------------------------------------------- sugestões
    print("\nsugestão pendente não sai do documento")
    doc = local.get(f"/api/documentos/{did}").json()
    bl = B.ler(doc["corpo"])
    alvo = next(b for b in bl if "Rua D" in b.texto)
    B.reescrever(alvo, alvo.texto.replace("Rua D", "Rua E"), "s1", 1)
    local.post(f"/api/documentos/{did}", json={"corpo": B.html(bl)})
    doc = local.get(f"/api/documentos/{did}").json()
    checar(doc["pendentes"] == 1, "uma sugestão pendente", doc["pendentes"])
    for rota in (f"/api/documentos/{did}/pdf", f"/api/documentos/{did}/docx"):
        r = local.get(rota)
        checar(r.status_code == 409 and "sugest" in r.json()["detail"], f"{rota.split('/')[-1]}: 409 com sugestão pendente")
    r = local.post(f"/api/documentos/{did}/concluir", json={})
    checar(r.status_code == 409, "concluir também espera a decisão")
    r = local.post(f"/api/documentos/{did}/sugestoes", json={"acao": "aceitar"})
    checar(r.status_code == 200 and r.json()["pendentes"] == 0 and "Rua E" in r.json()["corpo"], "aceitar todas")
    checar(local.get(f"/api/documentos/{did}/pdf").status_code == 200, "e o PDF sai")
    r = local.post(f"/api/documentos/{did}/biblioteca", json={"formato": "pdf"})
    checar(r.status_code == 409 and "conclua" in r.json()["detail"], "o rascunho da conversa não vai ao Acervo sem concluir")

    # ------------------------------------------------------ concluir
    print("\nconcluir")
    doc = local.get(f"/api/documentos/{alfa}").json()
    bl = B.ler(doc["corpo"])
    preco = next(b for b in bl if "5.000,00 (cinco mil reais)" in b.texto)
    B.reescrever(preco, "O preço total é de R$ 6.000,00 (seis mil reais).", "s9", 1)
    bl = B.aceito(bl)
    local.post(f"/api/documentos/{alfa}", json={"corpo": B.html(bl)})
    r = local.post(f"/api/documentos/{alfa}/concluir", json={})
    checar(r.status_code == 200, "conclui", r.text[:300])
    feito = r.json()
    destino = Path(feito["arquivo"])
    checar(destino.parent == original_docx.parent and destino.name == "Contrato Alfa - revisado.docx",
           "ao lado do original, com outro nome", destino)
    checar(file_sha1(original_docx) == sha_original, "o original não muda")
    checar(feito["como"] == "docx remendado do original", "o .docx é o original remendado", feito["como"])
    from docx import Document

    novo = Document(io.BytesIO(destino.read_bytes()))
    textos = [p.text for p in novo.paragraphs]
    checar("O preço total é de R$ 6.000,00 (seis mil reais)." in textos and not any("5.000,00 (cinco" in x for x in textos),
           "com a mudança aceita", textos[-3:])
    checar(novo.sections[0].header.paragraphs[0].text == "SILVA ADVOGADOS ASSOCIADOS" and
           novo.sections[0].footer.paragraphs[0].text.startswith("Rua da Bahia"), "cabeçalho e rodapé intactos")
    checar(any(p.style.name == "List Number" for p in novo.paragraphs) and len(novo.tables) == 1,
           "a numeração do Word e o quadro ficam")
    doc = local.get(f"/api/documentos/{alfa}").json()
    checar(doc["situacao"] == "concluido" and doc["arquivo"] == str(destino), "concluído, com o arquivo")
    checar(alfa in lista_do_editor(), "concluído aparece na lista do Editor")
    checar(any(d.name == destino.name for d in api.estado.searcher.documents), "e no Acervo")
    r = local.post(f"/api/documentos/{alfa}", json={"corpo": doc["corpo"].replace("6.000", "7.000")})
    checar(r.status_code == 409 and "reabra" in r.json()["detail"], "concluído não aceita gravação")
    r = local.post(f"/api/documentos/{alfa}/concluir", json={})
    checar(r.status_code == 409 and r.json().get("conflitos"), "concluir de novo com o mesmo nome pergunta", r.text[:200])
    r = local.post(f"/api/documentos/{alfa}/reabrir")
    checar(r.json()["situacao"] == "conversa", "reabrir volta a ser rascunho")
    r = local.post(f"/api/documentos/{alfa}/concluir",
                   json={"formato": "pdf", "decisoes": {}, "nome": "Contrato Alfa - assinatura"})
    checar(r.status_code == 200 and r.json()["nome"] == "Contrato Alfa - assinatura.pdf" and
           Path(r.json()["arquivo"]).read_bytes()[:4] == b"%PDF", "concluir em PDF, com outro nome", r.text[:200])
    r = local.post(f"/api/documentos/{did}/concluir", json={})
    checar(r.status_code == 200 and Path(r.json()["arquivo"]).parent.name == "Editor" and r.json()["formato"] == "docx",
           "o redigido conclui em DOCX, em Acervo › Editor", r.text[:200])

    # ---------------------------------------------------- ler_html
    print("\na leitura para o PDF")
    h = ('<p data-b="a">O prazo é de <del data-s="s1" data-r="1">30</del><ins data-s="s1" data-r="1">60</ins> dias.</p>'
         '<p data-b="b" data-s="s2" data-sug="del" data-r="1">Parágrafo que sai.</p>'
         '<p data-b="c" data-num="1.1">Item numerado.</p><p data-b="d"><del>tachado de verdade</del></p>')
    textos = [b.texto for b in D.ler_html(h)]
    checar(textos[0] == "O prazo é de 60 dias." and "Parágrafo que sai." not in textos, "o que sai não é lido", textos)
    checar("1.1 Item numerado." in textos, "o número do Word vai na frente", textos)
    tachado = next(b for b in D.ler_html(h) if "tachado" in b.texto)
    checar(tachado.trechos[0].riscado, "o tachado comum continua tachado")

    # ------------------------------------------------ lado e lixeira
    print("\no editor aberto ao lado e a lixeira")
    local.post(f"/api/trabalhos/{tid}/lado", json={"editor": beta["id"]})
    checar(local.get(f"/api/trabalhos/{tid}/documentos").json()["aberto"] == beta["id"], "a conversa lembra o editor aberto")
    r = local.delete(f"/api/trabalhos/{tid}")
    checar(r.status_code == 200, "apagar a conversa")
    checar(api.estado.documentos.obter(beta["id"]) is None, "o rascunho vai junto para a lixeira")
    concluido = api.estado.documentos.obter(alfa)
    checar(concluido is not None and concluido["trabalho_id"] == "", "o concluído fica, sem conversa")
    print()
    if _falhas:
        print(f"{len(_falhas)} falha(s): " + "; ".join(_falhas))
        return 1
    print("todos os testes passaram")
    return 0


if __name__ == "__main__":
    sys.exit(main())
