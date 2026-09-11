"""
Testes da folha, da papelada e do fechamento do mês.

Aqui o dano de errar não é uma tela feia — é um recibo com o valor errado, uma
nota que some do controle ou um mês que fecha dizendo que está tudo certo
quando não está. Então os testes cobrem:

  - a folha de agosto não muda quando o salário de hoje muda
  - estagiário não entra em "salários": bolsa não é salário
  - o recibo diz a natureza certa do pagamento — "bolsa de estágio", não "CLT"
  - nota registrada some da lista de "a registrar", e não antes
  - "contrato concluído" sai dos lançamentos, e some se abrir cobrança nova
  - o fechamento só diz "pronto" quando não há par faltando

    python tests/test_escritorio.py
"""

from __future__ import annotations

import sys
import tempfile
from datetime import date, timedelta
from pathlib import Path

RAIZ = Path(__file__).parent.parent
sys.path.insert(0, str(RAIZ / "src"))

import escritorio  # noqa: E402
from base import Base  # noqa: E402
from cadastros import Cadastros  # noqa: E402
from financeiro import Financeiro  # noqa: E402

_falhas: list[str] = []


def checar(condicao: bool, descricao: str, detalhe: str = "") -> None:
    if condicao:
        print(f"  ok   {descricao}")
    else:
        print(f"  FALHA {descricao}")
        if detalhe:
            print(f"         {detalhe}")
        _falhas.append(descricao)


def _dia(delta: int = 0) -> str:
    return (date.today() + timedelta(days=delta)).isoformat()


def _mundo(tmp: Path, nome: str):
    base = Base(tmp / f"{nome}.db")
    return base, Cadastros(base), Financeiro(base), escritorio.Folha(base), escritorio.PapeisFiscais(base)


def test_dias_para_fechar() -> None:
    print("\nquantos dias faltam para o mês fechar")
    hoje = date.today()
    faltam = escritorio.dias_para_fechar(hoje.strftime("%Y-%m"))
    checar(0 <= faltam <= 31, f"o mês de hoje devolve um número plausível ({faltam})")

    passado = (hoje.replace(day=1) - timedelta(days=1)).strftime("%Y-%m")
    checar(escritorio.dias_para_fechar(passado) == 0, "mês que já passou devolve zero, não negativo")
    checar(escritorio.dias_para_fechar("não é mês") == 0, "texto que não é mês não quebra")

    checar(escritorio.ultimo_dia("2026-02") == "2026-02-28", "fevereiro comum acaba no 28")
    checar(escritorio.ultimo_dia("2024-02") == "2024-02-29", "e o bissexto, no 29")


def test_folha(tmp: Path) -> None:
    print("\na folha do mês")
    base, cad, _, folha, _ = _mundo(tmp, "folha")
    try:
        try:
            folha.montar()
            checar(False, "sem ninguém com vínculo, recusa montar")
        except ValueError as exc:
            checar("vínculo" in str(exc), f"sem ninguém com vínculo, recusa montar ({exc})")

        cad.salvar({"tipo": "colaborador", "nome": "Ana", "vinculo": "clt",
                    "salario_centavos": 520000, "encargos_centavos": 113000})
        cad.salvar({"tipo": "colaborador", "nome": "Eva", "vinculo": "estagio",
                    "salario_centavos": 60000})
        # Sem vínculo definido: não entra, e não entrar é o certo.
        cad.salvar({"tipo": "colaborador", "nome": "Sem vínculo"})
        cad.salvar({"tipo": "cliente", "nome": "Um cliente qualquer"})

        f = folha.montar()
        checar(f["quantos"] == 2, f"entra quem tem vínculo, e só ({f['quantos']})")
        checar(f["salarios"] == 520000, f"salários somam só os não-estagiários ({f['salarios']})")
        checar(f["estagios"] == 60000, "a bolsa de estágio conta separado")
        checar(f["encargos"] == 113000, "e os encargos, também")
        checar(f["total"] == 693000, f"o total é a soma dos três ({f['total']})")
        checar(f["total_texto"] == "R$ 6.930,00", f"formatado em reais ({f['total_texto']})")

        # A razão de a folha ser cópia: o mês já montado não pode mudar sozinho.
        ana = base.um("SELECT id FROM cadastros WHERE nome = 'Ana'")["id"]
        cad.salvar({"tipo": "colaborador", "nome": "Ana", "vinculo": "clt",
                    "salario_centavos": 900000, "encargos_centavos": 200000}, ana)
        depois = folha.do_mes()
        checar(depois["total"] == 693000,
               f"aumentar o salário hoje NÃO reescreve o mês já montado ({depois['total']})")

        refeita = folha.montar()
        checar(refeita["total"] == 1160000,
               f"mas remontar de propósito pega os valores novos ({refeita['total']})")

        folha.ligar_ao_lancamento(refeita["mes"], 42)
        try:
            folha.montar()
            checar(False, "folha já lançada não se remonta por acidente")
        except ValueError as exc:
            checar("lançada" in str(exc) or "pagar" in str(exc),
                   f"folha já lançada não se remonta por acidente ({exc})")
    finally:
        base.fechar()


def test_recibos(tmp: Path) -> None:
    print("\nos recibos")
    base, cad, _, folha, _ = _mundo(tmp, "recibos")
    try:
        cad.salvar({"tipo": "colaborador", "nome": "Ana Advogada", "vinculo": "clt",
                    "salario_centavos": 520000, "encargos_centavos": 113000})
        cad.salvar({"tipo": "colaborador", "nome": "Eva Estagiária", "vinculo": "estagio",
                    "salario_centavos": 60000})
        cad.salvar({"tipo": "socio", "nome": "Sócio Fulano", "vinculo": "prolabore",
                    "salario_centavos": 1500000})
        f = folha.montar()

        destino = tmp / "recibos"
        feitos = escritorio.gerar_recibos(f, destino, "setembro de 2026", "Escritório de Teste")
        checar(len(feitos) == 3, f"um recibo por pessoa ({len(feitos)})")
        checar(all(Path(r["arquivo"]).exists() for r in feitos), "os três arquivos existem")

        from pypdf import PdfReader
        textos = {}
        for r in feitos:
            textos[r["nome"]] = PdfReader(r["arquivo"]).pages[0].extract_text().replace("\n", " ")

        # Chamar bolsa de estágio de "salário" num recibo é problema
        # trabalhista, não problema de texto.
        checar("bolsa de estágio" in textos["Eva Estagiária"],
               "o recibo do estagiário fala em bolsa de estágio")
        checar("salário" not in textos["Eva Estagiária"].lower(),
               "e não fala em salário")
        checar("salário" in textos["Ana Advogada"], "o da CLT fala em salário")
        checar("pró-labore" in textos["Sócio Fulano"], "o do sócio fala em pró-labore")

        checar("R$ 6.330,00" in textos["Ana Advogada"],
               "o valor do recibo é salário mais encargos")
        checar("Escritório de Teste" in textos["Ana Advogada"], "e diz quem pagou")
        checar("sem assinatura" in textos["Ana Advogada"],
               "o recibo avisa que sai sem assinatura — fingir assinatura seria pior")
    finally:
        base.fechar()


def test_notas_e_boletos(tmp: Path) -> None:
    print("\nnotas e boletos")
    base, cad, fin, _, papeis = _mundo(tmp, "papeis")
    try:
        cliente = cad.salvar({"tipo": "cliente", "nome": "Cooperativa"})
        recebido = fin.salvar({"tipo": "recebimento", "descricao": "Honorários",
                               "centavos": 1200000, "cadastro_id": cliente,
                               "vencimento": _dia(-3)})
        fin.liquidar(recebido)

        faltando = papeis.a_emitir()
        checar(len(faltando) == 1 and faltando[0]["id"] == recebido,
               "recebimento sem nota aparece em 'a registrar'")

        papeis.salvar({"tipo": "nota", "numero": "0142", "cadastro_id": cliente,
                       "lancamento_id": recebido, "centavos": 1200000, "data": _dia(-3)})
        checar(papeis.a_emitir() == [], "e some da lista quando a nota é registrada")

        notas = papeis.listar("nota", date.today().strftime("%Y-%m"))
        checar(len(notas) == 1 and notas[0]["cliente"] == "Cooperativa",
               "a nota sai com o nome do cliente junto")
        checar(notas[0]["valor"] == "R$ 12.000,00", "e com o valor formatado")

        try:
            papeis.salvar({"tipo": "recibo", "centavos": 100})
            checar(False, "tipo que não existe é recusado")
        except ValueError as exc:
            checar("nota" in str(exc) or "boleto" in str(exc),
                   f"tipo que não existe é recusado ({exc})")
        try:
            papeis.salvar({"tipo": "boleto", "centavos": 0})
            checar(False, "boleto de zero é recusado")
        except ValueError as exc:
            checar("valor" in str(exc), f"boleto de zero é recusado ({exc})")

        bid = papeis.salvar({"tipo": "boleto", "cadastro_id": cliente,
                             "centavos": 540000, "data": _dia(20)})
        abertos = papeis.listar("boleto")
        checar(abertos[0]["situacao"] == "aberto", "boleto nasce em aberto")
        checar(abertos[0]["situacao_rotulo"] == "na fila", "e a tela chama de 'na fila'")

        papeis.marcar_pago(bid)
        checar(papeis.listar("boleto")[0]["situacao"] == "pago", "baixar marca como pago")

        checar(papeis.apagar(bid), "apagar o registro funciona")
        checar(fin.obter(recebido) is not None,
               "e apagar o registro NÃO apaga o lançamento — são coisas diferentes")
    finally:
        base.fechar()


def test_concluidos_e_fechamento(tmp: Path) -> None:
    print("\ncontratos concluídos e o fechamento")
    base, cad, fin, folha, papeis = _mundo(tmp, "fecha")
    try:
        cliente = cad.salvar({"tipo": "cliente", "nome": "Condomínio 402"})
        outro = cad.salvar({"tipo": "cliente", "nome": "Fornecedor A"})

        quitado = fin.salvar({"tipo": "recebimento", "descricao": "Assembleia",
                              "centavos": 320000, "cadastro_id": cliente, "vencimento": _dia(-5)})
        fin.liquidar(quitado)
        fin.salvar({"tipo": "recebimento", "descricao": "Em aberto", "centavos": 850000,
                    "cadastro_id": outro, "vencimento": _dia(-12)})

        feitos = escritorio.contratos_concluidos(base)
        checar([c["nome"] for c in feitos] == ["Condomínio 402"],
               f"só quem não tem nada em aberto aparece ({[c['nome'] for c in feitos]})")
        checar(feitos[0]["recebido_texto"] == "R$ 3.200,00", "com o que recebeu no mês")

        # Abrir cobrança nova tira o cliente da lista, sem ninguém desmarcar
        # nada: é isso que um campo "encerrado" não faria.
        fin.salvar({"tipo": "recebimento", "descricao": "Novo trabalho", "centavos": 100000,
                    "cadastro_id": cliente, "vencimento": _dia(10)})
        checar(escritorio.contratos_concluidos(base) == [],
               "abrir cobrança nova tira da lista sozinho")

        despesa = fin.salvar({"tipo": "despesa", "descricao": "Aluguel", "centavos": 480000,
                              "vencimento": _dia(-4)})
        fin.liquidar(despesa)

        f = escritorio.fechamento(base, folha, papeis)
        oq = {x["o_que"] for x in f["faltas"]}
        checar("comprovante" in oq, f"despesa paga sem comprovante entra no fechamento ({oq})")
        checar("nota" in oq, "recebimento sem nota também")
        checar(not f["pronto"], "e o mês não está pronto")

        fin.anexar(despesa, "comprovante.pdf")
        papeis.salvar({"tipo": "nota", "numero": "1", "lancamento_id": quitado,
                       "centavos": 320000, "data": _dia(-5)})
        f2 = escritorio.fechamento(base, folha, papeis)
        checar(f2["pronto"], f"com os pares completos, o mês fecha ({[x['o_que'] for x in f2['faltas']]})")

        # Boleto vencido e ainda aberto tem que aparecer: e um vencido no
        # futuro, nao.
        papeis.salvar({"tipo": "boleto", "centavos": 100, "data": _dia(-2)})
        papeis.salvar({"tipo": "boleto", "centavos": 100, "data": _dia(30)})
        f3 = escritorio.fechamento(base, folha, papeis)
        vencidos = [x for x in f3["faltas"] if x["o_que"] == "boleto"]
        checar(len(vencidos) == 1 and vencidos[0]["quantos"] == 1,
               f"só o boleto vencido entra, não o que ainda vai vencer ({vencidos})")
    finally:
        base.fechar()


def test_exportar(tmp: Path) -> None:
    print("\na planilha do mês")
    base, cad, fin, folha, papeis = _mundo(tmp, "exporta")
    try:
        cliente = cad.salvar({"tipo": "cliente", "nome": "Cooperativa"})
        cad.salvar({"tipo": "colaborador", "nome": "Ana", "vinculo": "clt",
                    "salario_centavos": 520000, "encargos_centavos": 113000})
        folha.montar()
        rid = fin.salvar({"tipo": "recebimento", "descricao": "Honorários", "centavos": 1200000,
                          "cadastro_id": cliente, "vencimento": _dia(-2)})
        fin.liquidar(rid)
        papeis.salvar({"tipo": "nota", "numero": "0142", "cadastro_id": cliente,
                       "centavos": 1200000, "data": _dia(-2)})

        mes = date.today().strftime("%Y-%m")
        alvo = escritorio.exportar_mes(tmp / "mes.xlsx", mes,
                                       extrato=fin.extrato(mes), folha=folha.do_mes(mes),
                                       notas=papeis.listar("nota", mes),
                                       boletos=papeis.listar("boleto"))
        checar(alvo.exists() and alvo.stat().st_size > 3000, "a planilha sai com conteúdo")

        from openpyxl import load_workbook
        livro = load_workbook(alvo)
        checar(livro.sheetnames == ["Extrato", "Folha", "Notas e boletos"],
               f"uma aba por assunto ({livro.sheetnames})")

        extrato = livro["Extrato"]
        checar(extrato.max_row == 2, f"o lançamento do mês está lá ({extrato.max_row - 1})")

        # Planilha que chega com "R$ 1.200,00" em celula de texto nao soma do
        # outro lado - e o contador soma.
        valor = extrato.cell(row=2, column=6).value
        checar(isinstance(valor, (int, float)), f"o valor vai como número, não como texto ({valor!r})")
        checar(abs(valor - 12000.0) < 0.001, f"e com o valor certo ({valor})")

        da_folha = livro["Folha"]
        checar(da_folha.cell(row=2, column=5).value == 6330.0,
               "a folha vai com o total por pessoa")
    finally:
        base.fechar()


def main() -> int:
    print("=" * 55)
    print("PAULUS - o escritório por dentro")
    print("=" * 55)

    test_dias_para_fechar()
    with tempfile.TemporaryDirectory() as bruto:
        tmp = Path(bruto)
        test_folha(tmp)
        test_recibos(tmp)
        test_notas_e_boletos(tmp)
        test_concluidos_e_fechamento(tmp)
        test_exportar(tmp)

    print("\n" + "=" * 55)
    if _falhas:
        print(f"  {len(_falhas)} FALHA(S):")
        for f in _falhas:
            print(f"    - {f}")
        return 1
    print("  todos os testes passaram")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
