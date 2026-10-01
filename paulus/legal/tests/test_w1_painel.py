"""
W1 - o painel do PAVLVS no Word, o pareamento e a porta do suplemento
(src/word_suplemento.py, src/word_instalar.py, word/painel.*).

  - sem token: 401 em todas as rotas do suplemento; token revogado: 401 na
    hora; rota de suplemento fora da tabela: 403 (padrão "nega");
  - o pareamento só se completa com o Permitir da janela local (pedido deste
    computador) ou, de fora, com a sessão e o código do autenticador; o token
    sai uma vez e só o resumo fica guardado;
  - token de colaborador não alcança rota de titular nem módulo fechado;
  - token deste computador não vale pelo túnel;
  - toda chamada vai para a auditoria com o documento e os caracteres;
  - instalar: certificado só-localhost, porta HTTPS fixa, manifesto PAVLVS;
  - o painel no Edge, pela porta HTTPS, com o Office simulado: conectar,
    tema escuro, Word antigo, fora do Word, 320 px, e nenhuma requisição fora
    da CSP.

    PYTHONIOENCODING=utf-8 venv/Scripts/python.exe tests/test_w1_painel.py
"""

from __future__ import annotations

import json
import os
import sys
import tempfile
import time
from pathlib import Path

RAIZ = Path(__file__).parent.parent
TMP = Path(tempfile.mkdtemp(prefix="paulus-w1-"))
os.environ["PAULUS_DADOS"] = str(TMP / "dados")
os.environ["PAULUS_SEM_AVISOS"] = "1"
sys.path.insert(0, str(RAIZ / "src"))
sys.path.insert(0, str(RAIZ / "tests"))

_falhas: list[str] = []


def checar(cond, nome: str, detalhe=None) -> None:
    print(("  ok   " if cond else "  FALHOU  ") + nome + ("" if cond or detalhe is None else f"  -> {detalhe!r}"))
    if not cond:
        _falhas.append(nome)


FORA = {"Cf-Connecting-IP": "200.9.9.9"}


def test_guardados() -> None:
    print("\nos guardados: token uma vez, só o resumo no disco")
    import word_suplemento as w

    c = w.Conexoes(TMP / "u")
    conexao, token = c.criar(pessoa="Ana", email="", papel="titular", conta_id=0, origem="local", word="16.0")
    checar(token.startswith(w.PREFIXO_TOKEN) and "resumo" not in conexao, "o token volta só na criação")
    bruto = (TMP / "u" / "word" / "conexoes.json").read_text(encoding="utf-8")
    checar(token not in bruto and w._resumo(token) in bruto, "no disco, só o resumo")
    checar(c.do_token(token)["id"] == conexao["id"] and c.do_token(token + "x") is None and c.do_token("") is None,
           "o token certo acha a conexão; outro não")
    c.revogar(conexao["id"])
    checar(c.do_token(token) is None, "revogado, não acha mais")

    p = w.Pedidos()
    a = p.abrir(origem="local", endereco="127.0.0.1", word="16.0")
    checar(len(a["codigo"]) == 9 and a["codigo"][4] == "-" and all(ch in w.ALFABETO for ch in a["codigo"].replace("-", "")),
           "código XXXX-XXXX, sem letra que se confunda", a["codigo"])
    checar(p.trocar(a["id"], "errado")["estado"] == "vencido", "segredo errado não troca")
    checar(p.trocar(a["id"], a["segredo"])["estado"] == "esperando", "sem o Permitir, espera")
    checar(p.do_codigo(a["codigo"].lower().replace("-", " "), "local")["id"] == a["id"], "o código digitado sem traço e em minúsculas acha")
    checar(p.do_codigo(a["codigo"], "fora") is None, "o código de um pedido deste computador não serve de fora")
    p.decidir(a["id"], conexao={"pessoa": "Ana", "papel": "titular"}, token="t")
    r = p.trocar(a["id"], a["segredo"])
    checar(r["estado"] == "permitido" and r["token"] == "t", "permitido: o token sai")
    checar(p.trocar(a["id"], a["segredo"])["estado"] == "vencido", "e só uma vez")
    relogio = [1000.0]
    q = w.Pedidos(relogio=lambda: relogio[0])
    b = q.abrir(origem="local", endereco="x", word="")
    relogio[0] += w.PEDIDO_VALE_S + 1
    checar(q.trocar(b["id"], b["segredo"])["estado"] == "vencido", "depois de 10 minutos, vence")
    for _ in range(w.PEDIDOS_POR_ORIGEM):
        q.abrir(origem="local", endereco="y", word="")
    try:
        q.abrir(origem="local", endereco="y", word="")
        checar(False, "pedidos demais do mesmo endereço: recusa")
    except PermissionError:
        checar(True, "pedidos demais do mesmo endereço: recusa")


def test_http() -> None:
    print("\no suplemento com o app de verdade")
    import segredos

    if not segredos.disponivel():
        print("  pulado: sem DPAPI (fora do Windows)")
        return
    from fastapi.testclient import TestClient

    import api
    import word_suplemento as w
    from acesso.contas import codigo_totp

    local = TestClient(api.app, headers=api.cabecalho_local(), client=("127.0.0.1", 50000))
    # O painel não tem a chave da janela: é um cliente de 127.0.0.1 sem ela.
    painel = TestClient(api.app, base_url="https://localhost", client=("127.0.0.1", 50001))

    print("  desligado de fábrica")
    checar(api.estado.prefs.dados["word"]["ligado"] is False, "word.ligado vem desligado")
    checar(painel.get("/word/painel.html").status_code == 404, "desligado: o painel não abre")
    checar(painel.post("/api/word/parear", json={}).status_code == 404, "desligado: não pareia")
    r = local.post("/api/word/ligar", json={"ligado": True})
    checar(r.status_code == 200 and r.json()["ligado"] is True, "ligar pela janela", r.text[:200])

    print("  a página do painel")
    r = painel.get("/word/painel.html")
    csp = r.headers.get("content-security-policy", "")
    checar(r.status_code == 200 and "PAVLVS" in r.text, "o painel abre sem a chave da janela", r.status_code)
    checar("script-src 'self' https://appsforoffice.microsoft.com;" in csp and "connect-src 'self';" in csp,
           "CSP: script só daqui e do office.js; conexão só daqui", csp)
    checar("frame-ancestors 'self' https://*.officeapps.live.com" in csp and r.headers.get("x-frame-options") is None,
           "o Word na web pode pôr o painel em moldura (e só ele)")
    checar(painel.get("/word/prova/maquete.html").status_code == 404 and painel.get("/word/../src/api.py").status_code in (403, 404),
           "só os arquivos do painel saem")
    r = painel.get("/word/icones/conferir-80.png")
    checar(r.headers.get("content-type") == "image/png" and "no-store" not in r.headers.get("cache-control", ""),
           "os ícones saem, e podem ficar no cache do Office", r.headers.get("cache-control"))
    checar(painel.get("/api/status").status_code == 403, "sem a chave, o resto do PAULUS continua fechado")

    print("  sem token: 401 em todas as rotas do suplemento")
    for (metodo, caminho) in sorted(w.ROTAS):
        r = painel.request(metodo, caminho, json={})
        checar(r.status_code == 401, f"{metodo} {caminho} sem token: 401", r.status_code)
    r = painel.get("/api/word/s/eu", headers={"Authorization": "Bearer paulus_word_inventado"})
    checar(r.status_code == 401, "token inventado: 401", r.status_code)
    r = local.get("/api/word/s/eu")
    checar(r.status_code == 401, "nem a janela local passa sem token", r.status_code)

    print("  o pareamento deste computador")
    r = painel.post("/api/word/parear", json={"word": "16.0.14334.20918"})
    checar(r.status_code == 200 and r.json()["origem"] == "local", "o painel pede e recebe um código", r.text[:200])
    ped = r.json()
    r = painel.post("/api/word/parear/trocar", json={"pedido": ped["pedido"], "segredo": ped["segredo"]})
    checar(r.json() == {"estado": "esperando"}, "sem Permitir: espera", r.text)
    r = local.get("/api/word/pedidos")
    checar([p["codigo"] for p in r.json()["pedidos"]] == [ped["codigo"]], "a janela local vê o pedido com o código", r.text[:200])
    r = painel.post(f"/api/word/pedidos/{ped['pedido']}/permitir")
    checar(r.status_code == 403, "o próprio painel não se permite (sem a chave da janela)", r.status_code)
    r = local.post(f"/api/word/pedidos/{ped['pedido']}/permitir")
    checar(r.status_code == 200 and len(r.json()["conexoes"]) == 1, "a janela local permite", r.text[:200])
    r = painel.post("/api/word/parear/trocar", json={"pedido": ped["pedido"], "segredo": ped["segredo"]})
    checar(r.json().get("estado") == "permitido" and r.json().get("papel") == "titular", "o painel troca pelo token (titular)", r.text[:200])
    token = r.json()["token"]
    r = painel.post("/api/word/parear/trocar", json={"pedido": ped["pedido"], "segredo": ped["segredo"]})
    checar(r.json() == {"estado": "vencido"}, "o token sai uma vez só", r.text)
    bruto = (Path(os.environ["PAULUS_DADOS"]) / "word" / "conexoes.json").read_text(encoding="utf-8")
    checar(token not in bruto, "o token não está no disco")

    com_token = {"Authorization": "Bearer " + token, "X-PAULUS-Documento": "Peti%C3%A7%C3%A3o%20inicial.docx"}
    r = painel.get("/api/word/s/eu", headers=com_token)
    checar(r.status_code == 200 and r.json()["papel"] == "titular" and r.json()["origem"] == "local", "com token: quem sou", r.text[:200])
    r = painel.put("/api/word/s/preferencias", headers=com_token, json={"controlada": False})
    checar(r.status_code == 200 and r.json()["preferencias"]["controlada"] is False, "a preferência fica na conexão (por pessoa)")
    r = painel.get("/api/word/s/eu", headers={**com_token, **FORA})
    checar(r.status_code in (401, 403), "token deste computador não vale pelo túnel", r.status_code)

    print("  padrão nega")
    @api.app.get("/api/word/s/teste-sem-declaracao")
    def _sem() -> dict:
        return {"ok": True}

    @api.app.get("/api/word/s/teste-titular")
    def _tit() -> dict:
        return {"ok": True}

    @api.app.get("/api/word/s/teste-agenda")
    def _agenda() -> dict:
        return {"ok": True}

    w.declarar("GET /api/word/s/teste-titular", papel="titular")
    w.declarar("GET /api/word/s/teste-agenda", modulo="agenda", nivel="ver")
    r = painel.get("/api/word/s/teste-sem-declaracao", headers=com_token)
    checar(r.status_code == 403, "rota de suplemento fora da tabela: 403, mesmo com token", r.status_code)
    r = painel.get("/api/word/s/teste-titular", headers=com_token)
    checar(r.status_code == 200, "titular alcança rota de titular", r.text[:200])

    print("  a auditoria")
    linhas = [json.loads(l) for l in (Path(os.environ["PAULUS_DADOS"]) / "acesso" / "acessos.jsonl").read_text(encoding="utf-8").splitlines()]
    word = [l for l in linhas if l["acao"] == "word"]
    checar(any("Petição inicial.docx" in l["alvo"] and "caracteres" in l["alvo"] for l in word),
           "cada chamada: a ação, o documento e os caracteres", [l["alvo"] for l in word][-3:])
    put = [l["alvo"] for l in word if "PUT /api/word/s/preferencias" in l["alvo"]]
    checar(put and put[-1].endswith(f"{len(json.dumps({'controlada': False}, separators=(',', ':')))} caracteres"),
           "os caracteres contados são os que chegaram", put)
    checar(any(l["acao"] == "word_pareado" for l in linhas) and any(l["acao"] == "word_recusado" for l in linhas),
           "pareamento e recusa também ficam")

    print("  revogar vale na hora")
    cid = local.get("/api/word").json()["conexoes"][0]["id"]
    checar(local.delete(f"/api/word/conexoes/{cid}").status_code == 200, "a janela revoga")
    checar(painel.get("/api/word/s/eu", headers=com_token).status_code == 401, "token revogado: 401 na hora")

    print("  de fora (montagem C): sessão + autenticador")
    servico = api.estado.acesso_de_fora
    servico.conferir_turnstile = lambda token, ip="": "ok" if token == "ok" else "recusado"
    api.estado.prefs.dados["acesso_remoto"]["ligado"] = True
    api.estado.prefs.dados["acesso_remoto"]["so_google"] = False
    passo = int(time.time() // 30) - 1

    def entrar(nome, email, papel):
        c = servico.contas.criar(nome, email, papel, "senha-forte-" + nome.lower())
        servico.contas.confirmar_totp(c["conta"]["id"], codigo_totp(c["segredo"], passo))
        f = TestClient(api.app, base_url="https://x.paulus.ia.br", headers=FORA, client=("127.0.0.1", 50002))
        pend = f.post("/api/acesso/entrar", json={"email": email, "senha": "senha-forte-" + nome.lower(), "turnstile": "ok"}).json()["pendente"]
        r = f.post("/api/acesso/entrar/codigo", json={"pendente": pend, "codigo": codigo_totp(c["segredo"], passo + 1)})
        f.headers["X-PAULUS-CSRF"] = r.json()["csrf"]
        return c["conta"]["id"], c["segredo"], f

    tita_id, _tita_seg, tita = entrar("Tita", "tita@x.com", "titular")
    caio_id, caio_seg, caio = entrar("Caio", "caio@x.com", "colaborador")
    painel_fora = TestClient(api.app, base_url="https://x.paulus.ia.br", headers=FORA, client=("127.0.0.1", 50003))
    r = painel_fora.get("/word/painel.html")
    checar(r.status_code == 200, "o painel abre pelo túnel sem sessão (montagem C)", r.status_code)
    r = painel_fora.post("/api/word/parear", json={"word": "16.0 web"})
    checar(r.json().get("origem") == "fora", "o pedido pelo túnel é de fora", r.text[:200])
    pf = r.json()
    checar(local.get("/api/word/pedidos").json()["pedidos"] == [], "a janela local não lista (nem permite) pedido de fora")
    r = local.post(f"/api/word/pedidos/{pf['pedido']}/permitir")
    checar(r.status_code == 403, "a janela local não permite pedido de fora", r.status_code)
    checar(caio.get("/api/word/pedidos").json()["pedidos"] == [], "de fora, a lista de pedidos volta vazia")
    r = caio.post(f"/api/word/pedidos/{pf['pedido']}/permitir")
    checar(r.status_code == 403, "de fora, a rota da janela é bloqueada", r.status_code)
    r = painel_fora.post("/api/word/permitir-de-fora", json={"codigo": pf["codigo"], "codigo_autenticador": "000000"})
    checar(r.status_code in (401, 403), "sem sessão, não permite", r.status_code)
    r = caio.post("/api/word/permitir-de-fora", json={"codigo": pf["codigo"], "codigo_autenticador": "123456"})
    checar(r.status_code == 403 and "autenticador" in r.text, "com sessão e autenticador errado, não permite", r.text[:160])
    # O autenticador vem antes do código do painel: sem ele não dá nem para
    # saber se um código existe. (Cada código do autenticador vale uma vez.)
    r = caio.post("/api/word/permitir-de-fora", json={"codigo": pf["codigo"], "codigo_autenticador": codigo_totp(caio_seg, passo + 2)})
    checar(r.status_code == 200, "sessão + autenticador certo: permite", r.text[:200])
    r = painel_fora.post("/api/word/parear/trocar", json={"pedido": pf["pedido"], "segredo": pf["segredo"]})
    checar(r.json().get("papel") == "colaborador" and r.json().get("pessoa") == "Caio", "o token tem o papel de quem permitiu", r.text[:200])
    tok_caio = {"Authorization": "Bearer " + r.json()["token"]}
    checar(painel_fora.get("/api/word/s/eu", headers=tok_caio).status_code == 200, "o token de fora vale pelo túnel")
    r = painel_fora.get("/api/word/s/teste-titular", headers=tok_caio)
    checar(r.status_code == 403 and "titular" in r.text, "token de colaborador não alcança rota de titular", r.text[:160])
    checar(painel_fora.get("/api/word/s/teste-agenda", headers=tok_caio).status_code == 200, "módulo aberto para o colaborador: passa")
    local.put(f"/api/acesso/contas/{caio_id}/permissoes", json={"niveis": {"agenda": "nao"}})
    r = painel_fora.get("/api/word/s/teste-agenda", headers=tok_caio)
    checar(r.status_code == 403 and "Agenda" in r.text, "o titular fecha o módulo e vale no Word na hora", r.text[:160])
    api.estado.acesso_de_fora.contas.remover(caio_id)
    checar(painel_fora.get("/api/word/s/eu", headers=tok_caio).status_code == 401, "conta removida: o token cai")
    api.estado.prefs.dados["acesso_remoto"]["ligado"] = False
    checar(painel_fora.get("/word/painel.html").status_code == 403, "acesso de fora desligado: o túnel não abre o painel")
    del tita_id, tita

    print("  instalar (confiança e registro simulados)")
    import word_instalar

    inst = api.estado.word_instalacao
    orig = word_instalar.confiado
    pedidos_de_confianca = []
    try:
        word_instalar.confiado = lambda der: bool(pedidos_de_confianca)
        passos = inst.instalar(pedir=lambda arq: pedidos_de_confianca.append(arq) or 0, registrar=lambda m: None)
    finally:
        word_instalar.confiado = orig
    nomes = [p["passo"] for p in passos]
    checar(nomes == ["certificado", "confiança", "porta", "manifesto", "registro", "servidor"], "os passos, em ordem", passos)
    checar(len(pedidos_de_confianca) == 1, "pediu a confiança ao Windows uma vez")
    porta = inst.porta()
    checar(46300 <= porta < 46400 and passos[-1]["resultado"] == "no ar", "porta fixa e o servidor HTTPS no ar", passos[-1])
    xml = inst.manifesto_local().read_text(encoding="utf-8")
    checar(f"https://localhost:{porta}/word/painel.html?tela=inicio" in xml and '<DisplayName DefaultValue="PAVLVS"/>' in xml
           and '<bt:String id="Aba" DefaultValue="PAVLVS"/>' in xml, "manifesto com o endereço e o título PAVLVS")
    checar(inst.manifesto("local") == inst.manifesto("local") and inst._id("local") != inst._id("fora"),
           "o Id é fixo por instalação e diferente de fora")
    from cryptography import x509

    ac = x509.load_der_x509_certificate((inst.pasta_cert / "autoridade.cer").read_bytes())
    nc = ac.extensions.get_extension_for_class(x509.NameConstraints)
    checar(nc.critical and {str(n.value) for n in nc.value.permitted_subtrees} == {"localhost", "127.0.0.1/32"},
           "a autoridade só vale para localhost e 127.0.0.1")
    checar(sorted(p.name for p in inst.pasta_cert.iterdir()) == ["autoridade.cer", "certificado.json", "servidor.key", "servidor.pem"],
           "a chave da autoridade não fica no disco")
    sit = local.get("/api/word").json()["instalacao"]
    checar(sit["carregado"] is False, "instalado, mas ainda não diz 'abriu no Word'")
    checar(sit["instalado"] is True, "a situação diz instalado")

    test_documento_leva_o_pavlvs(local, inst)
    return porta


def test_documento_leva_o_pavlvs(local, inst) -> None:
    """
    No Word 2021 a aba só aparece no documento que traz o PAVLVS (medido em
    01/10/2026): o PAULUS abre o Word com um documento seu e todo .docx que
    ele gera leva o PAVLVS.
    """
    print("  o documento leva o PAVLVS")
    import io
    import zipfile

    import docx

    import api
    import word_instalar as wi

    def novo(txt: str) -> bytes:
        b = io.BytesIO()
        d = docx.Document()
        d.add_paragraph(txt)
        d.save(b)
        return b.getvalue()

    id_local = inst._id("local")
    original = novo("Petição de teste.")
    com = wi.com_pavlvs(original, id_=id_local, versao="0.9.23", abrir_painel=True)
    z = zipfile.ZipFile(io.BytesIO(com))
    checar(wi.tem_pavlvs(com, id_local) and "word/webextensions/taskpanes.xml" in z.namelist(), "com_pavlvs põe a referência e o painel")
    checar(z.read("word/document.xml") == zipfile.ZipFile(io.BytesIO(original)).read("word/document.xml"),
           "o texto do documento não muda")
    checar(docx.Document(io.BytesIO(com)).paragraphs[0].text == "Petição de teste.", "o .docx continua abrindo")
    checar(wi.com_pavlvs(com, id_=id_local, versao="0.9.23", abrir_painel=True) == com, "documento que já traz o PAVLVS volta igual")
    outro = wi.com_pavlvs(original, id_="outro-suplemento", versao="1", abrir_painel=True)
    checar(wi.com_pavlvs(outro, id_=id_local, versao="0.9.23", abrir_painel=True) == outro,
           "documento com outro suplemento não é mexido")
    checar(wi.com_pavlvs(b"nao e zip", id_=id_local, versao="1", abrir_painel=True) == b"nao e zip", "arquivo que não é .docx volta igual")

    abertos: list[str] = []
    api.estado.word_abrir_no_windows = lambda caminho: abertos.append(str(caminho))
    inst_precisa = inst.precisa_fechar_o_word
    try:
        inst.precisa_fechar_o_word = lambda **kw: False
        r = local.post("/api/word/abrir")
        checar(r.status_code == 200 and r.json()["aberto"] is True and abertos and abertos[-1].endswith("PAVLVS - comece aqui.docx"),
               "Abrir o Word: abre o 'comece aqui'", r.text[:200])
        checar(wi.tem_pavlvs(Path(abertos[-1]).read_bytes(), id_local), "o 'comece aqui' traz o PAVLVS")

        # o Editor: o documento vai para o Acervo com o PAVLVS e abre no Word
        r = local.post("/api/documentos", json={"titulo": "Contrato de teste W1", "tipo": "texto"})
        doc_id = r.json().get("id") or (r.json().get("documento") or {}).get("id")
        checar(bool(doc_id), "um documento no Editor", r.text[:200])
        r = local.post(f"/api/word/abrir-documento/{doc_id}")
        checar(r.status_code == 200 and r.json()["aberto"] is True and r.json()["caminho"].endswith(".docx"),
               "Abrir no Word (Editor): grava no Acervo e abre", r.text[:200])
        caminho = Path(r.json()["caminho"])
        checar(wi.tem_pavlvs(caminho.read_bytes(), id_local) and abertos[-1] == str(caminho),
               "o arquivo aberto traz o PAVLVS")
        r = local.get(f"/api/documentos/{doc_id}/docx")
        checar(wi.tem_pavlvs(r.content, id_local), "Baixar DOCX também leva o PAVLVS")
        r = local.post(f"/api/documentos/{doc_id}/biblioteca", json={"formato": "docx", "automatico": True})
        checar(r.status_code == 200 and wi.tem_pavlvs(Path(r.json()["caminho"]).read_bytes(), id_local),
               "Guardar no Acervo em .docx também leva o PAVLVS", r.text[:160])

        # Word aberto desde antes da instalação: pede para fechar, não abre
        marco = datetime_do(api.estado.prefs.dados["word"]["instalado_em"])
        checar(inst_precisa(desde=lambda: marco - 60) is True and inst_precisa(desde=lambda: marco + 60) is False
               and inst_precisa(desde=lambda: None) is False, "Word aberto antes da instalação: precisa fechar; depois ou fechado: não")
        inst.precisa_fechar_o_word = lambda **kw: True
        antes = len(abertos)
        r = local.post("/api/word/abrir")
        checar(r.json() == {"aberto": False, "precisa_fechar": True} and len(abertos) == antes,
               "com o Word antigo aberto, não abre: avisa para fechar", r.text[:160])
        checar(local.get("/api/word/word-aberto").json() == {"precisa_fechar": True}, "a tela consulta até o Word fechar")
    finally:
        inst.precisa_fechar_o_word = inst_precisa

    test_atalhos_e_botao_direito(local, inst, id_local, novo, abertos)

    # desligado: o .docx sai como sempre saiu
    api.estado.prefs.atualizar({"word": {"ligado": False}})
    r = local.get(f"/api/documentos/{doc_id}/docx")
    checar(not wi.tem_pavlvs(r.content, id_local), "com o PAVLVS desligado, o .docx sai sem ele")
    checar(local.post("/api/word/abrir").status_code == 409, "desligado: Abrir o Word recusa")
    api.estado.prefs.atualizar({"word": {"ligado": True}})


def test_atalhos_e_botao_direito(local, inst, id_local, novo, abertos) -> None:
    """O atalho "Word com PAVLVS" (--word) e o botão direito dos .docx (--word-abrir)."""
    print("  o atalho e o botão direito")
    import io
    import zipfile

    import api
    import word_atalhos as wa
    import word_instalar as wi

    checar(wa.pedido_do_word(["--word"]) == ("novo", "") and wa.pedido_do_word(["--perguntar", "x"]) is None,
           "--word pede um documento novo; os outros pedidos não são do Word")
    arquivo = TMP / "Petição do cliente.docx"
    checar(wa.pedido_do_word(["--word-abrir", str(arquivo)]) == ("abrir", str(arquivo.resolve())), "--word-abrir leva o arquivo")

    # o documento-base: sem "Modo de Compatibilidade"
    base = wi.documento_base([("", "x")])
    checar(b'w:name="compatibilityMode"' in zipfile.ZipFile(io.BytesIO(base)).read("word/settings.xml")
           and b'w:val="15"' in zipfile.ZipFile(io.BytesIO(base)).read("word/settings.xml"),
           "o documento que o PAULUS gera não abre em Modo de Compatibilidade")
    modelo = wi.documento_base(modelo=True)
    checar(b"template.main+xml" in zipfile.ZipFile(io.BytesIO(modelo)).read("[Content_Types].xml"), "o modelo é um .dotx")

    inst_precisa = inst.precisa_fechar_o_word
    try:
        inst.precisa_fechar_o_word = lambda **kw: False
        r = local.post("/api/word/externo", json={"acao": "novo"})
        checar(r.json().get("aberto") is True and abertos[-1].endswith("PAVLVS.dotx")
               and wi.tem_pavlvs(Path(abertos[-1]).read_bytes(), id_local),
               "o atalho abre um documento novo pelo modelo com o PAVLVS", r.text[:160])

        arquivo.write_bytes(novo("Texto do cliente, que não pode mudar."))
        texto_antes = zipfile.ZipFile(arquivo).read("word/document.xml")
        r = local.post("/api/word/externo", json={"acao": "abrir", "caminho": str(arquivo)})
        checar(r.json() == {"aberto": True, "resultado": "posto"} and abertos[-1] == str(arquivo),
               "o botão direito põe o PAVLVS no arquivo e o abre", r.text[:160])
        checar(wi.tem_pavlvs(arquivo.read_bytes(), id_local) and zipfile.ZipFile(arquivo).read("word/document.xml") == texto_antes,
               "o arquivo passa a trazer o PAVLVS, e o texto não muda")
        checar(not list(TMP.glob(".*.pavlvs")), "nenhuma cópia temporária fica para trás")
        r = local.post("/api/word/externo", json={"acao": "abrir", "caminho": str(arquivo)})
        checar(r.json().get("resultado") == "ja_tinha", "da segunda vez, o arquivo fica como está")
        r = local.post("/api/word/externo", json={"acao": "abrir", "caminho": str(TMP / "nao-existe.docx")})
        checar(r.json().get("aberto") is False, "arquivo que não existe: não abre nada")

        # Word aberto desde antes da instalação: espera fechar e abre sozinho
        estado_word = {"aberto": True}
        inst.precisa_fechar_o_word = lambda **kw: estado_word["aberto"]
        antes = len(abertos)
        r = local.post("/api/word/externo", json={"acao": "novo"})
        checar(r.json().get("esperando") is True and len(abertos) == antes, "com o Word antigo aberto: espera, sem abrir", r.text[:160])
        estado_word["aberto"] = False
        fim = time.time() + 6
        while time.time() < fim and len(abertos) == antes:
            time.sleep(0.2)
        checar(len(abertos) == antes + 1 and abertos[-1].endswith("PAVLVS.dotx"), "o Word fechou: o PAULUS abre sozinho")
    finally:
        inst.precisa_fechar_o_word = inst_precisa

    # o botão direito, numa chave de teste (não mexe no Explorer de verdade)
    if sys.platform == "win32":
        import winreg

        base = r"Software\PAULUS-teste-w1\SystemFileAssociations"
        exe = TMP / "PAULUS.exe"
        exe.write_bytes(b"")
        try:
            wa.ligar_botao_direito(exe, base=base)
            with winreg.OpenKey(winreg.HKEY_CURRENT_USER, base + r"\.docx\shell" + "\\" + wa.VERBO) as k:
                rotulo, _ = winreg.QueryValueEx(k, "")
            checar(rotulo == "Abrir no Word com o PAVLVS" and wa.botao_direito_ligado(exe, base=base),
                   "o botão direito dos .docx, só no usuário (HKCU)")
            wa.desligar_botao_direito(base=base)
            checar(not wa.botao_direito_ligado(exe, base=base), "e sai quando o PAVLVS sai do Word")
        finally:
            for sub in (base + r"\.docx\shell" + "\\" + wa.VERBO + r"\command", base + r"\.docx\shell" + "\\" + wa.VERBO,
                        base + r"\.docx\shell", base + r"\.docx", base, r"Software\PAULUS-teste-w1"):
                try:
                    winreg.DeleteKey(winreg.HKEY_CURRENT_USER, sub)
                except OSError:
                    pass
        lnk = TMP / "atalhos" / "Word com PAVLVS.lnk"
        wa.criar_atalho(lnk, exe, "--word", wa.icone_do_word(), "teste")
        checar(lnk.exists() and lnk.stat().st_size > 0, "o atalho .lnk é criado (pasta de teste)")

    checar(api.estado.word_instalacao.atalhos()["disponivel"] is False,
           "no código-fonte (sem PAULUS.exe), nenhum atalho é criado")


def datetime_do(texto: str) -> float:
    from datetime import datetime

    return datetime.strptime(texto, "%Y-%m-%d %H:%M:%S").timestamp()


OFFICE_FALSO = r"""
(function () {
  var cfg = window.__officeCfg || {};
  var conjuntos = cfg.conjuntos || { "WordApi": "1.3", "AddinCommands": "1.1", "DialogApi": "1.1" };
  function maior(a, b) { var x = a.split("."), y = b.split("."); for (var i = 0; i < 2; i++) { if (+x[i] !== +y[i]) return +x[i] > +y[i]; } return true; }
  window.Office = {
    EventType: { OfficeThemeChanged: "officeThemeChanged" },
    onReady: function (cb) { setTimeout(function () { cb({ host: "Word", platform: "PC" }); }, 10); },
    context: {
      diagnostics: { version: cfg.versao || "16.0.14334.20918", platform: "PC" },
      officeTheme: cfg.tema || { bodyBackgroundColor: "#ffffff", bodyForegroundColor: "#000000", controlBackgroundColor: "#ffffff" },
      requirements: { isSetSupported: function (n, v) { return conjuntos[n] ? maior(conjuntos[n], v || "1.1") : false; } },
      document: { url: "C:\\Clientes\\Peti\u00e7\u00e3o inicial.docx", addHandlerAsync: function () { throw new Error("sem evento"); } },
    },
  };
})();
"""


def test_painel(porta: int) -> None:
    print("\no painel no Edge, pela porta HTTPS, com o Office simulado")
    try:
        from playwright.sync_api import sync_playwright
    except ImportError:
        print("  pulado: sem Playwright")
        return
    import api

    base = f"https://localhost:{porta}"
    with sync_playwright() as p:
        try:
            nav = p.chromium.launch(channel="msedge")
        except Exception as exc:  # noqa: BLE001
            print(f"  pulado: sem o Edge ({exc})")
            return
        ctx = nav.new_context(ignore_https_errors=True, viewport={"width": 320, "height": 640})
        pedidos_rede: list[str] = []
        violacoes: list[str] = []
        ctx.on("request", lambda r: pedidos_rede.append(r.url))
        ctx.route("https://appsforoffice.microsoft.com/**", lambda rota: rota.fulfill(
            status=200, content_type="text/javascript", body=OFFICE_FALSO))
        pag = ctx.new_page()
        erros: list[str] = []
        pag.on("pageerror", lambda e: erros.append(str(e)))
        pag.on("console", lambda m: violacoes.append(m.text) if "Content Security Policy" in m.text else None)

        pag.goto(base + "/word/painel.html?tela=inicio")
        pag.wait_for_selector("text=O PAULUS no seu Word", timeout=10000)
        checar(True, "primeira vez: o que faz, em 3 linhas, e Conectar")
        largura = pag.evaluate("[document.documentElement.scrollWidth, document.documentElement.clientWidth]")
        checar(largura[0] <= largura[1], "a 320 px, nada passa da largura", largura)
        pag.click("text=Conectar ao PAULUS")
        pag.wait_for_selector(".codigo", timeout=10000)
        codigo = pag.inner_text(".codigo").strip()
        pend = [x for x in api.estado.word_pedidos.abertos("local") if x["codigo"] == codigo]
        checar(len(pend) == 1, "o código do painel é o do pedido na janela", codigo)
        time.sleep(2)
        checar(pag.is_visible(".codigo"), "sem o Permitir, o painel continua esperando")
        from fastapi.testclient import TestClient

        TestClient(api.app, headers=api.cabecalho_local()).post(f"/api/word/pedidos/{pend[0]['id']}/permitir")
        pag.wait_for_selector("text=Documento aberto", timeout=10000)
        checar("Petição inicial.docx" in pag.inner_text("#painel"), "conectado: o início com o nome do documento")
        checar(pag.locator('.funcao[aria-disabled="true"]').count() == 7 and pag.locator("text=em breve").count() == 7,
               "as funções das próximas etapas aparecem desabilitadas, com 'em breve'")
        checar(pag.evaluate("document.documentElement.dataset.tema") == "claro", "tema claro do Office")
        pag.click('[data-acao="controlada"]')
        time.sleep(0.5)
        checar("não" in pag.inner_text(".chave"), "a escolha 'alteração controlada' muda e fica dita")

        # o tema do Office muda (o painel relê a cada 2 s)
        pag.evaluate("Office.context.officeTheme = {bodyBackgroundColor: '#262626', bodyForegroundColor: '#f0f0f0', isDarkTheme: true}")
        # (wait_for_function avalia texto, e a CSP do painel não deixa: lê em laço)
        fim = time.time() + 5
        while time.time() < fim and pag.evaluate("document.documentElement.dataset.tema") != "escuro":
            time.sleep(0.2)
        checar(pag.evaluate("document.documentElement.dataset.tema") == "escuro", "o tema do Office mudou para escuro e o painel acompanhou")

        # a 480 px, as funções em duas colunas
        pag.set_viewport_size({"width": 480, "height": 640})
        time.sleep(0.2)
        checar(pag.evaluate("document.body.classList.contains('largo')"), "a 480 px, o painel se adapta (duas colunas)")

        # revogar na janela: ao reabrir, pede para conectar de novo, dizendo o porquê
        cid = api.estado.word_conexoes.listar()[-1]["id"]
        TestClient(api.app, headers=api.cabecalho_local()).delete(f"/api/word/conexoes/{cid}")
        pag.reload()
        pag.wait_for_selector("text=foi revogada", timeout=10000)
        checar(True, "revogado: o painel diz e pede para conectar de novo")

        # Word antigo: aviso claro, nenhuma função
        pag2 = ctx.new_page()
        pag2.add_init_script("window.__officeCfg = {conjuntos: {WordApi: '1.2'}, versao: '15.0.5023'}")
        pag2.goto(base + "/word/painel.html")
        pag2.wait_for_selector("text=antigo demais", timeout=10000)
        checar(pag2.locator(".funcao").count() == 0, "Word sem a WordApi 1.3: o aviso e nenhuma função")

        # fora do Word (o office.js não carrega)
        ctx2 = nav.new_context(ignore_https_errors=True, viewport={"width": 320, "height": 640})
        ctx2.route("https://appsforoffice.microsoft.com/**", lambda rota: rota.abort())
        pag3 = ctx2.new_page()
        pag3.goto(base + "/word/painel.html")
        pag3.wait_for_selector("text=Abra pelo Word", timeout=10000)
        checar(True, "fora do Word: diz para abrir pelo Word")

        hosts = sorted({u.split("/")[2] for u in pedidos_rede if u.startswith("http")})
        checar(set(hosts) <= {f"localhost:{porta}", "appsforoffice.microsoft.com"},
               "o painel só pede à própria origem e ao office.js (a CSP)", hosts)
        checar(not violacoes, "nenhuma violação de CSP", violacoes[:3])
        checar(not erros, "nenhum erro de página", erros[:3])
        nav.close()

    sit = api.estado.word_instalacao.situacao()
    checar(sit["carregado"] and sit["carregou"]["word"].startswith("16.0"), "o painel avisou que abriu: a tela diz 'abriu no Word'", sit.get("carregou"))


def test_janela() -> None:
    print("\na janela do PAULUS no Edge: Configurações › Word e o pedido com Permitir")
    try:
        from playwright.sync_api import sync_playwright
    except ImportError:
        print("  pulado: sem Playwright")
        return
    import httpx
    from fastapi.testclient import TestClient

    import api
    from test_gravacoes import _subir_servidor

    porta = 47911
    _subir_servidor(porta)
    base = f"http://127.0.0.1:{porta}"
    with sync_playwright() as p:
        try:
            nav = p.chromium.launch(channel="msedge")
        except Exception as exc:  # noqa: BLE001
            print(f"  pulado: sem o Edge ({exc})")
            return
        pag = nav.new_page(viewport={"width": 1280, "height": 900})
        erros: list[str] = []
        pag.on("pageerror", lambda e: erros.append(str(e)))
        pag.add_init_script("try { localStorage.setItem('paulus.boasvindas', '1'); } catch (e) {}")
        # A novidade não aparece a navegador de teste (navigator.webdriver):
        # aqui ela é o assunto, então a página se passa por um navegador comum.
        pag.add_init_script("Object.defineProperty(navigator, 'webdriver', {get: () => false})")
        pag.goto(f"{base}/entrar-local?chave={api.estado.acesso.chave}")
        pag.wait_for_load_state("networkidle")
        pag.evaluate("mostrarConfig('word')")
        pag.wait_for_selector("text=Instalar no Word deste computador", timeout=10000)
        texto = pag.inner_text("#cfg-tela")
        checar("Words conectados" in texto and "Instalar no Word" in texto and "o Word abre sozinho com o PAVLVS" in texto
               and "Catálogos de Suplementos" not in texto,
               "Configurações › Word: ligar, instalar (e o Word abre sozinho), de fora e os conectados; sem passo manual", texto[:300])

        # um Word deste computador pede: a janela mostra o código e Permitir
        r = httpx.post(f"{base}/api/word/parear", json={"word": "16.0.14334.20918"})
        ped = r.json()
        pag.wait_for_selector("text=Um Word quer se conectar", timeout=10000)
        checar(ped["codigo"] in pag.inner_text(".word-codigo"), "o diálogo mostra o mesmo código do painel", ped.get("codigo"))
        antes = len(api.estado.word_conexoes.listar())
        pag.click('[data-dialogo="confirmar"]')
        fim = time.time() + 5
        while time.time() < fim and len(api.estado.word_conexoes.listar()) == antes:
            time.sleep(0.2)
        checar(len(api.estado.word_conexoes.listar()) == antes + 1, "Permitir na janela cria a conexão")
        r = httpx.post(f"{base}/api/word/parear/trocar", json={"pedido": ped["pedido"], "segredo": ped["segredo"]})
        checar(r.json().get("estado") == "permitido", "e o painel recebe o token", r.text[:120])

        # Recusar
        ped2 = httpx.post(f"{base}/api/word/parear", json={"word": "x"}).json()
        pag.wait_for_selector("text=Um Word quer se conectar", timeout=10000)
        pag.click('[data-dialogo="cancelar"]')
        time.sleep(1)
        r = httpx.post(f"{base}/api/word/parear/trocar", json={"pedido": ped2["pedido"], "segredo": ped2["segredo"]})
        checar(r.json().get("estado") == "recusado", "Recusar na janela: o painel fica sabendo", r.text[:120])
        pag.wait_for_selector(".word-conexao", timeout=10000)
        checar(pag.locator('[data-word-acao="revogar"]').count() >= 1, "a lista mostra o Word conectado com Revogar")
        # A novidade ao entrar no Editor, nas duas formas: "abrir" (instalado
        # e nunca aberto no Word - o caso da 0.9.23) e "ativar" (não instalado).
        import word_instalar

        word_instalar.word_no_computador = lambda: True
        sem_avisos = os.environ.pop("PAULUS_SEM_AVISOS", None)
        chamadas: list[str] = []
        inst = api.estado.word_instalacao
        instalar_de_verdade = inst.instalar
        precisa_de_verdade = inst.precisa_fechar_o_word
        abrir_de_verdade = api.estado.word_abrir_no_windows
        inst.instalar = lambda **kw: chamadas.append("instalar") or [{"passo": "teste", "resultado": "ok"}]
        inst.precisa_fechar_o_word = lambda **kw: False
        api.estado.word_abrir_no_windows = lambda caminho: chamadas.append("abriu " + Path(caminho).name)
        instalado_em = api.estado.prefs.dados["word"]["instalado_em"]
        chave = "paulus.novidade.word.2"
        try:
            (inst.pasta / "carregou.json").unlink(missing_ok=True)
            pag.evaluate(f"try {{ localStorage.removeItem('{chave}'); }} catch (e) {{}}")
            pag.evaluate("mostrarEditor()")
            pag.wait_for_selector("text=O PAVLVS já está no seu Word", timeout=10000)
            texto = pag.inner_text(".dialogo")
            checar("Abrir o Word com o PAVLVS" in texto and "Word com PAVLVS" in texto and "botão direito" in texto and "em breve" in texto,
                   "instalado e nunca aberto: a novidade oferece abrir o Word, o atalho e o botão direito", texto[:200])
            pag.click('[data-dialogo="confirmar"]')
            fim = time.time() + 5
            while time.time() < fim and not chamadas:
                time.sleep(0.2)
            checar(chamadas == ["abriu PAVLVS - comece aqui.docx"], "e abre o Word com o 'comece aqui', sem reinstalar", chamadas)

            chamadas.clear()
            api.estado.prefs.atualizar({"word": {"instalado_em": ""}})
            pag.evaluate(f"try {{ localStorage.removeItem('{chave}'); }} catch (e) {{}}")
            pag.evaluate("mostrarConfig('word')")
            pag.evaluate("mostrarEditor()")
            pag.wait_for_selector("text=Novidade: o PAULUS dentro do Word", timeout=10000)
            texto = pag.inner_text(".dialogo")
            checar("em breve" in texto and "clique em Sim" in texto and "o Word abre sozinho" in texto and "Ativar no Word" in texto,
                   "não instalado: a novidade, o que esta versão faz e o botão Ativar", texto[:200])
            pag.click('[data-dialogo="cancelar"]')
            time.sleep(0.5)
            pag.evaluate("mostrarConfig('word')")
            pag.evaluate("mostrarEditor()")
            time.sleep(2)
            checar(pag.locator("text=Novidade: o PAULUS dentro do Word").count() == 0, "“Agora não”: não aparece de novo")
            pag.evaluate(f"try {{ localStorage.removeItem('{chave}'); }} catch (e) {{}}")
            pag.evaluate("mostrarConfig('word')")
            pag.evaluate("mostrarEditor()")
            pag.wait_for_selector("text=Novidade: o PAULUS dentro do Word", timeout=10000)
            api.estado.prefs.atualizar({"word": {"instalado_em": instalado_em}})  # o instalar falso não grava
            pag.click('[data-dialogo="confirmar"]')
            fim = time.time() + 6
            while time.time() < fim and len(chamadas) < 2:
                time.sleep(0.2)
            checar(chamadas == ["instalar", "abriu PAVLVS - comece aqui.docx"] and api.estado.prefs.dados["word"]["ligado"] is True,
                   "Ativar no Word: liga, instala e abre o Word com o PAVLVS", chamadas)
        finally:
            inst.instalar = instalar_de_verdade
            inst.precisa_fechar_o_word = precisa_de_verdade
            api.estado.word_abrir_no_windows = abrir_de_verdade
            api.estado.prefs.atualizar({"word": {"instalado_em": instalado_em}})
            if sem_avisos is not None:
                os.environ["PAULUS_SEM_AVISOS"] = sem_avisos
        r = TestClient(api.app, headers=api.cabecalho_local(), client=("127.0.0.1", 1)).get("/api/word")
        checar(r.json()["novidade"] is False, "servidor de teste (PAULUS_SEM_AVISOS): a novidade não aparece")
        checar(not erros, "nenhum erro de página", erros[:3])
        nav.close()


def main() -> int:
    print("=" * 55)
    print("  W1 — o painel do PAVLVS no Word")
    print("=" * 55)
    test_guardados()
    porta = test_http()
    if porta:
        test_painel(porta)
        test_janela()
    try:
        import api

        api.estado.word_instalacao.fechar()
    except Exception:  # noqa: BLE001
        pass
    print()
    if _falhas:
        print(f"FALHARAM {len(_falhas)}: " + "; ".join(_falhas))
        return 1
    print("todos os testes passaram")
    return 0


if __name__ == "__main__":
    sys.exit(main())
