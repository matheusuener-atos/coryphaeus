"""
Portão da N6 do emissor de NFS-e — contador e conferência fiscal.

  - o relatório do mês soma igual à soma dos XMLs (lidos de novo, do .zip);
  - o .zip tem todos os XMLs do período (DPS, NFS-e, pedidos e eventos) e
    nada de fora do mês, mais a planilha-resumo e o leia-me;
  - as conferências acham os casos plantados: recebimento sem nota, nota sem
    recebimento, valor divergente, retenção configurada e não aplicada,
    competência de outro mês e certificado vencendo;
  - "Mandar ao contador" vira e-mail em Aprovações, com o .zip;
  - guarda: os XMLs vão no backup cifrado; nota emitida não se apaga pela tela.

    PYTHONIOENCODING=utf-8 venv/Scripts/python.exe tests/test_n6_contador.py
"""

from __future__ import annotations

import io
import os
import shutil
import sys
import tempfile
import zipfile
from pathlib import Path

RAIZ = Path(__file__).parent.parent
TMP = Path(tempfile.mkdtemp(prefix="paulus-n6-"))
os.environ["PAULUS_DADOS"] = str(TMP / "dados")
os.environ["PAULUS_SEM_AVISOS"] = "1"
os.environ["PAULUS_NFSE_SEM_FILA"] = "1"
sys.path.insert(0, str(RAIZ / "src"))
sys.path.insert(0, str(RAIZ / "tests"))

from _nfse_comum import CNPJ_TOMADOR, CPF_TOMADOR, GOIANIA, certificado_a1, checar, falhas  # noqa: E402

SENHA = "senha-do-a1-n6"


def preparar(api):
    from _sefin_simulada import SefinSimulada

    e = api.estado.nfse
    e.ligar(True)
    e.prestador.gravar({
        "documento": "11222333000181", "razao_social": "Escritório de Teste", "inscricao_municipal": "123456",
        "municipio": GOIANIA, "opcao_simples": "1",
        "servico": {"ctribnac": "171401", "nbs": "113012000", "descricao": "Honorários advocatícios", "aliquota_iss_pct": "5"},
        "retencoes": {"irrf": {"quando": "tomador_pj", "aliquota_pct": "1,5", "minimo_reais": "50,00"},
                      "pis": {"quando": "nunca"}, "cofins": {"quando": "nunca"}, "csll": {"quando": "nunca"},
                      "cp": {"quando": "nunca"}, "iss": {"quando": "nunca"}},
        "ibscbs": {"enviar": True, "cst": "200", "cclasstrib": "200052", "cindop": "100301"},
        "contador": {"nome": "Carla Contadora", "email": "carla@contabil.com.br"},
    }, quem="teste")
    certificado_a1(TMP / "a1.pfx", senha=SENHA)
    e.instalar_certificado("a1.pfx", (TMP / "a1.pfx").read_bytes(), SENHA, True)
    api.estado.base.escrever("INSERT OR REPLACE INTO nfse_municipio (cmun, ambiente, situacao, prazo_cancelamento_dias, "
                             "consultado_em) VALUES (?, 'producao_restrita', 'conveniado', 30, datetime('now'))", (GOIANIA,))
    certificado_a1(TMP / "sefin.pfx", senha="sefin", nome="SEFIN NACIONAL TESTE")
    e.transporte = SefinSimulada(TMP / "sefin.json", "sucesso", (TMP / "sefin.pfx").read_bytes(), "sefin")
    return api.estado.cadastros.salvar({"nome": "ACME Ltda", "documento": CNPJ_TOMADOR, "end_logradouro": "Rua 1",
                                        "end_numero": "10", "end_bairro": "Centro", "end_cep": "74000000",
                                        "end_cmun": GOIANIA, "email_nota": "fiscal@acme.com.br"})


def nota(api, local, dados: dict) -> dict:
    n = local.post("/api/nfse/notas", json={"origem": "teste", "dados": dados}).json()
    assert not n["erros"], n["erros"]
    local.post(f"/api/nfse/notas/{n['id']}/pedir-aprovacao")
    pedido = next(p for p in api.estado.fila.pendentes if p.acao == "nfse.emitir" and p.dados.get("nota_id") == n["id"])
    local.post("/api/aprovacoes/decidir", json={"ids": [pedido.id], "aprovar": True})
    return api.estado.nfse.notas.obter(n["id"])


def recebimento(api, cid, centavos: int, quando: str, descricao: str = "Honorários") -> int:
    lid = api.estado.financeiro.salvar({"tipo": "recebimento", "descricao": descricao, "centavos": centavos,
                                        "cadastro_id": cid, "vencimento": quando})
    api.estado.financeiro.liquidar(lid, quando)
    return lid


def main() -> int:
    try:
        import api
        from fastapi.testclient import TestClient

        local = TestClient(api.app, headers=api.cabecalho_local())
        api.estado.contas.salvar_conta({"email": "escritorio@teste.com.br", "nome": "Escritório"})
        cid = preparar(api)
        pf = api.estado.cadastros.salvar({"nome": "Maria Pessoa Física", "documento": CPF_TOMADOR, "end_logradouro": "Rua 2",
                                          "end_numero": "20", "end_bairro": "Setor Oeste", "end_cep": "74000000",
                                          "end_cmun": GOIANIA})

        print("\nnotas de setembro (e uma de agosto, que não pode entrar)")
        l1 = recebimento(api, cid, 500000, "2026-09-10")
        n1 = nota(api, local, {"lancamento_id": l1})                                     # casada com o recebimento
        l2 = recebimento(api, cid, 300000, "2026-09-12")
        n2 = nota(api, local, {"lancamento_id": l2, "valor": "2.900,00"})                # valor divergente
        n3 = nota(api, local, {"cadastro_id": pf, "valor": "1.000,00", "competencia": "2026-09-20"})  # sem recebimento
        l4 = recebimento(api, cid, 100000, "2026-09-25")
        n4 = nota(api, local, {"lancamento_id": l4})                                     # IRRF abaixo do mínimo
        n5 = nota(api, local, {"cadastro_id": cid, "valor": "700,00", "competencia": "2026-09-05"})
        local.post(f"/api/nfse/notas/{n5['id']}/cancelar", json={"motivo": "1", "texto": "Emitida para o cliente errado"})
        pedido = next(p for p in api.estado.fila.pendentes if p.acao == "nfse.cancelar")
        local.post("/api/aprovacoes/decidir", json={"ids": [pedido.id], "aprovar": True})
        recebimento(api, cid, 450000, "2026-09-28", "Honorários sem nota")                # recebimento sem nota
        agosto = nota(api, local, {"cadastro_id": cid, "valor": "999,00", "competencia": "2026-08-20"})
        checar(api.estado.nfse.notas.obter(n5["id"])["estado"] == "cancelada", "a de 700 foi cancelada")

        print("\no relatório soma igual aos XMLs")
        r = local.get("/api/nfse/contador?mes=2026-09").json()
        rel = r["relatorio"]
        checar(rel["contagem"] == {"emitida": 4, "cancelada": 1, "substituida": 0}, "4 emitidas e 1 cancelada em setembro", rel["contagem"])
        r2 = local.post("/api/nfse/contador/exportar", json={"mes": "2026-09"}).json()
        zip_path = Path(r2["caminho"])
        checar(zip_path.exists() and zip_path.parent.name == "Contador", "o .zip vai para o Acervo, em Notas fiscais/Contador", str(zip_path))
        from nfse import contador

        with zipfile.ZipFile(zip_path) as z:
            nomes = z.namelist()
            nfses = [n for n in nomes if n.endswith("-nfse.xml")]
            somas = {k: 0 for k in contador.COLUNAS}
            for nome in nfses:
                xml = z.read(nome)
                chave = nome.split("/")[-1].replace("-nfse.xml", "")
                if chave == api.estado.nfse.notas.obter(n5["id"])["chave"]:
                    continue  # a cancelada não entra no total
                v = contador.valores_do_xml(xml)
                for k in somas:
                    somas[k] += v[k]
            planilha = z.read("resumo-2026-09.xlsx")
            leia = z.read("LEIA-ME.txt").decode("utf-8")
        checar(somas == rel["totais"], "os totais do relatório = a soma dos XMLs do .zip", (somas, rel["totais"]))
        esperado = {n1["id"], n2["id"], n3["id"], n4["id"], n5["id"]}
        chaves_mes = {api.estado.nfse.notas.obter(i)["chave"] for i in esperado}
        checar({n.split("/")[-1].replace("-nfse.xml", "") for n in nfses} == chaves_mes,
               "o .zip tem as NFS-e de setembro, todas", len(nfses))
        checar(agosto["chave"] not in " ".join(nomes), "e nenhuma de agosto")
        checar(sum(1 for n in nomes if n.endswith("-dps.xml")) == 5 and any(n.endswith("-pedido.xml") for n in nomes)
               and any(n.endswith("-evento.xml") for n in nomes), "as DPS, o pedido e o evento do cancelamento", nomes)
        checar("LEIA-ME.txt" in nomes and "PRODUÇÃO RESTRITA" in leia and "COMPETÊNCIA" in leia, "o leia-me explica o que vem")
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(planilha))
        ws = wb["Notas"]
        checar(ws.max_row == 1 + 5 + 1, "a planilha tem uma linha por nota e o total", ws.max_row)
        total_planilha = ws.cell(ws.max_row, 8).value
        from decimal import Decimal

        checar(Decimal(str(total_planilha)) * 100 == rel["totais"]["v_serv"], "o total da planilha confere", total_planilha)
        checar(rel["recebido"] == 500000 + 300000 + 100000 + 450000 and rel["diferenca"] == rel["recebido"] - rel["totais"]["v_serv"],
               "recebido e diferença recebido − faturado", (rel["recebido"], rel["diferenca"]))

        print("\nas conferências acham os casos plantados")
        tipos = {}
        for a in r["achados"]:
            tipos.setdefault(a["tipo"], []).append(a)
        checar(any("Honorários sem nota" in a["detalhe"] or "ACME" in a["titulo"] for a in tipos.get("recebimento_sem_nota", [])),
               "recebimento sem nota", tipos.get("recebimento_sem_nota"))
        checar(any(a.get("nota_id") == n3["id"] for a in tipos.get("nota_sem_recebimento", [])), "nota sem recebimento")
        checar(any(a.get("nota_id") == n2["id"] for a in tipos.get("valor_divergente", [])), "valor divergente (R$ 2.900 × R$ 3.000)")
        checar(any(a.get("nota_id") == n4["id"] and "IRRF" in a["detalhe"] for a in tipos.get("retencao_nao_aplicada", [])),
               "retenção configurada e não aplicada (IRRF abaixo do mínimo)", tipos.get("retencao_nao_aplicada"))
        checar(not any(a.get("nota_id") == n1["id"] for a in tipos.get("retencao_nao_aplicada", [])),
               "a nota com IRRF retido não é achado")
        checar(any(a.get("nota_id") == n1["id"] for a in tipos.get("competencia_fora_do_mes", [])),
               "competência de setembro emitida em outro mês", tipos.get("competencia_fora_do_mes"))
        certificado_a1(TMP / "vence.pfx", senha=SENHA, dias=5)
        api.estado.nfse.instalar_certificado("vence.pfx", (TMP / "vence.pfx").read_bytes(), SENHA, True)
        achados = local.get("/api/nfse/contador?mes=2026-09").json()["achados"]
        checar(any(a["tipo"] == "certificado" for a in achados), "certificado vencendo", [a["tipo"] for a in achados])

        print("\nmandar ao contador: e-mail em Aprovações, com o .zip")
        r3 = local.post("/api/nfse/contador/enviar", json={"mes": "2026-09"})
        checar(r3.status_code == 200, "pedido criado", r3.text[:200])
        email = api.estado.fila.obter(r3.json()["pedido"])
        checar(email.acao == "correio.enviar" and email.dados["para"] == "carla@contabil.com.br"
               and email.dados["anexos"][0].endswith(".zip") and Path(email.dados["anexos"][0]).exists(),
               "para a contadora, com o arquivo do mês", email.dados.get("para"))

        print("\nguarda: backup e nada de apagar nota emitida")
        import backup

        saida = TMP / "backup"
        feito = backup.fazer(Path(api.DADOS_DIR), saida, "senha-do-backup-1")
        arquivo = Path(feito.get("arquivo") or feito.get("caminho") or next(saida.glob("*.paulusbak")))
        aberto = TMP / "aberto.zip"
        backup.decifrar_arquivo(arquivo, aberto, "senha-do-backup-1")
        with zipfile.ZipFile(aberto) as z:
            dentro = z.namelist()
        n1_nome = Path(api.estado.nfse.notas.obter(n1["id"])["xml_nfse"]).name
        checar(any(d.endswith(n1_nome) and "nfse/xml" in d.replace("\\", "/") for d in dentro), "o XML da NFS-e vai no backup cifrado")
        checar(any(d.endswith("paulus.db") for d in dentro), "e a base com as notas e a numeração")
        rotas = {(m, getattr(rt, "path", "")) for rt in api.app.routes for m in getattr(rt, "methods", [])}
        checar(not any(m == "DELETE" and p.startswith("/api/nfse") for m, p in rotas), "não existe rota que apague nota")
        r4 = local.post(f"/api/nfse/notas/{n1['id']}/descartar")
        checar(r4.status_code == 400 and "cancela" in r4.json()["detail"], "descartar a emitida é recusado: cancela-se", r4.text[:200])
        papel = api.estado.base.um("SELECT id FROM papeis_fiscais WHERE nfse_nota_id = ?", (n1["id"],))
        r5 = local.delete(f"/api/financeiro/papeis/{papel['id']}")
        checar(r5.status_code == 400, "tirar do registro é recusado", r5.status_code)
    finally:
        shutil.rmtree(TMP, ignore_errors=True)
    if falhas:
        print(f"\n{len(falhas)} falha(s): {falhas}")
        return 1
    print("\ntodos os testes passaram")
    return 0


if __name__ == "__main__":
    sys.exit(main())
