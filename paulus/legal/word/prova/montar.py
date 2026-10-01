"""Prova da W0: manifesto por montagem e um .docx que abre o painel sozinho.

O .docx leva as partes webextension/taskpanes com
Office.AutoShowTaskpaneWithDocument, como o office-addin-dev-settings faz,
para o Word abrir o painel ao abrir o documento (sem clique na faixa).

Uso: python montar.py <pasta-saida> <letra> <url-da-pagina>
"""

from __future__ import annotations

import io
import sys
import uuid
import zipfile
from pathlib import Path
from xml.sax.saxutils import escape

import docx

ESPACO = uuid.UUID("6d1f0c52-5c1b-4f39-9a3b-0e7a2c6b9a10")


def id_da_montagem(letra: str) -> str:
    return str(uuid.uuid5(ESPACO, f"paulus-word-prova-{letra}"))


def manifesto(letra: str, url: str) -> str:
    origem = url.split("/", 3)
    origem = origem[0] + "//" + origem[2]
    return f"""<?xml version="1.0" encoding="UTF-8"?>
<OfficeApp xmlns="http://schemas.microsoft.com/office/appforoffice/1.1"
           xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance"
           xsi:type="TaskPaneApp">
  <Id>{id_da_montagem(letra)}</Id>
  <Version>1.0.0.0</Version>
  <ProviderName>PAULUS</ProviderName>
  <DefaultLocale>pt-BR</DefaultLocale>
  <DisplayName DefaultValue="PAULUS prova {letra}"/>
  <Description DefaultValue="Prova de conceito da montagem {letra}"/>
  <AppDomains><AppDomain>{escape(origem)}</AppDomain></AppDomains>
  <Hosts><Host Name="Document"/></Hosts>
  <Requirements><Sets DefaultMinVersion="1.1"><Set Name="WordApi" MinVersion="1.1"/></Sets></Requirements>
  <DefaultSettings><SourceLocation DefaultValue="{escape(url)}"/></DefaultSettings>
  <Permissions>ReadWriteDocument</Permissions>
</OfficeApp>
"""


def documento(letra: str, store: str, store_type: str) -> bytes:
    base = io.BytesIO()
    d = docx.Document()
    d.add_paragraph(f"Documento de prova do PAULUS (montagem {letra}).")
    d.add_paragraph("Art. 18 da Lei 8.078/1990 e Súmula 297 do STJ.")
    d.save(base)

    we = f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<we:webextension xmlns:we="http://schemas.microsoft.com/office/webextensions/webextension/2010/11" id="{{{str(uuid.uuid4()).upper()}}}">
  <we:reference id="{id_da_montagem(letra)}" version="1.0.0.0" store="{escape(store)}" storeType="{store_type}"/>
  <we:alternateReferences/>
  <we:properties><we:property name="Office.AutoShowTaskpaneWithDocument" value="true"/></we:properties>
  <we:bindings/>
  <we:snapshot xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"/>
</we:webextension>"""
    tp = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<wetp:taskpanes xmlns:wetp="http://schemas.microsoft.com/office/webextensions/taskpanes/2010/11">
  <wetp:taskpane dockstate="right" visibility="1" width="380" row="4">
    <wetp:webextensionref xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" r:id="rId1"/>
  </wetp:taskpane>
</wetp:taskpanes>"""
    tp_rels = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.microsoft.com/office/2011/relationships/webextension" Target="webextension1.xml"/></Relationships>"""

    saida = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(base.getvalue())) as zin, zipfile.ZipFile(saida, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            dados = zin.read(item.filename)
            if item.filename == "[Content_Types].xml":
                dados = dados.replace(b"</Types>", (
                    b'<Override PartName="/word/webextensions/taskpanes.xml" ContentType="application/vnd.ms-office.webextensiontaskpanes+xml"/>'
                    b'<Override PartName="/word/webextensions/webextension1.xml" ContentType="application/vnd.ms-office.webextension+xml"/>'
                    b"</Types>"))
            elif item.filename == "_rels/.rels":
                dados = dados.replace(b"</Relationships>", (
                    b'<Relationship Id="rIdPaulusWe" Type="http://schemas.microsoft.com/office/2011/relationships/webextensiontaskpanes" Target="word/webextensions/taskpanes.xml"/>'
                    b"</Relationships>"))
            zout.writestr(item, dados)
        zout.writestr("word/webextensions/taskpanes.xml", tp)
        zout.writestr("word/webextensions/_rels/taskpanes.xml.rels", tp_rels)
        zout.writestr("word/webextensions/webextension1.xml", we)
    return saida.getvalue()


if __name__ == "__main__":
    pasta, letra, url = Path(sys.argv[1]), sys.argv[2], sys.argv[3]
    pasta.mkdir(parents=True, exist_ok=True)
    arq = pasta / f"manifesto-{letra}.xml"
    arq.write_text(manifesto(letra, url), encoding="utf-8")
    # Carregado pelo registro de desenvolvedor (HKCU\...\WEF\Developer): o
    # documento aponta para o manifesto por "Registry".
    (pasta / f"documento-{letra}.docx").write_bytes(documento(letra, "developer", "Registry"))
    print(arq, id_da_montagem(letra))
