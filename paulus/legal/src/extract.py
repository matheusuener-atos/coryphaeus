"""
PAULUS Legal - Extracao de texto de contratos.

Suporta PDF, DOCX, TXT e MD. Mantem um cache em JSON para nao reprocessar
arquivos que nao mudaram entre execucoes (extracao de PDF grande e lenta).
"""

from __future__ import annotations

import hashlib
import json
from dataclasses import asdict, dataclass, replace
from pathlib import Path

SUPPORTED_SUFFIXES = {".pdf", ".docx", ".txt", ".md"}


@dataclass
class Document:
    """Um contrato ja extraido para texto puro."""

    name: str
    path: str
    text: str
    pages: int = 0
    sha1: str = ""

    @property
    def chars(self) -> int:
        return len(self.text)

    def to_dict(self) -> dict:
        return asdict(self)

    @classmethod
    def from_dict(cls, data: dict) -> "Document":
        return cls(
            name=data["name"],
            path=data["path"],
            text=data["text"],
            pages=data.get("pages", 0),
            sha1=data.get("sha1", ""),
        )


# --------------------------------------------------------------------------
# Extratores por formato
# --------------------------------------------------------------------------


def extract_pdf(path: Path) -> tuple[str, int]:
    """Extrai texto de um PDF. Retorna (texto, numero_de_paginas)."""
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    if reader.is_encrypted:
        # Muitos PDFs juridicos vem "protegidos" apenas contra edicao,
        # com senha vazia. Tentamos abrir antes de desistir.
        try:
            reader.decrypt("")
        except Exception as exc:  # pragma: no cover - depende do arquivo
            raise RuntimeError(f"PDF protegido por senha: {path.name}") from exc

    partes: list[str] = []
    for i, page in enumerate(reader.pages, start=1):
        try:
            texto = page.extract_text() or ""
        except Exception:
            texto = ""
        if texto.strip():
            partes.append(f"[pagina {i}]\n{texto.strip()}")

    return "\n\n".join(partes), len(reader.pages)


def extract_docx(path: Path) -> tuple[str, int]:
    """Extrai texto de um DOCX, incluindo o conteudo das tabelas."""
    import docx

    doc = docx.Document(str(path))
    partes = [p.text.strip() for p in doc.paragraphs if p.text.strip()]

    for tabela in doc.tables:
        for linha in tabela.rows:
            celulas = [c.text.strip() for c in linha.cells if c.text.strip()]
            if celulas:
                partes.append(" | ".join(celulas))

    return "\n".join(partes), 0


def extract_txt(path: Path) -> tuple[str, int]:
    return path.read_text(encoding="utf-8", errors="replace"), 0


def extract_file(path: Path) -> tuple[str, int]:
    """Dispatch por extensao. Levanta ValueError se o formato nao e suportado."""
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return extract_pdf(path)
    if suffix == ".docx":
        return extract_docx(path)
    if suffix in {".txt", ".md"}:
        return extract_txt(path)
    raise ValueError(f"Formato nao suportado: {suffix}")


# --------------------------------------------------------------------------
# Indexacao da pasta + cache
# --------------------------------------------------------------------------


def file_sha1(path: Path) -> str:
    h = hashlib.sha1()
    with open(path, "rb") as f:
        for bloco in iter(lambda: f.read(1 << 20), b""):
            h.update(bloco)
    return h.hexdigest()


def _load_cache(cache_path: Path) -> dict[str, Document]:
    if not cache_path.exists():
        return {}
    try:
        dados = json.loads(cache_path.read_text(encoding="utf-8"))
    except (json.JSONDecodeError, OSError):
        return {}
    return {d["sha1"]: Document.from_dict(d) for d in dados.get("documents", []) if d.get("sha1")}


def _save_cache(cache_path: Path, docs: list[Document]) -> None:
    cache_path.parent.mkdir(parents=True, exist_ok=True)
    payload = {"version": 1, "documents": [d.to_dict() for d in docs]}
    cache_path.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")


def index_all_contracts(
    folder: Path,
    cache_path: Path | None = None,
    *,
    force: bool = False,
    verbose: bool = True,
) -> list[Document]:
    """
    Extrai todos os contratos de `folder` (recursivamente).

    Arquivos ja extraidos com o mesmo conteudo (mesmo sha1) sao lidos do cache.
    Passe force=True para reprocessar tudo.
    """
    folder = Path(folder)
    if not folder.exists():
        return []

    cache = {} if (force or cache_path is None) else _load_cache(Path(cache_path))
    docs: list[Document] = []

    arquivos = sorted(
        p for p in folder.rglob("*") if p.is_file() and p.suffix.lower() in SUPPORTED_SUFFIXES
    )

    for path in arquivos:
        sha = file_sha1(path)
        if sha in cache:
            # Copia, e nao o objeto do cache. Dois arquivos iguais byte a byte
            # - "contrato.pdf" e "contrato (1).pdf" - tem o mesmo sha1 e caem
            # nesta linha os dois. Mexer no objeto guardado fazia o segundo
            # sobrescrever o caminho e o nome do primeiro, e a biblioteca
            # passava a mostrar a copia duas vezes, com o original sumido da
            # lista. O texto e o mesmo; o caminho e o nome nao sao.
            doc = replace(cache[sha], path=str(path), name=path.name)
            docs.append(doc)
            if verbose:
                print(f"  = {path.name} (cache)")
            continue

        try:
            texto, paginas = extract_file(path)
        except Exception as exc:
            if verbose:
                print(f"  ! {path.name}: {exc}")
            continue

        if not texto.strip():
            if verbose:
                print(f"  ! {path.name}: sem texto extraivel (PDF escaneado? precisa de OCR)")
            continue

        doc = Document(name=path.name, path=str(path), text=texto, pages=paginas, sha1=sha)
        docs.append(doc)
        if verbose:
            print(f"  + {path.name} ({doc.chars:,} chars)".replace(",", "."))

    if cache_path is not None:
        _save_cache(Path(cache_path), docs)

    return docs


if __name__ == "__main__":
    import sys

    pasta = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(__file__).parent.parent / "data" / "test_contracts"
    documentos = index_all_contracts(pasta)
    print(f"\n{len(documentos)} documento(s) extraido(s) de {pasta}")
