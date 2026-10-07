"""
Para o contador: o relatório do mês, os XMLs num .zip e a conferência por regra.

O relatório soma o que está nos XMLs das notas — o valor que vale é o que a
Sefin calculou e devolveu, não a previsão do cartão. O mês é o da
competência (o mês do fato gerador), que é como a apuração é feita; a data de
emissão vai junto.

A conferência acha, por regra, o que costuma dar trabalho no fechamento:
recebimento sem nota, nota sem recebimento, valor divergente, retenção
configurada e não aplicada, nota emitida em mês diferente da competência e o
certificado da nota perto de vencer. Cada achado vira aviso.

Guarda: os XMLs moram na pasta de dados (que vai inteira no backup cifrado) e
uma cópia no Acervo. Nota emitida não se apaga pela tela: cancela-se. O
prazo mínimo é o do CTN (art. 195, parágrafo único: até a prescrição dos
créditos, arts. 173 e 174); o PAULUS guarda sem prazo para apagar.
"""

from __future__ import annotations

import io
import json
import zipfile
from datetime import date, datetime
from pathlib import Path

from lxml import etree

from .dinheiro import centavos_do_xml, reais
from .notas import CANCELADA, EMITIDA, SUBSTITUIDA

NS = {"n": "http://www.sped.fazenda.gov.br/nfse"}
SITUACOES = (EMITIDA, CANCELADA, SUBSTITUIDA)
COLUNAS = ("v_serv", "v_desc", "v_bc_iss", "v_iss", "v_irrf", "v_csll_pcc", "v_cp", "v_total_ret", "v_liq", "v_ibs", "v_cbs")
ROTULOS = {"v_serv": "Valor do serviço", "v_desc": "Desconto incondicionado", "v_bc_iss": "Base do ISS", "v_iss": "ISS",
           "v_irrf": "IRRF retido", "v_csll_pcc": "PIS/COFINS/CSLL retidos", "v_cp": "CP retida",
           "v_total_ret": "Total retido", "v_liq": "Valor líquido", "v_ibs": "IBS", "v_cbs": "CBS"}


def valores_do_xml(xml: bytes) -> dict:
    """Os números da nota, lidos do XML da NFS-e (centavos)."""
    raiz = etree.fromstring(xml, etree.XMLParser(resolve_entities=False, no_network=True))
    inf = raiz.find("n:infNFSe", NS)
    dps = inf.find("n:DPS/n:infDPS", NS)

    def c(no, caminho):
        t = (no.findtext(caminho, default="", namespaces=NS) or "").strip() if no is not None else ""
        return centavos_do_xml(t) if t else 0

    return {
        "v_serv": c(dps, "n:valores/n:vServPrest/n:vServ"),
        "v_desc": c(dps, "n:valores/n:vDescCondIncond/n:vDescIncond"),
        "v_bc_iss": c(inf, "n:valores/n:vBC"),
        "v_iss": c(inf, "n:valores/n:vISSQN"),
        "v_irrf": c(dps, "n:valores/n:trib/n:tribFed/n:vRetIRRF"),
        "v_csll_pcc": c(dps, "n:valores/n:trib/n:tribFed/n:vRetCSLL"),
        "v_cp": c(dps, "n:valores/n:trib/n:tribFed/n:vRetCP"),
        "v_total_ret": c(inf, "n:valores/n:vTotalRet"),
        "v_liq": c(inf, "n:valores/n:vLiq"),
        "v_ibs": c(inf, "n:IBSCBS/n:totCIBS/n:gIBS/n:vIBSTot"),
        "v_cbs": c(inf, "n:IBSCBS/n:totCIBS/n:gCBS/n:vCBS"),
        "iss_retido": (dps.findtext("n:valores/n:trib/n:tribMun/n:tpRetISSQN", default="", namespaces=NS) == "2"),
    }


class Contador:
    def __init__(self, emissor) -> None:
        self.emissor = emissor
        self.base = emissor.base
        self.notas = emissor.notas

    def notas_do_mes(self, mes: str) -> list[dict]:
        return [n for n in self.notas.listar(SITUACOES, mes=mes, limite=5000) if n.get("xml_nfse")]

    # ------------------------------------------------------------ relatório

    def relatorio(self, mes: str) -> dict:
        notas = self.notas_do_mes(mes)
        linhas, totais, por_tomador = [], {k: 0 for k in COLUNAS}, {}
        contagem = {EMITIDA: 0, CANCELADA: 0, SUBSTITUIDA: 0}
        for n in sorted(notas, key=lambda x: (x.get("numero_nfse") or "").zfill(13)):
            v = valores_do_xml(Path(n["xml_nfse"]).read_bytes())
            contagem[n["estado"]] = contagem.get(n["estado"], 0) + 1
            linha = {"id": n["id"], "numero": n["numero_nfse"], "chave": n["chave"], "situacao": n["estado_rotulo"],
                     "estado": n["estado"], "tomador": n.get("tomador_nome") or "", "documento": n.get("tomador_documento") or "",
                     "competencia": n.get("competencia") or "", "emitida_em": (n.get("dh_proc") or "")[:10],
                     "ambiente": n["ambiente"], **{k: v[k] for k in COLUNAS}}
            linhas.append(linha)
            if n["estado"] == EMITIDA:
                # Só a nota que vale entra nos totais: a cancelada e a substituída não.
                for k in COLUNAS:
                    totais[k] += v[k]
                t = por_tomador.setdefault(linha["documento"] or linha["tomador"],
                                           {"tomador": linha["tomador"], "documento": linha["documento"], "notas": 0,
                                            **{k: 0 for k in COLUNAS}})
                t["notas"] += 1
                for k in COLUNAS:
                    t[k] += v[k]
        recebido = self._recebido(mes)
        return {"mes": mes, "linhas": linhas, "totais": totais, "por_tomador": sorted(por_tomador.values(), key=lambda t: -t["v_serv"]),
                "contagem": contagem, "recebido": recebido, "faturado": totais["v_serv"],
                "diferenca": recebido - totais["v_serv"], "rotulos": ROTULOS,
                "producao_restrita": any(l["ambiente"] != "producao" for l in linhas)}

    def _recebido(self, mes: str) -> int:
        linha = self.base.um("SELECT COALESCE(SUM(centavos),0) s FROM lancamentos WHERE tipo = 'recebimento' "
                             "AND liquidado_em != '' AND substr(liquidado_em,1,7) = ?", (mes,))
        return int(linha["s"] or 0)

    # ------------------------------------------------------------ exportação

    def arquivos_do_mes(self, mes: str) -> list[Path]:
        """Todos os XMLs do período: DPS assinada, NFS-e, pedidos e eventos."""
        arquivos: list[Path] = []
        for n in self.notas_do_mes(mes):
            for campo in ("xml_dps", "xml_nfse"):
                if n.get(campo) and Path(n[campo]).exists():
                    arquivos.append(Path(n[campo]))
            for ev in self.emissor.eventos.da_nota(n["id"]):
                for campo in ("xml_pedido", "xml_evento"):
                    if ev.get(campo) and Path(ev[campo]).exists():
                        arquivos.append(Path(ev[campo]))
        return arquivos

    def planilha(self, rel: dict) -> bytes:
        from openpyxl import Workbook
        from openpyxl.styles import Font

        wb = Workbook()
        ws = wb.active
        ws.title = "Notas"
        cab = ["Número", "Situação", "Emitida em", "Competência", "Tomador", "CPF/CNPJ", "Chave"] + [ROTULOS[k] for k in COLUNAS]
        ws.append(cab)
        for c in ws[1]:
            c.font = Font(bold=True)
        for l in rel["linhas"]:
            ws.append([l["numero"], l["situacao"], l["emitida_em"], l["competencia"], l["tomador"], l["documento"], l["chave"]]
                      + [_reais_planilha(l[k]) for k in COLUNAS])
        ws.append(["", "Total (só emitidas)", "", "", "", "", ""] + [_reais_planilha(rel["totais"][k]) for k in COLUNAS])
        for c in ws[ws.max_row]:
            c.font = Font(bold=True)
        ws2 = wb.create_sheet("Por tomador")
        ws2.append(["Tomador", "CPF/CNPJ", "Notas"] + [ROTULOS[k] for k in COLUNAS])
        for t in rel["por_tomador"]:
            ws2.append([t["tomador"], t["documento"], t["notas"]] + [_reais_planilha(t[k]) for k in COLUNAS])
        for folha in (ws, ws2):
            for coluna in folha.columns:
                folha.column_dimensions[coluna[0].column_letter].width = 16
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()

    def exportar(self, mes: str, destino: Path) -> dict:
        """O .zip do mês: os XMLs, a planilha-resumo e um leia-me. Devolve {caminho, arquivos}."""
        rel = self.relatorio(mes)
        arquivos = self.arquivos_do_mes(mes)
        prest = self.emissor.prestador.atual()["dados"]
        destino.mkdir(parents=True, exist_ok=True)
        alvo = destino / f"NFS-e {mes} - para o contador.zip"
        leia = _leia_me(mes, rel, prest, len(arquivos))
        with zipfile.ZipFile(alvo, "w", zipfile.ZIP_DEFLATED) as z:
            for a in arquivos:
                z.write(a, f"xml/{a.name}")
            z.writestr(f"resumo-{mes}.xlsx", self.planilha(rel))
            z.writestr("LEIA-ME.txt", leia)
        return {"caminho": str(alvo), "arquivos": len(arquivos), "relatorio": rel}

    # ------------------------------------------------------------ conferência

    def conferir(self, mes: str, hoje: date | None = None) -> list[dict]:
        """Os achados do mês, por regra. Cada um: {tipo, titulo, detalhe, id}."""
        hoje = hoje or date.today()
        achados: list[dict] = []
        emitidas = [n for n in self.notas_do_mes(mes) if n["estado"] == EMITIDA]

        for r in self.emissor.papeis_a_emitir(mes):
            achados.append({"tipo": "recebimento_sem_nota", "id": f"nfse:conf:sem_nota:{r['id']}",
                            "titulo": f"Recebimento sem nota — {r.get('cliente') or r.get('descricao') or ''}",
                            "detalhe": f"{reais(int(r['centavos']))} recebido em {str(r.get('liquidado_em') or '')[:10]}"})
        for n in emitidas:
            lanc = self.base.um("SELECT * FROM lancamentos WHERE id = ?", (n["lancamento_id"],)) if n.get("lancamento_id") else None
            if not lanc:
                achados.append({"tipo": "nota_sem_recebimento", "id": f"nfse:conf:sem_recebimento:{n['id']}",
                                "titulo": f"Nota sem recebimento — NFS-e {n['numero_nfse']} ({n.get('tomador_nome') or ''})",
                                "detalhe": f"{n['valor']}: ligue a um recebimento no Financeiro", "nota_id": n["id"]})
            else:
                if not lanc.get("liquidado_em"):
                    achados.append({"tipo": "nota_sem_recebimento", "id": f"nfse:conf:a_receber:{n['id']}",
                                    "titulo": f"Nota com recebimento em aberto — NFS-e {n['numero_nfse']}",
                                    "detalhe": f"{n['valor']} ainda não recebido", "nota_id": n["id"]})
                if int(lanc.get("centavos") or 0) != int(n["centavos"]):
                    achados.append({"tipo": "valor_divergente", "id": f"nfse:conf:divergente:{n['id']}",
                                    "titulo": f"Valor divergente — NFS-e {n['numero_nfse']}",
                                    "detalhe": f"nota {n['valor']} × recebimento {reais(int(lanc['centavos']))}", "nota_id": n["id"]})
            achados += self._retencoes_nao_aplicadas(n)
            emitida_mes = (n.get("dh_proc") or "")[:7]
            if emitida_mes and emitida_mes != (n.get("competencia") or "")[:7]:
                achados.append({"tipo": "competencia_fora_do_mes", "id": f"nfse:conf:competencia:{n['id']}",
                                "titulo": f"Competência de outro mês — NFS-e {n['numero_nfse']}",
                                "detalhe": f"competência {n['competencia']}, emitida em {emitida_mes}", "nota_id": n["id"]})
        cert = (self.emissor.certificado_para_tela() or {}).get("certificado") or {}
        dias = cert.get("dias_restantes")
        if dias is not None and dias <= 30:
            achados.append({"tipo": "certificado", "id": f"nfse:conf:certificado:{cert.get('valido_ate', '')}",
                            "titulo": "Certificado da nota " + ("vencido" if dias < 0 else f"vence em {dias} dia(s)"),
                            "detalhe": f"válido até {cert.get('valido_ate', '')}"})
        return achados

    def _retencoes_nao_aplicadas(self, n: dict) -> list[dict]:
        """A retenção que a configuração da nota mandava e o XML não traz."""
        prest = self.notas.prestador_da_nota(n)["dados"]
        tom = (n.get("rascunho") or {}).get("tomador") or {}
        pj = len("".join(ch for ch in str(tom.get("documento") or "") if ch.isalnum())) == 14
        v = valores_do_xml(Path(n["xml_nfse"]).read_bytes())
        faltam = []
        campo = {"irrf": "v_irrf", "cp": "v_cp"}
        for k, cfg in (prest.get("retencoes") or {}).items():
            quando = cfg.get("quando")
            aplica = quando == "sempre" or (quando == "tomador_pj" and pj)
            if not aplica or (k != "iss" and not int(cfg.get("aliquota_bp") or 0)):
                continue
            if k == "iss":
                tem = v["iss_retido"]
            elif k in campo:
                tem = v[campo[k]] > 0
            else:
                tem = v["v_csll_pcc"] > 0
            if not tem:
                faltam.append(k.upper())
        if not faltam:
            return []
        return [{"tipo": "retencao_nao_aplicada", "id": f"nfse:conf:retencao:{n['id']}",
                 "titulo": f"Retenção configurada e não aplicada — NFS-e {n['numero_nfse']}",
                 "detalhe": ", ".join(faltam) + " (abaixo do mínimo, ou a configuração mudou depois): confira com o contador",
                 "nota_id": n["id"]}]


def _reais_planilha(centavos: int):
    """A planilha pede número: o centavo vira reais em Decimal (exato), só na
    saída para o Excel do contador, e nunca volta para a conta."""
    from decimal import Decimal

    return Decimal(int(centavos)) / Decimal(100)


def _leia_me(mes: str, rel: dict, prest: dict, n_arquivos: int) -> str:
    t = rel["totais"]
    linhas = [
        f"NFS-e de {mes[5:7]}/{mes[:4]} — {prest.get('razao_social') or ''} ({prest.get('documento') or ''})",
        "",
        "Gerado pelo Paulus Legal em " + datetime.now().strftime("%d/%m/%Y %H:%M") + ".",
        "",
        "O que vem neste arquivo:",
        f"- xml/: {n_arquivos} XMLs — as DPS assinadas (…-dps.xml), as NFS-e devolvidas pela Sefin (…-nfse.xml),",
        "  os pedidos de evento (…-pedido.xml) e os eventos registrados (…-evento.xml);",
        f"- resumo-{mes}.xlsx: uma linha por nota e o total por tomador (só as notas emitidas entram nos totais);",
        "",
        f"Notas: {rel['contagem'].get('emitida', 0)} emitida(s), {rel['contagem'].get('cancelada', 0)} cancelada(s), "
        f"{rel['contagem'].get('substituida', 0)} substituída(s).",
        f"Faturado (valor do serviço das emitidas): {reais(t['v_serv'])}; recebido no mês: {reais(rel['recebido'])}.",
        f"ISS {reais(t['v_iss'])} · retenções {reais(t['v_total_ret'])} · IBS {reais(t['v_ibs'])} · CBS {reais(t['v_cbs'])} · "
        f"líquido {reais(t['v_liq'])}.",
        "",
        "O mês é o da COMPETÊNCIA (o mês do serviço). Os valores são os da NFS-e, calculados pela Sefin.",
    ]
    if rel.get("producao_restrita"):
        linhas += ["", "ATENÇÃO: há notas de PRODUÇÃO RESTRITA (homologação) neste arquivo: elas não têm valor fiscal."]
    return "\n".join(linhas) + "\n"
