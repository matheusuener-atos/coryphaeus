"""
O escritório por dentro: folha, papelada e o que fecha o mês.

O Financeiro já somava o que entra e o que sai. Falta a parte que não é
dinheiro passando — é o escritório funcionando: a folha das pessoas, os
comprovantes do que foi pago, as notas emitidas, os boletos na fila e o que
terminou no mês.

Três decisões moram aqui:

**A folha é uma cópia, não um cálculo.** O valor de cada pessoa fica gravado
no mês em que foi pago. Aumentar o salário hoje não pode reescrever a folha de
agosto — a folha de agosto foi aquela, e é por ela que se presta conta.

**Nota fiscal e boleto são registro, não emissão.** A nota sai no sistema da
prefeitura e o boleto sai do banco. Fingir que este programa emite qualquer um
dos dois seria mentir sobre o que ele faz. O que dá para fazer com verdade
nesta máquina é guardar o número, o valor e a data, e avisar do vencimento.

**Contrato concluído não é um campo — é uma conclusão.** Nada aqui pergunta
"este contrato acabou?". A resposta sai dos lançamentos: um cliente que
recebeu neste mês e não tem mais nada em aberto terminou. Um campo separado
viraria duas verdades sobre o mesmo cliente, e a errada é sempre a que não foi
atualizada.
"""

from __future__ import annotations

from calendar import monthrange
from datetime import date, datetime
from pathlib import Path

from financeiro import curto, em_reais

# Como a pessoa é paga. Muda o que sai na folha e o que se deve de encargo,
# e é o que separa "salários" de "estagiários" no resumo do mês.
VINCULOS = {
    "clt": "CLT",
    "estagio": "Estágio",
    "prolabore": "Pró-labore",
    "autonomo": "Autônomo",
}

# Estagiário não entra em "salários": a bolsa não é salário, e somar os dois
# esconde o número que o escritório precisa olhar.
VINCULO_ESTAGIO = "estagio"

TIPOS_PAPEL = {"nota": "Nota fiscal", "boleto": "Boleto"}

SITUACOES_BOLETO = {"aberto": "na fila", "pago": "pago", "cancelado": "cancelado"}

# Como o pagamento se chama no recibo. "Referente a CLT" não é português, e
# recibo é documento: bolsa de estágio não é salário, e pró-labore de sócio
# não é nenhum dos dois. O nome errado num recibo é um problema trabalhista,
# não um problema de texto.
NO_RECIBO = {
    "clt": "salário",
    "estagio": "bolsa de estágio",
    "prolabore": "pró-labore",
    "autonomo": "prestação de serviços",
}


def mes_de_hoje() -> str:
    return date.today().strftime("%Y-%m")


def dias_para_fechar(mes: str = "") -> int:
    """
    Quantos dias faltam para o mês acabar.

    O wireframe mostra "fechamento em 20 dias" no topo. É o último dia do mês
    menos hoje — número medido, não prazo inventado. Mês passado devolve zero:
    já fechou.
    """
    mes = mes or mes_de_hoje()
    try:
        ano, numero = (int(x) for x in mes.split("-"))
    except ValueError:
        return 0
    ultimo = date(ano, numero, monthrange(ano, numero)[1])
    return max(0, (ultimo - date.today()).days)


def _agora() -> str:
    return datetime.now().isoformat(timespec="seconds")


# ------------------------------------------------------------------ folha


class Folha:
    """A folha de pagamento, mês a mês."""

    def __init__(self, base) -> None:
        self.base = base

    def pessoas(self) -> list[dict]:
        """
        Quem entra na folha: quem tem vínculo definido no cadastro.

        Sem vínculo, a pessoa não entra — não por engano, mas porque não foi
        dito como ela é paga. Colocar todo colaborador na folha por um valor
        que ninguém digitou produziria uma folha errada com cara de certa.
        """
        return self.base.buscar(
            "SELECT id, nome, tipo, vinculo, salario_centavos, encargos_centavos "
            "FROM cadastros WHERE vinculo != '' ORDER BY nome COLLATE NOCASE"
        )

    def do_mes(self, mes: str = "") -> dict:
        """A folha gravada deste mês, com o resumo por natureza."""
        mes = mes or mes_de_hoje()
        linhas = self.base.buscar(
            "SELECT * FROM folha WHERE mes = ? ORDER BY nome COLLATE NOCASE", (mes,)
        )
        for l in linhas:
            l["vinculo_rotulo"] = VINCULOS.get(l["vinculo"], l["vinculo"])
            l["salario"] = em_reais(l["salario_centavos"])
            l["encargos"] = em_reais(l["encargos_centavos"])
            l["total_centavos"] = l["salario_centavos"] + l["encargos_centavos"]
            l["total"] = em_reais(l["total_centavos"])

        salarios = sum(l["salario_centavos"] for l in linhas if l["vinculo"] != VINCULO_ESTAGIO)
        estagios = sum(l["salario_centavos"] for l in linhas if l["vinculo"] == VINCULO_ESTAGIO)
        encargos = sum(l["encargos_centavos"] for l in linhas)
        total = salarios + estagios + encargos

        return {
            "mes": mes,
            "pessoas": linhas,
            "quantos": len(linhas),
            "salarios": salarios,
            "salarios_texto": em_reais(salarios),
            "encargos": encargos,
            "encargos_texto": em_reais(encargos),
            "estagios": estagios,
            "estagios_texto": em_reais(estagios),
            "total": total,
            "total_texto": em_reais(total),
            "total_curto": curto(total),
            "fechada": bool(linhas),
            "lancamento_id": linhas[0]["lancamento_id"] if linhas else None,
        }

    def montar(self, mes: str = "") -> dict:
        """
        Grava a folha do mês a partir dos cadastros, copiando os valores.

        Refazer é permitido e substitui: enquanto não foi paga, a folha é um
        rascunho. O que não se pode é recalcular sozinha depois de paga.
        """
        mes = mes or mes_de_hoje()
        pessoas = self.pessoas()
        if not pessoas:
            raise ValueError(
                "nenhuma pessoa com vínculo definido. Abra Cadastros, escolha a "
                "pessoa e diga como ela é paga e quanto recebe."
            )

        ja = self.do_mes(mes)
        if ja["lancamento_id"]:
            raise ValueError("a folha deste mês já virou conta a pagar; reabra o lançamento antes")

        self.base.escrever("DELETE FROM folha WHERE mes = ?", (mes,))
        for p in pessoas:
            self.base.escrever(
                "INSERT INTO folha (mes, cadastro_id, nome, vinculo, salario_centavos, "
                "encargos_centavos, criado_em) VALUES (?, ?, ?, ?, ?, ?, ?)",
                (mes, p["id"], p["nome"], p["vinculo"],
                 int(p["salario_centavos"] or 0), int(p["encargos_centavos"] or 0), _agora()),
            )
        return self.do_mes(mes)

    def ligar_ao_lancamento(self, mes: str, lancamento_id: int) -> None:
        self.base.escrever(
            "UPDATE folha SET lancamento_id = ? WHERE mes = ?", (lancamento_id, mes)
        )

    def apagar_mes(self, mes: str) -> int:
        return self.base.escrever("DELETE FROM folha WHERE mes = ?", (mes,))


# --------------------------------------------------------- notas e boletos


class PapeisFiscais:
    """
    Nota fiscal e boleto: o registro do que existe fora daqui.

    Os dois são a mesma forma — número, valor, data, a quem se refere — e a
    diferença está no que a data significa: na nota é quando foi emitida, no
    boleto é quando vence. Por isso moram na mesma tabela com um campo `tipo`:
    duas tabelas iguais divergiriam na primeira mudança.
    """

    def __init__(self, base) -> None:
        self.base = base

    def listar(self, tipo: str, mes: str = "") -> list[dict]:
        sql = "SELECT p.*, c.nome AS cliente FROM papeis_fiscais p " \
              "LEFT JOIN cadastros c ON c.id = p.cadastro_id WHERE p.tipo = ?"
        parametros: list = [tipo]
        if mes:
            sql += " AND substr(p.data,1,7) = ?"
            parametros.append(mes)
        sql += " ORDER BY p.data DESC, p.id DESC"

        linhas = self.base.buscar(sql, tuple(parametros))
        for l in linhas:
            l["valor"] = em_reais(l["centavos"])
            l["data_br"] = _br(l["data"])
            l["situacao_rotulo"] = SITUACOES_BOLETO.get(l["situacao"], l["situacao"])
            l["tipo_rotulo"] = TIPOS_PAPEL.get(l["tipo"], l["tipo"])
        return linhas

    def salvar(self, dados: dict, id_: int | None = None) -> int:
        tipo = str(dados.get("tipo", "")).strip()
        if tipo not in TIPOS_PAPEL:
            raise ValueError("é nota fiscal ou boleto")
        centavos = int(dados.get("centavos", 0) or 0)
        if centavos <= 0:
            raise ValueError("o valor precisa ser maior que zero")

        campos = (tipo, str(dados.get("numero", "")).strip(),
                  dados.get("cadastro_id") or None, dados.get("lancamento_id") or None,
                  centavos, str(dados.get("data", ""))[:10],
                  str(dados.get("situacao", "")).strip() or ("aberto" if tipo == "boleto" else "emitida"),
                  str(dados.get("observacao", "")).strip())

        if id_:
            self.base.escrever(
                "UPDATE papeis_fiscais SET tipo=?, numero=?, cadastro_id=?, lancamento_id=?, "
                "centavos=?, data=?, situacao=?, observacao=? WHERE id=?",
                campos + (id_,),
            )
            return id_
        return self.base.escrever(
            "INSERT INTO papeis_fiscais (tipo, numero, cadastro_id, lancamento_id, centavos, "
            "data, situacao, observacao, criado_em) VALUES (?,?,?,?,?,?,?,?,?)",
            campos + (_agora(),),
        )

    def apagar(self, id_: int) -> bool:
        return self.base.escrever("DELETE FROM papeis_fiscais WHERE id = ?", (id_,)) > 0

    def marcar_pago(self, id_: int, quando: str = "") -> bool:
        return self.base.escrever(
            "UPDATE papeis_fiscais SET situacao = 'pago', observacao = "
            "CASE WHEN observacao = '' THEN ? ELSE observacao END WHERE id = ? AND tipo = 'boleto'",
            (f"pago em {_br(quando or date.today().isoformat())}", id_),
        ) > 0

    def a_emitir(self, mes: str = "") -> list[dict]:
        """
        Recebimentos do mês que ainda não têm nota registrada.

        Não é uma lista de pendências inventada: é a diferença entre o que foi
        recebido e o que foi anotado como emitido. Se a nota foi emitida e não
        registrada aqui, aparece — e é justamente o que se quer ver.
        """
        mes = mes or mes_de_hoje()
        return self.base.buscar(
            "SELECT l.id, l.descricao, l.centavos, l.liquidado_em, c.nome AS cliente "
            "FROM lancamentos l LEFT JOIN cadastros c ON c.id = l.cadastro_id "
            "WHERE l.tipo = 'recebimento' AND l.liquidado_em != '' "
            "AND substr(l.liquidado_em,1,7) = ? "
            "AND NOT EXISTS (SELECT 1 FROM papeis_fiscais p "
            "                WHERE p.tipo = 'nota' AND p.lancamento_id = l.id) "
            "ORDER BY l.liquidado_em",
            (mes,),
        )


# ------------------------------------------------------- fechamento do mês


def contratos_concluidos(base, mes: str = "") -> list[dict]:
    """
    Quem terminou neste mês.

    Sai dos lançamentos, e não de um campo "encerrado": cliente que recebeu
    algo neste mês e não tem mais nada em aberto. Um campo separado seria uma
    segunda verdade sobre o mesmo cliente, e a errada é sempre a que ninguém
    lembrou de atualizar.
    """
    mes = mes or mes_de_hoje()
    linhas = base.buscar(
        "SELECT c.id, c.nome, "
        "       SUM(CASE WHEN substr(l.liquidado_em,1,7) = ? THEN l.centavos ELSE 0 END) recebido, "
        "       COUNT(CASE WHEN substr(l.liquidado_em,1,7) = ? THEN 1 END) quitados, "
        "       COUNT(CASE WHEN l.liquidado_em = '' THEN 1 END) abertos "
        "FROM cadastros c JOIN lancamentos l ON l.cadastro_id = c.id "
        "WHERE l.tipo = 'recebimento' "
        "GROUP BY c.id HAVING quitados > 0 AND abertos = 0 "
        "ORDER BY recebido DESC",
        (mes, mes),
    )
    for l in linhas:
        l["recebido_texto"] = em_reais(l["recebido"])
    return linhas


def comprovantes_do_mes(base, mes: str = "") -> list[dict]:
    mes = mes or mes_de_hoje()
    return base.buscar(
        "SELECT k.*, l.descricao, l.tipo, l.centavos, l.liquidado_em "
        "FROM comprovantes k JOIN lancamentos l ON l.id = k.lancamento_id "
        "WHERE substr(COALESCE(NULLIF(l.liquidado_em,''), l.vencimento),1,7) = ? "
        "ORDER BY k.criado_em DESC",
        (mes,),
    )


def sem_comprovante(base, mes: str = "") -> list[dict]:
    """
    O que saiu do caixa neste mês e não tem papel guardado.

    Só despesas liquidadas: cobrar comprovante de uma conta que ainda não foi
    paga seria pedir um papel que não existe.
    """
    mes = mes or mes_de_hoje()
    linhas = base.buscar(
        "SELECT l.id, l.descricao, l.centavos, l.liquidado_em FROM lancamentos l "
        "WHERE l.tipo = 'despesa' AND l.liquidado_em != '' "
        "AND substr(l.liquidado_em,1,7) = ? "
        "AND NOT EXISTS (SELECT 1 FROM comprovantes k WHERE k.lancamento_id = l.id) "
        "ORDER BY l.liquidado_em",
        (mes,),
    )
    for l in linhas:
        l["valor"] = em_reais(l["centavos"])
    return linhas


def fechamento(base, folha: Folha, papeis: PapeisFiscais, mes: str = "") -> dict:
    """
    O mês inteiro numa resposta: o que falta para fechar.

    Cada item é uma coisa que existe e está sem par — despesa paga sem
    comprovante, recebimento sem nota, boleto vencido. Nenhum deles é uma
    regra de escritório que inventei; todos são uma junção que não fechou.
    """
    mes = mes or mes_de_hoje()
    hoje = date.today().isoformat()

    faltas = []
    faltando_papel = sem_comprovante(base, mes)
    if faltando_papel:
        faltas.append({
            "o_que": "comprovante",
            "quantos": len(faltando_papel),
            "titulo": _quantos(len(faltando_papel), "despesa paga", "despesas pagas")
                      + " sem comprovante",
            "detalhe": ", ".join(l["descricao"] for l in faltando_papel[:3]),
        })

    notas = papeis.a_emitir(mes)
    if notas:
        faltas.append({
            "o_que": "nota",
            "quantos": len(notas),
            "titulo": _quantos(len(notas), "recebimento") + " sem nota registrada",
            "detalhe": ", ".join(l["descricao"] for l in notas[:3]),
        })

    vencidos = [b for b in papeis.listar("boleto")
                if b["situacao"] == "aberto" and b["data"] and b["data"] < hoje]
    if vencidos:
        faltas.append({
            "o_que": "boleto",
            "quantos": len(vencidos),
            "titulo": _quantos(len(vencidos), "boleto vencido", "boletos vencidos")
                      + " e ainda em aberto",
            "detalhe": ", ".join(b["cliente"] or b["numero"] or "sem nome" for b in vencidos[:3]),
        })

    da_folha = folha.do_mes(mes)
    if da_folha["quantos"] and not da_folha["lancamento_id"]:
        faltas.append({
            "o_que": "folha",
            "quantos": 1,
            "titulo": ("a folha de " + _quantos(da_folha["quantos"], "pessoa")
                       + " ainda não virou conta a pagar"),
            "detalhe": da_folha["total_texto"],
        })

    return {
        "mes": mes,
        "dias_para_fechar": dias_para_fechar(mes),
        "faltas": faltas,
        "pronto": not faltas,
    }


def _quantos(n: int, palavra: str, muitos: str = "") -> str:
    """"1 pessoa", "6 pessoas" — e nunca "6 pessoa(s)", que ninguém escreve."""
    return f"{n} {palavra if n == 1 else (muitos or palavra + 's')}"


def _br(iso: str) -> str:
    if not iso or len(iso) < 10:
        return iso or ""
    return f"{iso[8:10]}/{iso[5:7]}/{iso[0:4]}"


# -------------------------------------------------------- recibos e planilha


def ultimo_dia(mes: str) -> str:
    """O último dia do mês, em ISO. É quando a folha vence."""
    try:
        ano, numero = (int(x) for x in mes.split("-"))
    except ValueError:
        return ""
    return date(ano, numero, monthrange(ano, numero)[1]).isoformat()


def mes_por_extenso(mes: str) -> str:
    try:
        ano, numero = (int(x) for x in mes.split("-"))
        return f"{_MESES[numero - 1]} de {ano}"
    except (ValueError, IndexError):
        return mes


_MESES = ["janeiro", "fevereiro", "março", "abril", "maio", "junho",
          "julho", "agosto", "setembro", "outubro", "novembro", "dezembro"]


def gerar_recibos(da_folha: dict, destino: Path | str, mes_rotulo: str,
                  escritorio_nome: str = "") -> list[dict]:
    """
    Um recibo em PDF por pessoa da folha.

    O valor sai do que está gravado na folha daquele mês, e não do salário de
    hoje — é essa a razão de a folha ser cópia. Um recibo que muda de valor
    depois de assinado não é recibo.

    O recibo sai sem assinatura: assinar é da tela de assinar, com o
    certificado, e fingir aqui uma assinatura que não existe seria produzir um
    documento que parece valer e não vale.
    """
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

    pasta = Path(destino)
    pasta.mkdir(parents=True, exist_ok=True)

    corpo = ParagraphStyle("corpo", fontName="Times-Roman", fontSize=11, leading=17)
    titulo = ParagraphStyle("titulo", fontName="Times-Bold", fontSize=14, leading=20,
                            spaceAfter=18, alignment=1)
    linha = ParagraphStyle("linha", fontName="Times-Roman", fontSize=10, leading=15,
                           textColor="#555555")

    feitos: list[dict] = []
    for pessoa in da_folha["pessoas"]:
        arquivo = pasta / f"recibo-{_arquivavel(pessoa['nome'])}-{da_folha['mes']}.pdf"
        doc = SimpleDocTemplate(
            str(arquivo), pagesize=A4,
            leftMargin=3 * cm, rightMargin=2 * cm, topMargin=3 * cm, bottomMargin=2 * cm,
            title=f"Recibo - {pessoa['nome']} - {mes_rotulo}",
        )

        partes = [Paragraph("RECIBO DE PAGAMENTO", titulo)]
        if escritorio_nome:
            partes.append(Paragraph(escritorio_nome, linha))
            partes.append(Spacer(1, 10))

        natureza = NO_RECIBO.get(pessoa["vinculo"], "pagamento")
        partes.append(Paragraph(
            f"Recebi de {escritorio_nome or 'este escritório'} a importância de "
            f"<b>{em_reais(pessoa['total_centavos'])}</b>, referente a "
            f"{natureza} do mês de {mes_rotulo}.", corpo))
        partes.append(Spacer(1, 14))
        partes.append(Paragraph(
            f"{natureza.capitalize()}: {pessoa['salario']}<br/>"
            + (f"Encargos: {pessoa['encargos']}<br/>" if pessoa["encargos_centavos"] else "")
            + f"<b>Total: {pessoa['total']}</b>", corpo))
        partes.append(Spacer(1, 40))
        partes.append(Paragraph("_" * 46, corpo))
        partes.append(Paragraph(pessoa["nome"], corpo))
        partes.append(Spacer(1, 26))
        partes.append(Paragraph(
            "Documento gerado pelo PAULUS. Assinar com certificado digital é "
            "na tela Assinar documento — este PDF sai sem assinatura.", linha))

        doc.build(partes)
        feitos.append({"nome": pessoa["nome"], "arquivo": str(arquivo),
                       "valor": pessoa["total"]})

    return feitos


def _arquivavel(nome: str) -> str:
    import re
    import unicodedata

    plano = unicodedata.normalize("NFKD", nome or "")
    plano = "".join(c for c in plano if not unicodedata.combining(c))
    return re.sub(r"[^A-Za-z0-9]+", "-", plano).strip("-").lower() or "pessoa"


def exportar_mes(destino: Path | str, mes: str, *, extrato: dict, folha: dict,
                 notas: list[dict], boletos: list[dict]) -> Path:
    """
    O mês numa planilha, com uma aba por assunto.

    Existe para quem faz a contabilidade: o contador não abre este programa, e
    mandar print de tela não é prestação de contas. Valores vão como número, e
    não como texto — planilha que chega com "R$ 1.200,00" em célula de texto
    não soma do outro lado.
    """
    from openpyxl import Workbook

    livro = Workbook()
    aba = livro.active
    aba.title = "Extrato"
    aba.append(["Data", "Tipo", "Descrição", "Categoria", "Cliente",
                "Valor (R$)", "Saldo (R$)", "Comprovantes"])
    for item in extrato.get("linhas", []):
        aba.append([
            item.get("liquidado_em", ""),
            item.get("tipo_rotulo", item.get("tipo", "")),
            item.get("descricao", ""),
            item.get("categoria_rotulo", item.get("categoria", "")),
            item.get("cadastro_nome") or "",
            round(item.get("movimento", 0) / 100, 2),
            round(item.get("saldo", 0) / 100, 2),
            item.get("comprovantes", 0),
        ])

    se_folha = livro.create_sheet("Folha")
    se_folha.append(["Pessoa", "Vínculo", "Salário (R$)", "Encargos (R$)", "Total (R$)"])
    for pessoa in folha.get("pessoas", []):
        se_folha.append([
            pessoa["nome"], pessoa["vinculo_rotulo"],
            round(pessoa["salario_centavos"] / 100, 2),
            round(pessoa["encargos_centavos"] / 100, 2),
            round(pessoa["total_centavos"] / 100, 2),
        ])

    se_papeis = livro.create_sheet("Notas e boletos")
    se_papeis.append(["Tipo", "Número", "Cliente", "Data", "Valor (R$)", "Situação"])
    for papel in list(notas) + list(boletos):
        se_papeis.append([
            papel["tipo_rotulo"], papel.get("numero", ""), papel.get("cliente") or "",
            papel.get("data", ""), round(papel["centavos"] / 100, 2),
            papel.get("situacao_rotulo", ""),
        ])

    for folha_da_vez in (aba, se_folha, se_papeis):
        for coluna in folha_da_vez.columns:
            letra = coluna[0].column_letter
            largura = max((len(str(c.value or "")) for c in coluna), default=10)
            folha_da_vez.column_dimensions[letra].width = min(46, max(11, largura + 2))

    alvo = Path(destino)
    alvo.parent.mkdir(parents=True, exist_ok=True)
    livro.save(str(alvo))
    return alvo
