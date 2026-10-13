"""
PAULUS - O .docx de ida e volta (docs/PLANO-EDITOR.md, E1).

Abrir um .docx do Acervo no editor era ler o texto extraído - uma linha visual
virava um parágrafo, os quadros iam para o fim, sumiam títulos, negrito e a
numeração - e devolver um .docx novo, montado do zero, sem o cabeçalho, o
rodapé, a fonte e os estilos do escritório. O Word de ida não era o Word de
volta.

Aqui:

  ler(dados)      o .docx vira os blocos do editor (src/blocos.py), na ordem do
                  documento: títulos pelo estilo, negrito/itálico/sublinhado/
                  tachado de cada trecho, alinhamento, quadro no lugar dele, e a
                  numeração automática do Word como rótulo (`data-num`: "1.1",
                  "a)") - o número não é texto editável, como no Word. Cada
                  bloco leva `data-docx` com o lugar do parágrafo no arquivo.

  escrever(original, blocos)
                  o .docx de volta, REMENDANDO o original: o parágrafo que não
                  mudou volta exatamente como estava; o que mudou mantém o
                  estilo, o recuo e a numeração e troca só o texto; o novo copia
                  a cara do vizinho do mesmo tipo; o apagado sai. Cabeçalho,
                  rodapé, margens, fontes e estilos ficam - são do arquivo.
"""

from __future__ import annotations

import copy
import html as html_mod
import io
import re

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

import blocos as B

_TITULOS = {"title": "h1", "heading 1": "h1", "título 1": "h1", "titulo 1": "h1", "heading 2": "h2", "título 2": "h2",
            "titulo 2": "h2", "heading 3": "h3", "título 3": "h3", "titulo 3": "h3", "heading 4": "h4",
            "título 4": "h4", "titulo 4": "h4", "subtitle": "h2", "subtítulo": "h2"}
_CITACAO = {"quote", "citação", "citacao", "intense quote", "citação intensa"}
_ALINHAR = {"CENTER": "center", "RIGHT": "right", "JUSTIFY": "justify", "DISTRIBUTE": "justify"}


# ------------------------------------------------------------ elementos

def _elementos(doc) -> list[tuple[str, object]]:
    """
    Os parágrafos e quadros do corpo, na ordem, com a chave estável ("p12",
    "t3"). Desce em controle de conteúdo (w:sdt). Ler e escrever usam a mesma
    lista - é ela que casa o bloco do editor com o parágrafo do arquivo.
    """
    saida: list[tuple[str, object]] = []
    contagem = {"p": 0, "t": 0}

    def andar(pai) -> None:
        for el in pai.iterchildren():
            if el.tag == qn("w:p"):
                saida.append((f"p{contagem['p']}", el))
                contagem["p"] += 1
            elif el.tag == qn("w:tbl"):
                saida.append((f"t{contagem['t']}", el))
                contagem["t"] += 1
            elif el.tag == qn("w:sdt"):
                conteudo = el.find(qn("w:sdtContent"))
                if conteudo is not None:
                    andar(conteudo)

    andar(doc.element.body)
    return saida


# ------------------------------------------------------------ numeração

_FORMATOS = {
    "decimal": lambda n: str(n),
    "decimalZero": lambda n: f"{n:02d}",
    "lowerLetter": lambda n: _letra(n).lower(),
    "upperLetter": lambda n: _letra(n),
    "lowerRoman": lambda n: _romano(n).lower(),
    "upperRoman": lambda n: _romano(n),
    "ordinal": lambda n: f"{n}º",
}


def _letra(n: int) -> str:
    saida = ""
    while n > 0:
        n, r = divmod(n - 1, 26)
        saida = chr(65 + r) + saida
    return saida or "A"


def _romano(n: int) -> str:
    import esboco

    return esboco.para_romano(max(1, n))


class _Numeracao:
    """
    O número que o Word mostra na frente de um parágrafo de lista. O Word
    calcula isso na hora de desenhar - não está no texto -, então aqui se faz
    a mesma conta, no subconjunto que os documentos de escritório usam:
    decimal, letras, romanos, "%1.%2." e o reinício dos níveis de baixo.
    """

    def __init__(self, doc) -> None:
        self.niveis: dict[str, dict[int, dict]] = {}     # abstractNumId -> nivel -> {fmt, texto, inicio}
        self.de_num: dict[str, str] = {}                 # numId -> abstractNumId
        self.inicio_forcado: dict[tuple[str, int], int] = {}
        self.contas: dict[str, dict[int, int]] = {}
        self.estilos = doc.styles
        try:
            numeracao = doc.part.numbering_part.element
        except Exception:  # noqa: BLE001 - documento sem numeração
            return
        for absn in numeracao.findall(qn("w:abstractNum")):
            aid = absn.get(qn("w:abstractNumId"))
            niveis = {}
            for lvl in absn.findall(qn("w:lvl")):
                ilvl = int(lvl.get(qn("w:ilvl")) or 0)
                fmt = lvl.find(qn("w:numFmt"))
                txt = lvl.find(qn("w:lvlText"))
                ini = lvl.find(qn("w:start"))
                niveis[ilvl] = {"fmt": fmt.get(qn("w:val")) if fmt is not None else "decimal",
                                "texto": txt.get(qn("w:val")) if txt is not None else f"%{ilvl + 1}.",
                                "inicio": int(ini.get(qn("w:val"))) if ini is not None else 1}
            self.niveis[aid] = niveis
        for num in numeracao.findall(qn("w:num")):
            nid = num.get(qn("w:numId"))
            absn = num.find(qn("w:abstractNumId"))
            if absn is not None:
                self.de_num[nid] = absn.get(qn("w:val"))
            for ov in num.findall(qn("w:lvlOverride")):
                ini = ov.find(qn("w:startOverride"))
                if ini is not None:
                    self.inicio_forcado[(nid, int(ov.get(qn("w:ilvl")) or 0))] = int(ini.get(qn("w:val")))

    def _num_pr(self, p) -> tuple[str, int] | None:
        """(numId, nível) do parágrafo - dele mesmo ou do estilo dele."""
        ppr = p.find(qn("w:pPr"))
        num_pr = ppr.find(qn("w:numPr")) if ppr is not None else None
        estilo_id = None
        if ppr is not None and ppr.find(qn("w:pStyle")) is not None:
            estilo_id = ppr.find(qn("w:pStyle")).get(qn("w:val"))
        nid = ilvl = None
        if num_pr is not None:
            n = num_pr.find(qn("w:numId"))
            lv = num_pr.find(qn("w:ilvl"))
            nid = n.get(qn("w:val")) if n is not None else None
            ilvl = int(lv.get(qn("w:val"))) if lv is not None else None
        if nid is None and estilo_id:
            # A numeração pode vir do estilo ("List Number").
            visto = 0
            while estilo_id and visto < 8:
                visto += 1
                estilo = next((s for s in self.estilos if getattr(s, "style_id", None) == estilo_id), None)
                if estilo is None:
                    break
                sppr = estilo.element.find(qn("w:pPr"))
                snum = sppr.find(qn("w:numPr")) if sppr is not None else None
                if snum is not None:
                    n = snum.find(qn("w:numId"))
                    lv = snum.find(qn("w:ilvl"))
                    nid = n.get(qn("w:val")) if n is not None else None
                    if ilvl is None and lv is not None:
                        ilvl = int(lv.get(qn("w:val")))
                    break
                base = estilo.element.find(qn("w:basedOn"))
                estilo_id = base.get(qn("w:val")) if base is not None else None
        if not nid or nid == "0":
            return None
        return nid, ilvl or 0

    def rotulo(self, p) -> tuple[str, str]:
        """("1.2.", "decimal") do parágrafo, ou ("", "") se ele não é de lista."""
        achado = self._num_pr(p)
        if not achado:
            return "", ""
        nid, ilvl = achado
        aid = self.de_num.get(nid)
        niveis = self.niveis.get(aid or "", {})
        if ilvl not in niveis:
            return "", ""
        contas = self.contas.setdefault(aid, {})
        if ilvl in contas:
            contas[ilvl] += 1
        else:
            contas[ilvl] = self.inicio_forcado.get((nid, ilvl), niveis[ilvl]["inicio"])
        for mais_fundo in [k for k in contas if k > ilvl]:
            del contas[mais_fundo]
        nivel = niveis[ilvl]
        if nivel["fmt"] == "bullet":
            return "•", "bullet"
        if nivel["fmt"] == "none":
            return "", ""

        def trocar(m) -> str:
            k = int(m.group(1)) - 1
            fmt = niveis.get(k, {}).get("fmt", "decimal")
            valor = contas.get(k, niveis.get(k, {}).get("inicio", 1))
            return _FORMATOS.get(fmt, _FORMATOS["decimal"])(valor)

        return re.sub(r"%(\d)", trocar, nivel["texto"]).strip(), nivel["fmt"]


# ------------------------------------------------------------------ ler

def _estilo(p) -> str:
    try:
        return (p.style.name or "").lower()
    except Exception:  # noqa: BLE001
        return ""


def _efetivo(valor, *alternativas):
    if valor is not None:
        return bool(valor)
    for a in alternativas:
        if a is not None:
            return bool(a)
    return False


def _corridas(p: Paragraph, herdar_do_estilo: bool = True) -> list[B.Corrida]:
    """
    Os trechos do parágrafo com a marcação de cada um. No título, o negrito
    do estilo não vira marca: o título já é negrito na tela e no arquivo.
    """
    estilo_negrito = estilo_italico = None
    if herdar_do_estilo:
        try:
            estilo_negrito = p.style.font.bold
            estilo_italico = p.style.font.italic
        except Exception:  # noqa: BLE001
            pass
    saida: list[B.Corrida] = []
    for r in p.runs:
        partes = []
        for filho in r._r.iterchildren():
            if filho.tag == qn("w:t"):
                partes.append(filho.text or "")
            elif filho.tag in (qn("w:tab"), qn("w:ptab")):
                partes.append(" ")
            elif filho.tag in (qn("w:br"), qn("w:cr")):
                partes.append("\n")
        texto = "".join(partes)
        if not texto:
            continue
        estilo_run = getattr(r, "style", None)
        rf = getattr(estilo_run, "font", None)
        marcas = []
        if _efetivo(r.bold, getattr(rf, "bold", None), estilo_negrito):
            marcas.append("b")
        if _efetivo(r.italic, getattr(rf, "italic", None), estilo_italico):
            marcas.append("i")
        if r.underline:
            marcas.append("u")
        if r.font.strike:
            marcas.append("s")
        if r.font.superscript:
            marcas.append("sup")
        if r.font.subscript:
            marcas.append("sub")
        for k, pedaco in enumerate(texto.split("\n")):
            if k:
                saida.append(B.Corrida("\n", tuple(marcas)))
            if pedaco:
                saida.append(B.Corrida(pedaco, tuple(marcas)))
    return saida


def _alinhamento(p: Paragraph) -> str:
    a = p.alignment
    if a is None:
        try:
            a = p.style.paragraph_format.alignment
        except Exception:  # noqa: BLE001
            a = None
    nome = getattr(a, "name", "") or ""
    return _ALINHAR.get(nome.upper(), "")


def ler(dados: bytes) -> dict:
    """
    {"html": o corpo do editor, "blocos": a lista, "cabecalho": [linhas],
    "rodape": [linhas], "avisos": [...]}.
    """
    doc = Document(io.BytesIO(dados))
    numeracao = _Numeracao(doc)
    lista: list[B.Bloco] = []
    avisos: list[str] = []
    imagens = 0
    for chave, el in _elementos(doc):
        if chave.startswith("t"):
            t = Table(el, doc)
            linhas = []
            for linha in t.rows:
                vistos = set()
                celulas = []
                for c in linha.cells:
                    if id(c._tc) in vistos:     # célula mesclada aparece repetida
                        continue
                    vistos.add(id(c._tc))
                    celulas.append([False, " ".join(" ".join(par.text for par in c.paragraphs).split())])
                linhas.append(celulas)
            if linhas:
                for c in linhas[0]:
                    c[0] = True
            largura = max((len(x) for x in linhas), default=0)
            linhas = [x + [[False, ""]] * (largura - len(x)) for x in linhas]
            lista.append(B.Bloco(id="", tipo="table", linhas=linhas, docx=chave))
            continue
        p = Paragraph(el, doc)
        if el.findall(".//" + qn("w:drawing")) or el.findall(".//" + qn("w:pict")):
            imagens += 1
        estilo = _estilo(p)
        rotulo, fmt = numeracao.rotulo(el)
        tipo = _TITULOS.get(estilo, "")
        if not tipo and estilo in _CITACAO:
            tipo = "blockquote"
        b = B.Bloco(id="", tipo=tipo or "p", alinhar=_alinhamento(p), docx=chave)
        if fmt == "bullet":
            b.tipo, b.lista = "li", "ul"
            nivel = numeracao._num_pr(el)
            b.nivel = nivel[1] if nivel else 0
        elif rotulo:
            b.num = rotulo
        b.corridas = _corridas(p, herdar_do_estilo=b.tipo not in ("h1", "h2", "h3", "h4"))
        lista.append(b)
    blocos = B._dar_nomes(lista)
    if imagens:
        avisos.append(f"{imagens} parágrafo(s) com imagem: a imagem não aparece no editor, mas fica no arquivo "
                      "enquanto o parágrafo não mudar")
    return {"html": B.html(blocos), "blocos": blocos, "cabecalho": _linhas_de(doc, "header"),
            "rodape": _linhas_de(doc, "footer"), "avisos": avisos}


def _linhas_de(doc, qual: str) -> list[str]:
    linhas: list[str] = []
    for secao in doc.sections:
        parte = getattr(secao, qual, None)
        if parte is None or parte.is_linked_to_previous and linhas:
            continue
        for p in parte.paragraphs:
            t = " ".join(p.text.split())
            if t and t not in linhas:
                linhas.append(t)
    return linhas[:6]


# -------------------------------------------------------------- escrever

def _assinatura(b: B.Bloco) -> tuple:
    """O que, de um bloco, o .docx guarda: o texto com a marcação, o tipo e o alinhamento."""
    if b.tipo == "table":
        return ("table", tuple(tuple(c[1] for c in linha) for linha in b.linhas))
    corridas = [(c.texto, c.marcas) for c in b.corridas if c.sug != "del"]
    juntas: list = []
    for t, m in corridas:
        if juntas and juntas[-1][1] == m and t != "\n" and juntas[-1][0] != "\n":
            juntas[-1] = (juntas[-1][0] + t, m)
        else:
            juntas.append((t, m))
    return (b.tipo, b.alinhar, tuple(juntas))


def _rpr_base(p_el):
    """As propriedades do primeiro trecho do parágrafo (fonte, corpo, cor), sem negrito/itálico."""
    for r in p_el.findall(qn("w:r")):
        rpr = r.find(qn("w:rPr"))
        if rpr is not None:
            base = copy.deepcopy(rpr)
            for marca in ("w:b", "w:bCs", "w:i", "w:iCs", "w:u", "w:strike", "w:dstrike", "w:vertAlign"):
                for x in base.findall(qn(marca)):
                    base.remove(x)
            return base
    return None


def _w(tag: str):
    from docx.oxml import OxmlElement

    return OxmlElement(tag)


def _run(texto: str, marcas: tuple, base_rpr):
    r = _w("w:r")
    rpr = copy.deepcopy(base_rpr) if base_rpr is not None else _w("w:rPr")
    for m, tag in (("b", "w:b"), ("i", "w:i"), ("s", "w:strike")):
        if m in marcas:
            rpr.append(_w(tag))
    if "u" in marcas:
        u = _w("w:u")
        u.set(qn("w:val"), "single")
        rpr.append(u)
    if "sup" in marcas or "sub" in marcas:
        va = _w("w:vertAlign")
        va.set(qn("w:val"), "superscript" if "sup" in marcas else "subscript")
        rpr.append(va)
    if len(rpr):
        r.append(rpr)
    if texto == "\n":
        r.append(_w("w:br"))
        return r
    t = _w("w:t")
    t.text = texto
    t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    r.append(t)
    return r


def _trocar_texto(p_el, b: B.Bloco) -> None:
    """O parágrafo fica (pPr: estilo, recuo, numeração); os trechos são os do bloco."""
    base = _rpr_base(p_el)
    for filho in list(p_el):
        if filho.tag != qn("w:pPr"):
            p_el.remove(filho)
    for c in b.corridas:
        if c.sug == "del" or not c.texto:
            continue
        p_el.append(_run(c.texto, c.marcas, base))
    if b.alinhar or _alinhamento_do_el(p_el):
        _alinhar(p_el, b.alinhar)


def _alinhamento_do_el(p_el) -> str:
    ppr = p_el.find(qn("w:pPr"))
    jc = ppr.find(qn("w:jc")) if ppr is not None else None
    return jc.get(qn("w:val")) if jc is not None else ""


def _alinhar(p_el, alinhar: str) -> None:
    ppr = p_el.find(qn("w:pPr"))
    if ppr is None:
        ppr = _w("w:pPr")
        p_el.insert(0, ppr)
    for jc in ppr.findall(qn("w:jc")):
        ppr.remove(jc)
    valor = {"center": "center", "right": "right", "justify": "both"}.get(alinhar, "")
    if valor:
        jc = _w("w:jc")
        jc.set(qn("w:val"), valor)
        ppr.append(jc)


def _paragrafo_novo(b: B.Bloco, molde, manter_numeracao: bool):
    p_el = _w("w:p")
    base = None
    if molde is not None:
        ppr = molde.find(qn("w:pPr"))
        if ppr is not None:
            novo_ppr = copy.deepcopy(ppr)
            if not manter_numeracao:
                for np_ in novo_ppr.findall(qn("w:numPr")):
                    novo_ppr.remove(np_)
            for rpr in novo_ppr.findall(qn("w:rPr")):
                novo_ppr.remove(rpr)
            p_el.append(novo_ppr)
        base = _rpr_base(molde)
    for c in b.corridas:
        if c.sug == "del" or not c.texto:
            continue
        p_el.append(_run(c.texto, c.marcas, base))
    if b.alinhar or molde is not None:
        _alinhar(p_el, b.alinhar)
    return p_el


def _quadro_novo(doc, b: B.Bloco):
    colunas = max((len(x) for x in b.linhas), default=1)
    t = doc.add_table(rows=0, cols=colunas)
    try:
        t.style = "Table Grid"
    except Exception:  # noqa: BLE001
        pass
    for linha in b.linhas:
        celulas = t.add_row().cells
        for j in range(colunas):
            texto = linha[j][1] if j < len(linha) else ""
            run = celulas[j].paragraphs[0].add_run(texto)
            run.bold = bool(j < len(linha) and linha[j][0])
    el = t._tbl
    el.getparent().remove(el)
    return el


def _trocar_quadro(doc, el, b: B.Bloco):
    t = Table(el, doc)
    linhas_el = list(t.rows)
    if len(linhas_el) == len(b.linhas) and all(len({id(c._tc) for c in r.cells}) == len(b.linhas[i])
                                                 for i, r in enumerate(linhas_el)):
        for i, r in enumerate(linhas_el):
            vistos, k = set(), 0
            for c in r.cells:
                if id(c._tc) in vistos:
                    continue
                vistos.add(id(c._tc))
                pars = c.paragraphs
                p_el = pars[0]._p
                base = _rpr_base(p_el)
                for filho in list(p_el):
                    if filho.tag != qn("w:pPr"):
                        p_el.remove(filho)
                p_el.append(_run(b.linhas[i][k][1], ("b",) if b.linhas[i][k][0] and _tinha_negrito(pars[0]) else (), base))
                for extra in pars[1:]:
                    extra._p.getparent().remove(extra._p)
                k += 1
        return el
    return _quadro_novo(doc, b)


def _tinha_negrito(p: Paragraph) -> bool:
    return any(r.bold for r in p.runs)


def escrever(original: bytes, blocos_editor: list[B.Bloco]) -> bytes:
    """
    O .docx original remendado com o que está no editor (as sugestões
    pendentes contam como recusadas: só sai o que foi aceito).
    """
    doc = Document(io.BytesIO(original))
    elementos = dict(_elementos(doc))
    importados = {b.docx: _assinatura(b) for b in ler(original)["blocos"] if b.docx}
    corpo = doc.element.body
    sect = corpo.find(qn("w:sectPr"))

    aceitos = B.decidir(blocos_editor, False)
    novos: list = []
    molde_por_tipo: dict[str, object] = {}
    anterior_numerado = False
    for b in aceitos:
        el = elementos.get(b.docx) if b.docx else None
        if b.tipo == "table":
            if el is not None and el.tag == qn("w:tbl"):
                if importados.get(b.docx) != _assinatura(b):
                    el = _trocar_quadro(doc, el, b)
                novos.append(el)
            else:
                novos.append(_quadro_novo(doc, b))
            anterior_numerado = False
            continue
        chave_molde = _chave_molde(b)
        if el is not None and el.tag == qn("w:p"):
            if importados.get(b.docx) != _assinatura(b):
                _trocar_texto(el, b)
            novos.append(el)
            molde_por_tipo[chave_molde] = el
            anterior_numerado = bool(b.num)
            continue
        molde = molde_por_tipo.get(chave_molde)
        if molde is None:
            molde = next((elementos[x.docx] for x in aceitos if x.docx and _chave_molde(x) == chave_molde
                          and x.docx in elementos and elementos[x.docx].tag == qn("w:p")), None)
        novos.append(_paragrafo_novo(b, molde, manter_numeracao=anterior_numerado))
    # O corpo novo: os elementos na ordem do editor; o que não está mais no
    # editor sai. O sectPr (margens, cabeçalho, rodapé) fica no fim.
    for chave, el in _elementos(doc):
        pai = el.getparent()
        if pai is not None:
            pai.remove(el)
    for filho in list(corpo):
        if filho.tag == qn("w:sdt"):
            conteudo = filho.find(qn("w:sdtContent"))
            if conteudo is not None and not len(conteudo):
                corpo.remove(filho)
    for el in novos:
        if sect is not None:
            sect.addprevious(el)
        else:
            corpo.append(el)
    saida = io.BytesIO()
    doc.save(saida)
    return saida.getvalue()


def _chave_molde(b: B.Bloco) -> str:
    """O tipo de parágrafo para copiar a cara: título, cabeça de cláusula, item, corpo."""
    import esboco

    if b.tipo in ("h1", "h2", "h3", "h4", "blockquote", "li"):
        return b.tipo
    if esboco._numero_da_clausula(b.texto)[0] or esboco.RE_TITULO_NUMERADO.match(b.texto):
        return "cabeca"
    return "corpo"


def para_html_simples(texto: str) -> str:
    return "".join(f"<p>{html_mod.escape(p)}</p>" for p in texto.split("\n") if p.strip())
