"""
PAULUS Legal - Backend web (FastAPI).

Serve a interface e expoe a mesma logica do CLI por HTTP. Roda so em
localhost: o navegador precisa estar na mesma maquina que o Ollama, e nenhum
documento sai daqui.

    python src/api.py
    http://localhost:8000
"""

from __future__ import annotations

import errno
import json
import os
import re
import shutil
import tempfile
import subprocess
import sys
import queue
import threading
import time
import uuid
import unicodedata
from datetime import date, datetime
from collections.abc import Iterator
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))

from contextlib import asynccontextmanager

from fastapi import FastAPI, File, Form, HTTPException, Query, Request, UploadFile
from fastapi.concurrency import run_in_threadpool
from fastapi.responses import FileResponse, Response, StreamingResponse
from pydantic import BaseModel

import requests

import aprovacoes as fila_aprovacoes
import apoio
import assinatura
import traducao
import certificado
import correio
import correio_contas
import correio_oauth
import segredos
import destinos
import avisos
import bemestar
import conexoes
import documento
import financeiro
import extrato
import acervo
import citacao
import escritorio
import ferramentas
import intencao
import leis
import redacao
import ritmo as ritmo_mod
import planilha
import relatorios
import servicos as servicos_mod
import gravacoes as gravacoes_mod
import transcricao as transcricao_mod
import lixeira as lixeira_mod
import contextos as contextos_mod
from inteligencia import portas as inteligencia
from inteligencia.catalogo import Catalogo
from inteligencia.guarda import Biblioteca
import marca as marca_mod
import pastas
import recursos
import registro
from classify import Classificacao, ROTULOS
from agenda import Agenda, CAMPOS as CAMPOS_AGENDA, ONDES, TIPOS as TIPOS_AGENDA
from base import Base
from cadastros import Cadastros, TIPOS as TIPOS_CADASTRO
from config import Preferencias
import tarefas as tarefas_mod
from tarefas import Tarefas
from habilidade_base import (
    PRECISA_ASSISTENTE,
    PRECISA_DOCUMENTOS,
    Contexto,
)
from extract import SUPPORTED_SUFFIXES, extract_file, file_sha1, index_all_contracts
from jobs import AGUARDANDO, CONCLUIDO, EXECUTANDO, LIMITE_DE_NOME, PAUSADO, Etapa, Trabalhos, titular
from jobs import agora as jobs_agora
from llama_client import (
    DEFAULT_MODEL,
    LlamaClient,
    OllamaError,
    check_ollama,
    janela_para,
)
from organize import (
    PADROES_SUGERIDOS,
    aplicar_plano,
    desfazer,
    listar_diarios,
    montar_plano,
)
from scan import escanear
from search import ContractSearcher

BASE_DIR = Path(__file__).parent.parent
CONTRACTS_DIR = BASE_DIR / "data" / "test_contracts"
NOME_DA_PASTA_PADRAO = "Documentos do escritório"
CACHE_PATH = BASE_DIR / "data" / "extractions" / "index.json"
CLASSIFICACAO_PATH = BASE_DIR / "data" / "extractions" / "classificacao.json"
# Os nomes que a pessoa recusou na sugestao de cadastro (Cadastros).
SUGESTOES_IGNORADAS_PATH = BASE_DIR / "data" / "cadastros_ignorados.json"
DIARIOS_DIR = BASE_DIR / "data" / "diarios"
TRABALHOS_DIR = BASE_DIR / "data" / "trabalhos"
APROVACOES_PATH = BASE_DIR / "data" / "aprovacoes.json"
PREFERENCIAS_PATH = BASE_DIR / "data" / "preferencias.json"
BASE_PATH = BASE_DIR / "data" / "paulus.db"
CERTIFICADO_DIR = BASE_DIR / "data" / "certificado"
COFRE_PATH = CERTIFICADO_DIR / "cofre.json"
ASSINATURAS_PATH = BASE_DIR / "data" / "assinaturas.json"
CONTAS_EMAIL_PATH = BASE_DIR / "data" / "contas_email.json"
ENVIOS_PATH = BASE_DIR / "data" / "envios.json"
CLAUSULAS_PATH = BASE_DIR / "data" / "clausulas.json"
CONEXOES_PATH = BASE_DIR / "data" / "conexoes.json"
SESSOES_DIR = BASE_DIR / "data" / "sessoes"
COMPROVANTES_DIR = BASE_DIR / "data" / "comprovantes"
RECIBOS_DIR = BASE_DIR / "data" / "recibos"
EXPORTACOES_DIR = BASE_DIR / "data" / "exportacoes"
GRAVACOES_DIR = BASE_DIR / "data" / "gravacoes"
MODELOS_VOZ_DIR = BASE_DIR / "data" / "modelos" / "whisper"
LIXEIRA_DIR = BASE_DIR / "data" / "lixeira"
MARCA_DIR = BASE_DIR / "data" / "marca"
CONHECIMENTO_DIR = BASE_DIR / "data" / "conhecimento"
MAX_AUDIO_BYTES = 500 * 1024 * 1024
RITMO_PATH = BASE_DIR / "data" / "ritmo.json"
HABILIDADES_DIR = BASE_DIR / "habilidades"
FRONTEND_DIR = BASE_DIR / "frontend"
MAX_UPLOAD_BYTES = 50 * 1024 * 1024
MAX_CERTIFICADO_BYTES = 8 * 1024 * 1024


def _quem_sou() -> str:
    """O primeiro nome de quem usa esta maquina, para assinar a trilha."""
    nome = str((estado.prefs.dados.get("pessoa") or {}).get("nome", "")).strip()
    return nome.split(" ")[0] if nome else "você"


def _credenciais_oauth(provedor: str) -> dict:
    """
    Client ID (e, no Google, o secret) do aplicativo PAULUS - vem do codigo
    (src/oauth_app.py), igual em toda instalacao. Quem usa so faz login.
    """
    import oauth_app

    return oauth_app.credenciais(provedor)


class Estado:
    """Indice em memoria, compartilhado entre as requisicoes."""

    def __init__(self) -> None:
        self.searcher = ContractSearcher()
        self.pasta = CONTRACTS_DIR
        self.porta = 8000
        self.client = LlamaClient()
        # Organizador: resultado da ultima varredura/classificacao, por caminho.
        self.encontrados: list[dict] = []
        self.classificacoes: dict[str, Classificacao] = {}
        # O fim da ultima organizacao (movidos, falhas, diario para desfazer):
        # quando o mover passa pela fila, e daqui que a tela tira o desfecho.
        self.ultima_organizacao: dict | None = None
        # Conversas persistidas e o interruptor "ir devagar".
        self.trabalhos = Trabalhos(TRABALHOS_DIR)
        self.devagar = False
        # Levantado quando a pessoa pede para parar a leitura em andamento.
        self.cancelar = threading.Event()
        # O botao de parar da conversa: um sinal por resposta sendo escrita,
        # pelo id da conversa. Sai daqui quando a resposta termina.
        self.respondendo: dict[str, threading.Event] = {}
        # O andamento de verdade de cada resposta sendo escrita, para o cartao
        # de "Acontecendo agora": a fase, desde quando, a previsao de leitura
        # medida nesta maquina (quando ha) e as palavras ja escritas. So em
        # memoria - acabou a resposta, sai daqui.
        self.andamento: dict[str, dict] = {}
        # Habilidades carregadas da pasta do projeto, uma por arquivo.
        self.registro = registro.carregar(HABILIDADES_DIR)
        # Fila do que espera decisao humana, e as preferencias da casa.
        self.fila = fila_aprovacoes.Fila(APROVACOES_PATH)
        self.prefs = Preferencias(PREFERENCIAS_PATH)
        # Base local: cadastros, tarefas e o que vier depois.
        self.base = Base(BASE_PATH)
        self.cadastros = Cadastros(self.base, SUGESTOES_IGNORADAS_PATH)
        self.tarefas = Tarefas(self.base)
        self.agenda = Agenda(self.base)
        # Certificado digital: o arquivo, o selo e o historico de assinaturas.
        self.cofre = certificado.Cofre(COFRE_PATH, CERTIFICADO_DIR)
        self.assinaturas = assinatura.Registro(ASSINATURAS_PATH)
        # Contas de e-mail e o historico do que ja saiu daqui.
        self.contas = correio_contas.Contas(CONTAS_EMAIL_PATH)
        self.contas.credenciais_oauth = _credenciais_oauth
        self.envios = correio.RegistroEnvios(ENVIOS_PATH)
        # O login com Google/Microsoft em andamento (um por vez).
        self.entrada_oauth: correio_oauth.Entrada | None = None
        # Documentos de texto e planilhas, com historico de versoes.
        self.documentos = documento.Documentos(self.base)
        self.comentarios = documento.Comentarios(self.base)
        self.clausulas = documento.Clausulas(CLAUSULAS_PATH)
        # Os codigos oficiais em base local: citar vira consulta.
        self.leis = leis.Leis(self.base)
        # O que a pessoa marcou sobre cada documento da biblioteca.
        self.marcas = acervo.Marcas(self.base)
        # O que esta maquina ja mediu de si: quanto le e quanto escreve por
        # segundo. E daqui que sai o "leituras assim levaram ~70 s aqui".
        self.ritmo = ritmo_mod.Ritmo(RITMO_PATH)
        # Financeiro, bem-estar e conexoes; relatorios so le o que os outros gravaram.
        self.financeiro = financeiro.Financeiro(self.base)
        # A folha, as notas e os boletos: o escritorio por dentro.
        self.folha = escritorio.Folha(self.base)
        self.papeis = escritorio.PapeisFiscais(self.base)
        self.bem_estar = bemestar.BemEstar(self.base)
        # Avisos do Windows: olha os lembretes e o ciclo de foco no servidor,
        # para o aviso chegar com a janela minimizada ou noutra tela.
        self.vigia = avisos.Vigia(self.bem_estar, self.prefs, agenda=self.agenda)
        # Os avisos de evento (resposta, aprovacao, transcricao) leem daqui o
        # que esta ligado em Configuracoes.
        avisos.configurar(self.prefs)
        self.fila.ao_pedir = lambda p: avisos.avisar(
            "aprovacao", "Aprovação pendente", f"{p.titulo} — espera a sua confirmação em Aprovações.")
        # Servicos: as pastas de trabalho (docs/ui, A15).
        self.servicos = servicos_mod.Servicos(self.base, _quem_sou, lambda: self.pasta)
        # Gravacoes de audio, guardadas nesta maquina (docs/ui, A16).
        self.gravacoes = gravacoes_mod.Gravacoes(self.base, GRAVACOES_DIR)
        # Transcricao: o modelo de voz desta maquina e uma fila de fundo, um
        # audio por vez, porque o Whisper ocupa metade dos nucleos.
        voz = self.prefs.dados.get("voz") or {}
        self.transcritor = transcricao_mod.Transcritor(MODELOS_VOZ_DIR, modelo=voz.get("modelo") or transcricao_mod.PADRAO)
        self.fila_voz: "queue.Queue[int]" = queue.Queue()
        self.progresso_voz: dict[int, float] = {}
        self.transcrevendo: int | None = None
        self.erro_voz = ""
        # O download do modelo do assistente pelo cartao "falta baixar".
        self.puxando: dict | None = None
        # Sessoes de transcricao ao vivo (uma por gravacao em andamento).
        self.ao_vivo: dict[str, transcricao_mod.SessaoAoVivo] = {}
        # A camada de inteligencia de documentos: o que ja foi entendido de
        # cada documento, para nao entender de novo a cada pergunta. Com a
        # chave desligada ela devolve fallback na hora e nada muda.
        self.catalogo = Catalogo.carregar()
        self.analisando = False
        self.saber = inteligencia.Saber(
            Biblioteca(CONHECIMENTO_DIR, self.base), self.catalogo,
            ligada=bool(self.prefs.dados.get("inteligencia", True)))
        # O que o escritorio ensinou com as proprias palavras (docs/ui, A13).
        self.contextos = contextos_mod.Contextos(self.base)
        # A foto de quem usa e a logo do escritorio (docs/ui, A13).
        self.marca = marca_mod.Marca(MARCA_DIR)
        # A lixeira: apagar guarda por 30 dias; o que venceu some ao abrir.
        self.lixeira = lixeira_mod.Lixeira(self.base, LIXEIRA_DIR)
        self.lixeira.esvaziar_vencidos()
        for id_ in self.gravacoes.pendentes():
            self.fila_voz.put(id_)
        threading.Thread(target=self._trabalhar_voz, name="voz", daemon=True).start()
        self.conexoes = conexoes.Conexoes(CONEXOES_PATH, SESSOES_DIR)
        self.relatorios = relatorios.Relatorios(
            self.base, fila=self.fila, assinaturas=self.assinaturas,
            envios=self.envios, bem_estar=self.bem_estar,
        )

    def _trabalhar_voz(self) -> None:
        """A fila de transcricao: pega uma gravacao, transcreve, guarda, avisa a trilha."""
        while True:
            id_ = self.fila_voz.get()
            # Enquanto alguem grava com transcricao ao vivo, a fila espera: o
            # texto que chega enquanto a pessoa fala vale mais que o de fundo.
            while any(not s.fechada for s in self.ao_vivo.values()):
                time.sleep(1)
            self.transcrevendo = id_
            self.progresso_voz[id_] = 0.0
            try:
                caminho = self.gravacoes.caminho(id_)
                if not caminho:
                    self.gravacoes.marcar_transcricao(id_, "erro", erro="o áudio não está mais no disco")
                    continue
                self.gravacoes.marcar_transcricao(id_, "transcrevendo")
                r = self.transcritor.transcrever(
                    caminho, progresso=lambda f, i=id_: self.progresso_voz.__setitem__(i, f))
                self.gravacoes.marcar_transcricao(id_, "pronta", trechos=r["trechos"], modelo=r["modelo"], tempo=r["tempo"])
                g = self.gravacoes.obter(id_)
                if g and g.get("servico_id"):
                    self.servicos.trilha(int(g["servico_id"]), "Gravação transcrita: " + g["titulo"], "Assistente")
                if g:
                    avisos.avisar("gravacao", "Transcrição pronta", g.get("titulo") or "Gravação")
            except Exception as exc:  # noqa: BLE001 - a fila nao pode morrer por um audio
                self.gravacoes.marcar_transcricao(id_, "erro", erro=str(exc)[:300])
            finally:
                self.progresso_voz.pop(id_, None)
                self.transcrevendo = None
                self.fila_voz.task_done()

    def pastas_do_acervo(self) -> list[Path]:
        """A pasta do programa e as que o Organizar encheu - o que o Acervo le."""
        extras = [Path(p) for p in self.prefs.dados.get("pastas_acervo") or []]
        return [Path(self.pasta)] + [p for p in extras if p.is_dir()]

    def incluir_no_acervo(self, pasta: str | Path) -> bool:
        """Passa a ler `pasta` no Acervo. Pasta ja lida (ou dentro de uma) nao entra de novo."""
        alvo = Path(pasta).resolve()
        for ja in self.pastas_do_acervo():
            ja = ja.resolve()
            if alvo == ja or ja in alvo.parents:
                return False
        extras = list(self.prefs.dados.get("pastas_acervo") or [])
        self.prefs.atualizar({"pastas_acervo": extras + [str(alvo)]})
        return True

    def recarregar(self, *, force: bool = False) -> int:
        docs = index_all_contracts(self.pastas_do_acervo(), CACHE_PATH, force=force, verbose=False)
        searcher = ContractSearcher()
        searcher.add_contracts(docs)
        searcher.build()
        self.searcher = searcher

        # A janela do modelo acompanha o acervo. Com a janela fixa e pequena, o
        # programa lia um terco dos documentos e respondia "nao encontrei essa
        # informacao" sobre os outros dois tercos - resposta errada com cara de
        # certa, que e o pior tipo.
        if self.client is not None:
            self.client.num_ctx = janela_para(searcher.caracteres())

        # HOOK 1: o que entrou vai ser entendido uma vez, em segundo plano.
        # Nao bloqueia a indexacao e nao altera arquivo nenhum; se falhar, o
        # programa continua exatamente como era - documento sem metadata
        # responde pelo caminho de sempre.
        self.analisar_em_segundo_plano(docs)
        return len(docs)

    def analisar_em_segundo_plano(self, docs) -> None:
        if not self.saber.ligada or self.analisando:
            return

        def trabalhar() -> None:
            self.analisando = True
            try:
                for doc in docs:
                    try:
                        inteligencia.analisar_documento(
                            self.saber.biblioteca, self.catalogo, doc.path,
                            texto=doc.text, paginas=doc.pages, sha1=doc.sha1, titulo=doc.name)
                    except Exception:
                        continue   # um documento torto nao para o acervo
            finally:
                self.analisando = False

        threading.Thread(target=trabalhar, daemon=True).start()


estado = Estado()


@asynccontextmanager
async def lifespan(app: FastAPI):
    estado.pasta.mkdir(parents=True, exist_ok=True)
    total = estado.recarregar()
    estado.vigia.comecar()
    print(f"\n  PAULUS Legal - abra http://localhost:{estado.porta}")
    print(f"  {total} contrato(s) carregado(s) de {estado.pasta}\n")
    yield
    estado.vigia.parar()


app = FastAPI(title="PAULUS Legal", docs_url="/api/docs", lifespan=lifespan)


# ------------------------------------------------------------------ modelos


class Pergunta(BaseModel):
    pergunta: str
    top: int = 6
    # Os documentos que a tela diz estarem em foco - as pilulas do compositor,
    # postas pelo "/", por anexar, ou por ter perguntado sobre eles. Vem da
    # tela porque e o que a pessoa esta VENDO; adivinhar pelo texto quando a
    # tela ja mostra a resposta seria trocar uma certeza por um palpite.
    #
    # Lista porque anexar dois arquivos e perguntar sobre "eles" e o caso
    # comum - quem arrasta um contrato e o aditivo quer os dois lidos juntos.
    apenas: list[str] = []
    # "olha o acervo inteiro", dito pelo botao em vez de pela frase.
    tudo: bool = False
    # A pilula da caixa em "pergunto onde procurar": sem documento nomeado, a
    # conversa pergunta onde antes de sair lendo o acervo inteiro.
    sem_anexo: bool = False
    # O botao Retomar do cartao "Parado": a mesma pergunta de novo, sem
    # repeti-la no historico e trocando a resposta que parou no meio.
    retomar: bool = False


class Busca(BaseModel):
    termo: str
    top: int = 8


# ------------------------------------------------------------------- rotas


SEM_CACHE = {"Cache-Control": "no-cache"}


@app.get("/")
def index() -> FileResponse:
    return FileResponse(FRONTEND_DIR / "index.html", headers=SEM_CACHE)


def _arquivo_do_frontend(pasta: str, arquivo: str, tipo: str) -> FileResponse:
    """
    Um arquivo de frontend/css ou frontend/js. O nome vem da propria pagina;
    ainda assim, nada de subir pastas. "no-cache" faz o navegador embutido
    conferir a data a cada abertura: sem isso, um CSS trocado so aparecia
    depois de limpar o WebView2.
    """
    alvo = (FRONTEND_DIR / pasta / Path(arquivo).name).resolve()
    if alvo.parent != (FRONTEND_DIR / pasta).resolve() or not alvo.exists():
        raise HTTPException(status_code=404, detail="arquivo nao encontrado")
    return FileResponse(alvo, media_type=tipo, headers=SEM_CACHE)


@app.get("/css/{arquivo}")
def css_da_pagina(arquivo: str) -> FileResponse:
    return _arquivo_do_frontend("css", arquivo, "text/css")


@app.get("/js/{arquivo}")
def js_da_pagina(arquivo: str) -> FileResponse:
    return _arquivo_do_frontend("js", arquivo, "text/javascript")


@app.get("/fontes.css")
def fontes_css() -> FileResponse:
    return FileResponse(FRONTEND_DIR / "fontes.css", media_type="text/css")


@app.get("/fontes/{arquivo}")
def fonte(arquivo: str) -> FileResponse:
    # Nome vem do proprio CSS que servimos; ainda assim, nada de subir pastas.
    alvo = (FRONTEND_DIR / "fontes" / Path(arquivo).name).resolve()
    if alvo.parent != (FRONTEND_DIR / "fontes").resolve() or not alvo.exists():
        raise HTTPException(status_code=404, detail="fonte nao encontrada")
    return FileResponse(alvo, media_type="font/woff2")


@app.get("/img/{arquivo}")
def imagem(arquivo: str) -> FileResponse:
    alvo = (FRONTEND_DIR / "img" / Path(arquivo).name).resolve()
    if alvo.parent != (FRONTEND_DIR / "img").resolve() or not alvo.exists():
        raise HTTPException(status_code=404, detail="imagem nao encontrada")
    return FileResponse(alvo)


# ------------------------------------------------------------------- cache


class LimparCache(BaseModel):
    classificacao: bool = False


@app.post("/api/cache/limpar")
def cache_limpar(payload: LimparCache) -> dict:
    """
    Apaga o que foi guardado para nao refazer trabalho.

    Sao dois caches, e eles custam coisas diferentes para refazer. O da
    extracao e o texto tirado de cada PDF: volta sozinho na proxima leitura,
    so demora. O da classificacao e o que o MODELO decidiu sobre cada
    documento - tipo, partes, datas -, e refazer aquilo e rodar o modelo de
    novo em tudo, o que nesta maquina leva minutos por documento. Por isso o
    segundo so vai embora quando a pessoa marca, e a tela diz o preco.
    """
    sumiu, apagados = 0, []
    alvos = [("extração", CACHE_PATH)]
    if payload.classificacao:
        alvos.append(("classificação", CLASSIFICACAO_PATH))
    for rotulo, caminho in alvos:
        if caminho.exists():
            sumiu += caminho.stat().st_size
            try:
                caminho.unlink()
            except OSError as exc:
                raise HTTPException(status_code=500, detail=f"nao consegui apagar o cache de {rotulo}: {exc}") from exc
            apagados.append(rotulo)

    if not apagados:
        return {"apagados": [], "mb": 0, "aviso": "o cache já estava vazio"}

    mb = round(sumiu / 1024 / 1024, 1)
    aviso = "cache de " + " e de ".join(apagados) + " apagado"
    aviso += f" · {mb} MB liberados" if mb >= 0.1 else " · era menos de 0,1 MB"
    return {"apagados": apagados, "mb": mb, "aviso": aviso,
            "refaz": "os documentos são lidos de novo no próximo Reindexar"}


# --------------------------------------------------------------- novidades

NOVIDADES_PATH = BASE_DIR / "docs" / "NOVIDADES.md"


@app.get("/api/novidades")
def novidades() -> dict:
    """
    O que mudou no programa, escrito para quem usa.

    Sai de um arquivo do proprio repositorio, e nao de um servidor: o
    programa nao busca nada na internet, e a lista de novidades de uma versao
    e da versao - ela chega junto com o programa, nao depois dele.
    """
    if not NOVIDADES_PATH.exists():
        return {"blocos": [], "aviso": "ainda não há lista de novidades nesta instalação"}

    blocos: list[dict] = []
    intro: list[str] = []
    for linha in NOVIDADES_PATH.read_text(encoding="utf-8").splitlines():
        crua = linha.rstrip()
        if crua.startswith("## "):
            blocos.append({"titulo": crua[3:].strip(), "itens": []})
        elif crua.startswith("- ") and blocos:
            blocos[-1]["itens"].append(crua[2:].strip())
        elif crua.startswith("  ") and crua.strip() and blocos and blocos[-1]["itens"]:
            # Continuacao da linha anterior: o arquivo quebra em 80 colunas.
            blocos[-1]["itens"][-1] += " " + crua.strip()
        elif crua and not crua.startswith("#") and not blocos:
            intro.append(crua.strip())
    return {"blocos": blocos, "intro": " ".join(intro)}


# --------------------------------------------------------------- contextos


class NovoContexto(BaseModel):
    id: int | None = None
    titulo: str = ""
    texto: str = ""
    gaveta: str = ""


@app.get("/api/contextos")
def contextos_listar() -> dict:
    return estado.contextos.para_tela()


@app.post("/api/contextos")
def contextos_salvar(payload: NovoContexto) -> dict:
    try:
        id_ = estado.contextos.salvar(payload.model_dump(), payload.id)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    return {"id": id_, **estado.contextos.para_tela()}


# Arquivo de onde se aprende: ate aqui de tamanho, e so o que da para ler.
MAX_ENSINAR_BYTES = 20 * 1024 * 1024
FORMATOS_ENSINAR = (".pdf", ".docx", ".txt", ".md")


@app.post("/api/contextos/ler")
async def contextos_ler_arquivo(arquivo: UploadFile = File(...)) -> dict:
    """
    Le um arquivo e propoe o lembrete - sem guardar nada.

    A tela promete "eu leio, resumo em contextos curtos e MOSTRO antes de
    guardar", e e o que acontece: o que volta daqui vai para os campos do
    formulario, onde a pessoa corrige e confirma. O arquivo nao entra no
    Acervo nem fica no disco; o que sobra dele, se a pessoa guardar, sao as
    quatro linhas que ela leu antes.
    """
    nome = Path(arquivo.filename or "").name
    if not nome:
        raise HTTPException(status_code=400, detail="arquivo sem nome")
    if Path(nome).suffix.lower() not in FORMATOS_ENSINAR:
        raise HTTPException(status_code=400,
                            detail="dá para ler PDF, DOCX, TXT e MD; imagem ainda não")
    dados = await arquivo.read(MAX_ENSINAR_BYTES + 1)
    if len(dados) > MAX_ENSINAR_BYTES:
        raise HTTPException(status_code=413, detail="esse arquivo passa de 20 MB")

    def ler() -> tuple[str, int]:
        # Em pasta temporaria: os leitores trabalham sobre caminho, e o
        # arquivo nao tem por que ficar no computador depois disto.
        with tempfile.TemporaryDirectory() as tmp:
            alvo = Path(tmp) / nome
            alvo.write_bytes(dados)
            return extract_file(alvo)

    try:
        texto, paginas = await run_in_threadpool(ler)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    except Exception as exc:
        raise HTTPException(status_code=400, detail=f"não consegui ler esse arquivo: {exc}") from exc

    texto = " ".join((texto or "").split())
    if len(texto) < 40:
        raise HTTPException(
            status_code=422,
            detail="esse arquivo não tem texto para ler — se for um PDF digitalizado, passe o OCR antes")

    disponivel, motivo = check_ollama(estado.client.model)
    if not disponivel:
        # Sem modelo, ainda da para ensinar: o comeco do documento vai para o
        # campo e a pessoa escreve a regra com as palavras dela. Dizer isso e
        # melhor que devolver erro e deixar o arquivo sem serventia.
        return {"titulo": contextos_mod.titulo_do_arquivo(nome), "texto": texto[:contextos_mod.MAX_TEXTO],
                "de": nome, "paginas": paginas, "pelo_modelo": False,
                "aviso": f"{motivo} — trouxe o começo do arquivo para você resumir com as suas palavras"}

    try:
        resposta = await run_in_threadpool(
            lambda: estado.client.ask(contextos_mod.INSTRUCAO_DO_ARQUIVO,
                                      f"Documento “{nome}”:\n{texto[:contextos_mod.LEITURA_MAX]}",
                                      sistema=contextos_mod.INSTRUCAO_DO_ARQUIVO))
    except OllamaError as exc:
        raise HTTPException(status_code=503, detail=str(exc)) from exc

    sugestao = " ".join(_limpar_sugestao(resposta).split())[:contextos_mod.MAX_TEXTO]
    if not sugestao:
        sugestao = texto[:contextos_mod.MAX_TEXTO]
    return {"titulo": contextos_mod.titulo_do_arquivo(nome), "texto": sugestao,
            "de": nome, "paginas": paginas, "pelo_modelo": True,
            "aviso": "li o arquivo e escrevi isto — confira antes de guardar"}


@app.delete("/api/contextos/{id_}")
def contextos_apagar(id_: int) -> dict:
    item = estado.contextos.obter(id_)
    if not item:
        raise HTTPException(status_code=404, detail="esse lembrete nao existe")
    entrada = estado.lixeira.apagar_linha("contexto", id_, item["titulo"], item["gaveta"])
    return {**_foi_para_lixeira(entrada, "contexto", id_), **estado.contextos.para_tela()}


# ------------------------------------------------- o que ja foi lido


@app.get("/api/inteligencia")
def inteligencia_situacao() -> dict:
    """
    Quanto do acervo ja foi entendido, e quanto isso esta economizando.

    Sem medicao o retrofit e fe: a tela mostra quantos documentos tem
    metadata, o que cada secao rendeu e a proporcao de perguntas respondidas
    sem abrir documento. Tudo medido nesta maquina, nada estimado.
    """
    situacao = estado.saber.biblioteca.situacao()
    total = len(estado.searcher.documents)
    return {
        "ligada": estado.saber.ligada,
        "analisando": estado.analisando,
        "documentos": situacao["documentos"],
        "no_acervo": total,
        "secoes": situacao["secoes"],
        "medicao": estado.saber.medicao(),
        "catalogo": {
            "arquivo": str(estado.catalogo.caminho.name),
            "problemas": estado.catalogo.problemas(),
            "extratores": [
                {"secao": e.secao, "id": e.id, "nivel": e.nivel, "modelo": e.modelo}
                for e in (estado.catalogo.para_secao(s)
                          for s in estado.catalogo.secoes_declaradas())
                if e is not None
            ],
        },
        "pasta": str(CONHECIMENTO_DIR),
    }


# ------------------------------------------------------------------- marca


@app.get("/marca/{tipo}.png")
def marca_imagem(tipo: str) -> FileResponse:
    """
    A foto ou a logo, para a tela desenhar.

    Sem cache: a imagem troca no lugar, com o mesmo endereco. O `?v=` que a
    tela manda ja resolveria, mas quem abre o endereco direto tambem tem de
    ver a atual.
    """
    alvo = estado.marca.caminho(tipo)
    if not alvo:
        raise HTTPException(status_code=404, detail="essa imagem nao foi enviada")
    return FileResponse(alvo, media_type="image/png", headers=SEM_CACHE)


@app.post("/api/marca/{tipo}")
async def marca_enviar(tipo: str, arquivo: UploadFile = File(...)) -> dict:
    """A imagem escolhida vira um PNG pequeno em data/marca."""
    dados = await arquivo.read(marca_mod.MAX_BYTES + 1)
    try:
        item = await run_in_threadpool(estado.marca.guardar, tipo, dados)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    return {"tipo": tipo, **item, "marca": estado.marca.info()}


@app.delete("/api/marca/{tipo}")
def marca_tirar(tipo: str) -> dict:
    if tipo not in marca_mod.TIPOS:
        raise HTTPException(status_code=404, detail="imagem desconhecida")
    tirou = estado.marca.tirar(tipo)
    return {"tirou": tirou, "marca": estado.marca.info()}


@app.get("/api/status")
def status() -> dict:
    ok, mensagem = check_ollama(estado.client.model)
    return {
        "ollama": ok,
        "mensagem": mensagem,
        "motor": _motor(mensagem),
        "modelo": estado.client.model,
        # A interface nunca mostra o nome do modelo (regra de linguagem do
        # manual): mostra "Assistente local - X GB na sua maquina".
        "tamanho_gb": _tamanho_do_modelo(),
        "pasta": str(estado.pasta),
        "contratos": len(estado.searcher.documents),
        "trechos": len(estado.searcher.chunks),
    }


# Tamanho aproximado dos modelos que o programa sugere, para o cartao de
# "falta baixar" dizer quanto vem antes de a pessoa clicar.
TAMANHOS_MODELO = {"llama3.2:3b": "2,0 GB", "llama3.2:1b": "1,3 GB", "llama3.1:8b": "4,9 GB", "qwen2.5:3b": "1,9 GB", "gemma2:2b": "1,6 GB"}


def _ollama_exe() -> str | None:
    """O executavel do Ollama: no PATH ou onde o instalador do Windows o poe."""
    achado = shutil.which("ollama")
    if achado:
        return achado
    local = Path(os.environ.get("LOCALAPPDATA", "")) / "Programs" / "Ollama" / "ollama.exe"
    return str(local) if local.exists() else None


def _motor(mensagem: str = "") -> dict:
    """
    Em que pe esta o motor, para o cartao de erro (docs/ui/05) dizer a coisa
    certa: nao instalado, desligado, ou ligado sem o modelo.
    """
    try:
        modelos = estado.client.list_models()
        rodando = True
    except Exception:  # noqa: BLE001 - sem servidor e um estado, nao um erro
        modelos, rodando = [], False
    nome = estado.client.model
    presente = nome in modelos or any(m.split(":")[0] == nome.split(":")[0] and ":" not in nome for m in modelos)
    return {
        "instalado": bool(_ollama_exe()) or rodando,
        "rodando": rodando,
        "modelo_presente": presente,
        "modelo": nome,
        "tamanho": TAMANHOS_MODELO.get(nome, ""),
        "mensagem": mensagem,
        "puxando": estado.puxando,
    }


@app.post("/api/ollama/ligar")
def ollama_ligar() -> dict:
    """
    "Ligar agora" do cartao de erro: abre o servidor do Ollama daqui, sem
    terminal, e espera ate 20 s por ele. Se ja estiver de pe, so confere.
    """
    exe = _ollama_exe()
    if not exe:
        raise HTTPException(status_code=503, detail="O Ollama não está instalado nesta máquina")
    ok, _ = check_ollama(estado.client.model)
    if not ok:
        try:
            estado.client.list_models()
        except Exception:  # noqa: BLE001 - e isso que estamos tratando
            try:
                flags = getattr(subprocess, "CREATE_NO_WINDOW", 0) | getattr(subprocess, "DETACHED_PROCESS", 0)
                subprocess.Popen([exe, "serve"], stdin=subprocess.DEVNULL, stdout=subprocess.DEVNULL,
                                 stderr=subprocess.DEVNULL, creationflags=flags)
            except OSError as exc:
                raise HTTPException(status_code=503, detail="não consegui abrir o Ollama: " + str(exc)) from exc
            limite = time.time() + 20
            while time.time() < limite:
                try:
                    estado.client.list_models()
                    break
                except Exception:  # noqa: BLE001 - ainda subindo
                    time.sleep(0.5)
    ok, mensagem = check_ollama(estado.client.model)
    return {"ligado": ok, "mensagem": mensagem, "motor": _motor(mensagem)}


@app.post("/api/ollama/puxar")
def ollama_puxar() -> dict:
    """Baixa o modelo do assistente em segundo plano (`ollama pull`), com o andamento em /api/ollama/puxar."""
    exe = _ollama_exe()
    if not exe:
        raise HTTPException(status_code=503, detail="O Ollama não está instalado nesta máquina")
    if estado.puxando and estado.puxando.get("andando"):
        return estado.puxando
    estado.puxando = {"modelo": estado.client.model, "progresso": 0, "andando": True, "erro": "", "pronto": False, "linha": ""}
    andamento = estado.puxando

    def puxar() -> None:
        try:
            flags = getattr(subprocess, "CREATE_NO_WINDOW", 0)
            p = subprocess.Popen([exe, "pull", andamento["modelo"]], stdin=subprocess.DEVNULL, stdout=subprocess.PIPE,
                                 stderr=subprocess.STDOUT, creationflags=flags, text=True, encoding="utf-8", errors="replace")
            pedaco = ""
            while True:
                ch = p.stdout.read(1)
                if not ch:
                    break
                if ch in "\r\n":
                    linha = pedaco.strip()
                    pedaco = ""
                    if linha:
                        andamento["linha"] = linha[:120]
                        achado = re.search(r"(\d{1,3})%", linha)
                        if achado:
                            andamento["progresso"] = int(achado.group(1))
                else:
                    pedaco += ch
            if p.wait() == 0:
                andamento["progresso"] = 100
                andamento["pronto"] = True
            else:
                andamento["erro"] = "o download parou: " + (andamento["linha"] or "sem detalhe")
        except Exception as exc:  # noqa: BLE001 - o erro vai para a tela
            andamento["erro"] = str(exc)[:200]
        finally:
            andamento["andando"] = False

    threading.Thread(target=puxar, name="ollama-pull", daemon=True).start()
    return andamento


@app.get("/api/ollama/puxar")
def ollama_puxando() -> dict:
    return estado.puxando or {"andando": False, "progresso": 0, "pronto": False, "erro": ""}


def _tamanho_do_modelo() -> float | None:
    try:
        resp = requests.get(f"{estado.client.host}/api/tags", timeout=5)
        resp.raise_for_status()
        for modelo in resp.json().get("models", []):
            if modelo.get("name") == estado.client.model:
                return round(modelo.get("size", 0) / 1024**3, 1)
    except (requests.RequestException, ValueError):
        return None
    return None


def _sse(tipo: str, dados: dict) -> str:
    """Um evento no formato que o navegador consome (Server-Sent Events)."""
    return f"event: {tipo}\ndata: {json.dumps(dados, ensure_ascii=False)}\n\n"


def _disponibilidade() -> dict:
    ok, _ = check_ollama(estado.client.model)
    return {
        PRECISA_DOCUMENTOS: len(estado.searcher.documents) > 0,
        PRECISA_ASSISTENTE: ok,
    }


def _contexto(registrar=None, parar=None) -> Contexto:
    """O que as habilidades enxergam da aplicacao."""
    return Contexto(
        parar=parar,
        searcher=estado.searcher,
        client=estado.client,
        pasta=estado.pasta,
        cache_classificacao=CLASSIFICACAO_PATH,
        diarios=DIARIOS_DIR,
        recarregar=estado.recarregar,
        registrar=registrar,
        cancelado=estado.cancelar.is_set,
        antes_de_cada=lambda: recursos.esperar_maquina_livre(estado.devagar, limite_s=10),
        ritmo=estado.ritmo,
        # Os lembretes de Configuracoes > Aprendizado entram em toda resposta.
        ensinado=estado.contextos.bloco(),
        # O que ja foi lido uma vez, para nao ler de novo (HOOK 2).
        saber=estado.saber,
    )


@app.get("/api/destinos")
def listar_destinos() -> dict:
    """Menu lateral na ordem oficial do manual do sistema."""
    return {"grupos": destinos.por_grupo(), "contagem": destinos.contagem()}


@app.get("/api/habilidades")
def listar_habilidades() -> dict:
    """
    Catalogo do que o programa sabe fazer, com o que falta para cada coisa.

    A disponibilidade e checada agora: nao adianta oferecer "perguntar sobre os
    documentos" quando nao ha documento aberto.
    """
    return {
        "grupos": estado.registro.por_grupo(_disponibilidade()),
        "contagem": estado.registro.contagem(),
        "pasta": estado.registro.pasta,
    }


@app.post("/api/habilidades/recarregar")
def recarregar_habilidades() -> dict:
    """Rele a pasta de modulos sem fechar o programa."""
    estado.registro = registro.carregar(HABILIDADES_DIR)
    return {"contagem": estado.registro.contagem()}


@app.post("/api/habilidades/{id_}")
def executar_habilidade(id_: str, parametros: dict | None = None):
    """
    Porta unica para rodar qualquer habilidade.

    Se o modulo devolve um dict, sai como JSON. Se e um gerador, sai como
    stream de eventos - e a mesma porta serve para a busca instantanea e para
    uma leitura de uma hora.
    """
    habilidade = estado.registro.obter(id_)
    if not habilidade:
        raise HTTPException(status_code=404, detail="habilidade desconhecida")
    if not habilidade.executavel:
        raise HTTPException(
            status_code=400,
            detail=habilidade.problema or "essa habilidade ainda não faz nada",
        )

    faltando = [p for p in habilidade.precisa if not _disponibilidade().get(p, False)]
    if faltando:
        from habilidade_base import ROTULOS_PRECISA

        nomes = ", ".join(ROTULOS_PRECISA.get(p, p) for p in faltando)
        raise HTTPException(status_code=400, detail=f"precisa de: {nomes}")

    try:
        saida = habilidade.executar(_contexto(), **(parametros or {}))
    except TypeError as exc:
        raise HTTPException(status_code=400, detail=f"parâmetros inválidos: {exc}") from exc

    if not hasattr(saida, "__next__"):
        return saida

    def gerar() -> Iterator[str]:
        try:
            for tipo, dados in saida:
                yield _sse(tipo, dados)
        except Exception as exc:
            yield _sse("erro", {"mensagem": str(exc)})

    return StreamingResponse(
        gerar(),
        media_type="text/event-stream",
        headers={"Cache-Control": "no-cache", "X-Accel-Buffering": "no"},
    )


# ---------------------------------------------------------------- biblioteca


class Remocao(BaseModel):
    sha1: str


class AbrirPasta(BaseModel):
    caminho: str


def _itens_da_biblioteca() -> list[dict]:
    """
    Cada documento com tudo o que o programa sabe dele.

    Junta tres coisas que moram separadas: o indice de leitura (paginas,
    trechos), o cache de classificacao (tipo, cliente) e as marcas da pessoa
    (fixado, quando foi analisado). Sem juntar, a tela mostraria nome e
    tamanho - que e o que o Explorer ja faz.
    """
    from classify import CacheClassificacao

    cache = CacheClassificacao(CLASSIFICACAO_PATH).dados
    marcas = estado.marcas.de_todos()
    trechos_por_nome: dict[str, int] = {}
    for c in estado.searcher.chunks:
        trechos_por_nome[c.doc_name] = trechos_por_nome.get(c.doc_name, 0) + 1

    # A pasta padrao tem nome de programador no disco (test_contracts); na tela
    # ela e o acervo do escritorio. So o nome mostrado muda: o caminho segue o
    # mesmo, porque e por ele que o indice acha cada documento. Subpasta sai
    # contada da raiz que o Acervo le - "Acervo PAULUS › Cliente › Tipo" -,
    # que e o que a tela agrupa; so o ultimo nome repetia "Locacao" para
    # cada cliente.
    padrao = Path(estado.pasta).resolve()
    raizes = [(padrao, NOME_DA_PASTA_PADRAO)] + [
        (p.resolve(), p.name or str(p)) for p in estado.pastas_do_acervo()[1:]
    ]
    nomes: dict[Path, str] = {}

    def nome_da_pasta(pasta: Path) -> str:
        if pasta not in nomes:
            real = pasta.resolve()
            nome = pasta.name or str(pasta)
            for raiz, rotulo in raizes:
                if real == raiz or raiz in real.parents:
                    nome = " › ".join([rotulo, *real.relative_to(raiz).parts])
                    break
            nomes[pasta] = nome
        return nomes[pasta]

    itens: list[dict] = []
    for doc in estado.searcher.documents:
        caminho = Path(doc.path)
        try:
            info = caminho.stat()
            tamanho = info.st_size
            modificado = datetime.fromtimestamp(info.st_mtime).isoformat(timespec="seconds")
            existe = True
        except OSError:
            tamanho, modificado, existe = 0, "", False

        conhecido = cache.get(doc.sha1) or {}
        marca = marcas.get(doc.sha1) or {}

        item = {
            "sha1": doc.sha1,
            "nome": doc.name,
            "caminho": str(caminho),
            "pasta": str(caminho.parent),
            "pasta_curta": nome_da_pasta(caminho.parent),
            "existe": existe,
            "bytes": tamanho,
            "paginas": doc.pages,
            "caracteres": doc.chars,
            "trechos": trechos_por_nome.get(doc.name, 0),
            "modificado_em": modificado,
            "modificado": acervo.quando(modificado),
            "aberto_em": modificado[:10],
            "fixado": bool(marca.get("fixado")),
            "visto_em": marca.get("visto_em", ""),
            "tipo": conhecido.get("tipo", ""),
            "tipo_rotulo": ROTULOS.get(conhecido.get("tipo", ""), ""),
            "cliente": conhecido.get("cliente", ""),
            "data": conhecido.get("data", ""),
            "valor": conhecido.get("valor", ""),
        }
        item["assinado"] = _tem_assinatura_digital(caminho, tamanho, modificado) if existe else False
        item["analise"] = acervo.estado_da_analise(item)
        itens.append(item)

    return itens


_ASSINATURA_NO_ARQUIVO: dict[tuple, bool] = {}
LIMITE_PARA_PROCURAR_ASSINATURA = 40 * 1024 * 1024


def _tem_assinatura_digital(caminho: Path, tamanho: int, modificado: str) -> bool:
    """
    Se o PDF traz uma assinatura digital - so a presenca, nao a validade.

    Toda assinatura de PDF grava um /ByteRange (o trecho do arquivo que ela
    cobre); sem ele nao ha assinatura. Procurar a palavra no arquivo e barato,
    e o resultado fica guardado ate o arquivo mudar. Se ela vale, quem diz e
    /api/assinaturas/conferir, que le a assinatura de verdade.
    """
    if caminho.suffix.lower() != ".pdf" or not tamanho or tamanho > LIMITE_PARA_PROCURAR_ASSINATURA:
        return False
    chave = (str(caminho), tamanho, modificado)
    if chave not in _ASSINATURA_NO_ARQUIVO:
        try:
            _ASSINATURA_NO_ARQUIVO[chave] = b"/ByteRange" in caminho.read_bytes()
        except OSError:
            return False
    return _ASSINATURA_NO_ARQUIVO[chave]


@app.get("/api/documentos-abertos")
def documentos_abertos() -> dict:
    """
    So os nomes dos documentos abertos, para o "/" do compositor.

    Existe separada de /api/biblioteca porque aquela junta classificacao,
    marcas e estatistica de arquivo - trabalho demais para uma lista que a
    tela filtra a cada tecla digitada.
    """
    return {
        "documentos": [
            {"nome": d.name, "paginas": getattr(d, "pages", 0) or 0}
            for d in estado.searcher.documents
        ]
    }


@app.get("/api/biblioteca")
def biblioteca(termo: str = "", filtro: str = "todos", ordem: str = "modificacao",
               limite: int = 0) -> dict:
    """
    O que o PAULUS tem aberto, ja procurado, filtrado e ordenado.

    O corte acontece aqui e nao na tela porque procurar inclui o conteudo dos
    documentos, e o conteudo nao esta no navegador. A conta de quantos sobraram
    vem junto: "mostrando 8 de 128" e a diferenca entre "achei pouco" e "tem
    pouco".
    """
    todos = _itens_da_biblioteca()
    trechos = acervo.indexar_trechos(estado.searcher.chunks) if termo else {}

    achados = acervo.procurar(todos, termo, trechos)
    achados = acervo.filtrar(achados, filtro)
    achados = acervo.ordenar(achados, ordem)

    mostrados = achados[:limite] if limite else achados
    return {
        "documentos": mostrados,
        "achados": len(achados),
        "total": len(todos),
        "fixados": sum(1 for i in todos if i["fixado"]),
        "sem_analise": sum(1 for i in todos if not i["tipo"]),
        "filtros": acervo.FILTROS,
        "ordens": acervo.ORDENS,
        "pasta": str(estado.pasta),
        "total_bytes": sum(i["bytes"] for i in todos),
        "total_trechos": len(estado.searcher.chunks),
    }


class MarcarDocumento(BaseModel):
    sha1: str
    fixado: bool = True


class OndeCitou(BaseModel):
    nome: str = ""
    trecho: str = ""
    pergunta: str = ""


@app.post("/api/biblioteca/citacao")
def biblioteca_citacao(payload: OndeCitou) -> dict:
    """
    Em que pagina do documento este trecho esta, e onde marcar.

    A conversa ja dizia de onde tirou cada informacao - mas dizia em texto, e
    conferir exigia abrir o PDF por fora e procurar o paragrafo com o olho.
    Aqui a pagina volta com as coordenadas da marca, e a tela mostra o
    documento de verdade com o trecho destacado.
    """
    doc = next((d for d in estado.searcher.documents if d.name == payload.nome), None)
    if not doc:
        raise HTTPException(status_code=404, detail="esse documento nao esta aberto")

    alvo = Path(doc.path)
    lugar = citacao.onde_esta(alvo, payload.trecho)
    tamanho = alvo.stat().st_size if alvo.exists() else 0
    e_pdf = alvo.suffix.lower() == ".pdf"

    # Word e texto tambem tem pagina: o mesmo gerador do editor monta o PDF a
    # partir do texto lido - e o que o leitor da conversa ja mostra. Sem
    # trecho para achar, o total ainda precisa vir, senao o visor nao anda.
    if alvo.exists() and not lugar.get("total"):
        try:
            pdf = alvo.read_bytes() if e_pdf else ferramentas.pdf_do_documento(doc)
            lugar = {**lugar, "total": documento.paginas_de(pdf)}
        except Exception:  # noqa: BLE001 - sem contar, o visor mostra a primeira
            pass

    return {
        **lugar,
        "nome": doc.name,
        "caminho": str(alvo),
        "bytes": tamanho,
        "desenhavel": alvo.exists(),
        # Nao e a diagramacao do Word: e o texto do arquivo na folha do
        # PAULUS. A tela diz isso, em vez de apresentar como o original.
        "convertido": not e_pdf,
        "porque": citacao.por_que_este_trecho(payload.pergunta, payload.trecho),
    }


@app.get("/api/biblioteca/pagina")
def biblioteca_pagina(nome: str, numero: int = 1, largura: int = 1000):
    """A pagina do documento desenhada, para conferir sem sair daqui."""
    from fastapi.responses import Response

    doc = next((d for d in estado.searcher.documents if d.name == nome), None)
    if not doc:
        raise HTTPException(status_code=404, detail="esse documento nao esta aberto")
    if not Path(doc.path).exists():
        raise HTTPException(status_code=404, detail="o arquivo saiu do lugar")

    # PDF e desenhado do arquivo; Word e o resto, do PDF que o gerador do
    # editor monta a partir do texto lido - o visor da conversa mostra paginas
    # para qualquer formato.
    try:
        png = ferramentas.pagina_png(doc, numero, largura)
    except Exception as exc:
        raise HTTPException(status_code=500, detail=f"nao consegui desenhar: {exc}") from exc
    return Response(png, media_type="image/png",
                    headers={"Cache-Control": "max-age=120"})


@app.post("/api/biblioteca/fixar")
def biblioteca_fixar(payload: MarcarDocumento) -> dict:
    if not any(d.sha1 == payload.sha1 for d in estado.searcher.documents):
        raise HTTPException(status_code=404, detail="documento nao esta na biblioteca")
    estado.marcas.fixar(payload.sha1, payload.fixado)
    return {"sha1": payload.sha1, "fixado": payload.fixado}


class LoteDocumentos(BaseModel):
    # Caminhos, e nao SHA-1. Dois arquivos iguais byte a byte - "contrato.pdf"
    # e "contrato (1).pdf" - tem o mesmo SHA-1, e escolher por hash faria o
    # lote agir sobre a copia que a pessoa nao marcou. Apagar um e levar dois
    # e exatamente o estrago que a fila existe para impedir.
    caminhos: list[str] = []
    destino: str = ""


def _selecionados(caminhos: list[str]) -> list[dict]:
    escolhidos = set(caminhos)
    itens = [i for i in _itens_da_biblioteca() if i["caminho"] in escolhidos]
    if not itens:
        raise HTTPException(status_code=400, detail="nenhum documento selecionado")
    return itens


@app.post("/api/biblioteca/analisar")
def biblioteca_analisar(payload: LoteDocumentos) -> StreamingResponse:
    """
    Tomar vista de novo: le o documento outra vez e refaz a classificacao.

    Existe porque o arquivo muda depois de analisado - a tela marca esses como
    "mudou desde entao" justamente para que este botao tenha para que servir.

    Vai por evento e nao por resposta unica porque ler com o modelo leva perto
    de um minuto por documento nesta maquina. Oito documentos sao oito minutos
    de tela parada, e o andamento que aparece e o numero de documentos prontos
    - contado, nao estimado.
    """
    from classify import CacheClassificacao

    itens = _selecionados(payload.caminhos)

    habilidade = estado.registro.obter("classificar")
    if not habilidade or not habilidade.executavel:
        raise HTTPException(status_code=503, detail="a habilidade de classificar nao carregou")

    # Sem isto a releitura devolveria o que esta no cache, que e exatamente o
    # resultado que a pessoa pediu para refazer.
    cache = CacheClassificacao(CLASSIFICACAO_PATH)
    for item in itens:
        cache.dados.pop(item["sha1"], None)
    cache.salvar()

    estado.cancelar.clear()
    caminhos = [i["caminho"] for i in itens]

    def gerar() -> Iterator[str]:
        try:
            for tipo, dados in habilidade.executar(_contexto(), caminhos=caminhos):
                if tipo == "resultados":
                    for doc in dados.get("documentos", []):
                        if doc.get("sha1"):
                            estado.marcas.marcar_visto(doc["sha1"])
                    estado.recarregar(force=True)
                yield _sse(tipo, dados)
        except Exception as exc:
            yield _sse("erro", {"mensagem": str(exc)})

    return StreamingResponse(
        gerar(),
        media_type="text/event-stream",
        headers={"Cache-Control": "no-cache", "X-Accel-Buffering": "no"},
    )


@app.post("/api/biblioteca/lote/mover")
def biblioteca_lote_mover(payload: LoteDocumentos) -> dict:
    """
    Mover muitos de uma vez - pela fila, com o plano a vista.

    Reaproveita o plano do organizador de propósito: e o mesmo executor, o
    mesmo diario e o mesmo desfazer. Mover em lote e a operacao que mais
    estraga quando erra, e nao era hora de inventar um caminho novo para ela.
    """
    if not payload.destino:
        raise HTTPException(status_code=400, detail="escolha a pasta de destino")
    destino = Path(payload.destino)
    if not destino.is_dir():
        raise HTTPException(status_code=400, detail="essa pasta nao existe")

    itens = _selecionados(payload.caminhos)
    plano = acervo.plano_de_mover(itens, destino)
    if not plano["movimentos"]:
        return {"pedido": None, "impedidos": plano["impedidos"],
                "aviso": "nenhum dos escolhidos pode ser movido"}

    pedido = estado.fila.pedir(
        f"Mover {len(plano['movimentos'])} documento(s) para {destino.name}",
        "arquivo",
        acao="organizar.mover",
        resumo=acervo.resumo_do_lote("mover", itens, str(destino)),
        etiquetas=["dá para desfazer"],
        dados={"movimentos": plano["movimentos"], "destino": str(destino),
               "padrao": "escolha na biblioteca", "ignorados": []},
    )
    return {"pedido": pedido.to_dict(), "impedidos": plano["impedidos"],
            **estado.fila.para_tela()}


@app.post("/api/biblioteca/lote/apagar")
def biblioteca_lote_apagar(payload: LoteDocumentos) -> dict:
    """Tirar muitos da biblioteca de uma vez - pela fila, porque nao desfaz."""
    itens = _selecionados(payload.caminhos)
    pasta = Path(estado.pasta).resolve()

    dentro, fora = [], []
    for item in itens:
        if pasta in Path(item["caminho"]).resolve().parents:
            dentro.append(item)
        else:
            fora.append({"nome": item["nome"],
                         "motivo": "está fora da pasta do programa"})

    if not dentro:
        return {"pedido": None, "impedidos": fora,
                "aviso": "nenhum dos escolhidos está na pasta do programa"}

    pedido = estado.fila.pedir(
        f"Tirar {len(dentro)} documento(s) da biblioteca",
        "arquivo",
        acao="acervo.apagar",
        resumo=acervo.resumo_do_lote("apagar", dentro),
        etiquetas=["não dá para desfazer"],
        dados={"caminhos": [i["caminho"] for i in dentro],
               "nomes": [i["nome"] for i in dentro]},
    )
    return {"pedido": pedido.to_dict(), "impedidos": fora, **estado.fila.para_tela()}


@app.post("/api/biblioteca/lote/exportar")
def biblioteca_lote_exportar(payload: LoteDocumentos) -> dict:
    """Copiar muitos para uma pasta de fora - pela fila, porque sai daqui."""
    if not payload.destino:
        raise HTTPException(status_code=400, detail="escolha a pasta de destino")
    destino = Path(payload.destino)
    if not destino.is_dir():
        raise HTTPException(status_code=400, detail="essa pasta nao existe")

    itens = _selecionados(payload.caminhos)
    pedido = estado.fila.pedir(
        f"Copiar {len(itens)} documento(s) para {destino.name}",
        "arquivo",
        acao="acervo.exportar",
        resumo=acervo.resumo_do_lote("exportar", itens, str(destino)),
        etiquetas=["dá para desfazer"],
        dados={"caminhos": [i["caminho"] for i in itens], "destino": str(destino)},
    )
    return {"pedido": pedido.to_dict(), **estado.fila.para_tela()}


@app.post("/api/biblioteca/remover")
def biblioteca_remover(payload: Remocao) -> dict:
    """
    Tira um documento da biblioteca, apagando a copia que o programa guarda.

    So mexe em arquivo dentro da pasta do programa: o original de onde o
    documento veio nao e tocado.
    """
    alvo = next((d for d in estado.searcher.documents if d.sha1 == payload.sha1), None)
    if not alvo:
        raise HTTPException(status_code=404, detail="documento nao esta na biblioteca")

    caminho = Path(alvo.path).resolve()
    pasta = Path(estado.pasta).resolve()
    if pasta not in caminho.parents:
        raise HTTPException(
            status_code=400,
            detail="esse arquivo está fora da pasta do programa; remova por lá",
        )

    try:
        caminho.unlink(missing_ok=True)
    except OSError as exc:
        raise HTTPException(status_code=500, detail=f"não consegui remover: {exc}") from exc

    return {"removido": alvo.name, "documentos": estado.recarregar(force=True)}


@app.post("/api/biblioteca/abrir-pasta")
def biblioteca_abrir_pasta(payload: AbrirPasta) -> dict:
    """Abre a pasta no Explorer, para a pessoa ver o arquivo onde ele esta."""
    import os

    alvo = Path(payload.caminho)
    if not alvo.is_dir():
        raise HTTPException(status_code=400, detail="pasta nao encontrada")
    try:
        os.startfile(str(alvo))  # noqa: S606 - abre o gerenciador do proprio Windows
    except (OSError, AttributeError) as exc:
        raise HTTPException(status_code=500, detail=f"não consegui abrir: {exc}") from exc
    return {"aberta": str(alvo)}


@app.get("/api/documents")
def documentos() -> dict:
    return {
        "documentos": [
            {"nome": d.name, "chars": d.chars, "paginas": d.pages}
            for d in estado.searcher.documents
        ]
    }


@app.post("/api/reindex")
def reindexar() -> dict:
    try:
        return {"contratos": estado.recarregar(force=True)}
    except OSError as exc:
        # Disco cheio no meio da leitura (docs/ui/05): para, o cache do que
        # ja foi lido fica, e a tela diz o que fazer.
        if exc.errno == errno.ENOSPC:
            raise HTTPException(status_code=507, detail="O disco encheu no meio da leitura. O que já foi lido ficou guardado; libere espaço e reindexe de novo.") from exc
        raise


@app.post("/api/buscar-agora")
def buscar_agora(payload: Busca) -> dict:
    """Atalho para a habilidade `buscar`, mantido para nao quebrar chamadas antigas."""
    habilidade = estado.registro.obter("buscar")
    if not habilidade or not habilidade.executavel:
        raise HTTPException(status_code=503, detail="a habilidade de busca nao carregou")
    return habilidade.executar(_contexto(), termo=payload.termo, top=payload.top)


@app.post("/api/upload")
async def upload(arquivos: list[UploadFile]) -> dict:
    salvos: list[str] = []
    recusados: list[dict] = []

    for arquivo in arquivos:
        nome = Path(arquivo.filename or "").name  # descarta qualquer caminho
        if not nome:
            continue
        if Path(nome).suffix.lower() not in SUPPORTED_SUFFIXES:
            recusados.append({"nome": nome, "motivo": "formato nao suportado"})
            continue

        conteudo = await arquivo.read()
        if len(conteudo) > MAX_UPLOAD_BYTES:
            recusados.append({"nome": nome, "motivo": "arquivo maior que 50 MB"})
            continue

        (estado.pasta / nome).write_bytes(conteudo)
        salvos.append(nome)

    total = estado.recarregar()
    return {"salvos": salvos, "recusados": recusados, "contratos": total}


class AnexarCaminhos(BaseModel):
    caminhos: list[str]


@app.post("/api/anexar/caminhos")
def anexar_caminhos(payload: AnexarCaminhos) -> dict:
    """
    Anexar a partir do computador, pelo caminho do arquivo.

    E o "Meu computador" do pop-up de anexar: a tela navega pelas pastas desta
    maquina e manda os caminhos escolhidos, e aqui o arquivo e copiado para o
    acervo e lido - o mesmo que /api/upload faz com o que chega pelo seletor
    do Windows, sem o arquivo precisar atravessar o navegador. Arquivo que ja
    esta na pasta do acervo nao e copiado sobre ele mesmo.
    """
    salvos: list[str] = []
    recusados: list[dict] = []
    for bruto in payload.caminhos:
        origem = Path(bruto)
        nome = origem.name
        if not origem.is_file():
            recusados.append({"nome": nome or bruto, "motivo": "arquivo nao encontrado"})
            continue
        if origem.suffix.lower() not in SUPPORTED_SUFFIXES:
            recusados.append({"nome": nome, "motivo": "formato nao suportado"})
            continue
        try:
            if origem.stat().st_size > MAX_UPLOAD_BYTES:
                recusados.append({"nome": nome, "motivo": "arquivo maior que 50 MB"})
                continue
            destino = estado.pasta / nome
            if destino.resolve() != origem.resolve():
                shutil.copy2(origem, destino)
        except OSError as exc:
            recusados.append({"nome": nome, "motivo": f"nao consegui copiar: {exc.strerror or exc}"})
            continue
        salvos.append(nome)
    total = estado.recarregar() if salvos else len(estado.searcher.documents)
    return {"salvos": salvos, "recusados": recusados, "contratos": total}


# ------------------------------------------------------ trabalhos (conversas)


class NovoTrabalho(BaseModel):
    pedido: str = ""
    tipo: str = "conversa"


class Ajuste2(BaseModel):
    devagar: bool


@app.get("/api/trabalhos")
def trabalhos_listar() -> dict:
    dados = estado.trabalhos.listar()
    dados["pendencias"] = estado.trabalhos.pendencias
    return dados


@app.post("/api/trabalhos")
def trabalhos_criar(payload: NovoTrabalho) -> dict:
    titulo = titular(payload.pedido) if payload.pedido.strip() else "Nova conversa"
    trabalho = estado.trabalhos.criar(titulo, tipo=payload.tipo)
    return trabalho.to_dict()


@app.get("/api/trabalhos/{id_}")
def trabalhos_obter(id_: str) -> dict:
    trabalho = estado.trabalhos.obter(id_)
    if not trabalho:
        raise HTTPException(status_code=404, detail="trabalho nao encontrado")
    return trabalho.to_dict()


class Renomear(BaseModel):
    titulo: str


class MoverGrupo(BaseModel):
    grupo: str


class RenomearGrupo(BaseModel):
    de: str
    para: str = ""


class ExportarConversa(BaseModel):
    caminho: str


FORMATOS_DE_CONVERSA = {"md": "text/markdown", "txt": "text/plain", "docx":
                        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"}


def _conversa_exportada(trabalho, formato: str) -> bytes:
    """
    A conversa num arquivo, para levar para fora do programa.

    Markdown e o padrao: abre legivel em qualquer editor e vira documento
    formatado onde houver quem o leia. Texto e o mesmo sem marcacao. Word sai
    com o titulo, quem falou em negrito e o texto em paragrafos. Em todos, a
    resposta que a pessoa parou no meio diz isso, e a que citou trechos diz
    de quais documentos.
    """
    from datetime import datetime

    def quando(iso: str) -> str:
        try:
            return datetime.fromisoformat(iso).strftime("%d/%m/%Y %H:%M")
        except (TypeError, ValueError):
            return ""

    falas = []
    for m in trabalho.mensagens:
        quem = "Você" if m.autor == "pessoa" else "PAULUS"
        notas = []
        if getattr(m, "interrompida", False):
            notas.append("resposta parada no meio")
        documentos = sorted({f.get("documento", "") for f in (m.fontes or []) if f.get("documento")})
        if documentos:
            notas.append("trechos de " + ", ".join(documentos))
        falas.append((quem, quando(m.em), (m.texto or "").strip(), notas))
    agora = datetime.now().strftime("%d/%m/%Y %H:%M")

    if formato == "docx":
        import io

        from docx import Document
        from docx.shared import Pt

        doc = Document()
        doc.add_heading(trabalho.titulo, level=1)
        rodape = doc.add_paragraph()
        rodape.add_run(f"Exportado do PAULUS Legal em {agora}").italic = True
        for quem, em, texto, notas in falas:
            cabeca = doc.add_paragraph()
            cabeca.add_run(quem).bold = True
            if em:
                cabeca.add_run(f"  ·  {em}").font.size = Pt(9)
            for paragrafo in (texto.split("\n") if texto else [""]):
                doc.add_paragraph(paragrafo)
            for nota in notas:
                doc.add_paragraph().add_run(nota).italic = True
        saida = io.BytesIO()
        doc.save(saida)
        return saida.getvalue()

    linhas: list[str] = []
    if formato == "md":
        linhas += [f"# {trabalho.titulo}", "", f"*Exportado do PAULUS Legal em {agora}*", ""]
        for quem, em, texto, notas in falas:
            linhas += [f"**{quem}**" + (f" · {em}" if em else ""), "", texto, ""]
            linhas += [f"> _{nota}_" for nota in notas]
            if notas:
                linhas.append("")
    else:
        linhas += [trabalho.titulo, f"Exportado do PAULUS Legal em {agora}", ""]
        for quem, em, texto, notas in falas:
            linhas += [quem + (f" - {em}" if em else ""), texto]
            linhas += [f"({nota})" for nota in notas]
            linhas.append("")
    return "\n".join(linhas).rstrip("\n").encode("utf-8") + b"\n"


@app.post("/api/trabalhos/{id_}/renomear")
def trabalhos_renomear(id_: str, payload: Renomear) -> dict:
    trabalho = estado.trabalhos.obter(id_)
    if not trabalho:
        raise HTTPException(status_code=404, detail="conversa nao encontrada")

    titulo = " ".join(payload.titulo.split())[:LIMITE_DE_NOME]
    if not titulo:
        raise HTTPException(status_code=400, detail="o nome nao pode ficar vazio")

    trabalho.titulo = titulo
    estado.trabalhos.salvar(trabalho)
    return trabalho.resumo()


@app.post("/api/trabalhos/{id_}/grupo")
def trabalhos_grupo(id_: str, payload: MoverGrupo) -> dict:
    trabalho = estado.trabalhos.obter(id_)
    if not trabalho:
        raise HTTPException(status_code=404, detail="conversa nao encontrada")

    trabalho.grupo = " ".join(payload.grupo.split())[:LIMITE_DE_NOME]
    trabalho.atualizado_em = jobs_agora()
    estado.trabalhos.salvar(trabalho)
    return trabalho.resumo()


@app.get("/api/trabalhos/{id_}/exportar")
def trabalhos_exportar_baixar(id_: str, formato: str = "md") -> Response:
    """A conversa como download - o caminho de quem abre pelo navegador."""
    trabalho = estado.trabalhos.obter(id_)
    if not trabalho:
        raise HTTPException(status_code=404, detail="conversa nao encontrada")
    formato = formato if formato in FORMATOS_DE_CONVERSA else "md"
    nome = re.sub(r'[\\/:*?"<>|]+', "-", trabalho.titulo or "conversa").strip() or "conversa"
    from urllib.parse import quote

    return Response(
        content=_conversa_exportada(trabalho, formato),
        media_type=FORMATOS_DE_CONVERSA[formato],
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(nome + '.' + formato)}"},
    )


@app.post("/api/trabalhos/{id_}/exportar")
def trabalhos_exportar_gravar(id_: str, payload: ExportarConversa) -> dict:
    """
    A conversa gravada no caminho que a pessoa escolheu no "Salvar como".

    Na janela do programa, download de navegador nao chega a lugar nenhum: a
    tela pede o caminho ao Windows e o servidor grava. O formato sai da
    extensao escolhida; sem extensao, e Markdown.
    """
    trabalho = estado.trabalhos.obter(id_)
    if not trabalho:
        raise HTTPException(status_code=404, detail="conversa nao encontrada")
    caminho = Path(payload.caminho)
    formato = caminho.suffix.lower().lstrip(".")
    if formato not in FORMATOS_DE_CONVERSA:
        formato = "md"
        caminho = caminho.with_name(caminho.name + ".md")
    try:
        caminho.parent.mkdir(parents=True, exist_ok=True)
        caminho.write_bytes(_conversa_exportada(trabalho, formato))
    except OSError as exc:
        raise HTTPException(status_code=500, detail=f"nao consegui gravar o arquivo: {exc}") from exc
    return {"caminho": str(caminho), "nome": caminho.name, "formato": formato}


@app.post("/api/grupos/renomear")
def grupos_renomear(payload: RenomearGrupo) -> dict:
    """
    Renomeia um grupo de conversas - ou o desfaz, com `para` vazio.

    Desfazer o grupo nao apaga conversa nenhuma: elas vao para "Sem grupo".
    A resposta traz os ids, para a tela oferecer Desfazer (que e renomear de
    volta so aquelas conversas).
    """
    de = " ".join(payload.de.split())
    para = " ".join(payload.para.split())[:LIMITE_DE_NOME]
    if not de:
        raise HTTPException(status_code=400, detail="diga qual grupo")
    ids = estado.trabalhos.renomear_grupo(de, para)
    if not ids:
        raise HTTPException(status_code=404, detail="grupo nao encontrado")
    return {"ids": ids, "de": de, "para": para}


@app.post("/api/trabalhos/{id_}/duplicar")
def trabalhos_duplicar(id_: str) -> dict:
    copia = estado.trabalhos.duplicar(id_)
    if not copia:
        raise HTTPException(status_code=404, detail="conversa nao encontrada")
    return copia.to_dict()


@app.get("/api/agora")
def acontecendo_agora() -> dict:
    """
    O que esta acontecendo neste instante, para os cartoes da tela inicial.

    Tudo com numero real: cartao que diz "processando" sem dizer quanto falta
    e enfeite, nao informacao.
    """
    abertos = [
        estado.trabalhos.obter(t["id"])
        for t in estado.trabalhos.listar()["em_andamento"]
    ]
    abertos = [t for t in abertos if t]

    executando = []
    esperando = []
    for trabalho in abertos:
        if trabalho.aprovacao:
            esperando.append({
                "id": trabalho.id,
                "titulo": trabalho.titulo,
                "pergunta": trabalho.aprovacao.pergunta,
            })
            continue
        if trabalho.estado == EXECUTANDO:
            atual = next((e for e in trabalho.etapas if e.estado == EXECUTANDO), None)
            item = {
                "id": trabalho.id,
                "titulo": trabalho.titulo,
                "etapa": atual.titulo if atual else "",
                "feitos": atual.feitos if atual else 0,
                "total": atual.total if atual else 0,
                "progresso": trabalho.progresso or 0,
            }
            # A resposta andando diz em que fase esta e ha quanto tempo - e a
            # barra so existe onde ha medida: leitura com previsao desta
            # maquina. A conta de etapas (1 de 2) deixava a barra parada em
            # 50% a leitura inteira.
            a = estado.andamento.get(trabalho.id)
            if a:
                agora = time.time()
                item["andamento"] = {
                    "fase": a["fase"],
                    "decorrido_s": round(agora - a["inicio"], 1),
                    "fase_s": round(agora - a["desde"], 1),
                    "previsao_s": a["previsao_s"],
                    "palavras": a["palavras"],
                    "documentos": a["documentos"],
                }
            executando.append(item)

    return {
        "executando": executando,
        "esperando": esperando,
        "biblioteca": {
            "documentos": len(estado.searcher.documents),
            "trechos": len(estado.searcher.chunks),
        },
        "pausados": sum(1 for t in abertos if t.estado == "pausado"),
    }


@app.delete("/api/trabalhos/{id_}")
def trabalhos_remover(id_: str) -> dict:
    trabalho = estado.trabalhos.obter(id_)
    if not trabalho:
        raise HTTPException(status_code=404, detail="trabalho nao encontrado")
    caminho = estado.trabalhos.caminho_de(id_)
    movidos = [estado.lixeira.mover_arquivo(caminho)] if caminho.exists() else []
    entrada = estado.lixeira.guardar("conversa", id_, trabalho.titulo, "Assistente", {"id": id_}, movidos)
    estado.trabalhos.remover(id_)
    return _foi_para_lixeira(entrada, "removido", id_)


def _em_foco(trabalho) -> list[str]:
    """
    Os documentos que esta conversa vinha lendo.

    Aceita o formato antigo - um nome so, em texto - porque conversa gravada
    antes do escopo virar lista continua no disco.
    """
    guardado = trabalho.contexto.get("documento_em_foco") or []
    if isinstance(guardado, str):
        return [guardado] if guardado else []
    return [n for n in guardado if n]


def _escopo_da_pergunta(trabalho, pergunta: str, payload: Pergunta) -> tuple[list, list]:
    """
    Quais documentos ler, e quais a frase referiu de proposito.

    A regra mora em `intencao.escopo` - e da mesma familia das outras regras
    de conversa, e la da para testar sem subir o programa inteiro. Aqui fica
    so a contabilidade: ler o foco do trabalho e grava-lo de volta.
    """
    escolhidos, explicito = intencao.escopo(
        pergunta,
        abertos=estado.searcher.documents,
        em_foco=_em_foco(trabalho),
        pedidos=payload.apenas,
        tudo=payload.tudo,
        herdar_foco=not payload.sem_anexo,
    )
    trabalho.contexto["documento_em_foco"] = escolhidos
    return escolhidos, explicito


@app.post("/api/trabalhos/{id_}/parar")
def trabalhos_parar(id_: str) -> dict:
    """
    O botao de parar da caixa de pedido.

    Com resposta sendo escrita, levanta o sinal dela: a leitura larga o modelo
    em ate um quarto de segundo, guarda o que ja saiu e fecha a conversa como
    parada. Sem resposta nenhuma andando - a conversa ficou "trabalhando" de
    uma sessao que ja nao existe -, a propria rota a devolve como parada.
    """
    trabalho = estado.trabalhos.obter(id_)
    if not trabalho:
        raise HTTPException(status_code=404, detail="conversa nao encontrada")
    sinal = estado.respondendo.get(id_)
    if sinal:
        sinal.set()
        return {"parando": True, "estado": trabalho.estado}
    if trabalho.estado == EXECUTANDO:
        for etapa in trabalho.etapas:
            if etapa.estado == EXECUTANDO:
                etapa.estado = PAUSADO
        trabalho.estado = PAUSADO
        estado.trabalhos.salvar(trabalho)
    return {"parando": False, "estado": trabalho.estado}


@app.post("/api/trabalhos/{id_}/perguntar")
def trabalhos_perguntar(id_: str, payload: Pergunta) -> StreamingResponse:
    """
    Pergunta dentro de um trabalho.

    A logica de responder mora em habilidades/perguntar.py. Aqui fica so a
    contabilidade do trabalho: gravar a conversa, mexer nas etapas e salvar.
    """
    trabalho = estado.trabalhos.obter(id_)
    if not trabalho:
        raise HTTPException(status_code=404, detail="trabalho nao encontrado")
    if not payload.pergunta.strip():
        raise HTTPException(status_code=400, detail="pergunta vazia")

    habilidade = estado.registro.obter("perguntar")
    if not habilidade or not habilidade.executavel:
        raise HTTPException(status_code=503, detail="a habilidade de perguntar nao carregou")

    pergunta = payload.pergunta.strip()
    if trabalho.titulo == "Nova conversa" and not trabalho.mensagens:
        trabalho.titulo = titular(pergunta)

    ultima = trabalho.mensagens[-1] if trabalho.mensagens else None
    # Retomar troca a resposta parada - e tambem o cartao "onde eu procuro?",
    # que a pessoa respondeu escolhendo onde.
    if (payload.retomar and ultima and ultima.autor == "paulus"
            and (ultima.interrompida or (ultima.proposta or {}).get("tipo") == "escopo")):
        trabalho.mensagens.pop()
        ultima = trabalho.mensagens[-1] if trabalho.mensagens else None
    if not (payload.retomar and ultima and ultima.autor == "pessoa" and ultima.texto.strip() == pergunta):
        trabalho.dizer("pessoa", pergunta)

    # Antes de sair procurando: o que a pessoa pediu? A conversa tinha um
    # caminho so, e "anote uma reuniao no calendario" virava busca pela
    # palavra "reuniao" dentro dos contratos - resposta certa para a pergunta
    # errada. Ler a intencao e por regra: instantaneo e repetivel.
    # Sobre qual documento e esta conversa. Duas fontes, nesta ordem: o nome
    # que a frase escreve, e - quando ela diz "o referido documento" - o que
    # ficou em foco na conversa. Sem isso, "Exiba o referido documento aqui"
    # nao nomeia arquivo nenhum, cai na busca e e respondido pelo acervo
    # inteiro: foi assim que a pergunta sobre um comprovante de viagem voltou
    # falando de contrato de compra e venda.
    foco_antes = _em_foco(trabalho)
    citado, explicito = _escopo_da_pergunta(trabalho, pergunta, payload)

    lido = intencao.ler(pergunta, documentos=estado.searcher.documents,
                        cadastros=[f["nome"] for f in estado.cadastros.listar()])
    # "Exiba o referido documento" so vira acao de abrir quando se sabe qual e.
    # `explicito`, e nao `citado`: com o foco herdado, "mostre o valor do
    # adiantamento" abriria o arquivo em vez de responder a pergunta.
    # "Quero que voce abra o documento", com o anexo na caixa: o pedido e o
    # arquivo. Antes ia para o modelo, que respondia copiando o texto inteiro
    # na conversa. Agora vem o cartao - mostrar aqui, editar ou abrir no
    # Windows -, a nao ser que a frase peca o conteudo.
    if lido.tipo == "documentos" and intencao.quer_abrir(pergunta):
        alvo = explicito if len(explicito) == 1 else (
            citado if citado and intencao.pede_so_o_documento(pergunta) else [])
        if len(alvo) == 1:
            lido = intencao.Intencao(
                tipo="abrir", titulo=alvo[0], campos={"nome": alvo[0]},
                porque="“" + pergunta.strip().rstrip(".!?") + "”",
            )
        elif alvo:
            lido = intencao.Intencao(
                tipo="exibir", titulo=alvo[0], campos={"nome": alvo[0], "nomes": alvo[:4]},
                porque="você pediu para abrir os documentos anexados",
            )
    if lido.tipo in ("agenda", "tarefa", "sobre", "abrir", "servico", "cadastro", "nota", "exibir"):
        return _responder_sem_documentos(trabalho, lido, pergunta)

    # Tirou o anexo e perguntou sem nomear documento: antes de ler os
    # dezessete, pergunta onde - so no que a conversa vinha lendo, ou no
    # acervo inteiro. Ler tudo leva minutos, e nao foi o que a pessoa disse.
    if (payload.sem_anexo and not citado and not payload.tudo
            and not intencao.quer_todo_o_acervo(pergunta)):
        trabalho.contexto["documento_em_foco"] = foco_antes
        return _perguntar_onde_procurar(trabalho, pergunta, foco_antes)

    trabalho.etapas = [
        Etapa("Procurar nos documentos", estado=EXECUTANDO),
        Etapa("Ler os trechos e responder"),
    ]
    trabalho.estado = EXECUTANDO
    estado.trabalhos.salvar(trabalho)

    def registrar(texto: str) -> None:
        trabalho.registrar(texto)

    parar = threading.Event()
    estado.respondendo[id_] = parar

    def pausar_o_que_executava() -> None:
        for etapa in trabalho.etapas:
            if etapa.estado == EXECUTANDO:
                etapa.estado = PAUSADO
        trabalho.estado = PAUSADO

    def gerar() -> Iterator[str]:
        """
        A resposta e o que sempre garante que a conversa nao fica presa.

        Se quem esta do outro lado some no meio - a janela fechou, a pagina
        recarregou -, o gerador e fechado sem chegar ao fim, e a conversa
        ficava "trabalhando" para sempre: foi o que apareceu com 345 s no
        relogio. O `finally` a devolve como parada.
        """
        try:
            yield from _gerar()
        finally:
            estado.andamento.pop(id_, None)
            if estado.respondendo.get(id_) is parar:
                del estado.respondendo[id_]
            if trabalho.estado == EXECUTANDO:
                pausar_o_que_executava()
                estado.trabalhos.salvar(trabalho)

    def _gerar() -> Iterator[str]:
        import time

        inicio = time.time()
        partes: list[str] = []
        cobertura: dict = {}
        fontes: list[dict] = []
        andamento = {"fase": "procurando", "inicio": inicio, "desde": inicio,
                     "previsao_s": 0, "palavras": 0, "documentos": 0, "caracteres": 0}
        estado.andamento[id_] = andamento

        def fase(nome: str) -> None:
            andamento["fase"] = nome
            andamento["desde"] = time.time()
        medida: dict = {}
        lido_chars = 0
        # O nivel com que a camada respondeu, quando ela respondeu. Nulo quer
        # dizer o caminho de sempre.
        nivel: int | None = None
        inferencia = False

        passos = habilidade.executar(
            _contexto(registrar, parar=parar.is_set), pergunta=pergunta, top=payload.top,
            apenas=citado,
        )
        try:
            for tipo, dados in passos:
                if parar.is_set():
                    break
                if tipo == "fontes":
                    andamento["documentos"] = len(dados.get("consultados") or [])
                    nivel = dados.get("nivel", nivel)
                    inferencia = bool(dados.get("inferencia", inferencia))
                    cobertura = {
                        "consultados": dados["consultados"],
                        "ignorados": dados["ignorados"],
                        "total_contratos": dados["total_contratos"],
                        "apenas": dados.get("apenas") or [],
                    }
                    fontes = dados["trechos"]
                    trabalho.etapas[0].estado = CONCLUIDO
                    trabalho.etapas[0].detalhe = (
                        _quantos(dados["total_contratos"], "documento") + ", "
                        + _quantos(len(fontes), "trecho"))
                    # Sem 25/25 aqui: os dois numeros chegavam juntos, e uma
                    # barra que nasce cheia nao e progresso - e enfeite.
                    yield _sse("fontes", dados)
                    yield _sse("etapas", {"etapas": [asdict_etapa(e) for e in trabalho.etapas]})
                elif tipo == "lendo":
                    fase("lendo")
                    previsao = dados.get("previsao") or {}
                    andamento["previsao_s"] = previsao.get("segundos", 0) if previsao.get("sabe") else 0
                    andamento["caracteres"] = dados.get("caracteres", 0)
                    lido_chars = dados["caracteres"]
                    trabalho.etapas[1].titulo = "Lendo os documentos"
                    trabalho.etapas[1].estado = EXECUTANDO
                    trabalho.etapas[1].detalhe = f"{dados['caracteres']:,}".replace(",", ".") + " caracteres"
                    yield _sse("lendo", dados)
                    yield _sse("etapas", {"etapas": [asdict_etapa(e) for e in trabalho.etapas]})
                elif tipo == "escrevendo":
                    fase("escrevendo")
                    # A leitura acabou de verdade: a primeira palavra saiu.
                    trabalho.etapas[1].estado = CONCLUIDO
                    trabalho.etapas[1].detalhe = f"{dados['lendo_segundos']} s"
                    trabalho.etapas.append(Etapa("Escrevendo a resposta", estado=EXECUTANDO))
                    yield _sse("escrevendo", dados)
                    yield _sse("etapas", {"etapas": [asdict_etapa(e) for e in trabalho.etapas]})
                elif tipo == "medida":
                    medida = dados
                    yield _sse("medida", dados)
                elif tipo == "token":
                    partes.append(dados["t"])
                    # Contar a cada dez pedacos basta para o cartao, e nao
                    # refaz a conta do texto inteiro a cada palavra.
                    if len(partes) % 10 == 0:
                        andamento["palavras"] = len("".join(partes).split())
                    yield _sse("token", dados)
                elif tipo == "vazio":
                    trabalho.etapas[0].estado = CONCLUIDO
                    trabalho.etapas[1].estado = CONCLUIDO
                    trabalho.estado = CONCLUIDO
                    trabalho.dizer("paulus", dados["mensagem"])
                    estado.trabalhos.salvar(trabalho)
                    yield _sse("vazio", dados)
                    return
                elif tipo == "fim":
                    pass  # o fechamento e daqui: o modulo nao sabe de trabalho
        except Exception as exc:
            trabalho.etapas[1].estado = "falhou"
            trabalho.estado = "falhou"
            estado.trabalhos.salvar(trabalho)
            yield _sse("erro", {"mensagem": str(exc)})
            return

        segundos = round(time.time() - inicio, 1)

        # Parou no meio: guarda o que o modelo ja tinha escrito, marcado como
        # interrompido, e a conversa fica "parada" - da para perguntar de novo.
        # Nada disso entra no historico de ritmo: uma leitura cortada nao e
        # medida de quanto esta maquina leva.
        if parar.is_set():
            passos.close()
            pausar_o_que_executava()
            escrito = "".join(partes).strip()
            if escrito:
                trabalho.dizer(
                    "paulus", escrito,
                    fontes=fontes, cobertura=cobertura, segundos=segundos,
                    nivel=nivel, inferencia=inferencia, interrompida=True,
                )
            estado.trabalhos.salvar(trabalho)
            yield _sse("parado", {"segundos": segundos, "titulo": trabalho.titulo, "escreveu": bool(escrito)})
            return

        # O que acabou de acontecer entra no historico da maquina: e dele que
        # sai o "leituras deste tamanho levaram ~70 s aqui" da proxima vez.
        if medida and lido_chars:
            estado.ritmo.anotar(
                modelo=estado.client.model,
                caracteres=lido_chars,
                # A espera, e nao o tempo de calculo do modelo: a previsao
                # existe para dizer quanto a pessoa vai esperar, e a espera
                # inclui carregar o modelo quando ele esta frio.
                segundos_lendo=medida.get("esperou_segundos") or medida.get("lendo_segundos", 0),
                palavras=len("".join(partes).split()),
                segundos_escrevendo=medida.get("escrevendo_segundos", 0),
            )

        for etapa in trabalho.etapas:
            if etapa.estado == EXECUTANDO:
                etapa.estado = CONCLUIDO
        trabalho.estado = CONCLUIDO
        # Quer ver o documento? A oferta sai por regra dos trechos usados, e
        # fica guardada na propria resposta - e assim que a conversa sabe que
        # ja ofereceu e nao oferece o mesmo arquivo de novo.
        oferta = ferramentas.oferta_de_exibir(
            fontes, estado.searcher.documents, _documentos_ja_oferecidos(trabalho))
        trabalho.dizer(
            "paulus", "".join(partes).strip(),
            fontes=fontes, cobertura=cobertura, segundos=segundos,
            nivel=nivel, inferencia=inferencia, proposta=oferta or {},
        )
        estado.trabalhos.salvar(trabalho)
        # Uma resposta leva cerca de um minuto nesta maquina: quem foi para
        # outra janela esperar fica sabendo que acabou.
        resumo = " ".join("".join(partes).split())
        avisos.avisar("resposta", "Resposta pronta · " + (trabalho.titulo or "Conversa")[:60],
                      resumo[:140] + ("…" if len(resumo) > 140 else ""))
        yield _sse("fim", {"segundos": segundos, "titulo": trabalho.titulo})
        if oferta:
            yield _sse("oferta", oferta)

    return StreamingResponse(
        gerar(),
        media_type="text/event-stream",
        headers={"Cache-Control": "no-cache", "X-Accel-Buffering": "no"},
    )


def _documentos_ja_oferecidos(trabalho) -> set[str]:
    """
    Os arquivos que esta conversa ja ofereceu mostrar, ja mostrou, ou que a
    pessoa pediu para abrir - o cartao deles continua na conversa, e oferecer
    de novo embaixo da resposta seguinte seria o mesmo convite duas vezes.
    """
    vistos: set[str] = set()
    for m in trabalho.mensagens:
        proposta = m.proposta or {}
        if proposta.get("tipo") == "exibir":
            vistos.update(proposta.get("nomes") or (proposta.get("campos") or {}).get("nomes") or [])
        if proposta.get("tipo") == "abrir" and (proposta.get("campos") or {}).get("nome"):
            vistos.add(proposta["campos"]["nome"])
        if (m.feito or {}).get("tipo") in ("exibir", "abrir", "editar") and (m.feito or {}).get("nome"):
            vistos.add(m.feito["nome"])
    return vistos


class PropostaConfirmada(BaseModel):
    tipo: str = ""
    campos: dict = {}


@app.post("/api/trabalhos/{id_}/fazer")
def trabalhos_fazer(id_: str, payload: PropostaConfirmada) -> dict:
    """
    Grava o que a pessoa confirmou na conversa.

    A rota existe separada de propósito: entender a frase e gravar sao dois
    passos, e o segundo so acontece depois do sim. Os campos que chegam aqui
    sao os que estavam na tela - a pessoa pode ter corrigido a data ou a hora
    antes de confirmar, e vale o que ela deixou.
    """
    trabalho = estado.trabalhos.obter(id_)
    if not trabalho:
        raise HTTPException(status_code=404, detail="trabalho nao encontrado")

    campos = dict(payload.campos or {})
    pendente = False
    try:
        if payload.tipo == "abrir":
            nome = str(campos.get("nome", ""))
            doc = next((d for d in estado.searcher.documents if d.name == nome), None)
            if not doc:
                raise HTTPException(status_code=404, detail="esse documento nao esta mais aberto")
            alvo = Path(doc.path)
            if not alvo.exists():
                raise HTTPException(status_code=404, detail="o arquivo saiu do lugar")
            import os

            os.startfile(str(alvo))  # noqa: S606 - abre no programa do proprio Windows
            novo = 0
            feito = {"nome": nome, "caminho": str(alvo)}
            resumo = f"Abri “{nome}” no programa padrao do Windows"
            onde = "biblioteca"
        elif payload.tipo in ferramentas.POR_PROPOSTA:
            # Agenda, cadastro e nota fiscal: as ferramentas do catalogo. O
            # catalogo confere os campos do jeito que o modelo e a tela os
            # mandam, e so entao grava - ou, na NFS-e, so confere.
            feito_ferramenta = ferramentas.executar(ferramentas.POR_PROPOSTA[payload.tipo], campos, estado)
            novo = feito_ferramenta["id"]
            feito = feito_ferramenta["registro"]
            resumo = feito_ferramenta["resumo"]
            onde = feito_ferramenta["onde"]
            pendente = bool(feito_ferramenta.get("pendente"))
        elif payload.tipo == "tarefa":
            novo = estado.tarefas.salvar(campos)
            feito = estado.tarefas.obter(novo)
            resumo = f"Criei a tarefa “{feito['titulo']}”"
            if feito.get("prazo"):
                resumo += f", com prazo em {escritorio._br(feito['prazo'])}"
            onde = "tarefas"
        elif payload.tipo == "servico":
            # A pasta de trabalho nasce da conversa (docs/ui, A15). O cliente
            # e ligado pelo nome como esta em Cadastros; sem ficha, a pasta
            # abre sem cliente e o resumo diz isso.
            nome_cliente = " ".join(str(campos.get("cliente", "")).split())
            ficha = None
            if nome_cliente:
                ficha = next((f for f in estado.cadastros.listar()
                              if " ".join(f["nome"].split()).lower() == nome_cliente.lower()), None)
            novo = estado.servicos.salvar({
                "nome": str(campos.get("nome", "")), "cadastro_id": ficha["id"] if ficha else None,
                "descricao": str(campos.get("descricao", "")),
            })
            estado.servicos.trilha(novo, "Serviço aberto a partir da conversa “" + trabalho.titulo + "”")
            feito = estado.servicos.obter(novo)
            resumo = f"Abri o serviço “{feito['nome']}”"
            if feito.get("cliente_nome"):
                resumo += f" para {feito['cliente_nome']}"
            elif nome_cliente:
                resumo += f" sem cliente ligado — “{nome_cliente}” não está em Cadastros"
            onde = "servicos"
        else:
            raise HTTPException(status_code=400, detail="nao sei fazer isso")
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc

    trabalho.dizer("paulus", resumo + ".", feito={"tipo": payload.tipo, "id": novo, "onde": onde,
                                                  "pendente": pendente, "nome": str(campos.get("nome", ""))})
    estado.trabalhos.salvar(trabalho)
    return {"id": novo, "resumo": resumo, "onde": onde, "registro": feito, "pendente": pendente}


def _responder_sem_documentos(trabalho, lido, pergunta: str) -> StreamingResponse:
    """
    O que nao precisa ler documento nenhum.

    Duas coisas caem aqui. Uma acao sobre a agenda ou as tarefas vira uma
    PROPOSTA, com os campos a vista - entender nao e fazer, e nada entra no
    calendario de alguem por interpretacao de frase. E pergunta sobre o proprio
    programa e respondida do registro de habilidades, sem modelo: a lista do
    que ele faz esta declarada no codigo, e passar isso por um modelo de 3
    bilhoes de parametros so acrescentaria chance de erro.
    """
    def gerar() -> Iterator[str]:
        if lido.tipo == "sobre":
            texto = _o_que_eu_faco()
            trabalho.etapas = [Etapa("Responder", estado=CONCLUIDO)]
            trabalho.estado = CONCLUIDO
            trabalho.dizer("paulus", texto)
            estado.trabalhos.salvar(trabalho)
            yield _sse("token", {"t": texto})
            yield _sse("fim", {"segundos": 0, "titulo": trabalho.titulo})
            return

        # A regra achou a acao mas nao um campo obrigatorio, e a frase tem com
        # o que preencher: so aqui o modelo e chamado. E um minuto no
        # computador sem placa de video, entao a etapa aparece andando.
        ajuda: list[str] = []
        ferramenta = ferramentas.POR_PROPOSTA.get(lido.tipo, "")
        if lido.precisa_modelo and ferramenta:
            import time

            agora = time.time()
            trabalho.etapas = [Etapa("Entender o pedido", estado=EXECUTANDO)]
            trabalho.estado = EXECUTANDO
            estado.trabalhos.salvar(trabalho)
            estado.andamento[trabalho.id] = {"fase": "entendendo", "inicio": agora, "desde": agora,
                                             "previsao_s": 0, "palavras": 0, "documentos": 0,
                                             "caracteres": 0}
            yield _sse("etapas", {"etapas": [asdict_etapa(e) for e in trabalho.etapas]})
            try:
                ajuda = ferramentas.completar_com_modelo(
                    lido, pergunta,
                    lambda instrucao, sistema: estado.client.ask_json(instrucao, sistema=sistema))
            finally:
                estado.andamento.pop(trabalho.id, None)

        proposta = {
            "tipo": lido.tipo,
            "titulo": lido.titulo,
            "campos": lido.campos,
            "porque": lido.porque,
            "falta": lido.falta,
            "pergunta": pergunta,
            "ferramenta": ferramenta,
            "disponivel": ferramentas.CATALOGO_FERRAMENTAS[ferramenta]["disponivel"] if ferramenta else True,
            "ajuda_do_modelo": ajuda,
        }
        trabalho.etapas = [Etapa("Entender o pedido", estado=CONCLUIDO)]
        trabalho.estado = CONCLUIDO
        trabalho.dizer("paulus", "", proposta=proposta)
        estado.trabalhos.salvar(trabalho)
        yield _sse("proposta", proposta)
        yield _sse("fim", {"segundos": 0, "titulo": trabalho.titulo})

    return StreamingResponse(
        gerar(),
        media_type="text/event-stream",
        headers={"Cache-Control": "no-cache", "X-Accel-Buffering": "no"},
    )


def _perguntar_onde_procurar(trabalho, pergunta: str, foco: list[str]) -> StreamingResponse:
    """
    O cartao "onde eu procuro?": so no que a conversa vinha lendo, ou no
    acervo inteiro. Fica guardado como proposta; escolher refaz a pergunta
    com o Retomar, que troca o cartao pela resposta.
    """
    abertos = {d.name for d in estado.searcher.documents}
    proposta = {
        "tipo": "escopo", "titulo": "", "campos": {}, "porque": "", "falta": "",
        "pergunta": pergunta,
        "nomes": [n for n in foco if n in abertos][:3],
        "total": len(estado.searcher.documents),
    }

    def gerar() -> Iterator[str]:
        trabalho.etapas = [Etapa("Entender o pedido", estado=CONCLUIDO)]
        trabalho.estado = CONCLUIDO
        trabalho.dizer("paulus", "", proposta=proposta)
        estado.trabalhos.salvar(trabalho)
        yield _sse("proposta", proposta)
        yield _sse("fim", {"segundos": 0, "titulo": trabalho.titulo})

    return StreamingResponse(
        gerar(),
        media_type="text/event-stream",
        headers={"Cache-Control": "no-cache", "X-Accel-Buffering": "no"},
    )


def _o_que_eu_faco() -> str:
    """
    O que o programa faz, montado do proprio programa.

    Sai da lista de telas e do registro de habilidades - nao de um texto
    escrito a mao. Um texto a mao envelhece na primeira tela nova: passa a
    prometer o que nao existe, ou a esconder o que existe. E o que esta em pe
    e o que ainda nao esta saem separados, porque a diferenca e a informacao.
    """
    grupos = destinos.por_grupo()

    linhas = ["Eu trabalho com o que está nesta máquina. Nada sai daqui.", ""]
    faltando: list[str] = []

    for bloco in grupos:
        prontas = [d for d in bloco["destinos"] if d["pronta"]]
        faltando += [d["nome"] for d in bloco["destinos"] if not d["pronta"]]
        if not prontas:
            continue
        # Sem asterisco: a conversa mostra texto puro, e "**Documentos**"
        # apareceria com os asteriscos na tela.
        linhas.append(bloco["grupo"].upper())
        linhas += [f"  {d['nome']} — {d['resolve'][0].lower() + d['resolve'][1:]}"
                   for d in prontas]
        linhas.append("")

    linhas += [
        "AQUI NA CONVERSA",
        "  perguntar sobre os documentos abertos, com o trecho de origem à vista",
        "  anotar na agenda — “anote uma reunião dia 20/10 às 14h com lembrete "
        "30 minutos antes”",
        "  criar tarefa — “crie uma tarefa para revisar o contrato até sexta”",
        "  abrir um serviço — “abra um serviço para a Cooperativa: renovação "
        "do contrato de logística”",
        "",
        "Antes de gravar qualquer coisa eu mostro o que entendi, e você "
        "confirma. As outras telas estão no menu à esquerda.",
    ]

    if faltando:
        linhas += ["", "Ainda não faço: " + ", ".join(faltando) + "."]

    return "\n".join(linhas)


def _quantos(n: int, palavra: str, muitos: str = "") -> str:
    """"1 documento", "6 documentos" - nunca "6 documento(s)"."""
    return f"{n} {palavra if n == 1 else (muitos or palavra + 's')}"


def asdict_etapa(etapa: Etapa) -> dict:
    return {
        "titulo": etapa.titulo,
        "estado": etapa.estado,
        "feitos": etapa.feitos,
        "total": etapa.total,
        "detalhe": etapa.detalhe,
    }


# ------------------------------------------------------------------ maquina


@app.get("/api/recursos")
def maquina() -> dict:
    dados = recursos.ler()
    dados["devagar"] = estado.devagar
    return dados


@app.post("/api/config")
def configurar(payload: Ajuste2) -> dict:
    estado.devagar = payload.devagar
    return {"devagar": estado.devagar}


# ----------------------------------------------------------------- agenda


class FichaCompromisso(BaseModel):
    id: int | None = None
    dados: dict = {}


class NotaDia(BaseModel):
    dia: str
    texto: str = ""


def _documentos_com_data() -> list[dict]:
    from classify import CacheClassificacao

    return list(CacheClassificacao(CLASSIFICACAO_PATH).dados.values())


@app.get("/api/agenda")
def agenda_grade(de: str = "", ate: str = "", pessoa: int = 0) -> dict:
    """
    Tudo o que tem data no periodo: compromisso, prazo de tarefa e data de
    contrato na mesma grade. `pessoa` (id de Cadastros) e a agenda de alguem
    da equipe: so o que esta com essa pessoa.
    """
    from datetime import date, timedelta

    if not de or not ate:
        hoje = date.today()
        de = (hoje - timedelta(days=hoje.day - 1)).isoformat()
        ate = (date.fromisoformat(de) + timedelta(days=45)).isoformat()

    # "todas" sao as abertas; as concluidas que sao etapa de servico entram
    # para a etapa feita continuar no dia dela.
    tarefas = estado.tarefas.listar("todas") + [t for t in estado.tarefas.listar("concluidas") if t.get("servico_id")]
    grade = estado.agenda.grade(de, ate, tarefas, _documentos_com_data(), pessoa or None)
    grade["compromissos"] = estado.agenda.listar(de, ate)
    grade["tipos"] = [{"valor": k, "rotulo": v} for k, v in TIPOS_AGENDA.items()]
    grade["ondes"] = [{"valor": k, "rotulo": v} for k, v in ONDES.items() if k]
    grade["clientes"] = [{"id": f["id"], "nome": f["nome"]} for f in estado.cadastros.listar()]
    return grade


@app.get("/api/agenda/dia")
def agenda_dia(dia: str) -> dict:
    # As abertas e as concluidas do dia: a grade conta as duas, e o painel
    # mostra a feita riscada em vez de sumir com ela.
    tarefas = [t for t in estado.tarefas.listar("todas") + estado.tarefas.listar("concluidas") if t.get("prazo") == dia]
    return {
        "dia": dia,
        "compromissos": estado.agenda.listar(dia, dia),
        "tarefas": tarefas,
        "nota": estado.agenda.nota(dia),
    }


@app.post("/api/agenda")
def agenda_salvar(payload: FichaCompromisso) -> dict:
    try:
        id_ = estado.agenda.salvar(payload.dados, payload.id)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    return estado.agenda.obter(id_) or {}


@app.delete("/api/agenda/{id_}")
def agenda_apagar(id_: int) -> dict:
    c = estado.base.um("SELECT titulo, data, hora FROM compromissos WHERE id = ?", (id_,))
    if not c:
        raise HTTPException(status_code=404, detail="compromisso nao encontrado")
    entrada = estado.lixeira.apagar_linha("compromisso", id_, c["titulo"], "Agenda · " + c["data"] + " " + (c["hora"] or ""))
    return _foi_para_lixeira(entrada, "apagado", id_)


@app.post("/api/agenda/nota")
def agenda_nota(payload: NotaDia) -> dict:
    estado.agenda.gravar_nota(payload.dia, payload.texto)
    return {"dia": payload.dia, "texto": payload.texto}


@app.get("/api/agenda/livres")
def agenda_livres(duracao: int = 60, dia: str = "") -> dict:
    """Horarios em que cabe um compromisso desse tamanho."""
    regra = estado.prefs.dados.get("disponibilidade", {})
    if dia:
        return {"dia": dia, "horarios": estado.agenda.livres(dia, duracao, regra)}
    return {"proximos": estado.agenda.proximos_livres(duracao, regra)}


# -------------------------------------------------------------- cadastros


class FichaCadastro(BaseModel):
    id: int | None = None
    dados: dict = {}


class VinculoDoc(BaseModel):
    sha1: str
    nome: str = ""


@app.get("/api/cadastros")
def cadastros_listar(tipo: str = "", termo: str = "", ordem: str = "nome") -> dict:
    return {
        "fichas": estado.cadastros.listar(tipo, termo, ordem),
        "contagem": estado.cadastros.contagem(),
        "ordens": {"nome": "A–Z", "aberto": "em aberto", "atraso": "atraso"},
        "tipos": [{"valor": k, "rotulo": v} for k, v in TIPOS_CADASTRO.items()],
    }


# Quantas sugestoes a tela recebe de uma vez. A conta de verdade vai junto.
LIMITE_DE_SUGESTOES = 300


class NomeSugerido(BaseModel):
    nome: str


def _documentos_por_ler() -> tuple[int, list[str]]:
    """Quantos documentos o Acervo tem e quais ainda nao foram lidos."""
    from classify import CacheClassificacao

    lidos = CacheClassificacao(CLASSIFICACAO_PATH).dados
    documentos = list(estado.searcher.documents)
    faltam = [
        d.path for d in documentos
        if d.sha1 and d.sha1 not in lidos and Path(d.path).exists()
    ]
    return len(documentos), faltam


@app.get("/api/cadastros/sugestoes")
def cadastros_sugestoes() -> dict:
    """Quem ja aparece nos documentos e ainda nao tem ficha (nem foi recusado)."""
    todas = estado.cadastros.sugestoes(CLASSIFICACAO_PATH, limite=None)
    total, faltam = _documentos_por_ler()
    return {
        "sugestoes": todas[:LIMITE_DE_SUGESTOES],
        "total": len(todas),
        "ignorados": len(estado.cadastros.ignorados()),
        "levantamento": {"documentos": total, "faltam": len(faltam)},
    }


@app.post("/api/cadastros/sugestoes/ignorar")
def cadastros_sugestao_ignorar(payload: NomeSugerido) -> dict:
    """O nome recusado sai das sugestoes e nao volta."""
    try:
        quantos = estado.cadastros.ignorar(payload.nome)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    return {"nome": payload.nome, "ignorados": quantos}


@app.post("/api/cadastros/sugestoes/desfazer")
def cadastros_sugestao_desfazer(payload: NomeSugerido) -> dict:
    """Desfaz o ignorar: o nome volta a ser sugerido."""
    return {"nome": payload.nome, "voltou": estado.cadastros.voltar_a_sugerir(payload.nome)}


@app.post("/api/cadastros/levantamento")
def cadastros_levantamento() -> StreamingResponse:
    """
    Fazer levantamento: le os documentos do Acervo que ainda nao foram lidos,
    para achar quem assina e ainda nao tem ficha.

    E a mesma leitura do Acervo (habilidades/classificar.py), que guarda o
    resultado por conteudo - o que ja foi lido nao e lido de novo, e o que
    este levantamento ler serve ao Acervo tambem. Sem documento novo, nao
    ha o que ler: responde na hora com a conta de sugestoes de agora.
    """
    total, faltam = _documentos_por_ler()
    antes = len(estado.cadastros.sugestoes(CLASSIFICACAO_PATH, limite=None))

    habilidade = None
    if faltam:
        habilidade = estado.registro.obter("classificar")
        if not habilidade or not habilidade.executavel:
            raise HTTPException(status_code=503, detail="a leitura dos documentos não carregou")
        falta = [p for p in habilidade.precisa if not _disponibilidade().get(p, False)]
        if falta:
            from habilidade_base import ROTULOS_PRECISA

            nomes = ", ".join(ROTULOS_PRECISA.get(p, p) for p in falta)
            raise HTTPException(status_code=400, detail=f"para ler os documentos, precisa de: {nomes}")
        estado.cancelar.clear()

    def gerar() -> Iterator[str]:
        yield _sse("inicio", {"novos": len(faltam), "total": total})
        lidos, parado = 0, False
        try:
            if habilidade:
                for tipo, dados in habilidade.executar(_contexto(), caminhos=faltam):
                    if tipo == "progresso":
                        yield _sse("progresso", dados)
                    elif tipo == "resultados":
                        lidos = sum(1 for d in dados.get("documentos", []) if not d.get("erro"))
                        parado = bool(dados.get("parado"))
        except Exception as exc:
            yield _sse("erro", {"mensagem": str(exc)})
            return
        depois = len(estado.cadastros.sugestoes(CLASSIFICACAO_PATH, limite=None))
        yield _sse("fim", {
            "lidos": lidos, "total": total, "parado": parado, "novos": len(faltam),
            "sugestoes": depois, "novas": max(0, depois - antes),
        })

    return StreamingResponse(
        gerar(),
        media_type="text/event-stream",
        headers={"Cache-Control": "no-cache", "X-Accel-Buffering": "no"},
    )


@app.post("/api/cadastros")
def cadastros_salvar(payload: FichaCadastro) -> dict:
    try:
        id_ = estado.cadastros.salvar(payload.dados, payload.id)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    return estado.cadastros.obter(id_) or {}


@app.post("/api/cadastros/{id_}/vincular")
def cadastros_vincular(id_: int, payload: VinculoDoc) -> dict:
    if not estado.cadastros.obter(id_):
        raise HTTPException(status_code=404, detail="cadastro nao encontrado")
    estado.cadastros.vincular(id_, payload.sha1, payload.nome)
    return estado.cadastros.obter(id_) or {}


@app.delete("/api/cadastros/{id_}")
def cadastros_apagar(id_: int) -> dict:
    f = estado.base.um("SELECT nome, tipo FROM cadastros WHERE id = ?", (id_,))
    if not f:
        raise HTTPException(status_code=404, detail="cadastro nao encontrado")
    entrada = estado.lixeira.apagar_linha("cadastro", id_, f["nome"], "Cadastros · " + f["tipo"])
    return _foi_para_lixeira(entrada, "apagado", id_)


# ---------------------------------------------------------------- tarefas


class FichaTarefa(BaseModel):
    id: int | None = None
    dados: dict = {}


class MarcaTarefa(BaseModel):
    valor: bool = True


class NovaEtapa(BaseModel):
    titulo: str


@app.get("/api/tarefas")
def tarefas_listar(filtro: str = "meu_dia", lista: str = "") -> dict:
    """
    A visao Tarefas e compromissos da Agenda. O compromisso entra em Meu dia
    (os de hoje) e em Planejadas (os que vem); nas listas, em Importante e em
    Concluidas, nao - compromisso nao tem lista, estrela nem conclusao.
    """
    from datetime import date, timedelta

    hoje = date.today().isoformat()
    amanha = (date.today() + timedelta(days=1)).isoformat()
    compromissos: list[dict] = []
    if not lista and filtro == "meu_dia":
        compromissos = estado.agenda.listar(hoje, hoje)
    elif not lista and filtro == "planejadas":
        compromissos = estado.agenda.listar(amanha, "9999-12-31")
    contagens = estado.tarefas.contagens()
    contagens["compromissos_hoje"] = estado.agenda.base.contar("compromissos", "data = ?", (hoje,))
    contagens["compromissos_planejados"] = estado.agenda.base.contar("compromissos", "data > ?", (hoje,))
    return {
        "tarefas": estado.tarefas.listar(filtro, lista),
        "compromissos": compromissos,
        "contagens": contagens,
        "listas": estado.tarefas.listas(),
        "clientes": [
            {"id": f["id"], "nome": f["nome"]}
            for f in estado.cadastros.listar()
        ],
        "repeticoes": [
            {"valor": k, "rotulo": v} for k, v in tarefas_mod.REPETICOES.items()
        ],
    }


@app.get("/api/tarefas/sugestoes")
def tarefas_sugestoes() -> dict:
    """Prazos que os documentos ja lidos pedem para conferir."""
    from classify import CacheClassificacao

    cache = CacheClassificacao(CLASSIFICACAO_PATH).dados
    return {"sugestoes": estado.tarefas.sugerir(list(cache.values()))}


@app.post("/api/tarefas")
def tarefas_salvar(payload: FichaTarefa) -> dict:
    try:
        id_ = estado.tarefas.salvar(payload.dados, payload.id)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    return estado.tarefas.obter(id_) or {}


@app.post("/api/tarefas/{id_}/concluir")
def tarefas_concluir(id_: int, payload: MarcaTarefa) -> dict:
    """Concluir uma tarefa que repete ja deixa a proxima no lugar."""
    resultado = estado.tarefas.concluir(id_, payload.valor)
    saida = estado.tarefas.obter(id_) or {"apagada": True}
    if resultado.get("proxima"):
        proxima = estado.tarefas.obter(resultado["proxima"])
        saida["proxima"] = proxima
        saida["aviso_repeticao"] = (
            f"criei a próxima para {proxima['prazo']}" if proxima and proxima.get("prazo") else ""
        )
    return saida


@app.post("/api/tarefas/{id_}/importante")
def tarefas_importante(id_: int, payload: MarcaTarefa) -> dict:
    estado.tarefas.marcar_importante(id_, payload.valor)
    return estado.tarefas.obter(id_) or {}


@app.post("/api/tarefas/{id_}/meu-dia")
def tarefas_meu_dia(id_: int, payload: MarcaTarefa) -> dict:
    estado.tarefas.marcar_meu_dia(id_, payload.valor)
    return estado.tarefas.obter(id_) or {}


@app.post("/api/tarefas/{id_}/etapas")
def tarefas_nova_etapa(id_: int, payload: NovaEtapa) -> dict:
    try:
        estado.tarefas.nova_etapa(id_, payload.titulo)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    return estado.tarefas.obter(id_) or {}


@app.post("/api/etapas/{id_}")
def etapa_marcar(id_: int, payload: MarcaTarefa) -> dict:
    estado.tarefas.marcar_etapa(id_, payload.valor)
    return {"ok": True}


@app.get("/api/tarefas/{id_}/vinculos")
def tarefas_vinculos(id_: int) -> dict:
    """Os documentos ligados a tarefa, com o caminho de cada um se ainda existir."""
    from classify import CacheClassificacao

    cache = CacheClassificacao(CLASSIFICACAO_PATH).dados
    por_sha = {d.path: d for d in estado.searcher.documents}
    achados = []
    for v in estado.tarefas.vinculos_de(id_):
        doc = next((d for d in estado.searcher.documents if d.sha1 == v["sha1"]), None)
        conhecido = cache.get(v["sha1"]) or {}
        achados.append({
            **v,
            "existe": doc is not None,
            "caminho": doc.path if doc else "",
            "tipo_rotulo": ROTULOS.get(conhecido.get("tipo", ""), ""),
        })
    return {"vinculos": achados}


@app.post("/api/tarefas/{id_}/vincular")
def tarefas_vincular(id_: int, payload: VinculoDoc) -> dict:
    if not estado.tarefas.obter(id_):
        raise HTTPException(status_code=404, detail="tarefa não encontrada")
    estado.tarefas.vincular(id_, payload.sha1, payload.nome)
    return tarefas_vinculos(id_)


@app.delete("/api/vinculos/{id_}")
def vinculo_apagar(id_: int) -> dict:
    estado.tarefas.desvincular(id_)
    return {"removido": id_}


@app.delete("/api/tarefas/{id_}")
def tarefas_apagar(id_: int) -> dict:
    t = estado.base.um("SELECT titulo, prazo FROM tarefas WHERE id = ?", (id_,))
    if not t:
        raise HTTPException(status_code=404, detail="tarefa nao encontrada")
    entrada = estado.lixeira.apagar_linha("tarefa", id_, t["titulo"], "Agenda · Tarefas" + (" · prazo " + t["prazo"] if t["prazo"] else ""))
    return _foi_para_lixeira(entrada, "apagada", id_)


# ------------------------------------------------------------- aprovacoes


class Decisao(BaseModel):
    ids: list[str] = []
    aprovar: bool = True


@app.get("/api/aprovacoes")
def aprovacoes_listar() -> dict:
    return estado.fila.para_tela()


@app.post("/api/aprovacoes/decidir")
def aprovacoes_decidir(payload: Decisao) -> dict:
    """
    Aprova ou recusa. Aprovar executa a acao; recusar so encerra o pedido.

    A execucao nao mora na fila: cada acao e feita por quem sabe faze-la, e o
    resultado volta para o pedido. Assim a fila nao precisa entender de mover
    arquivo, assinar ou enviar.
    """
    feitos, falhas = [], []

    for id_ in payload.ids:
        pedido = estado.fila.decidir(id_, payload.aprovar)
        if not pedido:
            falhas.append({"id": id_, "motivo": "pedido nao esta mais na fila"})
            continue

        if not payload.aprovar:
            feitos.append({"id": id_, "estado": pedido.estado})
            continue

        executor = EXECUTORES.get(pedido.acao)
        if not executor:
            estado.fila.registrar_resultado(id_, "aprovado, sem nada a executar")
            feitos.append({"id": id_, "estado": pedido.estado})
            continue

        try:
            resultado = executor(pedido)
            # O executor devolve a frase do historico ou, quando ha para onde
            # levar a pessoa depois do sim, {"texto", "desfecho"}: a tela usa
            # o desfecho para abrir o que acabou de ser feito.
            desfecho = None
            if isinstance(resultado, dict):
                desfecho = resultado.get("desfecho")
                resultado = resultado.get("texto", "")
            estado.fila.registrar_resultado(id_, resultado)
            feito = {"id": id_, "estado": pedido.estado, "categoria": pedido.categoria, "resultado": resultado}
            if desfecho:
                feito["desfecho"] = desfecho
            feitos.append(feito)
        except Exception as exc:
            estado.fila.registrar_resultado(id_, f"nao consegui: {exc}", falhou=True)
            falhas.append({"id": id_, "motivo": str(exc)})

    return {"feitos": feitos, "falhas": falhas, **estado.fila.para_tela()}


def _executar_mover(pedido) -> str:
    """Executa um plano de organizacao que estava esperando aprovacao."""
    from organize import Movimento, Plano

    dados = pedido.dados
    plano = Plano(
        movimentos=[Movimento(**m) for m in dados.get("movimentos", [])],
        destino=dados.get("destino", ""),
        padrao=dados.get("padrao", ""),
        ignorados=dados.get("ignorados", []),
        operacao=dados.get("operacao", "mover"),
    )
    resultado = aplicar_plano(plano, DIARIOS_DIR)
    _guardar_organizacao(plano, resultado, pedido.id)
    estado.recarregar(force=True)
    verbo = "copiado(s)" if plano.operacao == "copiar" else "movido(s)"
    if resultado.falhas:
        return f"{resultado.movidos} {verbo}, {len(resultado.falhas)} falha(s)"
    return f"{resultado.movidos} arquivo(s) {verbo}"


def _executar_assinar(pedido) -> dict:
    """Assina o que estava esperando o sim na fila."""
    senha = estado.cofre.senha_agora()
    if not senha:
        raise RuntimeError("a senha do certificado expirou - abra o certificado e tente de novo")

    resultado = _assinar_de_fato(pedido.dados, senha)
    if resultado.erro:
        raise RuntimeError(resultado.erro)

    aviso = "" if resultado.icp_brasil else " (certificado fora da ICP-Brasil)"
    return {
        "texto": f"assinado, codigo {resultado.codigo}{aviso}",
        # Depois do sim, a tela abre o documento ja assinado.
        "desfecho": {"tipo": "assinatura", "origem": pedido.dados.get("arquivo", ""), **resultado.to_dict()},
    }


def _executar_assinar_lote(pedido) -> dict:
    """
    Assina os documentos de um lote que esperava o sim na fila.

    Um documento que falha nao derruba os outros: cada um tem a sua
    assinatura, e o desfecho diz quais sairam e quais nao. So quando nenhum
    sai o pedido conta como falho.
    """
    senha = estado.cofre.senha_agora()
    if not senha:
        raise RuntimeError("a senha do certificado expirou - abra o certificado e tente de novo")

    feitos, falhas = _assinar_varios(pedido.dados.get("itens") or [], senha)
    if not feitos:
        motivo = falhas[0]["motivo"] if falhas else "o lote estava vazio"
        raise RuntimeError(f"nenhum documento assinado: {motivo}")

    texto = f"{len(feitos)} de {len(feitos) + len(falhas)} documento(s) assinado(s)"
    if falhas:
        texto += f"; {len(falhas)} com problema"
    if not all(f.get("icp_brasil") for f in feitos):
        texto += " (certificado fora da ICP-Brasil)"
    return {"texto": texto, "desfecho": {"tipo": "assinatura_lote", "feitos": feitos, "falhas": falhas}}


def _executar_enviar(pedido) -> str:
    """Manda o e-mail que estava esperando o sim na fila."""
    resultado = _enviar_de_fato(pedido.dados)
    quem = ", ".join(resultado.get("para", []))
    anexos = resultado.get("anexos") or []
    return f"enviado para {quem}" + (f" com {len(anexos)} anexo(s)" if anexos else "")


def _executar_apagar_do_acervo(pedido) -> str:
    """
    Tira da biblioteca o que a pessoa aprovou tirar.

    So apaga dentro da pasta do programa, e confere de novo na hora de apagar:
    o pedido pode ter ficado na fila tempo suficiente para o arquivo ter sido
    movido para fora, e apagar fora daqui seria apagar o original de alguem.
    """
    import shutil

    pasta = Path(estado.pasta).resolve()
    apagados, recusados = 0, []

    for bruto in pedido.dados.get("caminhos", []):
        alvo = Path(bruto)
        try:
            resolvido = alvo.resolve()
        except OSError:
            recusados.append(alvo.name)
            continue
        if pasta not in resolvido.parents:
            recusados.append(alvo.name)
            continue
        try:
            resolvido.unlink(missing_ok=True)
            apagados += 1
        except OSError:
            recusados.append(alvo.name)

    estado.recarregar(force=True)
    if recusados:
        return f"{apagados} tirado(s); nao mexi em {len(recusados)}: {', '.join(recusados[:3])}"
    return f"{apagados} documento(s) tirado(s) da biblioteca"


def _executar_exportar(pedido) -> str:
    """Copia para fora da pasta do programa, sem passar por cima de nada."""
    import shutil

    destino = Path(pedido.dados.get("destino", ""))
    if not destino.is_dir():
        raise RuntimeError("a pasta de destino nao existe mais")

    copiados, falhas = 0, []
    for bruto in pedido.dados.get("caminhos", []):
        origem = Path(bruto)
        if not origem.exists():
            falhas.append(origem.name)
            continue
        alvo = destino / origem.name
        if alvo.exists():
            alvo = acervo._nome_livre(destino, origem.name, set())
        try:
            shutil.copy2(origem, alvo)
            copiados += 1
        except OSError:
            falhas.append(origem.name)

    if falhas:
        return f"{copiados} copiado(s); nao consegui {len(falhas)}: {', '.join(falhas[:3])}"
    return f"{copiados} documento(s) copiado(s) para {destino}"


EXECUTORES = {
    "organizar.mover": _executar_mover,
    "acervo.apagar": _executar_apagar_do_acervo,
    "acervo.exportar": _executar_exportar,
    "assinatura.assinar": _executar_assinar,
    "assinatura.lote": _executar_assinar_lote,
    "correio.enviar": _executar_enviar,
}


# ------------------------------------------------- folha, notas e boletos


class MesPedido(BaseModel):
    mes: str = ""


class FichaPapel(BaseModel):
    id: int | None = None
    dados: dict = {}


@app.get("/api/financeiro/folha")
def folha_ler(mes: str = "") -> dict:
    mes = mes or escritorio.mes_de_hoje()
    return {
        "folha": estado.folha.do_mes(mes),
        "candidatos": estado.folha.pessoas(),
        "vinculos": [{"valor": k, "rotulo": v} for k, v in escritorio.VINCULOS.items()],
    }


@app.post("/api/financeiro/folha/montar")
def folha_montar(payload: MesPedido) -> dict:
    """Copia os valores dos cadastros para a folha do mes."""
    try:
        return {"folha": estado.folha.montar(payload.mes)}
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc


@app.post("/api/financeiro/folha/pagar")
def folha_pagar(payload: MesPedido) -> dict:
    """
    A folha do mes vira uma conta a pagar.

    Um lancamento so, e nao um por pessoa: o que sai do caixa e a folha
    inteira, e o detalhamento por pessoa ja esta gravado na folha. Espalhar
    seis lancamentos no extrato esconderia o numero que importa.
    """
    mes = payload.mes or escritorio.mes_de_hoje()
    da_folha = estado.folha.do_mes(mes)
    if not da_folha["quantos"]:
        raise HTTPException(status_code=400, detail="monte a folha deste mes antes")
    if da_folha["lancamento_id"]:
        raise HTTPException(status_code=400, detail="a folha deste mes ja esta lancada")

    ultimo = escritorio.ultimo_dia(mes)
    id_ = estado.financeiro.salvar({
        "tipo": "despesa",
        "descricao": f"Folha de {escritorio.mes_por_extenso(mes)}",
        "centavos": da_folha["total"],
        "categoria": "folha",
        "vencimento": ultimo,
        "observacao": f"{da_folha['quantos']} pessoa(s)",
    })
    estado.folha.ligar_ao_lancamento(mes, id_)
    return {"folha": estado.folha.do_mes(mes), "lancamento": estado.financeiro.obter(id_)}


@app.post("/api/financeiro/folha/recibos")
def folha_recibos(payload: MesPedido) -> dict:
    """
    Um recibo em PDF por pessoa, na pasta do programa.

    O recibo sai com o valor que esta gravado na folha daquele mes - nao com o
    salario de hoje. E essa a razao de a folha ser copia.
    """
    mes = payload.mes or escritorio.mes_de_hoje()
    da_folha = estado.folha.do_mes(mes)
    if not da_folha["quantos"]:
        raise HTTPException(status_code=400, detail="monte a folha deste mes antes")

    # Quem paga e quem esta configurado em Preferencias. Sem isso o recibo
    # sairia dizendo "recebi de este escritorio", que nao serve de recibo.
    pessoa = estado.prefs.dados.get("pessoa", {})
    escritorio = estado.prefs.dados.get("escritorio", {})
    quem_paga = str(escritorio.get("nome", "")).strip() or str(pessoa.get("nome", "")).strip()

    destino = RECIBOS_DIR / mes
    try:
        feitos = escritorio.gerar_recibos(
            da_folha, destino, escritorio.mes_por_extenso(mes), quem_paga)
    except Exception as exc:
        raise HTTPException(status_code=500, detail=f"nao consegui gerar: {exc}") from exc

    return {"recibos": feitos, "pasta": str(destino)}


@app.get("/api/financeiro/papeis")
def papeis_listar(tipo: str = "", mes: str = "") -> dict:
    mes = mes or escritorio.mes_de_hoje()
    return {
        "notas": estado.papeis.listar("nota", mes),
        "boletos": estado.papeis.listar("boleto"),
        "a_emitir": estado.papeis.a_emitir(mes),
    }


@app.post("/api/financeiro/papeis")
def papeis_salvar(payload: FichaPapel) -> dict:
    try:
        id_ = estado.papeis.salvar(payload.dados, payload.id)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    return {"id": id_, **papeis_listar(mes=str(payload.dados.get("data", ""))[:7])}


@app.delete("/api/financeiro/papeis/{id_}")
def papeis_apagar(id_: int) -> dict:
    if not estado.papeis.apagar(id_):
        raise HTTPException(status_code=404, detail="nao achei esse registro")
    return {"apagado": id_}


@app.post("/api/financeiro/papeis/{id_}/pago")
def papeis_pago(id_: int) -> dict:
    if not estado.papeis.marcar_pago(id_):
        raise HTTPException(status_code=404, detail="nao achei esse boleto")
    return {"pago": id_}


@app.get("/api/financeiro/fechamento")
def financeiro_fechamento(mes: str = "") -> dict:
    return escritorio.fechamento(estado.base, estado.folha, estado.papeis, mes)


@app.post("/api/financeiro/comprovantes")
async def financeiro_comprovante(lancamento_id: int = Form(...),
                                 arquivo: UploadFile = File(...)) -> dict:
    """
    Guarda o comprovante e o liga ao lancamento.

    O arquivo fica na pasta do programa, com o mes no caminho: procurar o
    comprovante do DAS de setembro daqui a um ano nao pode depender de lembrar
    o nome que o banco deu ao PDF.
    """
    lancamento = estado.financeiro.obter(lancamento_id)
    if not lancamento:
        raise HTTPException(status_code=404, detail="nao achei esse lancamento")

    nome = Path(arquivo.filename or "").name
    if not nome:
        raise HTTPException(status_code=400, detail="arquivo sem nome")

    conteudo = await arquivo.read()
    if len(conteudo) > MAX_UPLOAD_BYTES:
        raise HTTPException(status_code=400, detail="arquivo maior que 50 MB")

    quando = lancamento.get("liquidado_em") or lancamento.get("vencimento") or ""
    pasta = COMPROVANTES_DIR / (quando[:7] or escritorio.mes_de_hoje())
    pasta.mkdir(parents=True, exist_ok=True)

    alvo = pasta / nome
    n = 2
    while alvo.exists():
        alvo = pasta / f"{Path(nome).stem} ({n}){Path(nome).suffix}"
        n += 1
    alvo.write_bytes(conteudo)

    import hashlib
    sha = hashlib.sha1(conteudo).hexdigest()
    id_ = estado.financeiro.anexar(lancamento_id, nome, str(alvo), sha)
    return {"id": id_, "nome": nome, "caminho": str(alvo),
            "comprovantes": estado.financeiro.comprovantes(lancamento_id)}


@app.delete("/api/financeiro/comprovantes/{id_}")
def financeiro_tirar_comprovante(id_: int) -> dict:
    """
    Desliga o comprovante do lancamento.

    O arquivo continua na pasta: apagar o papel do disco porque a ligacao
    estava errada seria destruir o comprovante por causa de um erro de
    digitacao.
    """
    if not estado.financeiro.tirar_comprovante(id_):
        raise HTTPException(status_code=404, detail="nao achei esse comprovante")
    return {"tirado": id_}


@app.get("/api/financeiro/exportar")
def financeiro_exportar(mes: str = "") -> FileResponse:
    """O mes inteiro numa planilha, para quem faz a contabilidade."""
    mes = mes or escritorio.mes_de_hoje()
    destino = EXPORTACOES_DIR / f"financeiro-{mes}.xlsx"
    try:
        escritorio.exportar_mes(
            destino, mes,
            extrato=estado.financeiro.extrato(mes),
            folha=estado.folha.do_mes(mes),
            notas=estado.papeis.listar("nota", mes),
            boletos=estado.papeis.listar("boleto"),
        )
    except Exception as exc:
        raise HTTPException(status_code=500, detail=f"nao consegui exportar: {exc}") from exc
    return FileResponse(
        destino,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers=_anexo(destino.name),
    )


# ----------------------------------------------------------- preferencias


@app.get("/api/preferencias")
def preferencias_ler() -> dict:
    dados = estado.prefs.para_tela()
    dados["modelos"] = _modelos_disponiveis()
    dados["modelo_atual"] = estado.client.model
    dados["pasta_acervo"] = str(estado.pasta)
    dados["marca"] = estado.marca.info()
    return dados


@app.post("/api/preferencias")
def preferencias_gravar(payload: dict) -> dict:
    estado.prefs.atualizar(payload)

    # O que a preferencia muda de verdade, agora: modelo e ritmo.
    modelo = estado.prefs.dados.get("modelo")
    if modelo and modelo != estado.client.model:
        # A janela vem junto. Sem esta linha, trocar de modelo devolvia o
        # cliente ao padrao de 8.192 e o programa voltava a ler um quarto do
        # acervo - calado, e so na proxima pergunta.
        estado.client = LlamaClient(
            model=modelo, num_ctx=janela_para(estado.searcher.caracteres())
        )
    estado.devagar = bool(estado.prefs.dados.get("devagar"))
    estado.saber.ligada = bool(estado.prefs.dados.get("inteligencia", True))
    voz = (estado.prefs.dados.get("voz") or {}).get("modelo")
    if voz in transcricao_mod.MODELOS and voz != estado.transcritor.modelo:
        estado.transcritor.escolher(voz)

    return preferencias_ler()


def _modelos_disponiveis() -> list[str]:
    try:
        return estado.client.list_models()
    except Exception:
        return []


# --------------------------------------------------------------- organizador


class Escaneamento(BaseModel):
    raizes: list[str]
    excluir: list[str] = []
    profundidade: int | None = None


class Ajuste(BaseModel):
    """Correcao feita pelo usuario na tabela, antes de montar o plano."""

    arquivo: str
    cliente: str | None = None
    tipo: str | None = None


class PedidoPlano(BaseModel):
    destino: str
    padrao: str
    ajustes: list[Ajuste] = []
    incluir_baixa_confianca: bool = True
    apenas: list[str] = []          # caminhos marcados; vazio = todos
    operacao: str = "mover"         # "mover" ou "copiar" (o original fica)


class PedidoDesfazer(BaseModel):
    diario: str


class PedidoLeitura(BaseModel):
    apenas: list[str] = []          # caminhos escolhidos na varredura; vazio = todos


def _guardar_organizacao(plano, resultado, pedido: str = "") -> dict:
    """
    O desfecho de um mover, na forma que a tela mostra e desfaz.

    O destino passa a ser lido pelo Acervo: organizar e trazer os documentos
    para ca. Antes, o que foi movido sumia da tela Documentos - estava no
    disco, fora da unica pasta que o indice lia.
    """
    if resultado.movidos and plano.destino:
        estado.incluir_no_acervo(plano.destino)
    estado.ultima_organizacao = {
        "pedido": pedido,
        "movidos": resultado.movidos,
        "falhas": resultado.falhas,
        "diario": resultado.diario,
        "destino": plano.destino,
        "operacao": plano.operacao,
        "pastas": plano.resumo_por_pasta(),
    }
    return estado.ultima_organizacao


@app.get("/api/organizar/opcoes")
def organizar_opcoes() -> dict:
    return {
        "padroes": PADROES_SUGERIDOS,
        "tipos": [{"valor": k, "rotulo": v} for k, v in ROTULOS.items()],
        "destino_sugerido": str(Path.home() / "Documentos" / "Acervo PAULUS"),
        "diarios": listar_diarios(DIARIOS_DIR),
        # O botao diz "Mover agora" ou "Enviar para aprovacao" por isto. Vem
        # aqui, e nao de /api/preferencias, que pergunta ao Ollama a lista
        # de modelos e fazia a tela esperar.
        "mover_sem_pedir": estado.prefs.pode("organizar_mover"),
    }


@app.get("/api/pastas")
def navegar_pastas(caminho: str = "", arquivos: bool = False) -> dict:
    """Um nivel do seletor de pastas. Sem caminho, mostra unidades e atalhos.
    Com `arquivos`, lista tambem os documentos que o programa sabe ler."""
    dados = pastas.listar(caminho, sufixos=SUPPORTED_SUFFIXES if arquivos else None)
    dados["migalhas"] = pastas.migalhas(caminho)
    if not caminho:
        # As pastas que o proprio PAULUS le, no topo: e por elas que se
        # reorganiza o acervo que ja esta no sistema.
        padrao = Path(estado.pasta).resolve()
        dados["acervo"] = [
            {"nome": NOME_DA_PASTA_PADRAO if p.resolve() == padrao else (p.name or str(p)),
             "caminho": str(p), "tipo": "acervo"}
            for p in estado.pastas_do_acervo() if p.is_dir()
        ]
    return dados


@app.post("/api/organizar/cancelar")
def organizar_cancelar() -> dict:
    estado.cancelar.set()
    return {"cancelando": True}


@app.post("/api/organizar/escanear")
def organizar_escanear(payload: Escaneamento) -> dict:
    if not payload.raizes:
        raise HTTPException(status_code=400, detail="nenhuma pasta escolhida")

    varredura = escanear(
        [Path(r) for r in payload.raizes],
        excluir=payload.excluir,
        profundidade=payload.profundidade,
    )
    estado.encontrados = [
        {"path": a.path, "nome": a.nome, "pasta": a.pasta, "suffix": a.suffix, "mb": round(a.mb, 2)}
        for a in varredura.arquivos
    ]

    # Quanto ja esta em cache define o tempo real da leitura: documento
    # conhecido sai em milissegundos, novo custa ~8 a 25 s. Cada arquivo leva
    # a marca, para a tela refazer a conta quando a pessoa escolhe quais ler.
    lidos = _em_cache([a["path"] for a in estado.encontrados])
    for a in estado.encontrados:
        a["lido"] = a["path"] in lidos
    conhecidos = len(lidos)
    novos = varredura.total - conhecidos

    return {
        "total": varredura.total,
        "em_cache": conhecidos,
        "novos": novos,
        "minutos_min": round(novos * 8 / 60, 1),
        "minutos_max": round(novos * 25 / 60, 1),
        "pastas_visitadas": varredura.pastas_visitadas,
        "sem_permissao": len(varredura.sem_permissao),
        "grandes": len(varredura.ignorados_por_tamanho),
        "arquivos": estado.encontrados[:500],
    }


def _em_cache(caminhos: list[str]) -> set[str]:
    """Quais desses documentos ja foram lidos antes (mesmo conteudo)."""
    from classify import CacheClassificacao
    from extract import file_sha1

    cache = CacheClassificacao(CLASSIFICACAO_PATH)
    if not cache.dados:
        return set()

    lidos: set[str] = set()
    for caminho in caminhos:
        try:
            if file_sha1(Path(caminho)) in cache.dados:
                lidos.add(caminho)
        except OSError:
            continue
    return lidos


@app.post("/api/organizar/classificar")
def organizar_classificar(payload: PedidoLeitura | None = None) -> StreamingResponse:
    """
    Le o que a varredura achou - ou so o que a pessoa escolheu dela.

    A leitura em si mora em habilidades/classificar.py. Aqui fica so o que e do
    servidor: guardar o resultado para os passos seguintes do organizador.
    """
    caminhos = [a["path"] for a in estado.encontrados]
    if payload and payload.apenas:
        # So caminhos que a varredura achou: a lista vem da tela, e ler um
        # caminho qualquer que chegasse aqui seria abrir o disco inteiro.
        escolhidos = set(payload.apenas)
        caminhos = [c for c in caminhos if c in escolhidos]
    if not caminhos:
        raise HTTPException(status_code=400, detail="nada para ler - faca a varredura antes")

    habilidade = estado.registro.obter("classificar")
    if not habilidade or not habilidade.executavel:
        raise HTTPException(status_code=503, detail="a habilidade de classificar nao carregou")

    estado.cancelar.clear()

    def gerar() -> Iterator[str]:
        try:
            for tipo, dados in habilidade.executar(_contexto(), caminhos=caminhos):
                if tipo == "resultados":
                    # O organizador precisa dos objetos, nao do JSON: os passos
                    # seguintes montam o plano a partir deles.
                    from classify import Classificacao as _C

                    estado.classificacoes = {
                        d["arquivo"]: _C(**{
                            k: v for k, v in d.items()
                            if k in _C.__dataclass_fields__
                        })
                        for d in dados["documentos"]
                    }
                yield _sse(tipo, dados)
        except Exception as exc:
            yield _sse("erro", {"mensagem": str(exc)})

    return StreamingResponse(
        gerar(),
        media_type="text/event-stream",
        headers={"Cache-Control": "no-cache", "X-Accel-Buffering": "no"},
    )


def _com_ajustes(payload: PedidoPlano) -> list[Classificacao]:
    """Aplica as correcoes do usuario sobre a ultima classificacao."""
    ajustes = {a.arquivo: a for a in payload.ajustes}
    selecao = set(payload.apenas)
    saida: list[Classificacao] = []

    for caminho, resultado in estado.classificacoes.items():
        if selecao and caminho not in selecao:
            continue

        ajuste = ajustes.get(caminho)
        if ajuste:
            # Copia rasa para nao gravar a correcao no cache de classificacao:
            # ela vale para este plano, nao para o documento.
            resultado = Classificacao(**{**resultado.__dict__})
            if ajuste.cliente is not None:
                resultado.cliente = ajuste.cliente
            if ajuste.tipo is not None:
                resultado.tipo = ajuste.tipo

        saida.append(resultado)

    return saida


@app.post("/api/organizar/plano")
def organizar_plano(payload: PedidoPlano) -> dict:
    if not estado.classificacoes:
        raise HTTPException(status_code=400, detail="classifique os documentos antes")
    if not payload.destino.strip():
        raise HTTPException(status_code=400, detail="informe a pasta de destino")

    plano = montar_plano(
        _com_ajustes(payload),
        Path(payload.destino),
        payload.padrao,
        incluir_baixa_confianca=payload.incluir_baixa_confianca,
        operacao=payload.operacao,
    )
    return plano.to_dict()


@app.post("/api/organizar/aplicar")
def organizar_aplicar(payload: PedidoPlano) -> dict:
    """
    Mover arquivo em lote e acao com efeito no disco do cliente.

    Se a permissao "mover arquivos sem pedir" estiver desligada - e ela vem
    desligada -, isto nao move nada: monta o pedido e devolve para a fila de
    aprovacao. Quem move e o executor, depois do sim.
    """
    if not estado.classificacoes:
        raise HTTPException(status_code=400, detail="classifique os documentos antes")

    plano = montar_plano(
        _com_ajustes(payload),
        Path(payload.destino),
        payload.padrao,
        incluir_baixa_confianca=payload.incluir_baixa_confianca,
        operacao=payload.operacao,
    )
    if not plano.movimentos:
        raise HTTPException(status_code=400, detail="nada a fazer: os marcados já estão no lugar")

    if not estado.prefs.pode("organizar_mover"):
        pastas_alvo = plano.resumo_por_pasta()
        pedido = estado.fila.pedir(
            ("Copiar" if plano.operacao == "copiar" else "Mover") + f" {plano.total} arquivo(s) para "
            + (NOME_DA_PASTA_PADRAO if Path(plano.destino).resolve() == Path(estado.pasta).resolve()
               else Path(plano.destino).name or plano.destino),
            "organizar",
            resumo=(
                f"{plano.total} arquivo(s) em {len(pastas_alvo)} pasta(s), "
                f"no padrao {plano.padrao}. Nada e apagado nem sobrescrito."
            ),
            etiquetas=["da para desfazer"],
            acao="organizar.mover",
            dados=plano.to_dict(),
            reversivel=True,
        )
        return {
            "aguardando_aprovacao": True,
            "pedido": pedido.to_dict(),
            "total": plano.total,
            "pastas": pastas_alvo,
        }

    resultado = aplicar_plano(plano, DIARIOS_DIR)
    _guardar_organizacao(plano, resultado)
    estado.classificacoes = {}
    estado.encontrados = []
    estado.recarregar(force=True)
    return {
        "aguardando_aprovacao": False,
        "destino": plano.destino,
        "operacao": plano.operacao,
        "pastas": plano.resumo_por_pasta(),
        "movidos": resultado.movidos,
        "falhas": resultado.falhas,
        "diario": resultado.diario,
        "diarios": listar_diarios(DIARIOS_DIR),
    }


@app.get("/api/organizar/ultima")
def organizar_ultima() -> dict:
    """O desfecho da ultima organizacao - a tela volta a ele depois da fila."""
    return estado.ultima_organizacao or {}


@app.post("/api/organizar/desfazer")
def organizar_desfazer(payload: PedidoDesfazer) -> dict:
    caminho = Path(payload.diario)
    # So diarios criados por este programa - nao aceita caminho arbitrario.
    if caminho.parent.resolve() != DIARIOS_DIR.resolve() or not caminho.exists():
        raise HTTPException(status_code=400, detail="diario desconhecido")

    resultado = desfazer(caminho)
    return {
        "revertidos": resultado.movidos,
        "falhas": resultado.falhas,
        "diarios": listar_diarios(DIARIOS_DIR),
    }



# -------------------------------------------------- certificado e assinatura


class SenhaCertificado(BaseModel):
    senha: str = ""
    guardar: bool = False


class AjusteCofre(BaseModel):
    minutos: int | None = None
    pedir_confirmacao: bool | None = None
    mostrar_documento: bool | None = None
    lote_sem_confirmar: bool | None = None


class DesenhoSelo(BaseModel):
    campo: str = "desenho"          # "desenho" ou "imagem"
    dados: str = ""                 # PNG em data URL, vindo do canvas


class PedidoAssinatura(BaseModel):
    arquivo: str
    paginas: str = "ultima"
    intervalo: str = ""
    posicao: str = "rodape_direita"
    x: float | None = None
    y: float | None = None
    # Largura do selo em pontos do PDF, quando a pessoa redimensionou na tela.
    # Vazio, o tamanho de sempre; a altura acompanha na mesma proporcao.
    tamanho: float | None = None
    senha_pdf: str = ""
    motivo: str = ""
    guardar_biblioteca: bool = True
    manter_original: bool = True
    senha_certificado: str = ""


class ItemDoLote(BaseModel):
    """Um documento do lote: onde o selo entra nele, posto pela pessoa na tela."""
    arquivo: str
    paginas: str = "ultima"
    intervalo: str = ""
    posicao: str = "rodape_direita"
    x: float | None = None
    y: float | None = None
    tamanho: float | None = None


class PedidoLote(BaseModel):
    """Varios documentos, uma senha: o que vale para todos fica fora dos itens."""
    itens: list[ItemDoLote] = []
    senha_pdf: str = ""
    motivo: str = ""
    guardar_biblioteca: bool = True
    manter_original: bool = True
    senha_certificado: str = ""


class AssinadosParaSalvar(BaseModel):
    arquivos: list[str] = []
    caminho: str = ""               # o .zip, escolhido no "Salvar como"
    pasta: str = ""                 # a pasta, para salvar um a um


MAX_LOTE = 200


@app.get("/api/certificado")
def certificado_ler() -> dict:
    dados = estado.cofre.para_tela()
    dados["registro"] = estado.assinaturas.para_tela(10)
    dados["pode_assinar_sozinho"] = estado.prefs.pode("assinar")
    return dados


@app.get("/api/certificado/windows")
def certificado_windows() -> dict:
    """
    Os certificados que ja estao no Windows.

    Listar sem oferecer assinatura direta e proposital: a chave de um e-CPF
    instalado costuma vir marcada como nao exportavel, e assinar por ela
    exigiria falar com a CryptoAPI. A tela mostra o que existe e explica que
    aqui se usa o arquivo .pfx.
    """
    return {"certificados": certificado.listar_windows()}


class CertificadoDoWindows(BaseModel):
    impressao: str
    senha: str


@app.post("/api/certificado/windows/usar")
def certificado_windows_usar(payload: CertificadoDoWindows) -> dict:
    """
    Pre-seleciona um certificado instalado no Windows. Nada e exportado: o
    cofre guarda so a parte publica e a marca da senha do PAULUS que a
    pessoa criou agora. Ao assinar, o PAULUS confere essa senha e pede a
    conta ao Windows, com a chave que nunca sai de la - e, se o
    certificado tiver protecao forte, o Windows pede a senha dele tambem.
    """
    if len(payload.senha) < 4:
        raise HTTPException(status_code=400, detail="crie uma senha de pelo menos 4 caracteres")
    escolhido = next((c for c in certificado.listar_windows() if c["impressao"] == payload.impressao.strip().upper()), None)
    if not escolhido:
        raise HTTPException(status_code=400, detail="esse certificado não está mais no Windows")
    if escolhido["vencido"]:
        raise HTTPException(status_code=400, detail="esse certificado está vencido")

    publico = certificado.publico_do_windows(escolhido["impressao"])
    if publico.get("erro"):
        raise HTTPException(status_code=400, detail=publico["erro"])
    estado.cofre.usar_windows(publico, payload.senha)
    return {**estado.cofre.para_tela(), "aviso": "certificado de " + escolhido["titular"] + " pronto para assinar"}


@app.post("/api/certificado/windows/abrir")
def certificado_windows_abrir() -> dict:
    """Abre a janela "Certificados" do Windows."""
    if not certificado.abrir_janela_do_windows():
        raise HTTPException(status_code=400, detail="não consegui abrir a janela do Windows")
    return {"aberta": True}


@app.post("/api/certificado/arquivo")
async def certificado_arquivo(arquivo: UploadFile) -> dict:
    nome = arquivo.filename or "certificado.pfx"
    if Path(nome).suffix.lower() not in (".pfx", ".p12"):
        raise HTTPException(status_code=400, detail="o certificado A1 e um arquivo .pfx ou .p12")

    conteudo = await arquivo.read()
    if len(conteudo) > MAX_CERTIFICADO_BYTES:
        raise HTTPException(status_code=400, detail="arquivo grande demais para um certificado")

    estado.cofre.guardar_arquivo(nome, conteudo)
    estado.cofre.esquecer_senha()
    return estado.cofre.para_tela()


@app.post("/api/certificado/senha")
def certificado_senha(payload: SenhaCertificado) -> dict:
    """
    Confere a senha e, se a pessoa pediu, guarda protegida pela conta Windows.

    A senha so vira memoria depois de abrir o certificado de verdade: guardar
    uma senha errada faria o programa falhar mais tarde, longe daqui.
    """
    if not estado.cofre.instalado:
        raise HTTPException(status_code=400, detail="nenhum certificado instalado")

    lido = estado.cofre.abrir(payload.senha)
    if lido.erro:
        raise HTTPException(status_code=400, detail=lido.erro)

    estado.cofre.lembrar(payload.senha)
    if payload.guardar and not estado.cofre.proteger(payload.senha):
        return {
            **estado.cofre.para_tela(),
            "aviso": "não consegui guardar a senha nesta máquina - vou perguntar a cada assinatura",
        }
    return estado.cofre.para_tela()


@app.post("/api/certificado/esquecer")
def certificado_esquecer() -> dict:
    """Apaga a senha guardada, mas mantem o certificado instalado."""
    estado.cofre.dados["senha_protegida"] = ""
    estado.cofre.dados["guardar_senha"] = False
    estado.cofre.esquecer_senha()
    estado.cofre.salvar()
    return estado.cofre.para_tela()


@app.delete("/api/certificado")
def certificado_remover() -> dict:
    estado.cofre.remover()
    return estado.cofre.para_tela()


@app.post("/api/certificado/opcoes")
def certificado_opcoes(payload: AjusteCofre) -> dict:
    for campo, valor in payload.model_dump(exclude_none=True).items():
        if campo == "minutos":
            valor = max(0, min(int(valor), 240))
        estado.cofre.dados[campo] = valor
    estado.cofre.salvar()
    return estado.cofre.para_tela()


@app.post("/api/certificado/selo")
def certificado_selo(payload: dict) -> dict:
    return {"selo": estado.cofre.gravar_selo(payload)}


@app.post("/api/certificado/selo/imagem")
def certificado_selo_imagem(payload: DesenhoSelo) -> dict:
    """Recebe o PNG do desenho ou do logotipo, vindo como data URL."""
    import base64

    cabeca, _, corpo = payload.dados.partition(",")
    if "image/png" not in cabeca or not corpo:
        raise HTTPException(status_code=400, detail="esperava uma imagem PNG")
    try:
        conteudo = base64.b64decode(corpo)
    except Exception as exc:
        raise HTTPException(status_code=400, detail="não consegui ler a imagem") from exc
    if len(conteudo) > 4 * 1024 * 1024:
        raise HTTPException(status_code=400, detail="imagem grande demais")

    try:
        nome = estado.cofre.guardar_imagem(payload.campo, conteudo)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    return {"selo": estado.cofre.dados["selo"], "arquivo": nome}


@app.post("/api/certificado/selo/restaurar")
def certificado_selo_restaurar() -> dict:
    return {"selo": estado.cofre.restaurar_selo()}


@app.delete("/api/certificado/selo/{campo}")
def certificado_selo_limpar(campo: str) -> dict:
    if campo not in ("desenho", "imagem"):
        raise HTTPException(status_code=400, detail="campo desconhecido")
    alvo = estado.cofre.caminho_do_selo(campo)
    if alvo:
        try:
            alvo.unlink()
        except OSError:
            pass
    return {"selo": estado.cofre.gravar_selo({campo: ""})}


@app.get("/api/certificado/selo/{campo}.png")
def certificado_selo_png(campo: str) -> FileResponse:
    alvo = estado.cofre.caminho_do_selo(campo) if campo in ("desenho", "imagem") else None
    if not alvo:
        raise HTTPException(status_code=404, detail="não há imagem gravada")
    return FileResponse(alvo, media_type="image/png")


@app.post("/api/certificado/teste")
def certificado_teste() -> dict:
    """
    Cria um certificado autoassinado para experimentar a assinatura.

    Existe para dar para conhecer a tela sem um e-CPF na mao. O que sai daqui
    assina de verdade do ponto de vista criptografico e NAO tem validade
    juridica - a tela repete isso em cada passo, e o registro grava assim.
    """
    if estado.cofre.instalado:
        raise HTTPException(
            status_code=400,
            detail="já existe um certificado instalado - remova antes de criar um de teste",
        )
    senha = "paulus"
    alvo = certificado.gerar_de_teste(CERTIFICADO_DIR / "certificado.pfx", senha, "CERTIFICADO DE TESTE")
    estado.cofre.dados["arquivo"] = str(alvo)
    estado.cofre.salvar()
    estado.cofre.lembrar(senha)
    return {**estado.cofre.para_tela(), "senha_do_teste": senha}


# ------------------------------------------------------------- assinar PDF


@app.get("/api/assinar/documento")
def assinar_documento(arquivo: str) -> dict:
    """Quantas paginas o PDF tem, e se ja vem assinado."""
    doc = assinatura.ler(arquivo)
    if doc.erro:
        raise HTTPException(status_code=400, detail=doc.erro)

    cofre = estado.cofre.para_tela()
    return {
        "documento": doc.to_dict(),
        "certificado": cofre["certificado"],
        "instalado": cofre["instalado"],
        "origem": cofre["origem"],
        "selo": cofre["selo"],
        "posicoes": cofre["posicoes"],
        "escolhas": [{"valor": k, "rotulo": v} for k, v in assinatura.ESCOLHAS_PAGINA.items()],
        "pede_confirmacao": cofre["pedir_confirmacao"],
        "precisa_senha": not bool(estado.cofre.senha_agora()),
    }


@app.get("/api/assinar/pagina")
def assinar_pagina(arquivo: str, numero: int = 1, largura: int = 1000):
    """A pagina desenhada como PNG, para a tela mostrar o documento."""
    from fastapi.responses import Response

    try:
        png = assinatura.pagina_png(arquivo, numero, largura)
    except Exception as exc:
        raise HTTPException(status_code=400, detail=f"não consegui desenhar a página: {exc}") from exc
    return Response(png, media_type="image/png", headers={"Cache-Control": "no-cache"})


@app.get("/api/assinar/pdfs")
def assinar_pdfs() -> dict:
    """Os PDFs que o programa ja conhece, para escolher sem procurar no disco."""
    achados = []
    for doc in estado.searcher.documents:
        caminho = Path(doc.path)
        if caminho.suffix.lower() != ".pdf" or not caminho.exists():
            continue
        achados.append({
            "path": str(caminho),
            "nome": caminho.name,
            "mb": round(caminho.stat().st_size / (1024 * 1024), 2),
            "paginas": doc.pages,
        })
    return {"pdfs": sorted(achados, key=lambda a: a["nome"].lower())}


@app.post("/api/assinar")
def assinar_agora(payload: PedidoAssinatura) -> dict:
    """
    Assina, ou poe o pedido na fila.

    Assinatura tem efeito juridico, entao segue a mesma regra do resto: sem a
    permissao "assinar sem revisar" ligada - e ela vem desligada -, isto nao
    assina nada, monta o pedido e devolve para a fila de aprovacao.
    """
    doc = assinatura.ler(payload.arquivo)
    if doc.erro:
        raise HTTPException(status_code=400, detail=doc.erro)

    if not estado.cofre.instalado:
        raise HTTPException(status_code=400, detail="nenhum certificado instalado")

    senha = payload.senha_certificado or estado.cofre.senha_agora()
    if not senha:
        raise HTTPException(status_code=400, detail="preciso da senha do certificado")

    cert = estado.cofre.abrir(senha)
    if cert.erro:
        raise HTTPException(status_code=400, detail=cert.erro)
    if payload.senha_certificado:
        estado.cofre.lembrar(payload.senha_certificado)

    alvos = assinatura.paginas_alvo(payload.paginas, doc.paginas, payload.intervalo)
    if not alvos:
        raise HTTPException(status_code=400, detail="nenhuma página escolhida")

    dados = {
        **payload.model_dump(exclude={"senha_certificado"}),
        "paginas_alvo": alvos,
        "titular": cert.titular,
        "icp_brasil": cert.icp_brasil,
    }

    if not estado.prefs.pode("assinar"):
        etiquetas = ["não dá para desfazer"]
        if not cert.icp_brasil:
            etiquetas.append("certificado fora da ICP-Brasil")
        pedido = estado.fila.pedir(
            f"Assinar {Path(payload.arquivo).name}",
            "assinatura",
            resumo=assinatura.resumo_em_portugues(cert, doc, alvos, payload.senha_pdf),
            etiquetas=etiquetas,
            acao="assinatura.assinar",
            dados=dados,
            reversivel=False,
        )
        return {"aguardando_aprovacao": True, "pedido": pedido.to_dict()}

    resultado = _assinar_de_fato(dados, senha)
    if resultado.erro:
        raise HTTPException(status_code=400, detail=resultado.erro)
    return {"aguardando_aprovacao": False, **resultado.to_dict()}


@app.post("/api/assinar/lote")
def assinar_lote(payload: PedidoLote) -> dict:
    """
    Assina varios documentos com uma senha so - ou poe o lote na fila.

    Cada documento leva o selo onde a pessoa pos na tela e recebe a SUA
    assinatura: o lote e so a forma de pedir, nao junta os PDFs. A regra de
    alcada e a mesma do documento avulso; quando ela pede aprovacao, o lote
    vira UM pedido, aprovado (ou recusado) de uma vez, com os arquivos
    listados nele. Documento que nao da para assinar (nao abre, sem pagina
    escolhida) volta em `falhas` com o motivo, sem parar os outros.
    """
    if not payload.itens:
        raise HTTPException(status_code=400, detail="nenhum documento no lote")
    if len(payload.itens) > MAX_LOTE:
        raise HTTPException(status_code=400, detail=f"o lote aceita até {MAX_LOTE} documentos de uma vez")
    vistos = set()
    for item in payload.itens:
        chave = str(Path(item.arquivo)).lower()
        if chave in vistos:
            raise HTTPException(status_code=400, detail=f"{Path(item.arquivo).name} aparece duas vezes no lote")
        vistos.add(chave)

    if not estado.cofre.instalado:
        raise HTTPException(status_code=400, detail="nenhum certificado instalado")
    senha = payload.senha_certificado or estado.cofre.senha_agora()
    if not senha:
        raise HTTPException(status_code=400, detail="preciso da senha do certificado")
    cert = estado.cofre.abrir(senha)
    if cert.erro:
        raise HTTPException(status_code=400, detail=cert.erro)
    if payload.senha_certificado:
        estado.cofre.lembrar(payload.senha_certificado)

    comum = payload.model_dump(exclude={"itens", "senha_certificado"})
    prontos, recusados = [], []
    for item in payload.itens:
        doc = assinatura.ler(item.arquivo)
        nome = Path(item.arquivo).name
        if doc.erro:
            recusados.append({"arquivo": item.arquivo, "nome": nome, "motivo": doc.erro})
            continue
        alvos = assinatura.paginas_alvo(item.paginas, doc.paginas, item.intervalo)
        if not alvos:
            recusados.append({"arquivo": item.arquivo, "nome": nome, "motivo": "nenhuma página escolhida"})
            continue
        prontos.append({**comum, **item.model_dump(), "paginas_alvo": alvos,
                        "titular": cert.titular, "icp_brasil": cert.icp_brasil})
    if not prontos:
        raise HTTPException(status_code=400, detail="nenhum documento do lote pode ser assinado: " +
                            "; ".join(f"{r['nome']} ({r['motivo']})" for r in recusados[:3]))

    if not estado.prefs.pode("assinar"):
        etiquetas = ["não dá para desfazer", "lote"]
        if not cert.icp_brasil:
            etiquetas.append("certificado fora da ICP-Brasil")
        nomes = [Path(d["arquivo"]).name for d in prontos]
        lista = ", ".join(nomes[:5]) + (f" e mais {len(nomes) - 5}" if len(nomes) > 5 else "")
        quem = cert.titular or "seu certificado"
        quantos = f"{len(prontos)} documentos" if len(prontos) > 1 else "1 documento"
        resumo = (f"Vou assinar {quantos} com o {getattr(cert, 'tipo', '') or 'certificado'} de {quem}: {lista}. "
                  "Cada um recebe a sua assinatura, com o selo onde você o pôs na tela.")
        if payload.senha_pdf:
            resumo += " Os arquivos saem protegidos por senha."
        resumo += " Nada é enviado para a internet."
        pedido = estado.fila.pedir(
            f"Assinar {len(prontos)} documentos" if len(prontos) > 1 else f"Assinar {nomes[0]}",
            "assinatura",
            resumo=resumo,
            etiquetas=etiquetas,
            acao="assinatura.lote",
            # `arquivos` e o que Aprovacoes lista em "O que vai sair".
            dados={"itens": prontos, "arquivos": [d["arquivo"] for d in prontos]},
            reversivel=False,
        )
        return {"aguardando_aprovacao": True, "pedido": pedido.to_dict(), "falhas": recusados}

    feitos, falhas = _assinar_varios(prontos, senha)
    return {"aguardando_aprovacao": False, "feitos": feitos, "falhas": recusados + falhas}


def _assinar_varios(itens: list[dict], senha: str) -> tuple[list[dict], list[dict]]:
    """
    Assina um por um, sem parar no primeiro erro, e rele o Acervo UMA vez no
    fim - reler a cada documento deixaria um lote de vinte muito lento.
    """
    feitos, falhas = [], []
    for dados in itens:
        origem = dados.get("arquivo", "")
        try:
            resultado = _assinar_de_fato(dados, senha, recarregar=False)
        except Exception as exc:  # noqa: BLE001 - um PDF ruim nao derruba o lote
            resultado = assinatura.Resultado(erro=str(exc) or "erro ao assinar")
        if resultado.erro:
            falhas.append({"arquivo": origem, "nome": Path(origem).name, "motivo": resultado.erro})
        else:
            feitos.append({"origem": origem, **resultado.to_dict()})
    if feitos and any(d.get("guardar_biblioteca", True) for d in itens):
        estado.recarregar(force=True)
    return feitos, falhas


def _assinar_de_fato(dados: dict, senha: str, recarregar: bool = True) -> "assinatura.Resultado":
    """O trabalho em si, chamado direto ou depois do sim na fila."""
    origem = Path(dados["arquivo"])
    destino = _nome_do_assinado(origem, dados.get("guardar_biblioteca", True))

    ponto = None
    if dados.get("x") is not None and dados.get("y") is not None:
        ponto = (float(dados["x"]), float(dados["y"]))

    do_windows = estado.cofre.origem == "windows"
    if do_windows:
        # A senha do PAULUS e a porta: sem ela certa, o Windows nem e chamado.
        conferido = estado.cofre.abrir(senha)
        if conferido.erro:
            return assinatura.Resultado(erro=conferido.erro)
    resultado = assinatura.assinar(
        origem,
        destino,
        arquivo_pfx="" if do_windows else str(estado.cofre.arquivo),
        windows=estado.cofre.windows if do_windows else None,
        senha=senha,
        selo=estado.cofre.dados.get("selo") or {},
        escolha_paginas=dados.get("paginas", "ultima"),
        intervalo=dados.get("intervalo", ""),
        posicao=dados.get("posicao", "rodape_direita"),
        ponto=ponto,
        tamanho=float(dados["tamanho"]) if dados.get("tamanho") else None,
        senha_pdf=dados.get("senha_pdf", ""),
        motivo=dados.get("motivo", ""),
        imagem=estado.cofre.caminho_do_selo("desenho") or estado.cofre.caminho_do_selo("imagem"),
    )
    if resultado.erro:
        return resultado

    estado.assinaturas.anotar(resultado, origem)

    if not dados.get("manter_original", True):
        try:
            origem.unlink()
        except OSError:
            pass

    if recarregar and dados.get("guardar_biblioteca", True):
        estado.recarregar(force=True)

    return resultado


def _nome_do_assinado(origem: Path, na_biblioteca: bool) -> Path:
    """
    Onde o assinado e gravado, sem nunca passar por cima do original.

    Sobrescrever o documento de origem com a versao assinada seria perder o
    original de um jeito que nao se desfaz.
    """
    pasta = estado.pasta if na_biblioteca else origem.parent
    base = f"{origem.stem} - assinado"
    destino = pasta / f"{base}.pdf"
    conta = 2
    while destino.exists():
        destino = pasta / f"{base} ({conta}).pdf"
        conta += 1
    return destino


CONFORMIDADE_DIR = BASE_DIR / "data" / "conformidade"
APOIO_DIR = BASE_DIR / "data" / "apoio"


def _pdf_conhecido(caminho: str) -> Path:
    """
    O PDF pedido, se for um que o programa conhece: do Acervo, de uma pasta
    que o Acervo le, ou um relatorio de conformidade gerado aqui. Um caminho
    qualquer vindo da tela nao abre nem baixa arquivo arbitrario do disco.
    """
    try:
        alvo = Path(caminho).resolve()
    except (OSError, ValueError):
        raise HTTPException(status_code=404, detail="arquivo não encontrado") from None
    if alvo.suffix.lower() != ".pdf" or not alvo.is_file():
        raise HTTPException(status_code=404, detail="arquivo não encontrado")
    raizes = [Path(p).resolve() for p in estado.pastas_do_acervo()] + [CONFORMIDADE_DIR.resolve(), APOIO_DIR.resolve()]
    conhecidos = {Path(d.path).resolve() for d in estado.searcher.documents}
    if alvo in conhecidos or any(r == alvo.parent or r in alvo.parents for r in raizes):
        return alvo
    raise HTTPException(status_code=403, detail="esse arquivo não é do Acervo")


@app.post("/api/arquivos/abrir")
def arquivos_abrir(payload: dict) -> dict:
    """O PDF no programa padrao do Windows."""
    import os

    alvo = _pdf_conhecido(str(payload.get("caminho", "")))
    try:
        os.startfile(str(alvo))  # noqa: S606 - abre no programa do proprio Windows
    except (OSError, AttributeError) as exc:
        raise HTTPException(status_code=500, detail=f"não consegui abrir: {exc}") from exc
    return {"aberto": str(alvo)}


@app.get("/api/arquivos/baixar")
def arquivos_baixar(caminho: str) -> FileResponse:
    alvo = _pdf_conhecido(caminho)
    return FileResponse(alvo, media_type="application/pdf", filename=alvo.name)


@app.post("/api/arquivos/salvar")
def arquivos_salvar(payload: dict) -> dict:
    """
    Uma copia do PDF onde a pessoa escolheu no "Salvar como" do Windows.

    Na janela do programa, download de navegador nao chega a lugar nenhum:
    quem salva e o servidor. Vale para PDF do Acervo, relatorio gerado aqui ou
    PDF assinado por este programa - nunca um caminho qualquer.
    """
    import shutil

    origem = str(payload.get("caminho", ""))
    try:
        alvo = _pdf_conhecido(origem)
    except HTTPException:
        alvo = _assinados_daqui([origem])[0]
    destino = Path(str(payload.get("destino", "")))
    if not str(destino) or not destino.parent.is_dir():
        raise HTTPException(status_code=400, detail="escolha uma pasta que exista")
    if destino.suffix.lower() != ".pdf":
        destino = destino.with_name(destino.name + ".pdf")
    if destino.resolve() == alvo.resolve():
        return {"nome": destino.name, "pasta": str(destino.parent)}
    # O seletor do PAULUS nao pergunta "substituir?": o que ja esta na pasta
    # fica, e a copia ganha "(2)" no nome.
    if destino.exists():
        destino = acervo._nome_livre(destino.parent, destino.name, set())
    try:
        shutil.copy2(alvo, destino)
    except OSError as exc:
        raise HTTPException(status_code=500, detail=f"não consegui salvar: {exc}") from exc
    return {"nome": destino.name, "pasta": str(destino.parent)}


@app.get("/api/arquivos/pagina")
def arquivos_pagina(caminho: str, numero: int = 1, largura: int = 700):
    """Uma pagina desenhada - a previa do relatorio antes de baixar."""
    from fastapi.responses import Response

    alvo = _pdf_conhecido(caminho)
    return Response(documento.pagina_png(alvo.read_bytes(), numero, largura),
                    media_type="image/png", headers={"Cache-Control": "no-cache"})


@app.post("/api/assinaturas/conformidade")
def assinaturas_conformidade(payload: dict) -> dict:
    """
    O "PAVLVS - Certificado de conformidade" do PDF: conferencia feita agora,
    gravada em data/conformidade para ver, baixar ou mandar.
    """
    import conformidade

    alvo = _pdf_conhecido(str(payload.get("arquivo", "")))
    assinaturas = assinatura.verificar(alvo, str(payload.get("senha", "")))
    relatorio = conformidade.gerar(alvo, assinaturas, CONFORMIDADE_DIR)
    tom, frase = conformidade.veredito(assinaturas)
    return {"caminho": str(relatorio), "nome": relatorio.name, "tom": tom, "veredito": frase,
            "paginas": documento.paginas_de(relatorio.read_bytes())}


@app.get("/api/assinar/baixar")
def assinar_baixar(arquivo: str) -> FileResponse:
    """Entrega o PDF assinado para o navegador salvar."""
    alvo = Path(arquivo)
    if not alvo.exists() or alvo.suffix.lower() != ".pdf":
        raise HTTPException(status_code=404, detail="arquivo não encontrado")
    return FileResponse(alvo, media_type="application/pdf", filename=alvo.name)


def _assinados_daqui(arquivos: list[str]) -> list[Path]:
    """
    Os PDFs pedidos, se todos foram assinados por este programa.

    O .zip e a copia para uma pasta so aceitam o que esta no registro de
    assinaturas: um caminho qualquer vindo da tela nao vira um jeito de
    empacotar arquivo arbitrario do disco.
    """
    if not arquivos:
        raise HTTPException(status_code=400, detail="nenhum documento para salvar")
    conhecidos = set()
    for item in estado.assinaturas.itens:
        try:
            conhecidos.add(Path(item.get("destino") or "").resolve())
        except (OSError, ValueError):
            continue
    achados, vistos = [], set()
    for bruto in arquivos:
        try:
            alvo = Path(bruto).resolve()
        except (OSError, ValueError):
            raise HTTPException(status_code=404, detail="arquivo não encontrado") from None
        if alvo in vistos:
            continue
        if alvo not in conhecidos:
            raise HTTPException(status_code=403, detail=f"{alvo.name} não foi assinado por aqui")
        if not alvo.is_file():
            raise HTTPException(status_code=404, detail=f"{alvo.name} não está mais no disco")
        vistos.add(alvo)
        achados.append(alvo)
    return achados


def _nomes_no_zip(arquivos: list[Path]) -> list[str]:
    """Um nome por arquivo, sem colisao: o segundo "x.pdf" vira "x (2).pdf"."""
    usados, nomes = set(), []
    for alvo in arquivos:
        nome = alvo.name
        conta = 2
        while nome.lower() in usados:
            nome = f"{alvo.stem} ({conta}){alvo.suffix}"
            conta += 1
        usados.add(nome.lower())
        nomes.append(nome)
    return nomes


def _zip_dos_assinados(arquivos: list[Path]) -> bytes:
    import io
    import zipfile

    memoria = io.BytesIO()
    # PDF ja vem comprimido por dentro: guardar sem recomprimir e mais rapido
    # e o arquivo sai do mesmo tamanho.
    with zipfile.ZipFile(memoria, "w", zipfile.ZIP_STORED) as z:
        for alvo, nome in zip(arquivos, _nomes_no_zip(arquivos)):
            z.write(alvo, nome)
    return memoria.getvalue()


def _nome_do_zip() -> str:
    return f"PDFs assinados {datetime.now():%Y-%m-%d %H%M}.zip"


@app.get("/api/assinar/zip")
def assinar_zip_baixar(arquivo: list[str] = Query(default=[])) -> Response:
    """Os assinados num .zip, para o navegador baixar (um `arquivo=` por PDF)."""
    from urllib.parse import quote

    conteudo = _zip_dos_assinados(_assinados_daqui(arquivo))
    return Response(conteudo, media_type="application/zip",
                    headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(_nome_do_zip())}"})


@app.post("/api/assinar/zip")
def assinar_zip_gravar(payload: AssinadosParaSalvar) -> dict:
    """
    Os assinados num .zip gravado onde a pessoa escolheu no "Salvar como".
    Na janela do programa, download de navegador nao chega a lugar nenhum.
    """
    arquivos = _assinados_daqui(payload.arquivos)
    if not payload.caminho:
        raise HTTPException(status_code=400, detail="diga onde gravar o .zip")
    caminho = Path(payload.caminho)
    if caminho.suffix.lower() != ".zip":
        caminho = caminho.with_name(caminho.name + ".zip")
    # A pasta vem do seletor do PAULUS, que nao pergunta "substituir?": um
    # .zip que ja esta la nao e sobrescrito, o novo ganha "(2)" no nome.
    if caminho.exists():
        caminho = acervo._nome_livre(caminho.parent, caminho.name, set())
    try:
        caminho.parent.mkdir(parents=True, exist_ok=True)
        caminho.write_bytes(_zip_dos_assinados(arquivos))
    except OSError as exc:
        raise HTTPException(status_code=500, detail=f"não consegui gravar o .zip: {exc}") from exc
    return {"caminho": str(caminho), "nome": caminho.name, "quantos": len(arquivos)}


@app.post("/api/assinar/copiar")
def assinar_copiar(payload: AssinadosParaSalvar) -> dict:
    """
    Os assinados, um a um, copiados para a pasta escolhida. Nunca passa por
    cima de um arquivo que ja esta la: o repetido ganha "(2)" no nome.
    """
    arquivos = _assinados_daqui(payload.arquivos)
    pasta = Path(payload.pasta) if payload.pasta else None
    if not pasta or not pasta.is_dir():
        raise HTTPException(status_code=400, detail="escolha uma pasta que exista")
    copiados, usados = [], set()
    for alvo in arquivos:
        destino = pasta / alvo.name
        if destino.exists() or str(destino).lower() in usados:
            destino = acervo._nome_livre(pasta, alvo.name, usados)
        try:
            shutil.copy2(alvo, destino)
        except OSError as exc:
            raise HTTPException(status_code=500, detail=f"não consegui copiar {alvo.name}: {exc}") from exc
        usados.add(str(destino).lower())
        copiados.append(destino.name)
    return {"pasta": str(pasta), "copiados": copiados}


@app.get("/api/assinaturas")
def assinaturas_registro(codigo: str = "") -> dict:
    if codigo:
        achado = estado.assinaturas.procurar(codigo)
        if not achado:
            raise HTTPException(status_code=404, detail="não achei esse código no registro")
        return {"assinatura": achado}
    return estado.assinaturas.para_tela()


@app.get("/api/assinaturas/conferir")
def assinaturas_conferir(arquivo: str, senha: str = "") -> dict:
    """
    O que as assinaturas de um PDF dizem, lidas do proprio arquivo.

    Offline nao da para checar revogacao nem a cadeia ate a raiz da ICP-Brasil.
    O que da para afirmar - o documento foi mexido depois de assinado, ou nao -
    e o que vale a pena responder, e e o que vai aqui.
    """
    import conformidade

    assinaturas = assinatura.verificar(arquivo, senha)
    tom, frase = conformidade.veredito(assinaturas)
    alvo = Path(arquivo)
    return {
        "assinaturas": assinaturas,
        "conferencia_limitada": True,
        "tom": tom,
        "veredito": frase,
        "arquivo": {"nome": alvo.name, **(conformidade.resumo_do_arquivo(alvo) if alvo.is_file() else {})},
    }

# ----------------------------------------------------------------- e-mail


class EnderecoEmail(BaseModel):
    email: str = ""


class FichaConta(BaseModel):
    dados: dict = {}
    senha: str = ""


class SenhaConta(BaseModel):
    id: str = ""
    senha: str = ""


class PedidoEnvio(BaseModel):
    conta_id: str = ""
    para: str = ""
    cc: str = ""
    cco: str = ""
    assunto: str = ""
    corpo: str = ""
    corpo_html: str = ""            # o mesmo texto com a formatacao da faixa de edicao
    anexos: list[str] = []          # caminhos de arquivos da biblioteca
    responder_a: str = ""           # Message-ID, quando e resposta
    senha: str = ""


def _credencial_ou_http(conta) -> str:
    """
    A senha, ou o access token renovado, da conta.

    401 quando falta algo que so a pessoa resolve (senha, ou entrar de novo
    pelo login); 502 quando o problema e de rede, que passa sozinho.
    """
    try:
        senha = estado.contas.credencial(conta)
    except correio_oauth.ErroOAuth as exc:
        codigo = 401 if exc.precisa_entrar else 502
        raise HTTPException(status_code=codigo, detail=str(exc)) from exc
    if not senha:
        raise HTTPException(status_code=401, detail=f"preciso da senha de {conta.email}")
    return senha


def _conta_e_senha(id_: str = "") -> tuple:
    """A conta pedida (ou a em uso) com a senha (ou o token) disponivel."""
    conta = estado.contas.obter(id_) if id_ else estado.contas.em_uso
    if not conta:
        raise HTTPException(status_code=400, detail="nenhuma conta de e-mail conectada")
    return conta, _credencial_ou_http(conta)


def _info_oauth() -> dict:
    """Quais logins este PAULUS oferece: os provedores que vem registrados."""
    import oauth_app

    return {
        "google": {"configurado": oauth_app.configurado("google")},
        "microsoft": {"configurado": oauth_app.configurado("microsoft")},
        "prazo_s": correio_oauth.PRAZO_LOGIN,
    }


def _emails_dos_cadastros() -> dict:
    """{e-mail: nome} das fichas, para marcar quem e cliente na caixa."""
    achados = {}
    for ficha in estado.cadastros.listar():
        alvo = (ficha.get("email") or "").strip().lower()
        if alvo:
            achados[alvo] = ficha.get("nome", "")
    return achados


@app.get("/api/email/contas")
def email_contas() -> dict:
    dados = _contas_para_tela()
    dados["registro"] = estado.envios.para_tela(10)
    dados["pode_enviar_sozinho"] = estado.prefs.pode("enviar_mensagem")
    return dados


def _contas_para_tela() -> dict:
    dados = estado.contas.para_tela()
    oauth = _info_oauth()
    dados["oauth"] = oauth
    if oauth["microsoft"]["configurado"]:
        dados["aviso_microsoft"] = correio_contas.AVISO_MICROSOFT_OAUTH
    return dados


@app.post("/api/email/detectar")
def email_detectar(payload: EnderecoEmail) -> dict:
    """Acha os servidores do endereco, pela tabela ou sondando o dominio."""
    oauth = _info_oauth()
    return correio_contas.detectar(
        payload.email, oauth={p: oauth[p]["configurado"] for p in ("google", "microsoft")},
    )


@app.post("/api/email/testar")
def email_testar(payload: FichaConta) -> dict:
    """
    Prova a conta antes de guardar.

    Testa entrada e saida separadamente: da para ler e da para enviar sao
    coisas diferentes, e guardar uma conta que so le - sem dizer - seria
    descobrir o problema na hora de mandar o e-mail que importava.
    """
    campos = {k: v for k, v in payload.dados.items()
              if k in correio_contas.Conta.__dataclass_fields__}
    campos.pop("senha_protegida", None)
    conta = correio_contas.Conta(**{"id": "teste", **campos})

    senha = payload.senha
    if not senha and payload.dados.get("id"):
        guardada = estado.contas.obter(str(payload.dados["id"]))
        if guardada and guardada.por_login:
            # A conta de login se testa com ela mesma: servidores e token dela.
            return correio.testar(guardada, _credencial_ou_http(guardada))
        senha = estado.contas.senha(guardada) if guardada else ""
    if not senha:
        raise HTTPException(status_code=400, detail="informe a senha para testar")

    return correio.testar(conta, senha)


@app.post("/api/email/contas")
def email_salvar_conta(payload: FichaConta) -> dict:
    try:
        conta = estado.contas.salvar_conta(payload.dados, payload.senha)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    if payload.senha:
        estado.contas.lembrar(conta.id, payload.senha)
    return _contas_para_tela()


@app.post("/api/email/contas/senha")
def email_senha_conta(payload: SenhaConta) -> dict:
    """Confere a senha contra o servidor antes de aceitar como boa."""
    conta = estado.contas.obter(payload.id)
    if not conta:
        raise HTTPException(status_code=404, detail="conta não encontrada")
    if conta.por_login and not payload.senha:
        raise HTTPException(
            status_code=400,
            detail=f"esta conta entra pelo login {correio_oauth.rotulo(conta.autenticacao)} - use o botão de entrar",
        )
    if conta.por_login:
        # Senha de app numa conta que entrava pelo login: prova com a senha,
        # nos servidores dela, antes de trocar o jeito de entrar.
        import dataclasses

        conta = dataclasses.replace(conta, autenticacao="senha")

    prova = correio.testar(conta, payload.senha)
    if not prova["entrada"]:
        estado.contas.marcar_erro(conta, prova["erro_entrada"])
        raise HTTPException(status_code=400, detail=prova["erro_entrada"] or "o servidor recusou a senha")

    estado.contas.lembrar(conta.id, payload.senha)
    guardada = estado.contas.salvar_conta({"id": conta.id, "email": conta.email}, payload.senha)
    estado.contas.marcar_ok(guardada)
    return _contas_para_tela()


@app.post("/api/email/contas/{id_}/usar")
def email_usar_conta(id_: str) -> dict:
    if not estado.contas.usar(id_):
        raise HTTPException(status_code=404, detail="conta não encontrada")
    return _contas_para_tela()


@app.delete("/api/email/contas/{id_}")
def email_apagar_conta(id_: str) -> dict:
    if not estado.contas.apagar(id_):
        raise HTTPException(status_code=404, detail="conta não encontrada")
    return _contas_para_tela()


@app.post("/api/email/esquecer-senhas")
def email_esquecer_senhas() -> dict:
    quantas = estado.contas.esquecer_senhas()
    return {**_contas_para_tela(), "apagadas": quantas}


# ------------------------------------------- login com Google e Microsoft


class PedidoEntrada(BaseModel):
    provedor: str = ""
    email: str = ""          # login_hint, opcional


@app.get("/api/email/oauth")
def email_oauth_info() -> dict:
    return _info_oauth()


def _concluir_entrada(provedor: str, tokens: dict, email: str, nome: str) -> dict:
    """Guarda a conta que entrou e prova a conexao, dizendo o que funcionou."""
    conta = estado.contas.ligar_oauth(provedor, email, nome, tokens)
    prova = correio.testar(conta, tokens["access_token"])
    if prova["entrada"]:
        estado.contas.marcar_ok(conta)
    else:
        estado.contas.marcar_erro(conta, prova["erro_entrada"])
    return {"conta": conta.to_dict(em_uso=estado.contas.em_uso is conta), "prova": prova}


@app.post("/api/email/oauth/entrar")
def email_oauth_entrar(payload: PedidoEntrada) -> dict:
    """
    Comeca o login: sobe o servidor de volta e abre o navegador.

    Um login por vez: comecar outro cancela o que estava esperando.
    """
    provedor = payload.provedor.strip().lower()
    if provedor not in correio_oauth.PROVEDORES:
        raise HTTPException(status_code=400, detail="provedor desconhecido")
    anterior = estado.entrada_oauth
    if anterior and not anterior.terminou:
        anterior.cancelar()
    try:
        entrada = correio_oauth.Entrada(
            provedor, _credenciais_oauth(provedor), _concluir_entrada,
            login_hint=payload.email.strip() if "@" in payload.email else "",
        )
        estado.entrada_oauth = entrada
        return entrada.iniciar()
    except correio_oauth.ErroOAuth as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    except OSError as exc:
        raise HTTPException(status_code=500, detail=f"não consegui abrir a porta local para o login: {exc}") from exc


@app.get("/api/email/oauth/andamento")
def email_oauth_andamento() -> dict:
    entrada = estado.entrada_oauth
    if not entrada:
        return {"fase": "nenhum"}
    dados = entrada.andamento()
    if entrada.fase == "pronto":
        dados["contas"] = _contas_para_tela()
    return dados


@app.post("/api/email/oauth/cancelar")
def email_oauth_cancelar() -> dict:
    entrada = estado.entrada_oauth
    if entrada and not entrada.terminou:
        entrada.cancelar()
    return entrada.andamento() if entrada else {"fase": "nenhum"}


@app.post("/api/email/oauth/reabrir")
def email_oauth_reabrir() -> dict:
    """Abre de novo a pagina do login que esta esperando - o mesmo pedido, nao outro."""
    entrada = estado.entrada_oauth
    return entrada.reabrir() if entrada else {"fase": "nenhum"}


# --------------------------------------------------------- caixa de entrada


@app.get("/api/email/caixa")
def email_caixa(conta_id: str = "", filtro: str = "tudo", busca: str = "",
                limite: int = 25, antes_de: str = "") -> dict:
    """
    Os cabecalhos das mensagens - nao as mensagens.

    O corpo fica no servidor ate alguem abrir. E a promessa do rodape da tela,
    e tambem o que faz uma caixa com milhares de mensagens abrir em segundos.
    """
    conta, senha = _conta_e_senha(conta_id)
    try:
        dados = correio.listar(
            conta, senha, filtro=filtro, busca=busca,
            limite=min(max(int(limite), 1), 60), antes_de=antes_de,
            clientes=_emails_dos_cadastros(),
        )
    except correio.ErroCorreio as exc:
        estado.contas.marcar_erro(conta, str(exc))
        raise HTTPException(status_code=502, detail=str(exc)) from exc

    estado.contas.marcar_ok(conta)
    dados["conta"] = conta.to_dict(em_uso=True)
    dados["filtros"] = [{"valor": k, "rotulo": v} for k, v in correio.FILTROS.items()]
    return dados


@app.get("/api/email/mensagem")
def email_mensagem(uid: str, conta_id: str = "") -> dict:
    conta, senha = _conta_e_senha(conta_id)
    try:
        msg = correio.abrir(conta, senha, uid)
    except correio.ErroCorreio as exc:
        raise HTTPException(status_code=502, detail=str(exc)) from exc

    dados = msg.to_dict()
    dados["pode_rascunhar"] = conta.pode_rascunhar
    return dados


@app.get("/api/email/anexo")
def email_anexo(uid: str, nome: str, conta_id: str = ""):
    """Entrega o anexo ao navegador, sem gravar nada no disco."""
    from fastapi.responses import Response

    conta, senha = _conta_e_senha(conta_id)
    try:
        achado, dados = correio.baixar_anexo(conta, senha, uid, nome)
    except correio.ErroCorreio as exc:
        raise HTTPException(status_code=502, detail=str(exc)) from exc

    import mimetypes

    tipo, _ = mimetypes.guess_type(achado)
    return Response(dados, media_type=tipo or "application/octet-stream")


@app.post("/api/email/anexo/guardar")
def email_guardar_anexo(payload: dict) -> dict:
    """Traz o anexo para a biblioteca, sem passar por cima de nada."""
    conta, senha = _conta_e_senha(str(payload.get("conta_id", "")))
    uid = str(payload.get("uid", ""))
    nome = str(payload.get("nome", ""))

    try:
        achado, dados = correio.baixar_anexo(conta, senha, uid, nome)
    except correio.ErroCorreio as exc:
        raise HTTPException(status_code=502, detail=str(exc)) from exc

    limpo = Path(achado).name
    if Path(limpo).suffix.lower() not in SUPPORTED_SUFFIXES:
        raise HTTPException(
            status_code=400,
            detail=f"não sei ler {Path(limpo).suffix or 'esse tipo'} - guardo PDF, DOCX, TXT e MD",
        )

    estado.pasta.mkdir(parents=True, exist_ok=True)
    destino = estado.pasta / limpo
    conta_repetida = 2
    while destino.exists():
        destino = estado.pasta / f"{Path(limpo).stem} ({conta_repetida}){Path(limpo).suffix}"
        conta_repetida += 1

    destino.write_bytes(dados)
    return {"guardado": destino.name, "documentos": estado.recarregar(force=True)}


@app.post("/api/email/arquivar")
def email_arquivar(payload: dict) -> dict:
    """Uma (`uid`) ou varias (`uids`, a selecao da lista) numa conexao so."""
    conta, senha = _conta_e_senha(str(payload.get("conta_id", "")))
    if isinstance(payload.get("uids"), list):
        uids = correio._uids_validos(payload["uids"])
        if not uids:
            raise HTTPException(status_code=400, detail="nenhuma mensagem escolhida")
        try:
            feito = correio.arquivar_varias(conta, senha, uids)
        except correio.ErroCorreio as exc:
            raise HTTPException(status_code=502, detail=str(exc)) from exc
        so_lidas = feito["so_lidas"]
        return {
            **feito,
            "aviso": ("" if not so_lidas else
                      "seu servidor não tem pasta de arquivo - marquei como lidas" if not feito["arquivadas"] else
                      f"{so_lidas} não foram copiadas para o arquivo - ficaram na caixa, marcadas como lidas"),
        }
    uid = str(payload.get("uid", ""))
    try:
        movida = correio.arquivar(conta, senha, uid)
    except correio.ErroCorreio as exc:
        raise HTTPException(status_code=502, detail=str(exc)) from exc
    return {
        "arquivada": movida,
        "aviso": "" if movida else "seu servidor não tem pasta de arquivo - marquei só como lida",
    }


@app.post("/api/email/estrela")
def email_estrela(payload: dict) -> dict:
    """Poe ou tira a estrela (\\Flagged) nas mensagens escolhidas."""
    conta, senha = _conta_e_senha(str(payload.get("conta_id", "")))
    uids = correio._uids_validos(payload.get("uids") or [])
    if not uids:
        raise HTTPException(status_code=400, detail="nenhuma mensagem escolhida")
    try:
        feitas = correio.sinalizar(conta, senha, uids, bool(payload.get("sim", True)))
    except correio.ErroCorreio as exc:
        raise HTTPException(status_code=502, detail=str(exc)) from exc
    return {"feitas": feitas}


@app.post("/api/email/excluir")
def email_excluir(payload: dict) -> dict:
    """Manda as mensagens escolhidas para a lixeira do servidor (nao apaga de vez)."""
    conta, senha = _conta_e_senha(str(payload.get("conta_id", "")))
    uids = correio._uids_validos(payload.get("uids") or [])
    if not uids:
        raise HTTPException(status_code=400, detail="nenhuma mensagem escolhida")
    try:
        feito = correio.excluir_varias(conta, senha, uids)
    except correio.ErroCorreio as exc:
        raise HTTPException(status_code=502, detail=str(exc)) from exc
    if feito["sem_lixeira"]:
        raise HTTPException(status_code=409, detail="seu servidor não tem lixeira - não excluí nada, para não apagar de vez")
    falta = len(uids) - feito["excluidas"]
    return {**feito, "aviso": f"{falta} não foram para a lixeira - ficaram na caixa" if falta else ""}


@app.post("/api/email/marcar")
def email_marcar(payload: dict) -> dict:
    """Marca como lidas (ou nao lidas) as mensagens escolhidas na lista."""
    conta, senha = _conta_e_senha(str(payload.get("conta_id", "")))
    uids = correio._uids_validos(payload.get("uids") or [])
    if not uids:
        raise HTTPException(status_code=400, detail="nenhuma mensagem escolhida")
    lido = bool(payload.get("lido", True))
    try:
        n = correio.marcar_lidas(conta, senha, uids, lido)
    except correio.ErroCorreio as exc:
        raise HTTPException(status_code=502, detail=str(exc)) from exc
    return {"marcadas": n, "lido": lido, "uids": uids}


# O resumo da caixa pelo modelo: em segundo plano, um por vez, guardado pela
# chave das mensagens. A regra responde na hora, no mesmo pedido.
_resumos_da_caixa = correio.ResumosDaCaixa()


def _cabecalhos_do_pedido(payload: dict) -> list[dict]:
    campos = ("uid", "de_nome", "de_email", "de_cadastro", "assunto", "prazo", "lido", "respondido", "tem_anexo")
    return [{k: m.get(k) for k in campos} for m in (payload.get("mensagens") or [])[:200] if isinstance(m, dict)]


@app.post("/api/email/caixa/resumo")
def email_caixa_resumo(payload: dict) -> dict:
    """
    O resumo da lista que a tela mostra: a regra agora, o modelo depois.

    Recebe os cabecalhos que a lista ja tem - nenhum corpo sai do servidor
    de e-mail para isto. O modelo roda numa thread; a resposta volta na hora
    com "resumindo" e a tela pergunta de novo em GET.
    """
    mensagens = _cabecalhos_do_pedido(payload)
    chave = correio.chave_do_resumo(str(payload.get("conta_id", "")), mensagens)
    regra = correio.resumo_por_regra(mensagens)
    if payload.get("de_novo"):
        _resumos_da_caixa.esquecer(chave)
    modelo = _resumos_da_caixa.estado(chave)
    if modelo.get("estado") in ("pronto", "resumindo", "na_fila") or not mensagens:
        return {"chave": chave, "regra": regra, "modelo": modelo if mensagens else {"estado": "nenhum"}}
    disponivel, motivo = check_ollama(estado.client.model)
    if not disponivel:
        return {"chave": chave, "regra": regra, "modelo": {"estado": "indisponivel", "motivo": motivo}}
    return {"chave": chave, "regra": regra, "modelo": _resumos_da_caixa.pedir(chave, mensagens, estado.client)}


_contextos_de_email = correio.ResumosDaCaixa(fazer=correio.contexto_da_mensagem, limite=60)


@app.post("/api/email/contexto")
def email_contexto(payload: dict) -> dict:
    """
    O "Contexto IA" da mensagem aberta, pelo modelo desta maquina.

    Recebe o texto que a tela ja tem (a pessoa abriu a mensagem). Roda numa
    thread, como o resumo da caixa: volta "resumindo" e a tela pergunta de
    novo em GET. Guardado por conta e UID - abrir de novo nao roda outra vez.
    """
    uid = str(payload.get("uid", "")).strip()
    if not uid:
        raise HTTPException(status_code=400, detail="qual mensagem?")
    chave = f"{payload.get('conta_id', '')}:{uid}"
    if payload.get("de_novo"):
        _contextos_de_email.esquecer(chave)
    feito = _contextos_de_email.estado(chave)
    if feito.get("estado") in ("pronto", "resumindo", "na_fila"):
        return {"chave": chave, "modelo": feito}
    disponivel, motivo = check_ollama(estado.client.model)
    if not disponivel:
        return {"chave": chave, "modelo": {"estado": "indisponivel", "motivo": motivo}}
    item = {"de": str(payload.get("de", ""))[:200], "assunto": str(payload.get("assunto", ""))[:300],
            "corpo": str(payload.get("corpo", ""))}
    return {"chave": chave, "modelo": _contextos_de_email.pedir(chave, [item], estado.client)}


@app.get("/api/email/contexto")
def email_contexto_estado(chave: str) -> dict:
    return {"chave": chave, "modelo": _contextos_de_email.estado(chave)}


@app.get("/api/email/caixa/resumo")
def email_caixa_resumo_estado(chave: str) -> dict:
    return {"chave": chave, "modelo": _resumos_da_caixa.estado(chave)}


@app.post("/api/email/reescrever")
def email_reescrever(payload: dict) -> dict:
    """
    O "Pedir aqui" do Escrever: o modelo reescreve o corpo do e-mail.

    Volta como sugestao - a tela troca o texto e oferece manter ou
    descartar. Leva cerca de um minuto nesta maquina.
    """
    pedido = str(payload.get("pedido", "")).strip()
    if not pedido:
        raise HTTPException(status_code=400, detail="diga o que mudar no e-mail")
    disponivel, motivo = check_ollama(estado.client.model)
    if not disponivel:
        raise HTTPException(status_code=503, detail=motivo)
    try:
        texto = correio.reescrever_email(estado.client, pedido,
                                         str(payload.get("assunto", "")), str(payload.get("corpo", "")))
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    except OllamaError as exc:
        raise HTTPException(status_code=503, detail=str(exc)) from exc
    return {"sugestao": texto, "pedido": pedido}


# ------------------------------------------------------ apoiar o projeto


@app.post("/api/apoio/pix")
def apoio_pix(payload: dict) -> dict:
    """O Pix de apoio: o site (Worker) cria no Mercado Pago e devolve o QR."""
    try:
        return apoio.criar_pix(float(payload.get("valor") or 0), str(payload.get("email", "")))
    except (ValueError, TypeError) as exc:
        raise HTTPException(status_code=400, detail="valor inválido") from exc
    except apoio.ErroDeApoio as exc:
        raise HTTPException(status_code=502, detail=str(exc)) from exc


@app.post("/api/apoio/pix/recuperar")
def apoio_pix_recuperar(payload: dict) -> dict:
    """Os Pix pagos que o historico desta maquina perdeu (versao antiga)."""
    try:
        return apoio.recuperar_pix([str(e) for e in payload.get("emails") or []], [str(d) for d in payload.get("datas") or []])
    except apoio.ErroDeApoio as exc:
        raise HTTPException(status_code=502, detail=str(exc)) from exc


@app.get("/api/apoio/pix/{id_}")
def apoio_pix_situacao(id_: str) -> dict:
    try:
        return apoio.situacao_do_pix(id_)
    except apoio.ErroDeApoio as exc:
        raise HTTPException(status_code=502, detail=str(exc)) from exc


@app.get("/api/apoio/assinatura/{id_}")
def apoio_assinatura_situacao(id_: str) -> dict:
    try:
        return apoio.situacao_da_assinatura(id_)
    except apoio.ErroDeApoio as exc:
        raise HTTPException(status_code=502, detail=str(exc)) from exc


@app.post("/api/apoio/extrato")
def apoio_extrato(payload: dict) -> dict:
    """
    O "PAULUS - Extrato de apoio" em PDF: os Pix confirmados nesta maquina
    (a tela manda) e as cobrancas do cartao, consultadas no Mercado Pago
    agora. Gravado em data/apoio, para baixar pela janela de sempre.
    """
    import extrato_apoio

    pix = [p for p in (payload.get("pix") or []) if isinstance(p, dict)][:500]
    cobrancas: list = []
    assinatura = None
    # A assinatura de agora e as ja interrompidas: as cobrancas de todas.
    for i, ass in enumerate([a for a in (payload.get("assinaturas") or []) if isinstance(a, dict)][:20]):
        if not (ass.get("id") and ass.get("chave")):
            continue
        try:
            cobrancas += apoio.pagamentos_da_assinatura(str(ass["id"]), str(ass["chave"])).get("pagamentos", [])
            if i == 0 and ass.get("atual"):
                situacao = apoio.situacao_da_assinatura(str(ass["id"]))
                assinatura = {"situacao": situacao.get("situacao", ""), "valor": ass.get("valor")}
        except apoio.ErroDeApoio as exc:
            raise HTTPException(status_code=502, detail=f"não consegui as cobranças do cartão: {exc}") from exc
    try:
        caminho = extrato_apoio.gerar(APOIO_DIR, nome=str(payload.get("nome", ""))[:80], email=str(payload.get("email", ""))[:120],
                                      pix=pix, cobrancas=cobrancas, assinatura=assinatura)
    except (OSError, ValueError) as exc:
        raise HTTPException(status_code=500, detail=f"não consegui montar o extrato: {exc}") from exc
    return {"caminho": str(caminho), "nome": caminho.name}


@app.post("/api/apoio/assinatura/{id_}/valor")
def apoio_assinatura_valor(id_: str, payload: dict) -> dict:
    try:
        return apoio.mudar_valor(id_, str(payload.get("chave", "")), float(payload.get("valor") or 0))
    except (ValueError, TypeError) as exc:
        raise HTTPException(status_code=400, detail="valor inválido") from exc
    except apoio.ErroDeApoio as exc:
        raise HTTPException(status_code=502, detail=str(exc)) from exc


@app.post("/api/apoio/assinatura/{id_}/interromper")
def apoio_assinatura_interromper(id_: str, payload: dict) -> dict:
    try:
        return apoio.interromper(id_, str(payload.get("chave", "")))
    except apoio.ErroDeApoio as exc:
        raise HTTPException(status_code=502, detail=str(exc)) from exc


@app.post("/api/apoio/assinatura")
def apoio_assinatura(payload: dict) -> dict:
    """A assinatura no cartao: devolve o link da pagina do Mercado Pago."""
    try:
        return apoio.criar_assinatura(float(payload.get("valor") or 0), str(payload.get("email", "")))
    except (ValueError, TypeError) as exc:
        raise HTTPException(status_code=400, detail="valor inválido") from exc
    except apoio.ErroDeApoio as exc:
        raise HTTPException(status_code=502, detail=str(exc)) from exc


@app.post("/api/email/traduzir")
def email_traduzir(payload: dict) -> dict:
    """
    A mensagem aberta em portugues, pelo tradutor desta maquina (src/traducao.py):
    o texto e, quando ha, o HTML com as tags intactas. Menos de um segundo.
    """
    texto = str(payload.get("texto", ""))
    html_email = str(payload.get("html", ""))
    lingua = traducao.adivinhar_lingua(texto) or str(payload.get("lingua", ""))
    if lingua not in traducao.PACOTES:
        raise HTTPException(status_code=400, detail="não reconheci a língua da mensagem")
    if not traducao.disponivel(lingua):
        raise HTTPException(status_code=503, detail="o tradutor desta máquina não está instalado (python src/traducao.py baixar)")
    try:
        return {
            "lingua": lingua,
            "texto": traducao.traduzir_texto(texto, lingua) if texto.strip() else "",
            "html": traducao.traduzir_html(html_email, lingua) if html_email.strip() else "",
        }
    except (ValueError, RuntimeError) as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc


@app.post("/api/email/anexos/conferir")
def email_anexos_conferir(payload: dict) -> dict:
    """Nome e tamanho dos arquivos escolhidos para anexar - e se passam do limite."""
    achados = []
    for bruto in (payload.get("caminhos") or [])[:40]:
        alvo = Path(str(bruto))
        existe = alvo.is_file()
        tamanho = alvo.stat().st_size if existe else 0
        achados.append({
            "path": str(alvo), "nome": alvo.name, "existe": existe,
            "mb": round(tamanho / (1024 * 1024), 2),
            "grande": tamanho > correio.MAX_ANEXO_BYTES,
        })
    return {"anexos": achados, "limite_mb": correio.MAX_ANEXO_BYTES // (1024 * 1024)}


@app.post("/api/email/rascunho")
def email_rascunho(payload: dict) -> dict:
    """
    Escreve um rascunho de resposta, quando alguem pede.

    Sob demanda de proposito: rodar o modelo em cada mensagem que chega custaria
    perto de um minuto por e-mail nesta maquina, e a caixa de entrada demoraria
    uma hora para abrir.
    """
    conta, senha = _conta_e_senha(str(payload.get("conta_id", "")))
    if not conta.pode_rascunhar:
        raise HTTPException(status_code=403, detail="esta conta não permite que eu escreva rascunhos")

    disponivel, motivo = check_ollama(estado.client.model)
    if not disponivel:
        raise HTTPException(status_code=503, detail=motivo)

    try:
        msg = correio.abrir(conta, senha, str(payload.get("uid", "")), marcar_lido=False)
    except correio.ErroCorreio as exc:
        raise HTTPException(status_code=502, detail=str(exc)) from exc

    quem = estado.prefs.dados.get("pessoa", {}).get("nome", "") or conta.nome
    try:
        texto = correio.sugerir_resposta(estado.client, msg, quem)
    except OllamaError as exc:
        raise HTTPException(status_code=503, detail=str(exc)) from exc

    return {
        "rascunho": texto,
        "para": [msg.de_email],
        "assunto": msg.assunto if msg.assunto.lower().startswith("re:") else f"Re: {msg.assunto}",
        "responder_a": "",
        "sobre": msg.assunto,
    }


# ---------------------------------------------------------------- enviar


def _montar_do_pedido(payload: PedidoEnvio):
    conta = estado.contas.obter(payload.conta_id) if payload.conta_id else estado.contas.em_uso
    if not conta:
        raise HTTPException(status_code=400, detail="nenhuma conta de e-mail conectada")

    para = correio.enderecos(payload.para)
    if not para:
        raise HTTPException(status_code=400, detail="informe pelo menos um destinatário")

    anexos = [Path(a) for a in payload.anexos]
    if anexos and not conta.pode_anexar:
        raise HTTPException(status_code=403, detail="esta conta não permite que eu anexe arquivos")
    for alvo in anexos:
        if not alvo.exists():
            raise HTTPException(status_code=400, detail=f"não achei o anexo {alvo.name}")

    try:
        msg = correio.montar_email(
            conta,
            para=para,
            assunto=payload.assunto,
            corpo=payload.corpo,
            corpo_html=payload.corpo_html,
            cc=correio.enderecos(payload.cc),
            anexos=anexos,
            responder_a=payload.responder_a,
        )
    except correio.ErroCorreio as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc

    return conta, msg, para


@app.post("/api/email/previa")
def email_previa(payload: PedidoEnvio) -> dict:
    """
    Como a mensagem vai chegar - sem nada ter saido.

    Montar e enviar sao passos separados justamente para isto existir.
    """
    conta, msg, para = _montar_do_pedido(payload)
    corpo = msg.get_body(preferencelist=("plain",))
    return {
        "de": str(msg.get("From", "")),
        "para": para,
        "cc": correio.enderecos(payload.cc),
        "cco": correio.enderecos(payload.cco),
        "assunto": str(msg.get("Subject", "")) or "(sem assunto)",
        "corpo": corpo.get_content().rstrip() if corpo else "",
        "html": correio.html_da_mensagem(msg),
        "anexos": [
            {"nome": p.get_filename(), "kb": round(len(p.get_payload(decode=True) or b"") / 1024, 1)}
            for p in msg.iter_attachments()
        ],
        "resumo": correio.resumo_do_envio(
            conta, para, [p.get_filename() for p in msg.iter_attachments()],
            not conta.pode_enviar_sem_confirmar,
        ),
    }


@app.post("/api/email/enviar")
def email_enviar(payload: PedidoEnvio) -> dict:
    """
    Envia, ou poe o pedido na fila.

    E-mail que sai nao volta. Sem a permissao "enviar sem confirmar" - que vem
    desligada, no programa e na conta -, isto nao manda nada: monta o pedido e
    devolve para a fila de aprovacao.
    """
    conta, msg, para = _montar_do_pedido(payload)

    if payload.senha and not conta.por_login:
        estado.contas.lembrar(conta.id, payload.senha)
    if not estado.contas.tem_credencial(conta):
        if conta.por_login:
            raise HTTPException(
                status_code=401,
                detail=f"é preciso entrar de novo com {correio_oauth.rotulo(conta.autenticacao)} em {conta.email}",
            )
        raise HTTPException(status_code=401, detail=f"preciso da senha de {conta.email}")

    anexos = [Path(a).name for a in payload.anexos]
    livre = estado.prefs.pode("enviar_mensagem") and conta.pode_enviar_sem_confirmar

    if not livre:
        pedido = estado.fila.pedir(
            f"Enviar e-mail para {', '.join(para[:2])}" + ("…" if len(para) > 2 else ""),
            "email",
            resumo=correio.resumo_do_envio(conta, para, anexos, False),
            etiquetas=["não dá para desfazer"] + (["com anexo"] if anexos else []),
            acao="correio.enviar",
            dados=payload.model_dump(exclude={"senha"}),
            reversivel=False,
        )
        return {"aguardando_aprovacao": True, "pedido": pedido.to_dict()}

    resultado = _enviar_de_fato(payload.model_dump(exclude={"senha"}))
    return {"aguardando_aprovacao": False, **resultado}


def _enviar_de_fato(dados: dict) -> dict:
    """O envio em si, chamado direto ou depois do sim na fila."""
    payload = PedidoEnvio(**{k: v for k, v in dados.items() if k in PedidoEnvio.model_fields})
    conta, msg, _ = _montar_do_pedido(payload)

    try:
        senha = estado.contas.credencial(conta)
    except correio_oauth.ErroOAuth as exc:
        raise RuntimeError(str(exc)) from exc
    if not senha:
        raise RuntimeError(f"a senha de {conta.email} não está mais disponível - entre na conta de novo")

    try:
        resultado = correio.enviar(conta, senha, msg, correio.enderecos(payload.cco))
    except correio.ErroCorreio as exc:
        raise RuntimeError(str(exc)) from exc

    estado.contas.marcar_ok(conta)
    return estado.envios.anotar(conta.email, resultado)


@app.get("/api/email/envios")
def email_envios() -> dict:
    return estado.envios.para_tela()


@app.get("/api/email/anexaveis")
def email_anexaveis() -> dict:
    """Arquivos da biblioteca que dao para anexar."""
    achados = []
    for doc in estado.searcher.documents:
        caminho = Path(doc.path)
        if not caminho.exists():
            continue
        achados.append({
            "path": str(caminho),
            "nome": caminho.name,
            "mb": round(caminho.stat().st_size / (1024 * 1024), 2),
        })
    return {"arquivos": sorted(achados, key=lambda a: a["nome"].lower())}

# ------------------------------------------------- documentos e planilha


class NovoDocumento(BaseModel):
    titulo: str = ""
    tipo: str = "texto"
    corpo: str = ""
    cadastro_id: int | None = None


class GravarDocumento(BaseModel):
    corpo: str = ""
    titulo: str = ""
    nota: str = ""


class PedidoAoAssistente(BaseModel):
    pedido: str = ""
    trecho: str = ""            # o que estava selecionado, quando havia
    selecao: str = ""           # celulas, na planilha
    # O editor aberto ao lado de uma conversa: o pedido e o que foi feito
    # ficam nela, como qualquer outra coisa dita ali.
    trabalho_id: str = ""


class CelulaPlanilha(BaseModel):
    aba: int = 0
    ref: str = ""
    dados: dict = {}


def _documento_ou_404(id_: int) -> dict:
    item = estado.documentos.obter(id_)
    if not item:
        raise HTTPException(status_code=404, detail="documento não encontrado")
    return item


@app.get("/api/documentos")
def documentos_listar(tipo: str = "") -> dict:
    return {
        "documentos": estado.documentos.listar(tipo),
        "contagem": estado.documentos.contagem(),
        "clientes": [{"id": f["id"], "nome": f["nome"]} for f in estado.cadastros.listar()],
    }


@app.post("/api/documentos")
def documentos_criar(payload: NovoDocumento) -> dict:
    corpo = payload.corpo
    if not corpo and payload.tipo == "planilha":
        corpo = planilha.para_json([planilha.Aba()])
    id_ = estado.documentos.criar(payload.titulo, payload.tipo, corpo, payload.cadastro_id)
    return estado.documentos.obter(id_) or {}


class ImportarParaEditar(BaseModel):
    caminho: str = ""
    nome: str = ""
    # Vindo de uma conversa: o rascunho fica ligado a ela, e abrir para
    # editar de novo volta ao mesmo rascunho em vez de criar outra copia.
    trabalho_id: str = ""


@app.post("/api/documentos/importar")
def documentos_importar(payload: ImportarParaEditar) -> dict:
    """
    Traz um documento da biblioteca para o editor, como rascunho.

    O arquivo original NAO e tocado. O que nasce aqui e uma copia editavel: a
    pessoa pediu para abrir para editar, e editar o .docx de origem no lugar
    seria mexer no documento que ja foi assinado, enviado ou protocolado.
    """
    alvo = Path(payload.caminho) if payload.caminho else None
    if not alvo or not alvo.exists():
        doc = next((d for d in estado.searcher.documents if d.name == payload.nome), None)
        if not doc:
            raise HTTPException(status_code=404, detail="nao achei esse documento")
        alvo = Path(doc.path)
    if not alvo.exists():
        raise HTTPException(status_code=404, detail="o arquivo saiu do lugar")

    lido = next((d for d in estado.searcher.documents if d.path == str(alvo)), None)
    if lido is None:
        raise HTTPException(status_code=400, detail="esse arquivo ainda nao foi lido")

    trabalho = estado.trabalhos.obter(payload.trabalho_id) if payload.trabalho_id else None
    if trabalho:
        for m in reversed(trabalho.mensagens):
            feito = m.feito or {}
            if feito.get("tipo") == "editar" and feito.get("nome") == lido.name:
                rascunho = estado.documentos.obter(int(feito.get("id") or 0))
                if rascunho:
                    return {"id": rascunho["id"], "titulo": rascunho["titulo"], "de": lido.name,
                            "paragrafos": 0, "reaberto": True}

    # Paragrafo por paragrafo: o texto extraido vem com quebras de linha, e
    # jogar tudo num <p> so daria um bloco unico impossivel de editar.
    paragrafos = [p.strip() for p in lido.text.splitlines() if p.strip()]
    corpo = "".join(f"<p>{_escapar(p)}</p>" for p in paragrafos) or "<p><br></p>"

    id_ = estado.documentos.criar(Path(lido.name).stem, "texto", corpo, None)
    if trabalho:
        trabalho.dizer("paulus", f"Abri “{lido.name}” para editar.",
                       feito={"tipo": "editar", "id": id_, "nome": lido.name, "onde": "editor"})
        estado.trabalhos.salvar(trabalho)
    return {"id": id_, "titulo": Path(lido.name).stem, "de": lido.name,
            "paragrafos": len(paragrafos), "reaberto": False}


def _escapar(texto: str) -> str:
    return (texto.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;"))


# A rota fixa vem ANTES da rota com parametro: o FastAPI casa na ordem de
# declaracao, e /api/documentos/{id_} engoliria /api/documentos/modelos.
@app.get("/api/documentos/modelos")
def documentos_modelos() -> dict:
    """
    As clausulas do escritorio e os codigos de lei disponiveis.

    Os codigos ficaram de fora ate a Etapa 11 porque o programa nao tinha o
    texto das leis, e pedir o artigo ao modelo produziria numero plausivel e
    errado dentro de um contrato. Com o texto oficial no disco, citar passa a
    ser consulta - e quando um codigo nao esta instalado, a tela diz isso em
    vez de chutar.
    """
    instalados = estado.leis.instalados()
    return {
        "clausulas": estado.clausulas.listar(),
        "codigos": instalados,
        "tem_codigo": any(c["instalado"] for c in instalados),
        "sem_codigo": {
            "titulo": "Nenhum código instalado",
            "porque": (
                "Sem o texto oficial no disco eu não cito artigo: o número sairia de um "
                "modelo pequeno, plausível e possivelmente errado dentro de um contrato."
            ),
            "onde": "Configurações · Códigos de lei",
        },
    }


@app.post("/api/documentos/modelos")
def documentos_gravar_modelo(payload: dict) -> dict:
    try:
        estado.clausulas.gravar(
            str(payload.get("titulo", "")), str(payload.get("texto", "")),
            str(payload.get("id", "")),
        )
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    return {"clausulas": estado.clausulas.listar()}


@app.delete("/api/documentos/modelos/{id_}")
def documentos_apagar_modelo(id_: str) -> dict:
    estado.clausulas.apagar(id_)
    return {"clausulas": estado.clausulas.listar()}


@app.get("/api/documentos/{id_}")
def documentos_obter(id_: int) -> dict:
    item = _documento_ou_404(id_)
    if item["tipo"] == "texto":
        blocos = documento.ler_html(item["corpo"])
        item["contagem"] = documento.contar(blocos)
        # Pagina medida, nao estimada: a conta antiga era palavras/450, e
        # errava em todo documento com titulo, lista ou paragrafo curto.
        item["paginacao"] = documento.mapa_de_paginas(
            blocos, timbre=_timbre_do_documento(item.get("formato")), formato=item.get("formato"))
        item["formato"] = documento.normalizar_formato(item.get("formato"))
        item["folha"] = _folha_para_tela(item)
    return item


@app.post("/api/documentos/{id_}")
def documentos_gravar(id_: int, payload: GravarDocumento) -> dict:
    try:
        resultado = estado.documentos.salvar(id_, payload.corpo, payload.titulo, payload.nota)
    except ValueError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc

    if payload.corpo and not payload.corpo.lstrip().startswith("{"):
        resultado["contagem"] = documento.contar(documento.ler_html(payload.corpo))
    return resultado


@app.delete("/api/documentos/{id_}")
def documentos_apagar(id_: int) -> dict:
    d = estado.base.um("SELECT titulo, tipo FROM documentos WHERE id = ?", (id_,))
    if not d:
        raise HTTPException(status_code=404, detail="documento não encontrado")
    entrada = estado.lixeira.apagar_linha("documento", id_, d["titulo"], "Documentos · " + ("planilha" if d["tipo"] == "planilha" else "texto"))
    return _foi_para_lixeira(entrada, "apagado", id_)


@app.get("/api/documentos/{id_}/versoes")
def documentos_versoes(id_: int) -> dict:
    _documento_ou_404(id_)
    return {"versoes": estado.documentos.versoes(id_)}


@app.get("/api/documentos/{id_}/comparar")
def documentos_comparar(id_: int, de: int = 0, ate: int = 0) -> dict:
    """
    O que mudou entre duas versoes.

    Compara o texto, nao o HTML: o editor troca <b> por <strong> sozinho, e
    listar isso como alteracao de contrato faria a comparacao perder utilidade
    justamente no dia em que ela importa.
    """
    _documento_ou_404(id_)
    ultima = estado.documentos.ultima_versao(id_)
    ate = ate or ultima
    de = de or max(1, ate - 1)

    antes = estado.documentos.corpo_da_versao(id_, de)
    depois = estado.documentos.corpo_da_versao(id_, ate)
    if antes is None or depois is None:
        raise HTTPException(status_code=404, detail="essa versão não existe")

    return {"de": de, "ate": ate, "mudancas": documento.comparar(antes, depois)}


@app.post("/api/documentos/{id_}/notas")
def documentos_notas(id_: int, payload: dict) -> dict:
    """
    As notas de margem deste documento, presas ao paragrafo de cada uma.

    Duas origens, e a tela diz qual e qual:

      - regra: o que `conferir` acha, instantaneo e sempre igual
      - assistente: o que o modelo comentou quando pediram

    A da regra nao precisa de modelo nem de rede, e por isso aparece sempre. A
    do modelo so existe onde alguem pediu.
    """
    item = _documento_ou_404(id_)
    if item["tipo"] != "texto":
        raise HTTPException(status_code=400, detail="isso é uma planilha")

    corpo = payload.get("corpo")
    blocos = documento.ler_html(item["corpo"] if corpo is None else corpo)

    notas = []
    for aviso in documento.conferir(blocos, estado.cadastros.listar()):
        notas.append({
            "id": 0,
            "origem": "regra",
            "grau": aviso["grau"],
            "bloco": aviso["bloco"] - 1,     # a tela conta de zero
            "titulo": aviso["titulo"],
            "texto": aviso["detalhe"],
        })
    for c in estado.comentarios.ancorar(id_, blocos):
        notas.append({
            "id": c["id"],
            "origem": c["origem"],
            "grau": "comentario",
            "bloco": c["bloco"],
            "titulo": "Comentário",
            "texto": c["texto"],
            "trecho": c["trecho"],
            "criado_em": c["criado_em"],
        })

    return {"notas": notas, "blocos": len(blocos)}


@app.post("/api/documentos/{id_}/comentar")
def documentos_comentar(id_: int, payload: dict) -> dict:
    """
    O assistente comentando um trecho - sem reescrever nada.

    E diferente de pedir uma alteracao: aqui a resposta e uma observacao para
    quem le decidir, e nao um texto para entrar no contrato. O trecho vai junto
    porque comentario sobre "o documento" nao gruda em lugar nenhum.
    """
    item = _documento_ou_404(id_)
    trecho = str(payload.get("trecho", "")).strip()
    if not trecho:
        raise HTTPException(
            status_code=400,
            detail="selecione no texto o trecho que você quer que eu comente")

    disponivel, motivo = check_ollama(estado.client.model)
    if not disponivel:
        raise HTTPException(status_code=503, detail=motivo)

    pedido = (
        "Escreva UMA observacao curta sobre o trecho abaixo, para quem vai "
        "revisar o contrato. Diga o que falta, o que esta ambiguo ou o que "
        "merece conferencia. Nao reescreva o trecho e nao proponha texto novo. "
        "No maximo tres frases."
    )
    try:
        resposta = estado.client.ask(pedido, f"Trecho do contrato:\n{trecho[:2000]}",
                                     sistema=SISTEMA_COMENTARIO)
    except OllamaError as exc:
        raise HTTPException(status_code=503, detail=str(exc)) from exc

    texto = _limpar_sugestao(resposta)
    if not texto:
        raise HTTPException(status_code=422, detail="o modelo não devolveu observação nenhuma")

    novo = estado.comentarios.criar(id_, trecho, texto, origem="assistente")
    return {"id": novo, "texto": texto, "trecho": trecho[:600],
            "aviso": ("Escrito por um modelo pequeno rodando nesta máquina. "
                      "É observação para conferir, não parecer.")}


@app.post("/api/documentos/{id_}/comentarios")
def comentarios_criar(id_: int, payload: dict) -> dict:
    """Um comentario escrito pela propria pessoa, preso a um trecho."""
    _documento_ou_404(id_)
    try:
        novo = estado.comentarios.criar(
            id_, payload.get("trecho", ""), payload.get("texto", ""), origem="pessoa")
    except ValueError as erro:
        raise HTTPException(status_code=400, detail=str(erro))
    return {"id": novo}


@app.post("/api/documentos/{id_}/comentarios/{comentario_id}/resolver")
def comentarios_resolver(id_: int, comentario_id: int, payload: dict) -> dict:
    _documento_ou_404(id_)
    achou = estado.comentarios.resolver(comentario_id, bool(payload.get("resolvido", True)))
    if not achou:
        raise HTTPException(status_code=404, detail="comentário não encontrado")
    return {"ok": True}


@app.delete("/api/documentos/{id_}/comentarios/{comentario_id}")
def comentarios_apagar(id_: int, comentario_id: int) -> dict:
    _documento_ou_404(id_)
    if not estado.comentarios.apagar(comentario_id):
        raise HTTPException(status_code=404, detail="comentário não encontrado")
    return {"ok": True}


@app.post("/api/documentos/{id_}/alteracoes")
def documentos_alteracoes(id_: int, payload: dict) -> dict:
    """
    O que mudou no texto que esta no editor agora, em relacao a uma versao.

    Controlar alteracoes no Word marca cada tecla enquanto se digita. Aqui a
    marca vem da comparacao com uma versao gravada, e nao de interceptar a
    digitacao - porque interceptar tecla dentro de um campo editavel e o
    caminho curto para perder texto de contrato, e texto de contrato perdido
    nao tem conserto do lado de ca.

    O resultado e o mesmo que interessa: o que entrou, o que saiu, o que mudou,
    e o caminho de volta para cada um.
    """
    item = _documento_ou_404(id_)
    if item["tipo"] != "texto":
        raise HTTPException(status_code=400, detail="isso é uma planilha")

    ultima = estado.documentos.ultima_versao(id_)
    desde = int(payload.get("desde") or ultima)
    antes = estado.documentos.corpo_da_versao(id_, desde)
    if antes is None:
        raise HTTPException(status_code=404, detail="essa versão não existe")

    corpo = payload.get("corpo")
    depois = item["corpo"] if corpo is None else corpo
    return {
        "desde": desde,
        "ultima": ultima,
        "mudancas": documento.comparar(antes, depois),
    }


@app.post("/api/documentos/{id_}/restaurar")
def documentos_restaurar(id_: int, payload: dict) -> dict:
    _documento_ou_404(id_)
    try:
        resultado = estado.documentos.restaurar(id_, int(payload.get("numero", 0)))
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    return {**resultado, "documento": estado.documentos.obter(id_)}


# ------------------------------------------------ sair: PDF, DOCX, avisos


def _timbre_do_escritorio() -> dict:
    """
    O timbre sai dos dados profissionais, e so quando o escritorio quer.

    Fica desligado por padrao: por um lado, o dado pode nao estar preenchido;
    por outro, nem todo documento sai em papel timbrado - uma minuta interna
    com timbre parece peca protocolada.
    """
    if not estado.prefs.dados.get("timbre_no_pdf"):
        return {}
    return _dados_do_escritorio()


def _dados_do_escritorio() -> dict:
    """O timbre montado dos dados profissionais, com a chave ligada ou nao."""
    pessoa = estado.prefs.dados.get("pessoa", {})
    # Sem nome nao ha timbre: o resto sozinho sairia como um endereco solto no
    # alto da folha. Quem ligou a chave precisa saber disso na tela de
    # Configuracoes, e nao ao abrir o PDF e nao ver nada.
    if not str(pessoa.get("nome", "")).strip():
        return {}
    logo = estado.marca.caminho("logo")
    return {
        "nome": pessoa.get("nome", ""),
        "oab": pessoa.get("oab", ""),
        "cpf": pessoa.get("cpf", ""),
        "endereco": pessoa.get("endereco", ""),
        "telefone": pessoa.get("telefone", ""),
        "email": pessoa.get("email", ""),
        # A logo do escritorio vai no alto do papel, acima do nome. Quando nao
        # ha logo, o timbre continua sendo so o texto, como antes.
        "logo": str(logo) if logo else "",
    }


def _timbre_do_documento(formato) -> dict | None:
    """
    O cabecalho que ESTE documento leva, pela folha dele: o timbre de
    Configuracoes (quando a chave esta ligada, ou sempre), o que a pessoa
    escreveu, ou nenhum.
    """
    folha = documento.normalizar_formato(formato)["folha"]
    modo = folha["cabecalho"]
    if modo == "nenhum":
        return None
    if modo == "padrao":
        base = _timbre_do_escritorio()
    elif modo == "escritorio":
        base = _dados_do_escritorio()
    else:
        logo = estado.marca.caminho("logo")
        base = {"linhas": folha["linhas"], "logo": str(logo) if logo else ""}
    if not base:
        return None
    if not folha["logo"]:
        base = {**base, "logo": ""}
    return base


def _folha_para_tela(item: dict) -> dict:
    """A folha do documento e o desenho dela, para o editor e o modo Folha."""
    formato = documento.normalizar_formato(item.get("formato"))
    escritorio = _dados_do_escritorio()
    return {
        **formato["folha"],
        "desenho": documento.desenho_da_folha(
            _timbre_do_documento(formato), formato, item.get("titulo", "")),
        "escritorio": {
            "linhas": documento._linhas_do_timbre(escritorio) if escritorio else [],
            "logo": bool(estado.marca.caminho("logo")),
            "timbre_ligado": bool(estado.prefs.dados.get("timbre_no_pdf")),
        },
    }


def _pdf_do_documento(item: dict, timbre: dict | None = None) -> bytes:
    blocos = documento.ler_html(item["corpo"])
    formato = documento.normalizar_formato(item.get("formato"))
    rodape = documento.texto_do_rodape(formato["folha"], item["titulo"])
    if timbre is None:
        timbre = _timbre_do_documento(formato)
    return documento.para_pdf(blocos, item["titulo"], rodape, timbre=timbre or None,
                              formato=item.get("formato"))


@app.get("/api/documentos/{id_}/pdf")
def documentos_pdf(id_: int):
    from fastapi.responses import Response

    item = _documento_ou_404(id_)
    if item["tipo"] != "texto":
        raise HTTPException(status_code=400, detail="isso é uma planilha - baixe em XLSX ou CSV")
    return Response(
        _pdf_do_documento(item), media_type="application/pdf",
        headers=_anexo(_arquivo(item["titulo"]) + ".pdf"),
    )


@app.get("/api/impressoras")
def impressoras_listar() -> dict:
    """As impressoras desta conta do Windows, para o dialogo de imprimir."""
    import impressao

    try:
        lista = impressao.listar()
    except Exception as exc:  # driver quebrado nao pode derrubar a tela
        return {"disponivel": impressao.disponivel(), "impressoras": [], "erro": str(exc)}
    return {"disponivel": impressao.disponivel(), "impressoras": lista}


@app.post("/api/documentos/{id_}/imprimir")
def documentos_imprimir(id_: int, payload: dict) -> dict:
    """O documento direto na impressora escolhida, sem o dialogo do navegador."""
    import impressao

    item = _documento_ou_404(id_)
    if item["tipo"] != "texto":
        raise HTTPException(status_code=400, detail="isso é uma planilha - baixe em XLSX para imprimir")
    try:
        return impressao.imprimir(
            _pdf_do_documento(item), str(payload.get("impressora", "")), item["titulo"],
            paginas=str(payload.get("paginas", "todas")), copias=int(payload.get("copias", 1) or 1),
            cor=bool(payload.get("cor", True)), frente_verso=bool(payload.get("frente_verso", False)))
    except impressao.ErroDeImpressao as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    except (TypeError, ValueError) as exc:
        raise HTTPException(status_code=400, detail="não entendi o pedido de impressão") from exc


class PdfProtegido(BaseModel):
    senha: str = ""


@app.post("/api/documentos/{id_}/pdf")
def documentos_pdf_com_senha(id_: int, payload: PdfProtegido):
    """
    O mesmo PDF, pedindo senha para abrir.

    E POST, e nao um `?senha=` no endereco, de proposito: endereco fica no
    historico do navegador e no registro do servidor, e uma senha nao tem por
    que morar em nenhum dos dois.
    """
    from fastapi.responses import Response

    item = _documento_ou_404(id_)
    if item["tipo"] != "texto":
        raise HTTPException(status_code=400, detail="isso é uma planilha - baixe em XLSX ou CSV")
    try:
        dados = documento.proteger_com_senha(_pdf_do_documento(item), payload.senha)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    return Response(dados, media_type="application/pdf",
                    headers=_anexo(_arquivo(item["titulo"]) + ".pdf"))


@app.get("/api/documentos/{id_}/docx")
def documentos_docx(id_: int):
    from fastapi.responses import Response

    item = _documento_ou_404(id_)
    if item["tipo"] != "texto":
        raise HTTPException(status_code=400, detail="isso é uma planilha - baixe em XLSX ou CSV")

    dados = documento.para_docx(documento.ler_html(item["corpo"]), item["titulo"],
                                formato=item.get("formato"))
    return Response(
        dados,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers=_anexo(_arquivo(item["titulo"]) + ".docx"),
    )


@app.get("/api/documentos/{id_}/pagina")
def documentos_pagina(id_: int, numero: int = 1, largura: int = 900, versao: int = 0):
    """
    A pagina desenhada, para a pre-visualizacao.

    E o PDF de verdade rasterizado, nao uma aproximacao em CSS: o que a tela
    mostra e o arquivo que vai sair.

    Com `versao`, desenha uma versao antiga - e o que permite ver as duas lado
    a lado na comparacao. Listar o que mudou em texto nao mostra como ficou.
    """
    from fastapi.responses import Response

    item = _documento_ou_404(id_)
    if item["tipo"] != "texto":
        raise HTTPException(status_code=400, detail="planilha não tem página de PDF")

    if versao:
        corpo = estado.documentos.corpo_da_versao(id_, versao)
        if corpo is None:
            raise HTTPException(status_code=404, detail="essa versão não existe")
        item = {**item, "corpo": corpo}

    pdf = _pdf_do_documento(item)
    return Response(
        documento.pagina_png(pdf, numero, largura),
        media_type="image/png", headers={"Cache-Control": "no-cache"},
    )


@app.post("/api/documentos/{id_}/paginacao")
def documentos_paginacao(id_: int, payload: dict) -> dict:
    """
    Em que pagina cai cada paragrafo do que esta no editor agora.

    O editor mostrava "paginas ~4": caracteres divididos por uma media, que
    erra com titulo, lista ou paragrafo curto. Aqui o numero e medido - o PDF
    e montado de verdade, o mesmo que a pre-visualizacao desenha.

    Recebe o corpo em vez de ler do banco porque a pergunta e sobre o texto
    que esta na tela, ainda nao gravado.
    """
    item = _documento_ou_404(id_)
    if item["tipo"] != "texto":
        raise HTTPException(status_code=400, detail="planilha não tem página de PDF")

    corpo = payload.get("corpo")
    blocos = documento.ler_html(item["corpo"] if corpo is None else corpo)
    # Com `formato`, mede uma folha proposta sem gravar nada.
    formato = item.get("formato")
    if isinstance(payload.get("formato"), dict):
        pedido = payload["formato"]
        if "folha" not in pedido:
            pedido = {**pedido, "folha": documento.normalizar_formato(formato)["folha"]}
        formato = documento.normalizar_formato(pedido)
    mapa = documento.mapa_de_paginas(
        blocos, timbre=_timbre_do_documento(formato), formato=formato)

    return {**mapa, "blocos": len(blocos), "formato": documento.normalizar_formato(formato),
            "folha": _folha_para_tela({**item, "formato": formato})}


@app.post("/api/documentos/{id_}/formato")
def documentos_formato(id_: int, payload: dict) -> dict:
    """A fonte, o corpo, o recuo e as entrelinhas da folha deste documento."""
    item = _documento_ou_404(id_)
    if item["tipo"] != "texto":
        raise HTTPException(status_code=400, detail="planilha não tem formato de folha")

    try:
        formato = estado.documentos.formatar(id_, payload.get("formato", payload))
    except ValueError as erro:
        raise HTTPException(status_code=404, detail=str(erro))

    corpo = payload.get("corpo")
    blocos = documento.ler_html(item["corpo"] if corpo is None else corpo)
    mapa = documento.mapa_de_paginas(
        blocos, timbre=_timbre_do_documento(formato), formato=formato)
    return {**mapa, "blocos": len(blocos), "formato": formato,
            "folha": _folha_para_tela({**item, "formato": formato})}


@app.post("/api/documentos/{id_}/folha/sugerir")
def documentos_folha_sugerir(id_: int, payload: dict) -> dict:
    """
    O papel timbrado que o pedido descreve ("cabecalho com meu nome e OAB,
    pagina 1 de 3 no rodape"). Devolve a proposta sem gravar.
    """
    item = _documento_ou_404(id_)
    pedido = str(payload.get("pedido", "")).strip()[:600]
    if not pedido:
        raise HTTPException(status_code=400, detail="diga como quer a folha")
    atual = payload.get("folha") or documento.normalizar_formato(item.get("formato"))["folha"]
    escritorio = _dados_do_escritorio()
    linhas = documento._linhas_do_timbre(escritorio) if escritorio else []
    return documento.sugerir_folha(
        pedido, atual, linhas, item.get("titulo", ""),
        lambda instrucao, sistema, esquema: estado.client.ask_json(
            instrucao, schema_hint=esquema, sistema=sistema))


@app.get("/api/documentos/{id_}/conferir")
def documentos_conferir(id_: int) -> dict:
    """
    O que conferir antes de o documento sair.

    Tudo regra, nada de modelo: lacuna de modelo que ficou, documento com
    digito verificador errado, valor com cifrao e sem numero, nome que aparece
    no texto sem bater com a ficha do cadastro. Roda em milissegundos e a
    resposta e sempre a mesma.
    """
    item = _documento_ou_404(id_)
    if item["tipo"] != "texto":
        return {"avisos": [], "paginas": 0}

    blocos = documento.ler_html(item["corpo"])
    avisos = documento.conferir(blocos, estado.cadastros.listar())
    pdf = _pdf_do_documento(item)

    return {
        "avisos": avisos,
        "impedem": sum(1 for a in avisos if a["grau"] == "impede"),
        "paginas": documento.paginas_de(pdf),
        "bytes": len(pdf),
        "contagem": documento.contar(blocos),
        "versao": item["versao"],
    }


@app.post("/api/documentos/{id_}/biblioteca")
def documentos_para_biblioteca(id_: int, payload: dict | None = None) -> dict:
    """
    Grava o PDF na biblioteca, para poder assinar e anexar em e-mail.

    Nunca passa por cima de arquivo existente: sai com numero no fim.
    """
    item = _documento_ou_404(id_)
    formato = str((payload or {}).get("formato", "pdf")).lower()
    senha = str((payload or {}).get("senha", ""))

    if item["tipo"] == "planilha":
        abas = planilha.de_dict(json.loads(item["corpo"] or "{}"))
        dados = planilha.para_xlsx(abas, [planilha.calcular_aba(a) for a in abas])
        sufixo = ".xlsx"
    elif formato == "docx":
        dados = documento.para_docx(documento.ler_html(item["corpo"]), item["titulo"])
        sufixo = ".docx"
    else:
        dados = _pdf_do_documento(item)
        # Com senha, o arquivo que vai para a biblioteca (e dali para o anexo
        # do e-mail) ja sai protegido: proteger so o download deixaria passar
        # justamente o caminho que sai desta maquina.
        if senha:
            try:
                dados = documento.proteger_com_senha(dados, senha)
            except ValueError as exc:
                raise HTTPException(status_code=400, detail=str(exc)) from exc
        sufixo = ".pdf"

    estado.pasta.mkdir(parents=True, exist_ok=True)
    base = _arquivo(item["titulo"])
    destino = estado.pasta / f"{base}{sufixo}"
    conta = 2
    while destino.exists():
        destino = estado.pasta / f"{base} ({conta}){sufixo}"
        conta += 1

    destino.write_bytes(dados)
    return {
        "guardado": destino.name,
        "caminho": str(destino),
        "documentos": estado.recarregar(force=True) if sufixo != ".xlsx" else estado.recarregar(),
    }


def _arquivo(titulo: str) -> str:
    """Titulo virando nome de arquivo que o Windows aceita."""
    limpo = re.sub(r'[<>:"/\\|?*]', "-", titulo or "documento").strip(" .")
    return (limpo or "documento")[:90]


def _anexo(nome: str) -> dict:
    """
    O cabecalho que manda o navegador salvar o arquivo com esse nome.

    Cabecalho HTTP e latin-1, e titulo de documento brasileiro tem acento e
    travessao - "Contrato de honorarios - Cooperativa" derrubava a rota com
    UnicodeEncodeError. A regra do proprio HTTP resolve: um nome sem acento
    para quem so entende o basico, e o nome de verdade em UTF-8 ao lado.
    """
    from urllib.parse import quote

    simples = unicodedata.normalize("NFKD", nome).encode("ascii", "ignore").decode()
    simples = (simples or "documento").replace(_ASPAS, "")
    return {
        "Content-Disposition":
            "attachment; filename=" + _ASPAS + simples + _ASPAS
            + "; filename*=UTF-8''" + quote(nome)
    }


_ASPAS = chr(34)



# --------------------------------------------------------- o assistente


INSTRUCAO_EDITOR = """Voce ajuda um advogado brasileiro a redigir. Responda em
portugues do Brasil, direto, sem preambulo.

Regras:
- devolva APENAS o texto pedido, pronto para entrar no documento
- nao invente numero de artigo, de lei, de sumula nem de processo
- nao invente nome, data, valor nem prazo que nao estejam no que foi dado
- se faltar informacao para escrever, escreva o texto com a lacuna marcada
  entre colchetes, por exemplo [VALOR]
- sem marcacao, sem asteriscos, sem titulo"""

# A instrucao de SISTEMA do editor. Sem ela valia a do acervo, que manda citar
# o nome do arquivo de onde a informacao saiu - e o modelo obedecia as duas
# ordens ao mesmo tempo: devolvia o texto pedido com "Arquivo: contrato.txt" na
# frente, com o nome do arquivo inventado.
SISTEMA_EDITOR = """Voce e o PAULUS, ajudando a redigir um documento juridico
brasileiro. Voce nao esta consultando arquivos: o que vem abaixo e o texto em
que se esta trabalhando agora.

Nunca cite nome de arquivo, nunca escreva cabecalho como "Arquivo:" ou
"Documento:", nunca repita o pedido antes de responder.

Nunca invente artigo de lei, sumula, processo, nome, data, valor ou prazo.

Voce NAO conhece a tela deste programa. Nunca diga onde clicar, nunca cite
botao, menu ou atalho."""

SISTEMA_COMENTARIO = SISTEMA_EDITOR + """

Aqui voce NAO reescreve nada: voce comenta. A resposta e uma observacao curta
para quem vai revisar decidir - o que falta, o que esta ambiguo, o que merece
conferencia. Nao proponha texto novo e nao repita o trecho."""


class PedidoDeRedacao(BaseModel):
    corpo: str = ""
    cadastro_id: int | None = None


@app.post("/api/documentos/{id_}/clausulas")
def documentos_renumerar(id_: int, payload: PedidoDeRedacao) -> dict:
    """
    Poe as clausulas em sequencia e leva as referencias cruzadas junto.

    Nao grava: devolve o texto novo e a lista do que mudaria. Renumerar um
    contrato sem mostrar o que mudou e pedir para a pessoa aceitar no escuro.
    """
    _documento_ou_404(id_)
    corpo = payload.corpo or ""
    resultado = redacao.renumerar(corpo)
    return {
        **resultado,
        "quebradas": redacao.referencias_quebradas(resultado["texto"]),
        "estilo": redacao.estilo_do_documento(corpo),
    }


@app.get("/api/documentos/{id_}/conferir-clausulas")
def documentos_conferir_clausulas(id_: int) -> dict:
    """O que esta fora de ordem ou apontando para o nada, sem mexer em nada."""
    item = _documento_ou_404(id_)
    corpo = item["corpo"] or ""
    achadas = redacao.achar_clausulas(corpo)
    numeros = [a["numero"] for a in achadas]
    return {
        "clausulas": len(achadas),
        "numeros": numeros,
        "fora_de_ordem": numeros != sorted(numeros) or numeros != list(range(1, len(numeros) + 1)),
        "quebradas": redacao.referencias_quebradas(corpo),
        "estilo": redacao.estilo_do_documento(corpo),
    }


@app.get("/api/redacao/qualificacao")
def redacao_qualificacao(cadastro_id: int) -> dict:
    """
    O paragrafo de qualificacao de uma parte, com o que o cadastro tem.

    O que falta vira lacuna entre colchetes - e sai declarado, porque e mais
    barato completar o cadastro do que cacar colchete dentro do contrato
    depois de assinado.
    """
    ficha = estado.cadastros.obter(cadastro_id)
    if not ficha:
        raise HTTPException(status_code=404, detail="nao achei esse cadastro")
    return {
        "texto": redacao.qualificar(ficha),
        "falta": redacao.o_que_falta(ficha),
        "nome": ficha["nome"],
        "cadastro_id": cadastro_id,
    }


@app.post("/api/documentos/{id_}/assistente")
def documentos_assistente(id_: int, payload: PedidoAoAssistente) -> dict:
    """
    O assistente escrevendo dentro do documento.

    O ganho desta tela e este - o editor em si o Word ja faz. O que volta e
    sugestao: entra no texto so quando a pessoa aceitar.
    """
    item = _documento_ou_404(id_)
    pedido = payload.pedido.strip()
    if not pedido:
        raise HTTPException(status_code=400, detail="diga o que você quer que eu escreva")

    disponivel, motivo = check_ollama(estado.client.model)
    if not disponivel:
        raise HTTPException(status_code=503, detail=motivo)

    if payload.trecho.strip():
        contexto = f"Trecho selecionado do documento:\n{payload.trecho[:2500]}"
    else:
        texto = documento.para_texto(documento.ler_html(item["corpo"]))
        contexto = f"Documento (início):\n{texto[:2500]}"
    instrucao = INSTRUCAO_EDITOR + f"\n\nPedido: {pedido}"

    # Vindo de uma conversa, o pedido e trabalho dela como qualquer pergunta:
    # a conversa fica "trabalhando" com as etapas a vista, o cartao
    # "Acontecendo agora" da tela inicial mostra o andamento - com barra, quando
    # esta maquina ja mediu quanto le e escreve -, e o botao de parar alcanca.
    trabalho = estado.trabalhos.obter(payload.trabalho_id) if payload.trabalho_id else None
    parar = threading.Event()
    if trabalho:
        agora = time.time()
        trabalho.dizer("pessoa", pedido)
        trabalho.etapas = [Etapa("Entender o pedido", estado=CONCLUIDO),
                           Etapa(f"Escrever em “{item['titulo']}”", estado=EXECUTANDO)]
        trabalho.estado = EXECUTANDO
        estado.trabalhos.salvar(trabalho)
        estado.respondendo[trabalho.id] = parar
        estado.andamento[trabalho.id] = {
            "fase": "documento", "inicio": agora, "desde": agora,
            "previsao_s": _previsao_de_escrita(len(instrucao) + len(contexto)),
            "palavras": 0, "documentos": 1, "caracteres": len(contexto),
        }

    def terminar(etapa: str, texto: str, **extras) -> None:
        if not trabalho:
            return
        estado.andamento.pop(trabalho.id, None)
        estado.respondendo.pop(trabalho.id, None)
        trabalho.etapas[-1].estado = etapa
        trabalho.estado = CONCLUIDO if etapa == CONCLUIDO else (PAUSADO if etapa == PAUSADO else "falhou")
        trabalho.dizer("paulus", texto, **extras)
        estado.trabalhos.salvar(trabalho)

    try:
        # No editor entram so as regras de redacao: como o escritorio escreve
        # muda o texto sugerido; o nome de um cliente nao tem o que fazer aqui.
        resposta = estado.client.ask(instrucao, contexto,
                                     sistema=SISTEMA_EDITOR,
                                     ensinado=estado.contextos.bloco(["Regras de redação"]),
                                     parar=parar.is_set)
    except OllamaError as exc:
        terminar("falhou", f"Não consegui escrever em “{item['titulo']}”: {exc}")
        raise HTTPException(status_code=503, detail=str(exc)) from exc
    except Exception:
        terminar("falhou", f"Não consegui escrever em “{item['titulo']}”.")
        raise

    if parar.is_set():
        terminar(PAUSADO, f"Parei antes de escrever em “{item['titulo']}”. O documento ficou como estava.")
        raise HTTPException(status_code=409, detail="parei antes de terminar a alteração")

    sugestao = _limpar_sugestao(resposta)
    inteiro = documento.para_texto(documento.ler_html(item["corpo"]))
    if _e_o_documento_de_volta(sugestao, inteiro):
        detalhe = ("o modelo devolveu o documento de volta em vez da alteração. "
                   "Selecione no texto o trecho que você quer mudar e peça de "
                   "novo — com o trecho à mão ele acerta.")
        terminar("falhou", "Não consegui: " + detalhe)
        raise HTTPException(status_code=422, detail=detalhe)

    onde = "o trecho selecionado" if payload.trecho.strip() else "o fim do documento"
    terminar(CONCLUIDO, f"Escrevi em “{item['titulo']}” ({onde}). A alteração ficou marcada "
                        "no documento, esperando você manter ou descartar.",
             feito={"tipo": "alteracao", "id": id_, "nome": item["titulo"], "onde": "editor"})
    avisos.avisar("resposta", "Alteração pronta · " + item["titulo"][:60],
                  f"Escrevi em {onde}. Está marcada no documento, esperando você manter ou descartar.")

    return {
        "sugestao": sugestao,
        "sobre": payload.trecho[:160],
        "aviso": (
            "Escrito por um modelo pequeno rodando nesta máquina. Confira nomes, "
            "datas, valores e qualquer artigo de lei antes de aceitar."
        ),
    }


def _previsao_de_escrita(caracteres: int) -> float:
    """
    Quanto uma alteracao no documento deve levar nesta maquina: ler o pedido
    e o trecho, e escrever uns 120 palavras. Zero quando nao ha medida - a
    tela mostra a barra so com numero de verdade atras.
    """
    leitura = estado.ritmo.previsao_de_leitura(estado.client.model, caracteres)
    if not leitura.get("sabe"):
        return 0
    por_segundo = estado.ritmo.palavras_por_segundo(estado.client.model)
    return round(leitura["segundos"] + (120 / por_segundo if por_segundo else 0), 1)


def _limpar_sugestao(texto: str) -> str:
    limpo = (texto or "").strip()
    limpo = re.sub(r"\*\*(.+?)\*\*", r"\1", limpo)
    limpo = re.sub(r"^\s*[*-]\s+", "", limpo, flags=re.M)
    limpo = re.sub(r"^\s*(sugest[aã]o|resposta|texto)\s*:\s*", "", limpo, flags=re.I)
    # "[Texto do documento]", "[Documento]": rotulo que o modelo copia do
    # cabecalho do contexto e entrega como se fosse parte do texto.
    limpo = re.sub(r"^\s*\[[^\]]{0,40}\]\s*", "", limpo)
    # "Arquivo: contrato.txt" na primeira linha: o modelo obedecendo a ordem de
    # citar a origem, com um nome de arquivo que ele mesmo inventou. A ordem ja
    # saiu do prompt do editor, mas o rotulo ainda escapa de vez em quando.
    limpo = re.sub(r"^\s*(arquivo|documento|origem|fonte)\s*:[^\n]{0,80}\n+", "",
                   limpo, flags=re.I)
    return limpo.strip()


def _e_o_documento_de_volta(sugestao: str, texto_do_documento: str) -> bool:
    """
    O modelo devolveu o documento em vez da alteracao?

    Acontece com pedido amplo - "deixe mais formal" sem trecho selecionado - e
    o estrago e silencioso: o texto entra no fim e o documento dobra de
    tamanho. Aconteceu de verdade: 355 palavras viraram 712.
    """
    a = re.sub(r"\s+", " ", sugestao or "").strip().lower()
    b = re.sub(r"\s+", " ", texto_do_documento or "").strip().lower()
    if len(a) < 400 or len(b) < 400:
        return False
    # Comeca igual, ou e quase do tamanho do documento e aparece dentro dele:
    # nos dois casos e copia, nao reescrita.
    return a[:300] == b[:300] or (len(a) > len(b) * 0.7 and a[:150] in b)


# ------------------------------------------------------------ planilha


def _abas_do(item: dict) -> list:
    try:
        return planilha.de_dict(json.loads(item["corpo"] or "{}"))
    except json.JSONDecodeError:
        return [planilha.Aba()]


def _resposta_planilha(id_: int, item: dict, abas: list) -> dict:
    calculados = [planilha.calcular_aba(a) for a in abas]
    return {
        "id": id_,
        "titulo": item["titulo"],
        "versao": item.get("versao", 1),
        "abas": [a.to_dict() for a in abas],
        "calculado": calculados,
        "resumos": [planilha.resumo(a, c) for a, c in zip(abas, calculados)],
        "formatos": [{"valor": k, "rotulo": v} for k, v in planilha.FORMATOS.items()],
        # Ate onde a grade cresce quando a pessoa rola.
        "limites": {"linhas": planilha.MAX_LINHAS, "colunas": planilha.MAX_COLUNAS},
        "funcoes": [
            {"nome": n, "exemplo": e, "explica": x} for n, e, x in planilha.AJUDA_FUNCOES
        ],
    }


@app.get("/api/planilha/{id_}")
def planilha_abrir(id_: int) -> dict:
    item = _documento_ou_404(id_)
    if item["tipo"] != "planilha":
        raise HTTPException(status_code=400, detail="isso não é uma planilha")
    return _resposta_planilha(id_, item, _abas_do(item))


@app.post("/api/planilha/{id_}/celula")
def planilha_celula(id_: int, payload: CelulaPlanilha) -> dict:
    item = _documento_ou_404(id_)
    abas = _abas_do(item)
    if not 0 <= payload.aba < len(abas):
        raise HTTPException(status_code=400, detail="aba não encontrada")
    if not planilha.RE_CELULA.match(payload.ref.upper()):
        raise HTTPException(status_code=400, detail="referência de célula inválida")

    abas[payload.aba].gravar(payload.ref, payload.dados)
    estado.documentos.salvar(id_, planilha.para_json(abas), nota=f"{payload.ref.upper()} alterada")
    return _resposta_planilha(id_, estado.documentos.obter(id_), abas)


@app.post("/api/planilha/{id_}/faixa")
def planilha_faixa(id_: int, payload: dict) -> dict:
    """
    Aplica formato, negrito, italico ou borda a uma selecao inteira.

    Uma celula de cada vez obrigava a repetir o clique por linha - e numa
    tabela de parcelas ninguem faz isso, simplesmente deixa sem.
    """
    item = _documento_ou_404(id_)
    abas = _abas_do(item)
    indice = int(payload.get("aba", 0))
    if not 0 <= indice < len(abas):
        raise HTTPException(status_code=400, detail="aba não encontrada")

    refs = planilha.refs_da_faixa(payload.get("faixa", ""))
    if not refs:
        raise HTTPException(status_code=400, detail="seleção vazia")
    if len(refs) > planilha.MAX_LINHAS * 4:
        raise HTTPException(status_code=400, detail="seleção grande demais")

    dados = {c: payload[c] for c in ("formato", "negrito", "italico", "borda", "valor")
             if c in payload}
    if not dados:
        raise HTTPException(status_code=400, detail="nada para aplicar")

    for ref in refs:
        abas[indice].gravar(ref, dados)

    estado.documentos.salvar(
        id_, planilha.para_json(abas),
        nota=f"{str(payload.get('faixa', '')).upper()}: {', '.join(dados)}")
    return _resposta_planilha(id_, estado.documentos.obter(id_), abas)


@app.post("/api/planilha/{id_}/grafico")
def planilha_grafico(id_: int, payload: dict) -> dict:
    """
    A selecao virando serie para desenhar.

    Quem escolhe qual coluna e valor e qual e rotulo e a propria selecao: a
    ultima coluna com numero e o valor, a coluna de texto a esquerda sao os
    rotulos. Pedir isso num formulario antes de ver o desenho seria trocar um
    grafico por um questionario.
    """
    item = _documento_ou_404(id_)
    abas = _abas_do(item)
    indice = int(payload.get("aba", 0))
    if not 0 <= indice < len(abas):
        raise HTTPException(status_code=400, detail="aba não encontrada")

    aba = abas[indice]
    return planilha.serie_da_faixa(aba, planilha.calcular_aba(aba), payload.get("faixa", ""))


def _aba_do_pedido(item: dict, payload: dict) -> tuple[list, int]:
    abas = _abas_do(item)
    try:
        indice = int(payload.get("aba", 0))
    except (TypeError, ValueError):
        indice = -1
    if not 0 <= indice < len(abas):
        raise HTTPException(status_code=400, detail="aba não encontrada")
    return abas, indice


@app.post("/api/planilha/{id_}/lote")
def planilha_lote(id_: int, payload: dict) -> dict:
    """
    Varias celulas numa chamada so: colar, preencher com a alca, limpar,
    recortar. Uma versao no Historico para a operacao inteira - que e o que a
    pessoa fez: "colei a tabela", e nao "mudei 200 celulas".
    """
    item = _documento_ou_404(id_)
    if item["tipo"] != "planilha":
        raise HTTPException(status_code=400, detail="isso não é uma planilha")
    abas, indice = _aba_do_pedido(item, payload)
    itens = payload.get("itens")
    if not isinstance(itens, list) or not itens:
        raise HTTPException(status_code=400, detail="nada para gravar")
    if len(itens) > planilha.MAX_LINHAS * 10:
        raise HTTPException(status_code=400, detail="seleção grande demais")

    feitas = planilha.gravar_lote(abas[indice], itens)
    nota = " ".join(str(payload.get("nota") or "").split())[:80] or "células alteradas"
    estado.documentos.salvar(id_, planilha.para_json(abas), nota=nota)
    return {**_resposta_planilha(id_, estado.documentos.obter(id_), abas), "feitas": feitas}


@app.post("/api/planilha/{id_}/estrutura")
def planilha_estrutura(id_: int, payload: dict) -> dict:
    """
    Inserir ou excluir linhas e colunas, com as formulas seguindo as celulas.

    `em` e o numero da linha (1, 2, ...) ou a letra/indice da coluna.
    """
    item = _documento_ou_404(id_)
    abas, indice = _aba_do_pedido(item, payload)
    eixo = str(payload.get("eixo", ""))
    em = payload.get("em", 1)
    if eixo == "colunas" and isinstance(em, str) and not em.strip().isdigit():
        em = planilha.indice_da_coluna(em.strip())
    try:
        feito = planilha.reestruturar(abas[indice], eixo, str(payload.get("acao", "")),
                                      int(em), int(payload.get("quantas", 1) or 1))
    except (TypeError, ValueError) as erro:
        raise HTTPException(status_code=400, detail=str(erro))

    if eixo == "linhas":
        onde = f"linha {feito['em']}" if feito["quantas"] == 1 else f"linhas {feito['em']}-{feito['em'] + feito['quantas'] - 1}"
    else:
        letra = planilha.letra_da_coluna(feito["em"])
        onde = f"coluna {letra}" if feito["quantas"] == 1 else (
            f"colunas {letra}-{planilha.letra_da_coluna(feito['em'] + feito['quantas'] - 1)}")
    estado.documentos.salvar(
        id_, planilha.para_json(abas),
        nota=("inseriu " if feito["acao"] == "inserir" else "excluiu ") + onde)
    return {**_resposta_planilha(id_, estado.documentos.obter(id_), abas), "feito": feito}


@app.post("/api/planilha/{id_}/layout")
def planilha_layout(id_: int, payload: dict) -> dict:
    """
    Largura de coluna, altura de linha e o que esta oculto.

    {"larguras": {"B": 180, "C": null}, "alturas": {"3": 40},
     "ocultar": {"colunas": ["D"], "linhas": [7]}, "reexibir": {...}}
    null volta a medida ao padrao.
    """
    item = _documento_ou_404(id_)
    abas, indice = _aba_do_pedido(item, payload)
    aba = abas[indice]
    larguras = payload.get("larguras") or {}
    alturas = payload.get("alturas") or {}
    ocultar = payload.get("ocultar") or {}
    reexibir = payload.get("reexibir") or {}
    if not all(isinstance(x, dict) for x in (larguras, alturas, ocultar, reexibir)):
        raise HTTPException(status_code=400, detail="formato inválido")
    try:
        aba.medir(larguras, alturas)
        aba.ocultar(ocultar.get("colunas") or [], ocultar.get("linhas") or [], True)
        aba.ocultar(reexibir.get("colunas") or [], reexibir.get("linhas") or [], False)
    except (TypeError, ValueError):
        raise HTTPException(status_code=400, detail="medida inválida")

    partes = []
    if larguras:
        partes.append("largura de " + ", ".join(sorted(str(k).upper() for k in larguras)))
    if alturas:
        partes.append("altura da linha " + ", ".join(sorted((str(k) for k in alturas), key=lambda x: int(x) if x.isdigit() else 0)))
    if ocultar:
        partes.append("ocultou")
    if reexibir:
        partes.append("reexibiu")
    estado.documentos.salvar(id_, planilha.para_json(abas), nota=("; ".join(partes) or "layout")[:80])
    return _resposta_planilha(id_, estado.documentos.obter(id_), abas)


@app.post("/api/planilha/{id_}/mesclar")
def planilha_mesclar(id_: int, payload: dict) -> dict:
    """
    Junta as celulas de uma faixa numa so, ou desfaz a juncao.

    Nao cobre celula com conteudo. O Excel junta e joga fora o que estava
    debaixo, avisando numa caixa que todo mundo clica em OK sem ler - e ali
    some um valor que ninguem mais procura.
    """
    item = _documento_ou_404(id_)
    abas = _abas_do(item)
    indice = int(payload.get("aba", 0))
    if not 0 <= indice < len(abas):
        raise HTTPException(status_code=400, detail="aba não encontrada")

    faixa = payload.get("faixa", "")
    try:
        feito = (planilha.separar(abas[indice], faixa) if payload.get("separar")
                 else planilha.mesclar(abas[indice], faixa))
    except ValueError as erro:
        raise HTTPException(status_code=400, detail=str(erro))

    estado.documentos.salvar(
        id_, planilha.para_json(abas),
        nota=("separou " if payload.get("separar") else "juntou ") + str(faixa).upper())
    return {**_resposta_planilha(id_, estado.documentos.obter(id_), abas), "feito": feito}


@app.post("/api/planilha/{id_}/congelar")
def planilha_congelar(id_: int, payload: dict) -> dict:
    """Prende a primeira linha no lugar - e vai junto para o XLSX."""
    item = _documento_ou_404(id_)
    abas = _abas_do(item)
    indice = int(payload.get("aba", 0))
    if not 0 <= indice < len(abas):
        raise HTTPException(status_code=400, detail="aba não encontrada")

    abas[indice].congelar_cabecalho = bool(payload.get("congelar"))
    estado.documentos.salvar(
        id_, planilha.para_json(abas),
        nota="congelou o cabeçalho" if abas[indice].congelar_cabecalho
        else "soltou o cabeçalho")
    return _resposta_planilha(id_, estado.documentos.obter(id_), abas)


@app.post("/api/planilha/{id_}/selecao")
def planilha_selecao(id_: int, payload: dict) -> dict:
    """
    O que a seleção tem dentro: soma, média, mínimo, máximo, vazias.

    Quem conta é a planilha, e não a tela: os mesmos números que as fórmulas
    usam. Uma soma calculada em JavaScript e outra em Python é a promessa de
    duas respostas diferentes para a mesma coluna.
    """
    item = _documento_ou_404(id_)
    abas = _abas_do(item)
    indice = int(payload.get("aba", 0))
    if not 0 <= indice < len(abas):
        raise HTTPException(status_code=400, detail="aba não encontrada")

    refs = planilha.refs_da_faixa(payload.get("faixa", ""))
    if not refs:
        raise HTTPException(status_code=400, detail="seleção vazia")

    calculado = planilha.calcular_aba(abas[indice])
    return {**planilha.resumo_selecao(abas[indice], calculado, refs),
            "faixa": str(payload.get("faixa", "")).upper()}


@app.post("/api/planilha/{id_}/ordenar")
def planilha_ordenar(id_: int, payload: dict) -> dict:
    """
    Ordena as linhas de uma faixa por uma coluna, com a linha inteira junto.

    Cria versao: ordenar muda o documento, e e a operacao desta tela com maior
    potencial de estrago que ninguem percebe olhando - cada celula continua com
    um numero plausivel, so que na linha errada. O "voltar para a v(n-1)" tem
    que existir.
    """
    item = _documento_ou_404(id_)
    abas = _abas_do(item)
    indice = int(payload.get("aba", 0))
    if not 0 <= indice < len(abas):
        raise HTTPException(status_code=400, detail="aba não encontrada")

    aba = abas[indice]
    calculado = planilha.calcular_aba(aba)
    try:
        feito = planilha.ordenar(
            aba, calculado, payload.get("faixa", ""), payload.get("coluna", ""),
            crescente=bool(payload.get("crescente", True)),
            com_cabecalho=bool(payload.get("com_cabecalho", False)),
        )
    except ValueError as erro:
        raise HTTPException(status_code=400, detail=str(erro))

    estado.documentos.salvar(
        id_, planilha.para_json(abas),
        nota=f"ordenou {str(payload.get('faixa', '')).upper()} por {payload.get('coluna', '')}")
    return {**_resposta_planilha(id_, estado.documentos.obter(id_), abas), "ordenou": feito}


@app.post("/api/planilha/{id_}/filtrar")
def planilha_filtrar(id_: int, payload: dict) -> dict:
    """
    Quais linhas esconder para ver so o que casa com o criterio.

    Nao grava nada: filtro e jeito de olhar. Um filtro gravado esconderia
    linhas de quem abrisse o arquivo depois sem saber que ha filtro, e uma
    tabela de parcelas com linhas faltando e um erro que ninguem percebe.
    """
    item = _documento_ou_404(id_)
    abas = _abas_do(item)
    indice = int(payload.get("aba", 0))
    if not 0 <= indice < len(abas):
        raise HTTPException(status_code=400, detail="aba não encontrada")

    aba = abas[indice]
    try:
        return planilha.filtrar(
            aba, planilha.calcular_aba(aba), payload.get("faixa", ""),
            payload.get("coluna", ""), payload.get("criterio", ""),
            com_cabecalho=bool(payload.get("com_cabecalho", True)),
        )
    except ValueError as erro:
        raise HTTPException(status_code=400, detail=str(erro))


@app.post("/api/planilha/{id_}/aba")
def planilha_nova_aba(id_: int, payload: dict) -> dict:
    item = _documento_ou_404(id_)
    abas = _abas_do(item)
    if len(abas) >= 6:
        raise HTTPException(status_code=400, detail="seis abas é o limite por planilha")

    nome = " ".join(str(payload.get("nome", "")).split())[:24] or f"Página {len(abas) + 1}"
    abas.append(planilha.Aba(nome=nome))
    estado.documentos.salvar(id_, planilha.para_json(abas), nota=f"aba {nome} criada")
    return _resposta_planilha(id_, estado.documentos.obter(id_), abas)


@app.delete("/api/planilha/{id_}/aba/{indice}")
def planilha_apagar_aba(id_: int, indice: int) -> dict:
    item = _documento_ou_404(id_)
    abas = _abas_do(item)
    if len(abas) <= 1:
        raise HTTPException(status_code=400, detail="a planilha precisa de pelo menos uma aba")
    if not 0 <= indice < len(abas):
        raise HTTPException(status_code=400, detail="aba não encontrada")

    fora = abas.pop(indice)
    estado.documentos.salvar(id_, planilha.para_json(abas), nota=f"aba {fora.nome} apagada")
    return _resposta_planilha(id_, estado.documentos.obter(id_), abas)


class TrazerParaPlanilha(BaseModel):
    de: str = ""          # "financeiro" ou "prazos"
    mes: str = ""         # AAAA-MM; vazio traz tudo o que existe
    categoria: str = "honorarios"


@app.post("/api/planilha/{id_}/trazer")
def planilha_trazer(id_: int, payload: TrazerParaPlanilha) -> dict:
    """
    Uma aba nova com o que ja esta no programa: honorarios ou prazos.

    O painel da planilha oferecia as duas coisas e as duas diziam "em breve".
    Sao dados que o programa ja tem - o que faltava era escreve-los em linhas
    e colunas. Vem sempre numa ABA NOVA, nunca por cima do que a pessoa
    escreveu: uma importacao que apaga trabalho alheio nao se desfaz.

    O que entra e valor, nao formula: a planilha e uma fotografia do dia em
    que foi feita, e o total e que e formula, para continuar certo se alguem
    corrigir uma linha.
    """
    item = _documento_ou_404(id_)
    if item["tipo"] != "planilha":
        raise HTTPException(status_code=400, detail="isso é um documento de texto")
    abas = _abas_do(item)
    if len(abas) >= 6:
        raise HTTPException(status_code=400, detail="seis abas é o limite por planilha")

    if payload.de == "financeiro":
        aba, quantos = _aba_do_financeiro(payload.mes, payload.categoria)
    elif payload.de == "prazos":
        aba, quantos = _aba_dos_prazos()
    else:
        raise HTTPException(status_code=400, detail="não sei trazer isso")

    if not quantos:
        raise HTTPException(
            status_code=404,
            detail=("não achei honorários lançados nesse período" if payload.de == "financeiro"
                    else "não há prazos lidos nos documentos do Acervo"))

    abas.append(aba)
    estado.documentos.salvar(id_, planilha.para_json(abas), nota=f"aba {aba.nome} criada")
    return {**_resposta_planilha(id_, estado.documentos.obter(id_), abas),
            "aba_nova": len(abas) - 1, "quantos": quantos,
            "aviso": f"{_quantos(quantos, 'linha')} na aba “{aba.nome}”"}


def _cabecalho_da_aba(aba, titulos: list[str]) -> None:
    for coluna, titulo in enumerate(titulos):
        aba.gravar(f"{planilha.letra_da_coluna(coluna)}1", {"valor": titulo, "negrito": True})
    aba.congelar_cabecalho = True


def _aba_do_financeiro(mes: str, categoria: str):
    """Os honorarios do Financeiro em linhas: quem, quando, quanto, pago ou nao."""
    lancamentos = [
        x for x in estado.financeiro.listar(tipo="recebimento", mes=mes)
        if not categoria or x.get("categoria") == categoria
    ]
    rotulo = financeiro.CATEGORIAS.get(categoria, "Recebimentos")
    aba = planilha.Aba(nome=(rotulo + (" " + escritorio.mes_por_extenso(mes) if mes else ""))[:24])
    _cabecalho_da_aba(aba, ["Cliente", "Descrição", "Vencimento", "Valor", "Situação"])

    for i, l in enumerate(lancamentos, start=2):
        aba.gravar(f"A{i}", {"valor": l.get("cadastro_nome") or "—"})
        aba.gravar(f"B{i}", {"valor": l.get("descricao") or ""})
        aba.gravar(f"C{i}", {"valor": _data_br(l.get("vencimento")), "formato": "data"})
        # O valor vai em reais com virgula, como se digita aqui; a planilha
        # entende isso como numero e soma.
        aba.gravar(f"D{i}", {"valor": financeiro.em_reais(l.get("centavos", 0)).replace("R$ ", ""),
                             "formato": "moeda"})
        aba.gravar(f"E{i}", {"valor": "recebido" if l.get("liquidado_em") else "a receber"})

    if lancamentos:
        fim = len(lancamentos) + 1
        aba.gravar(f"C{fim + 1}", {"valor": "Total", "negrito": True})
        aba.gravar(f"D{fim + 1}", {"valor": f"=SOMA(D2:D{fim})", "formato": "moeda", "negrito": True})
    return aba, len(lancamentos)


def _aba_dos_prazos():
    """Os prazos que os documentos do Acervo pedem para conferir."""
    from classify import CacheClassificacao

    cache = CacheClassificacao(CLASSIFICACAO_PATH).dados
    prazos = estado.tarefas.sugerir(list(cache.values()))
    aba = planilha.Aba(nome="Prazos do Acervo")
    _cabecalho_da_aba(aba, ["Documento", "Cliente", "Tipo", "Prazo", "Faltam (dias)"])

    hoje = date.today()
    for i, p in enumerate(prazos, start=2):
        aba.gravar(f"A{i}", {"valor": p.get("arquivo") or ""})
        aba.gravar(f"B{i}", {"valor": p.get("cliente") or "—"})
        aba.gravar(f"C{i}", {"valor": p.get("lista") or ""})
        aba.gravar(f"D{i}", {"valor": _data_br(p.get("prazo")), "formato": "data"})
        try:
            faltam = (date.fromisoformat(p["prazo"]) - hoje).days
        except (ValueError, KeyError, TypeError):
            faltam = ""
        aba.gravar(f"E{i}", {"valor": str(faltam), "formato": "numero"})
    return aba, len(prazos)


def _data_br(iso: str | None) -> str:
    """2026-09-20 vira 20/09/2026 - a planilha guarda como se digita."""
    if not iso:
        return ""
    try:
        return date.fromisoformat(str(iso)[:10]).strftime("%d/%m/%Y")
    except ValueError:
        return str(iso)

@app.post("/api/planilha/{id_}/importar")
async def planilha_importar(id_: int, arquivo: UploadFile) -> dict:
    """Traz um CSV ou XLSX para dentro da planilha, como abas novas."""
    item = _documento_ou_404(id_)
    nome = Path(arquivo.filename or "").name
    sufixo = Path(nome).suffix.lower()
    if sufixo not in (".csv", ".xlsx", ".txt"):
        raise HTTPException(status_code=400, detail="importo CSV e XLSX")

    dados = await arquivo.read()
    if len(dados) > 12 * 1024 * 1024:
        raise HTTPException(status_code=400, detail="arquivo grande demais")

    abas = _abas_do(item)
    try:
        if sufixo == ".xlsx":
            novas = planilha.de_xlsx(dados)
        else:
            texto = dados.decode("utf-8", errors="replace")
            if texto.count("�") > len(texto) * 0.01:
                texto = dados.decode("cp1252", errors="replace")
            novas = [planilha.de_csv(texto, Path(nome).stem[:24])]
    except Exception as exc:
        raise HTTPException(status_code=400, detail=f"não consegui ler o arquivo: {exc}") from exc

    abas = (abas + novas)[:6]
    estado.documentos.salvar(id_, planilha.para_json(abas), nota=f"importado {nome}")
    return _resposta_planilha(id_, estado.documentos.obter(id_), abas)


@app.get("/api/planilha/{id_}/exportar")
def planilha_exportar(id_: int, formato: str = "xlsx", aba: int = 0):
    from fastapi.responses import Response

    item = _documento_ou_404(id_)
    abas = _abas_do(item)
    nome = _arquivo(item["titulo"])

    if formato == "csv":
        if not 0 <= aba < len(abas):
            aba = 0
        texto = planilha.para_csv(abas[aba], planilha.calcular_aba(abas[aba]))
        return Response(
            texto.encode("utf-8-sig"), media_type="text/csv",
            headers=_anexo(nome + ".csv"),
        )

    dados = planilha.para_xlsx(abas, [planilha.calcular_aba(a) for a in abas])
    return Response(
        dados,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers=_anexo(nome + ".xlsx"),
    )


@app.post("/api/planilha/{id_}/assistente")
def planilha_assistente(id_: int, payload: PedidoAoAssistente) -> dict:
    """
    Pede uma formula em portugues e recebe a formula pronta.

    O modelo escreve a formula; quem calcula e o motor da planilha. Assim um
    erro do modelo vira #NOME? numa celula, nao numero errado num honorario.
    """
    item = _documento_ou_404(id_)
    if item["tipo"] != "planilha":
        raise HTTPException(status_code=400, detail="isso não é uma planilha")
    if not payload.pedido.strip():
        raise HTTPException(status_code=400, detail="diga o que você quer calcular")

    disponivel, motivo = check_ollama(estado.client.model)
    if not disponivel:
        raise HTTPException(status_code=503, detail=motivo)

    abas = _abas_do(item)
    aba = abas[0]
    calculado = planilha.calcular_aba(aba)

    linhas = []
    for ref in sorted(aba.celulas, key=lambda r: (planilha._posicao(r)[0], planilha._posicao(r)[1]))[:60]:
        linhas.append(f"{ref}: {(calculado.get(ref) or {}).get('texto', '')}")

    instrucao = (
        "Voce escreve formulas de planilha no padrao brasileiro: ponto e virgula separa "
        "argumento e virgula e o decimal. Funcoes em portugues: SOMA, MEDIA, SE, SOMASE, "
        "CONT.SE, MAXIMO, MINIMO, ARRED. Responda APENAS com a formula, comecando por =. "
        "Nao explique. Nao invente celula que nao esta na lista."
    )
    try:
        resposta = estado.client.ask(instrucao + f"\n\nPedido: {payload.pedido}",
                                     "Células com valor:\n" + "\n".join(linhas))
    except OllamaError as exc:
        raise HTTPException(status_code=503, detail=str(exc)) from exc

    formula = _so_a_formula(resposta)
    prova = _provar_formula(formula, aba)
    return {
        "formula": formula,
        "resultado": prova["texto"],
        "erro": prova["erro"],
        "aviso": "Confira a fórmula antes de inserir: o modelo pode citar a célula errada.",
    }


def _so_a_formula(bruto: str) -> str:
    texto = (bruto or "").strip().replace("`", "")
    achado = re.search(r"=[^\n]+", texto)
    return achado.group().strip() if achado else texto.split("\n")[0].strip()


def _provar_formula(formula: str, aba) -> dict:
    """Roda a formula antes de oferecer, para a tela ja mostrar o resultado."""
    prova = planilha.Aba(nome=aba.nome, celulas=dict(aba.celulas))
    prova.gravar("ZZ999", {"valor": formula})
    calculado = planilha.calcular_aba(prova)
    celula = calculado.get("ZZ999", {})
    return {"texto": celula.get("texto", ""), "erro": bool(celula.get("erro"))}

# --------------------------------------- financeiro, relatorios, bem-estar


class FichaLancamento(BaseModel):
    id: int | None = None
    dados: dict = {}


class Liquidacao(BaseModel):
    quando: str = ""


class FichaLembrete(BaseModel):
    id: int | None = None
    dados: dict = {}


class PedidoCiclo(BaseModel):
    tarefa: str = ""
    foco: int = 0
    pausa: int = 0


class MensagemWhats(BaseModel):
    telefone: str = ""
    nome: str = ""
    texto: str = ""
    anexo: str = ""


@app.get("/api/financeiro")
def financeiro_painel(mes: str = "") -> dict:
    """
    O painel do mes.

    Todo numero aqui e soma de lancamento que existe. Sem lancamento, o numero
    e zero e a tela diz que esta vazio - grafico bonito de dado que ninguem
    digitou e a forma mais rapida de o financeiro perder a confianca de quem usa.
    """
    dados = estado.financeiro.para_tela(mes)
    dados["clientes"] = [
        {"id": f["id"], "nome": f["nome"]} for f in estado.cadastros.listar()
    ]
    dados["onde_moram"] = (
        "Os lançamentos ficam no mesmo arquivo do resto do programa, nesta máquina. "
        "Não há senha separada para o financeiro: quem abre o Windows com a sua conta "
        "vê esta tela. Uma permissão por pessoa só faria sentido num PAULUS de equipe, "
        "que ainda não existe."
    )

    # O escritorio por dentro vem na mesma resposta: sao seis blocos da mesma
    # tela, e seis chamadas fariam a tela pintar em pedacos.
    mes_alvo = dados["painel"]["mes"]
    dados["folha"] = estado.folha.do_mes(mes_alvo)
    dados["candidatos_folha"] = estado.folha.pessoas()
    dados["notas"] = estado.papeis.listar("nota", mes_alvo)
    dados["notas_a_emitir"] = estado.papeis.a_emitir(mes_alvo)
    dados["boletos"] = estado.papeis.listar("boleto")
    dados["comprovantes"] = escritorio.comprovantes_do_mes(estado.base, mes_alvo)
    dados["sem_comprovante"] = escritorio.sem_comprovante(estado.base, mes_alvo)
    dados["concluidos"] = escritorio.contratos_concluidos(estado.base, mes_alvo)
    dados["fechamento"] = escritorio.fechamento(
        estado.base, estado.folha, estado.papeis, mes_alvo)
    dados["meses"] = _meses_com_movimento()
    return dados


def _meses_com_movimento() -> list[dict]:
    """
    Os meses que o seletor oferece: os que tem lancamento, mais o de hoje.

    Oferecer doze meses fixos daria meses vazios para escolher; oferecer so os
    que tem dado esconderia o mes corrente enquanto nada foi lancado nele.
    """
    linhas = estado.base.buscar(
        "SELECT DISTINCT substr(COALESCE(NULLIF(liquidado_em,''), vencimento, ''),1,7) m "
        "FROM lancamentos WHERE m != '' ORDER BY m DESC"
    )
    meses = [l["m"] for l in linhas]
    hoje = escritorio.mes_de_hoje()
    if hoje not in meses:
        meses.insert(0, hoje)
    return [{"valor": m, "rotulo": escritorio.mes_por_extenso(m)} for m in meses[:24]]


@app.get("/api/financeiro/lancamentos")
def financeiro_listar(tipo: str = "", mes: str = "", situacao: str = "") -> dict:
    return {"lancamentos": estado.financeiro.listar(tipo=tipo, mes=mes, situacao=situacao)}


@app.get("/api/financeiro/extrato")
def financeiro_extrato(mes: str = "") -> dict:
    return estado.financeiro.extrato(mes)


# --------------------------------------------------- o extrato do banco

MAX_EXTRATO_BYTES = 8 * 1024 * 1024


@app.post("/api/financeiro/banco")
async def financeiro_banco(arquivo: UploadFile = File(...)) -> dict:
    """
    Le o extrato do banco e diz o que parece ser o que - sem gravar nada.

    Conciliar propoe; quem da baixa e a pessoa, na tela, marcando. Dar baixa
    sozinho por causa de um valor igual seria escrever no livro-caixa de
    alguem por palpite, e valor igual acontece toda semana num escritorio que
    cobra honorario fixo.
    """
    nome = Path(arquivo.filename or "").name
    dados = await arquivo.read(MAX_EXTRATO_BYTES + 1)
    if len(dados) > MAX_EXTRATO_BYTES:
        raise HTTPException(status_code=413, detail="esse arquivo passa de 8 MB")

    try:
        movimentos = extrato.ler(dados, nome)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc

    lancamentos: list[dict] = []
    for mes in extrato.meses_do_extrato(movimentos):
        lancamentos += estado.financeiro.listar(mes=mes)
    vistos, unicos = set(), []
    for l in lancamentos:
        if l["id"] not in vistos:
            vistos.add(l["id"])
            unicos.append(l)

    pares = extrato.conciliar(movimentos, unicos)
    return {"arquivo": nome, "pares": [p.to_dict() for p in pares],
            "resumo": extrato.resumo(pares),
            "categorias": [{"valor": k, "rotulo": v} for k, v in financeiro.CATEGORIAS.items()]}


class ConciliarEscolha(BaseModel):
    # Os pares que a pessoa marcou: lancamento que recebe a baixa e a data do
    # banco. Movimento sem lancamento pode virar lancamento novo.
    baixas: list[dict] = []
    novos: list[dict] = []


@app.post("/api/financeiro/banco/aplicar")
def financeiro_banco_aplicar(payload: ConciliarEscolha) -> dict:
    """Da baixa no que foi marcado e lanca o que a pessoa mandou lancar."""
    baixados, criados, erros = 0, 0, []

    for item in payload.baixas:
        id_ = int(item.get("lancamento_id") or 0)
        quando = str(item.get("data") or "")[:10]
        if not id_:
            continue
        if estado.financeiro.liquidar(id_, quando):
            baixados += 1
        else:
            erros.append(f"lançamento {id_} não estava aberto")

    for item in payload.novos:
        dados = {
            "tipo": "recebimento" if int(item.get("centavos", 0)) >= 0 else "despesa",
            "descricao": str(item.get("descricao", ""))[:160] or "Do extrato do banco",
            "centavos": abs(int(item.get("centavos", 0))),
            "categoria": str(item.get("categoria") or "outros"),
            "cadastro_id": item.get("cadastro_id"),
            "vencimento": str(item.get("data") or "")[:10],
            "liquidado_em": str(item.get("data") or "")[:10],
            "observacao": "Lançado a partir do extrato do banco.",
        }
        try:
            estado.financeiro.salvar(dados)
            criados += 1
        except ValueError as exc:
            erros.append(str(exc))

    partes = []
    if baixados:
        partes.append(_quantos(baixados, "baixa"))
    if criados:
        partes.append(_quantos(criados, "lançamento") + " novo" + ("s" if criados > 1 else ""))
    aviso = " e ".join(partes) + " a partir do extrato" if partes else "nada foi marcado"
    return {"baixados": baixados, "criados": criados, "erros": erros, "aviso": aviso,
            **estado.financeiro.para_tela()}


@app.post("/api/financeiro/lancamentos")
def financeiro_salvar(payload: FichaLancamento) -> dict:
    try:
        id_ = estado.financeiro.salvar(payload.dados, payload.id)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    return estado.financeiro.obter(id_) or {}


@app.post("/api/financeiro/lancamentos/{id_}/liquidar")
def financeiro_liquidar(id_: int, payload: Liquidacao) -> dict:
    if not estado.financeiro.liquidar(id_, payload.quando):
        raise HTTPException(status_code=404, detail="lançamento não encontrado")
    return estado.financeiro.obter(id_) or {}


@app.post("/api/financeiro/lancamentos/{id_}/reabrir")
def financeiro_reabrir(id_: int) -> dict:
    if not estado.financeiro.reabrir(id_):
        raise HTTPException(status_code=404, detail="lançamento não encontrado")
    return estado.financeiro.obter(id_) or {}


@app.delete("/api/financeiro/lancamentos/{id_}")
def financeiro_apagar(id_: int) -> dict:
    l = estado.base.um("SELECT descricao, centavos, vencimento FROM lancamentos WHERE id = ?", (id_,))
    if not l:
        raise HTTPException(status_code=404, detail="lançamento não encontrado")
    reais = f"R$ {l['centavos'] / 100:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    entrada = estado.lixeira.apagar_linha("lancamento", id_, l["descricao"], "Financeiro · " + reais + (" · " + l["vencimento"] if l["vencimento"] else ""))
    return _foi_para_lixeira(entrada, "apagado", id_)


@app.post("/api/financeiro/lancamentos/{id_}/comprovante")
async def financeiro_comprovante(id_: int, arquivo: UploadFile) -> dict:
    """Guarda o comprovante na biblioteca e o liga ao lancamento."""
    if not estado.financeiro.obter(id_):
        raise HTTPException(status_code=404, detail="lançamento não encontrado")

    nome = Path(arquivo.filename or "").name
    if not nome:
        raise HTTPException(status_code=400, detail="arquivo sem nome")

    dados = await arquivo.read()
    if len(dados) > MAX_UPLOAD_BYTES:
        raise HTTPException(status_code=400, detail="arquivo grande demais")

    pasta = BASE_DIR / "data" / "comprovantes"
    pasta.mkdir(parents=True, exist_ok=True)
    destino = pasta / nome
    conta = 2
    while destino.exists():
        destino = pasta / f"{Path(nome).stem} ({conta}){Path(nome).suffix}"
        conta += 1
    destino.write_bytes(dados)

    import hashlib

    sha1 = hashlib.sha1(dados).hexdigest()
    estado.financeiro.anexar(id_, destino.name, str(destino), sha1)
    return estado.financeiro.obter(id_) or {}


@app.delete("/api/financeiro/comprovante/{id_}")
def financeiro_tirar_comprovante(id_: int) -> dict:
    estado.financeiro.tirar_comprovante(id_)
    return {"removido": id_}


@app.post("/api/financeiro/cobrar")
def financeiro_cobrar(payload: dict) -> dict:
    """
    Prepara a cobranca de um recebimento atrasado como rascunho de e-mail.

    Nao envia: monta o texto e devolve. O envio segue o caminho de sempre, com
    a fila de aprovacao no meio.
    """
    item = estado.financeiro.obter(int(payload.get("id", 0)))
    if not item:
        raise HTTPException(status_code=404, detail="lançamento não encontrado")

    ficha = estado.cadastros.obter(item["cadastro_id"]) if item["cadastro_id"] else None
    vencimento = item["vencimento"]
    atraso = item["dias_atraso"]

    corpo = (
        f"Prezados,\n\n"
        f"Consta em nosso controle o valor de {item['valor']} referente a "
        f"{item['descricao']}"
    )
    if vencimento:
        corpo += f", com vencimento em {financeiro._br(vencimento)}"
        if atraso:
            corpo += f" — portanto há {atraso} dia(s)"
    corpo += (
        ".\n\nCaso o pagamento já tenha sido efetuado, favor desconsiderar e nos "
        "encaminhar o comprovante.\n\nFico à disposição."
    )

    return {
        "para": (ficha or {}).get("email", ""),
        "assunto": f"Cobrança — {item['descricao']}",
        "corpo": corpo,
        "cliente": (ficha or {}).get("nome", ""),
        "sem_email": not (ficha or {}).get("email"),
    }


# ------------------------------------------------------------ relatorios


@app.get("/api/relatorios")
def relatorios_ver(quando: str = "") -> dict:
    return estado.relatorios.para_tela(quando)


@app.get("/api/relatorios/acoes")
def relatorios_acoes(limite: int = 40) -> dict:
    """O que o assistente fez com o seu sim, lido da fila de aprovacoes."""
    return {"acoes": estado.relatorios.acoes(limite=limite)}


@app.post("/api/relatorios/parecer")
def relatorios_parecer(payload: dict | None = None) -> dict:
    """
    O parecer em texto sobre o dia.

    O modelo recebe os numeros ja apurados e escreve sobre eles. Nao calcula:
    um erro de aritmetica de um modelo de 3 bilhoes de parametros viraria
    numero errado num parecer financeiro, e parecer com numero errado e pior
    que nenhum parecer.
    """
    quando = str((payload or {}).get("quando", ""))
    numeros = estado.relatorios.numeros_para_o_parecer(quando)

    disponivel, motivo = check_ollama(estado.client.model)
    if not disponivel:
        raise HTTPException(status_code=503, detail=motivo)

    try:
        texto = estado.client.ask(relatorios.INSTRUCAO_PARECER, numeros)
    except OllamaError as exc:
        raise HTTPException(status_code=503, detail=str(exc)) from exc

    return {
        "parecer": _limpar_sugestao(texto),
        "numeros": numeros,
        "aviso": "Escrito sobre os números acima, que foram calculados aqui. Confira antes de usar.",
    }


@app.post("/api/relatorios/pdf")
def relatorios_pdf(payload: dict | None = None):
    """O relatorio do dia como PDF, pelo mesmo gerador dos documentos."""
    from fastapi.responses import Response

    quando = str((payload or {}).get("quando", "")) or datetime.now().strftime("%Y-%m-%d")
    d = estado.relatorios.dia(quando)
    m = estado.relatorios.mes(quando[:7])

    html = [f"<h1>Relatório de {d['rotulo']}</h1>"]
    t = d["tarefas"]
    html.append(f"<h2>Tarefas</h2><p>{len(t['feitas'])} de {t['planejado']} concluída(s).</p>")
    if t["feitas"]:
        html.append("<ul>" + "".join(
            f"<li>{_escapar(x['titulo'])} — {x['hora']}</li>" for x in t["feitas"]) + "</ul>")
    if t["abertas"]:
        html.append("<p>Ficaram abertas:</p><ul>" + "".join(
            f"<li>{_escapar(x['titulo'])} ({_escapar(x['quando'])})</li>" for x in t["abertas"]) + "</ul>")

    if d["medido"]:
        html.append(f"<h2>Tempo</h2><p>Ativo por {d['medido']['tempo_ativo']}, "
                    f"{d['medido']['pausas']} pausa(s), {d['medido']['ciclos']} ciclo(s) de foco.</p>")

    if d["acoes"]:
        html.append("<h2>Ações aprovadas por você</h2><ul>" + "".join(
            f"<li>{_escapar(a['titulo'])} — {_escapar(a['resultado'] or a['estado'])} ({a['hora']})</li>"
            for a in d["acoes"]) + "</ul>")

    e = m["extrato"]
    html.append(f"<h2>Financeiro — {_escapar(m['rotulo'])}</h2>"
                f"<p>Entradas {m['entradas_texto']}, saídas {m['saidas_texto']}, "
                f"resultado {e['resultado_texto']}. Saldo em caixa: {e['saldo_final_texto']}.</p>")
    if e["linhas"]:
        html.append("<ul>" + "".join(
            f"<li>{l['dia']} — {_escapar(l['descricao'])}: {l['movimento_texto']}</li>"
            for l in e["linhas"]) + "</ul>")

    html.append("<p><i>Relatório montado nesta máquina. Nenhum dado saiu daqui.</i></p>")

    pdf = documento.para_pdf(
        documento.ler_html("".join(html)),
        f"Relatório de {quando}",
        "PAULUS · relatório gerado nesta máquina",
    )
    return Response(pdf, media_type="application/pdf",
                    headers=_anexo(f"Relatorio {quando}.pdf"))


def _escapar(texto: str) -> str:
    from xml.sax.saxutils import escape

    return escape(str(texto or ""))


# ------------------------------------------------------------ bem-estar


@app.get("/api/bemestar")
def bemestar_ver() -> dict:
    dados = estado.bem_estar.para_tela()
    dados["sugeridos_criados"] = estado.bem_estar.sugerir_lembretes()
    if dados["sugeridos_criados"]:
        dados["lembretes"] = estado.bem_estar.lembretes()
    return dados


@app.get("/api/avisos")
def avisos_ver() -> dict:
    """Se da para avisar no Windows, se a pessoa quer e em que horario."""
    d = estado.prefs.dados.get("disponibilidade") or {}
    nomes = ["seg", "ter", "qua", "qui", "sex", "sáb", "dom"]
    dias = sorted(d.get("dias") or [0, 1, 2, 3, 4])
    faixa = (f"{nomes[dias[0]]} a {nomes[dias[-1]]}" if dias == list(range(dias[0], dias[-1] + 1))
             else ", ".join(nomes[i] for i in dias))
    return {
        "disponivel": avisos.disponivel(),
        "ligado": estado.vigia.ligado(),
        "tipos": [{"chave": k, "rotulo": t["rotulo"], "explica": t["explica"], "so_fora": t["so_fora"]}
                  for k, t in avisos.TIPOS.items()],
        "horario": f"{faixa}, das {d.get('inicio') or '00:00'} às {d.get('fim') or '23:59'}",
    }


@app.post("/api/avisos/teste")
def avisos_teste() -> dict:
    """Um aviso de verdade, para a pessoa ver como chega - espera o Windows responder."""
    if not avisos.disponivel():
        raise HTTPException(status_code=400, detail="avisos do Windows só existem no Windows")
    ok = avisos.notificar("Avisos ligados", "É assim que os lembretes e o fim do ciclo de foco vão chegar.", esperar=True)
    return {"ok": ok}


@app.get("/api/bemestar/semana")
def bemestar_semana(ate: str = "") -> dict:
    """Uma semana qualquer, para comparar a atual com a anterior."""
    return estado.bem_estar.semana(ate)


@app.post("/api/bemestar/medir")
def bemestar_medir(payload: dict) -> dict:
    """
    Liga ou desliga o acompanhamento.

    Desligar e opcao de verdade: a thread para e nada mais e gravado. O ciclo
    de foco e os lembretes continuam funcionando sem ele.
    """
    if payload.get("ligar"):
        if not estado.bem_estar.ligar():
            raise HTTPException(
                status_code=400,
                detail="não consigo medir a atividade neste sistema - o ciclo de foco e os lembretes continuam funcionando",
            )
    else:
        estado.bem_estar.desligar()
    return estado.bem_estar.para_tela()


@app.post("/api/bemestar/ciclo")
def bemestar_ciclo(payload: PedidoCiclo) -> dict:
    return estado.bem_estar.comecar_ciclo(payload.tarefa, payload.foco, payload.pausa)


@app.post("/api/bemestar/pausa")
def bemestar_pausa() -> dict:
    return estado.bem_estar.ir_para_pausa()


@app.post("/api/bemestar/parar")
def bemestar_parar() -> dict:
    return estado.bem_estar.parar_ciclo()


@app.get("/api/bemestar/ciclo")
def bemestar_ciclo_estado() -> dict:
    return {**estado.bem_estar.estado_do_ciclo(), "alerta": estado.bem_estar.alerta()}


@app.post("/api/bemestar/lembretes")
def bemestar_salvar_lembrete(payload: FichaLembrete) -> dict:
    try:
        estado.bem_estar.salvar_lembrete(payload.dados, payload.id)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    return {"lembretes": estado.bem_estar.lembretes()}


@app.post("/api/bemestar/lembretes/restaurar")
def bemestar_restaurar_lembretes() -> dict:
    """Devolve a lista de fabrica para quem se perdeu editando."""
    quantos = estado.bem_estar.restaurar_lembretes()
    return {"lembretes": estado.bem_estar.lembretes(), "restaurados": quantos}


@app.post("/api/bemestar/lembretes/{id_}/feito")
def bemestar_marcar(id_: int) -> dict:
    if not estado.bem_estar.marcar_lembrete(id_):
        raise HTTPException(status_code=404, detail="lembrete não encontrado")
    return {"lembretes": estado.bem_estar.lembretes(), "hoje": estado.bem_estar.dia()}


@app.delete("/api/bemestar/lembretes/{id_}")
def bemestar_apagar_lembrete(id_: int) -> dict:
    estado.bem_estar.apagar_lembrete(id_)
    return {"lembretes": estado.bem_estar.lembretes()}


# ------------------------------------------------------------- servicos


class FichaServico(BaseModel):
    id: int | None = None
    dados: dict = {}


@app.get("/api/servicos")
def servicos_listar(filtro: str = "", termo: str = "") -> dict:
    return {
        "servicos": estado.servicos.listar(filtro, termo),
        "contagem": estado.servicos.contagem(),
        "status": [{"valor": k, "rotulo": v} for k, v in servicos_mod.STATUS.items()],
        "clientes": [{"id": f["id"], "nome": f["nome"], "tipo": f["tipo"], "observacao": f.get("observacao", "")}
                     for f in estado.cadastros.listar()],
    }


# Arquivos da pasta de um servico que o indice ja tentou ler e nao conseguiu
# (PDF escaneado, arquivo corrompido): caminho e data de modificacao. Sem
# isso, cada abertura da pasta mandaria reler o Acervo atras deles.
_TENTADOS_NA_PASTA: set[tuple[str, float]] = set()


def _chave_do_arquivo(p: Path) -> tuple[str, float]:
    try:
        return (os.path.normcase(str(p.resolve())), p.stat().st_mtime)
    except OSError:
        return (os.path.normcase(str(p)), 0.0)


def _documentos_por_caminho() -> dict:
    return {os.path.normcase(str(Path(d.path).resolve())): d for d in estado.searcher.documents}


def _sincronizar_pasta_do_servico(id_: int) -> Path | None:
    """
    O que esta na pasta do servico no disco entra no servico.

    Arquivo posto na pasta pelo Windows ainda nao foi lido: o Acervo e relido
    uma vez e o arquivo, ja com sha1, e ligado. Arquivo que alguem desligou
    de proposito nao volta sozinho.
    """
    pasta = estado.servicos.pasta_de(id_, criar=True)
    if not pasta or not pasta.is_dir():
        return pasta
    arquivos = [p for p in pasta.rglob("*") if p.is_file() and p.suffix.lower() in SUPPORTED_SUFFIXES]
    if not arquivos:
        return pasta
    por_caminho = _documentos_por_caminho()
    novos = [p for p in arquivos
             if _chave_do_arquivo(p)[0] not in por_caminho and _chave_do_arquivo(p) not in _TENTADOS_NA_PASTA]
    if novos:
        estado.recarregar()
        por_caminho = _documentos_por_caminho()
        _TENTADOS_NA_PASTA.update(_chave_do_arquivo(p) for p in novos if _chave_do_arquivo(p)[0] not in por_caminho)
    ligados = {a["sha1"] for a in estado.servicos.arquivos_de(id_)}
    desligados = estado.servicos.desligados(id_)
    for p in arquivos:
        doc = por_caminho.get(_chave_do_arquivo(p)[0])
        if doc and doc.sha1 not in ligados and doc.sha1 not in desligados:
            estado.servicos.vincular(id_, doc.sha1, doc.name, "Pasta do serviço")
            ligados.add(doc.sha1)
    return pasta


@app.post("/api/servicos")
def servicos_salvar(payload: FichaServico) -> dict:
    antes = estado.servicos.pasta_de(payload.id, criar=False) if payload.id else None
    try:
        id_ = estado.servicos.salvar(payload.dados, payload.id)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    # Renomear o servico renomeou a pasta: os caminhos no indice e os dos
    # audios que moram nela mudaram.
    depois = estado.servicos.pasta_de(id_, criar=False)
    if antes and depois and antes != depois and depois.is_dir():
        estado.gravacoes.trocar_pasta(antes, depois)
        if any(depois.iterdir()):
            estado.recarregar()
    return servicos_obter(id_)


@app.get("/api/servicos/{id_}")
def servicos_obter(id_: int) -> dict:
    if not estado.base.um("SELECT id FROM servicos WHERE id = ?", (id_,)):
        raise HTTPException(status_code=404, detail="serviço não encontrado")
    pasta = _sincronizar_pasta_do_servico(id_)
    s = estado.servicos.obter(id_)
    if not s:
        raise HTTPException(status_code=404, detail="serviço não encontrado")
    s["gravacoes"] = [g for g in estado.gravacoes.listar() if g.get("servico_id") == id_]
    s["pasta_caminho"] = str(pasta) if pasta else ""
    return s


@app.delete("/api/servicos/{id_}")
def servicos_apagar(id_: int) -> dict:
    s = estado.servicos.obter(id_)
    if not s:
        raise HTTPException(status_code=404, detail="serviço não encontrado")
    pasta = estado.servicos.pasta_de(id_, criar=False)
    entrada = estado.lixeira.apagar_linha("servico", id_, s["nome"], "Serviços" + (" · " + s["cliente_nome"] if s.get("cliente_nome") else ""))
    # A pasta vazia sai; com arquivos, fica no Acervo - os arquivos sao do
    # escritorio, nao do servico.
    servicos_mod.remover_pasta_vazia(pasta)
    return _foi_para_lixeira(entrada, "apagado", id_)


class AnexarAoServico(BaseModel):
    caminhos: list[str]


def _nome_livre_de_arquivo(pasta: Path, nome: str) -> Path:
    alvo = pasta / nome
    n = 2
    while alvo.exists():
        alvo = pasta / f"{Path(nome).stem} ({n}){Path(nome).suffix}"
        n += 1
    return alvo


@app.post("/api/servicos/{id_}/anexar")
def servicos_anexar(id_: int, payload: AnexarAoServico) -> dict:
    """
    Adicionar arquivos ao servico: copia cada um para a pasta do servico no
    Acervo e liga a copia.

    Vale para o que vem do Acervo e do computador. O original fica onde
    estava; o que ja esta na pasta nao e copiado sobre si mesmo, e o mesmo
    conteudo com o mesmo nome nao vira "(2)".
    """
    pasta = estado.servicos.pasta_de(id_, criar=True)
    if not pasta:
        raise HTTPException(status_code=404, detail="serviço não encontrado")
    destinos: list[Path] = []
    recusados: list[dict] = []
    for bruto in payload.caminhos:
        origem = Path(bruto)
        nome = origem.name
        if not origem.is_file():
            recusados.append({"nome": nome or bruto, "motivo": "arquivo não encontrado"})
            continue
        if origem.suffix.lower() not in SUPPORTED_SUFFIXES:
            recusados.append({"nome": nome, "motivo": "formato não suportado"})
            continue
        try:
            if origem.stat().st_size > MAX_UPLOAD_BYTES:
                recusados.append({"nome": nome, "motivo": "arquivo maior que 50 MB"})
                continue
            if origem.resolve().is_relative_to(pasta.resolve()):
                destino = origem
            else:
                destino = pasta / nome
                if destino.exists() and file_sha1(destino) != file_sha1(origem):
                    destino = _nome_livre_de_arquivo(pasta, nome)
                if not destino.exists():
                    shutil.copy2(origem, destino)
        except OSError as exc:
            recusados.append({"nome": nome, "motivo": f"não consegui copiar: {exc.strerror or exc}"})
            continue
        destinos.append(destino)
    if destinos:
        estado.recarregar()
    por_caminho = _documentos_por_caminho()
    ligados: list[str] = []
    for destino in destinos:
        doc = por_caminho.get(_chave_do_arquivo(destino)[0])
        if not doc:
            recusados.append({"nome": destino.name, "motivo": "sem texto para ler (PDF escaneado?)"})
            continue
        estado.servicos.vincular(id_, doc.sha1, doc.name)
        ligados.append(doc.name)
    return {"ligados": ligados, "recusados": recusados, "pasta": str(pasta),
            "arquivos": estado.servicos.arquivos_de(id_)}


@app.post("/api/servicos/{id_}/status")
def servicos_status(id_: int, payload: dict) -> dict:
    try:
        estado.servicos.mudar_status(id_, str(payload.get("status", "")))
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    return estado.servicos.obter(id_) or {}


@app.post("/api/servicos/{id_}/etapas")
def servicos_etapa_nova(id_: int, payload: dict) -> dict:
    try:
        return {"etapas": estado.servicos.etapa_adicionar(id_, str(payload.get("titulo", "")), str(payload.get("quando", "")),
                                                          payload.get("responsavel_id"))}
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc


@app.post("/api/servicos/{id_}/etapas/{indice}/editar")
def servicos_etapa_editar(id_: int, indice: int, payload: dict) -> dict:
    """O prazo, o nome ou quem cuida de uma etapa - so o que vier no corpo."""
    try:
        return {"etapas": estado.servicos.etapa_editar(id_, indice, payload)}
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc


@app.post("/api/servicos/{id_}/etapas/{indice}")
def servicos_etapa_alternar(id_: int, indice: int) -> dict:
    try:
        return {"etapas": estado.servicos.etapa_alternar(id_, indice, _quem_sou())}
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc


@app.delete("/api/servicos/{id_}/etapas/{indice}")
def servicos_etapa_remover(id_: int, indice: int) -> dict:
    try:
        return {"etapas": estado.servicos.etapa_remover(id_, indice)}
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc


@app.post("/api/servicos/{id_}/anotacoes")
def servicos_anotar(id_: int, payload: dict) -> dict:
    try:
        return {"anotacoes": estado.servicos.anotar(id_, str(payload.get("texto", "")), _quem_sou())}
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc


@app.post("/api/servicos/{id_}/anotacoes/{indice}")
def servicos_anotacao_editar(id_: int, indice: int, payload: dict) -> dict:
    try:
        return {"anotacoes": estado.servicos.anotacao_editar(id_, indice, str(payload.get("texto", "")), _quem_sou())}
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc


@app.delete("/api/servicos/{id_}/anotacoes/{indice}")
def servicos_anotacao_remover(id_: int, indice: int) -> dict:
    try:
        return {"anotacoes": estado.servicos.anotacao_remover(id_, indice)}
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc


@app.post("/api/servicos/{id_}/prazos/{origem}/{item_id}")
def servicos_prazo_editar(id_: int, origem: str, item_id: int, payload: dict) -> dict:
    """
    Trocar o titulo, a data ou a hora de um prazo da pasta, sem mexer no
    resto da ficha. O compromisso e a tarefa sao gravados inteiros (a Agenda
    e Tarefas regravam todos os campos): aqui a ficha atual e lida e so o que
    veio no corpo muda.
    """
    if not estado.servicos.obter(id_):
        raise HTTPException(status_code=404, detail="serviço não encontrado")
    titulo = " ".join(str(payload.get("titulo", "")).split()) if "titulo" in payload else None
    if titulo == "":
        raise HTTPException(status_code=400, detail="o prazo precisa de um título")
    try:
        if origem == "compromisso":
            atual = estado.agenda.obter(item_id)
            if not atual:
                raise HTTPException(status_code=404, detail="compromisso não encontrado")
            dados = {c: atual.get(c) for c in CAMPOS_AGENDA}
            if titulo is not None:
                dados["titulo"] = titulo
            if payload.get("data"):
                dados["data"] = servicos_mod.data_de_prazo(str(payload["data"])) or dados["data"]
            if payload.get("hora"):
                dados["hora"] = str(payload["hora"])[:5]
            estado.agenda.salvar(dados, item_id)
        elif origem == "tarefa":
            atual = estado.tarefas.obter(item_id)
            if not atual:
                raise HTTPException(status_code=404, detail="tarefa não encontrada")
            dados = {c: atual.get(c) for c in tarefas_mod.CAMPOS}
            if titulo is not None:
                dados["titulo"] = titulo
            if payload.get("data"):
                dados["prazo"] = servicos_mod.data_de_prazo(str(payload["data"])) or dados["prazo"]
            estado.tarefas.salvar(dados, item_id)
        else:
            raise HTTPException(status_code=400, detail="prazo desconhecido")
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    return {"prazos": (estado.servicos.obter(id_) or {}).get("prazos", [])}


@app.post("/api/servicos/{id_}/vincular")
def servicos_vincular(id_: int, payload: VinculoDoc) -> dict:
    if not estado.servicos.obter(id_):
        raise HTTPException(status_code=404, detail="serviço não encontrado")
    estado.servicos.vincular(id_, payload.sha1, payload.nome)
    return {"arquivos": estado.servicos.arquivos_de(id_)}


@app.delete("/api/servicos/{id_}/vinculos/{sha1}")
def servicos_desvincular(id_: int, sha1: str) -> dict:
    estado.servicos.desvincular(id_, sha1)
    return {"arquivos": estado.servicos.arquivos_de(id_)}


@app.post("/api/servicos/{id_}/resumo")
def servicos_resumo(id_: int) -> dict:
    """
    O resumo do servico, escrito pelo modelo sobre o que esta gravado.

    O modelo recebe o servico inteiro em texto - etapas, prazos, anotacoes,
    nomes dos arquivos - e escreve sobre isso. Nao le os arquivos aqui: isso
    e a conversa do Assistente, que tem o Acervo.
    """
    s = estado.servicos.obter(id_)
    if not s:
        raise HTTPException(status_code=404, detail="serviço não encontrado")
    disponivel, motivo = check_ollama(estado.client.model)
    if not disponivel:
        raise HTTPException(status_code=503, detail=motivo)
    try:
        texto = estado.client.ask(servicos_mod.INSTRUCAO_RESUMO, estado.servicos.texto_para_resumo(s))
    except OllamaError as exc:
        raise HTTPException(status_code=503, detail=str(exc)) from exc
    estado.servicos.guardar_resumo(id_, _limpar_sugestao(texto))
    return estado.servicos.obter(id_) or {}


@app.post("/api/servicos/{id_}/conversar")
def servicos_conversar(id_: int, payload: dict) -> dict:
    """
    Conversar sobre a pasta, da caixa de pedido dela.

    O modelo recebe o que esta gravado no servico - e as ultimas perguntas,
    para a conversa continuar -, responde, e a pergunta com a resposta entram
    no historico da pasta. Como o resumo, nao le os arquivos: isso e a
    conversa do Assistente, que tem o Acervo.
    """
    pergunta = " ".join(str(payload.get("pergunta", "")).split())
    if not pergunta:
        raise HTTPException(status_code=400, detail="escreva a pergunta")
    s = estado.servicos.obter(id_)
    if not s:
        raise HTTPException(status_code=404, detail="serviço não encontrado")
    disponivel, motivo = check_ollama(estado.client.model)
    if not disponivel:
        raise HTTPException(status_code=503, detail=motivo)
    try:
        resposta = estado.client.ask(servicos_mod.INSTRUCAO_CONVERSA + f"\n\nPergunta: {pergunta}",
                                     estado.servicos.texto_para_conversa(s))
    except OllamaError as exc:
        raise HTTPException(status_code=503, detail=str(exc)) from exc
    estado.servicos.conversar(id_, pergunta, _limpar_sugestao(resposta), _quem_sou())
    return estado.servicos.obter(id_) or {}


# ------------------------------------------------------------------ voz


def _situacao_da_voz() -> dict:
    s = estado.transcritor.situacao()
    s["fila"] = list(estado.fila_voz.queue)
    s["transcrevendo"] = estado.transcrevendo
    s["erro"] = estado.erro_voz
    return s


@app.get("/api/voz")
def voz_situacao() -> dict:
    return _situacao_da_voz()


@app.post("/api/voz/baixar")
def voz_baixar(payload: dict) -> dict:
    """Baixa o modelo de voz em segundo plano. Uma vez, com internet; depois nunca mais."""
    nome = str(payload.get("modelo") or estado.transcritor.modelo)
    if nome not in transcricao_mod.MODELOS:
        raise HTTPException(status_code=400, detail="modelo de voz desconhecido")
    if estado.transcritor.baixando:
        return _situacao_da_voz()
    if estado.transcritor.instalado(nome):
        return _situacao_da_voz()

    def baixar() -> None:
        estado.erro_voz = ""
        try:
            estado.transcritor.baixar(nome)
        except Exception as exc:  # noqa: BLE001 - o erro vai para a tela
            estado.erro_voz = "não consegui baixar: " + str(exc)[:200]

    estado.transcritor.baixando = nome
    threading.Thread(target=baixar, name="baixar-voz", daemon=True).start()
    return _situacao_da_voz()


@app.post("/api/voz/modelo")
def voz_modelo(payload: dict) -> dict:
    nome = str(payload.get("modelo") or "")
    try:
        estado.transcritor.escolher(nome)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    estado.prefs.atualizar({"voz": {"modelo": nome}})
    return _situacao_da_voz()


def _purgar_sessoes() -> None:
    """Sessao sem audio ha 15 minutos e uma gravacao que nao voltou: some."""
    agora = time.time()
    for sid in [s for s, v in estado.ao_vivo.items() if agora - v.ultima > 900]:
        estado.ao_vivo.pop(sid, None)


def _sessao_viva(sid: str) -> transcricao_mod.SessaoAoVivo:
    s = estado.ao_vivo.get(sid)
    if not s:
        raise HTTPException(status_code=404, detail="sessão de transcrição não encontrada")
    return s


@app.post("/api/voz/ao-vivo")
def voz_ao_vivo_abrir(payload: dict | None = None) -> dict:
    """Abre uma sessao para transcrever enquanto grava. O audio vem em pedacos por /audio."""
    _purgar_sessoes()
    if not estado.transcritor.instalado():
        raise HTTPException(status_code=503, detail="o modelo de voz não está baixado nesta máquina")
    sessao = transcricao_mod.SessaoAoVivo(uuid.uuid4().hex, estado.transcritor.modelo)
    estado.ao_vivo[sessao.id] = sessao
    s = sessao.situacao()
    s["rotulo"] = transcricao_mod.MODELOS[sessao.modelo]["rotulo"]
    return s


@app.post("/api/voz/ao-vivo/{sid}/audio")
async def voz_ao_vivo_audio(sid: str, request: Request) -> dict:
    """Um pedaco de audio (PCM 16 kHz, 16 bits, mono). Devolve os trechos novos, se houve pausa para transcrever."""
    sessao = _sessao_viva(sid)
    if sessao.fechada:
        raise HTTPException(status_code=409, detail="a sessão já foi encerrada")
    pcm = await request.body()
    if len(pcm) < 2:
        return {"trechos": [], **sessao.situacao()}
    try:
        novos = await run_in_threadpool(estado.transcritor.ao_vivo_receber, sessao, pcm[: len(pcm) - len(pcm) % 2])
    except Exception as exc:  # noqa: BLE001 - o erro vai para a tela, a sessao continua
        raise HTTPException(status_code=500, detail="a transcrição falhou neste pedaço: " + str(exc)[:200]) from exc
    return {"trechos": novos, **sessao.situacao()}


@app.post("/api/voz/ao-vivo/{sid}/fim")
async def voz_ao_vivo_fim(sid: str) -> dict:
    """A gravacao parou: transcreve o que sobrou e devolve tudo. A sessao fica ate a gravacao ser arquivada."""
    sessao = _sessao_viva(sid)
    await run_in_threadpool(estado.transcritor.ao_vivo_fim, sessao)
    return {"trechos": sessao.trechos, **sessao.situacao()}


@app.delete("/api/voz/ao-vivo/{sid}")
def voz_ao_vivo_descartar(sid: str) -> dict:
    sessao = estado.ao_vivo.pop(sid, None)
    if sessao:
        sessao.fechada = True
    return {"descartada": sid}


# ------------------------------------------------------------ gravacoes


def _com_progresso(g: dict) -> dict:
    if g.get("transcricao_estado") == "transcrevendo":
        g["progresso"] = round(estado.progresso_voz.get(int(g["id"]), 0.0), 3)
    elif g.get("transcricao_estado") == "fila":
        fila = list(estado.fila_voz.queue)
        g["posicao_na_fila"] = (fila.index(int(g["id"])) + 1) if int(g["id"]) in fila else 0
    return g


@app.get("/api/gravacoes")
def gravacoes_listar(tipo: str = "", termo: str = "") -> dict:
    return {
        "gravacoes": [_com_progresso(g) for g in estado.gravacoes.listar(tipo, termo)],
        "total_segundos": estado.gravacoes.total_segundos(),
        "tipos": [{"valor": k, "rotulo": v} for k, v in gravacoes_mod.TIPOS.items()],
        "clientes": [{"id": f["id"], "nome": f["nome"]} for f in estado.cadastros.listar("cliente")],
        "servicos": [{"id": s["id"], "nome": s["nome"]} for s in estado.servicos.listar("andamento")],
        "pessoas": _pessoas_para_gravacao(),
        "pasta": str(GRAVACOES_DIR),
    }


def _pessoas_para_gravacao() -> list[dict]:
    """
    Os nomes que o campo Participantes sugere: as fichas de Cadastros (equipe,
    clientes, fornecedores) e quem ja participou de outra gravacao. Nome
    repetido aparece uma vez, com o que se sabe dele.
    """
    rotulos = {"colaborador": "equipe", "socio": "sócio", "cliente": "cliente"}
    vistos: dict[str, dict] = {}
    for f in estado.cadastros.listar():
        nome = " ".join(str(f.get("nome", "")).split())
        if nome and nome.casefold() not in vistos:
            vistos[nome.casefold()] = {"nome": nome, "detalhe": f.get("observacao") or rotulos.get(f.get("tipo", ""), f.get("tipo", ""))}
    for g in estado.gravacoes.listar():
        for nome in g.get("participantes_lista", []):
            if nome.casefold() not in vistos:
                vistos[nome.casefold()] = {"nome": nome, "detalhe": "participou de uma gravação"}
    return sorted(vistos.values(), key=lambda p: p["nome"].casefold())


@app.post("/api/gravacoes")
async def gravacoes_guardar(
    arquivo: UploadFile,
    titulo: str = Form(""), tipo: str = Form("reuniao"), cadastro_id: str = Form(""),
    servico_id: str = Form(""), participantes: str = Form(""), duracao_s: str = Form("0"),
    origem: str = Form("gravada"), marcadores: str = Form("[]"), transcrever: str = Form("1"),
    sessao: str = Form(""),
) -> dict:
    """O audio entra por aqui, gravado no navegador ou importado de um arquivo."""
    conteudo = await arquivo.read()
    if not conteudo:
        raise HTTPException(status_code=400, detail="o áudio veio vazio")
    if len(conteudo) > MAX_AUDIO_BYTES:
        raise HTTPException(status_code=400, detail="áudio maior que 500 MB")
    try:
        lista = json.loads(marcadores or "[]")
    except json.JSONDecodeError:
        lista = []
    try:
        id_ = estado.gravacoes.guardar({
            "titulo": titulo, "tipo": tipo, "cadastro_id": int(cadastro_id) if cadastro_id.isdigit() else None,
            "servico_id": int(servico_id) if servico_id.isdigit() else None, "participantes": participantes,
            "duracao_s": int(duracao_s) if duracao_s.isdigit() else 0, "origem": origem, "marcadores": lista,
        }, conteudo, arquivo.filename or "")
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    g = estado.gravacoes.obter(id_) or {}
    if g.get("servico_id"):
        # Gravada ja ligada a um servico: o audio vai para a pasta dele.
        _mover_audio_da_gravacao(id_, int(g["servico_id"]))
        estado.servicos.trilha(int(g["servico_id"]), "Gravação arquivada: " + g["titulo"])
        g = estado.gravacoes.obter(id_) or g
    viva = estado.ao_vivo.pop(sessao, None) if sessao else None
    if viva is not None:
        # O que foi transcrito enquanto gravava ja serve; fecha o que sobrou.
        await run_in_threadpool(estado.transcritor.ao_vivo_fim, viva)
    if viva is not None and viva.trechos:
        estado.gravacoes.marcar_transcricao(id_, "pronta", trechos=viva.trechos, modelo=viva.modelo + " ao vivo", tempo=0.0)
        g = estado.gravacoes.obter(id_) or g
    elif estado.transcritor.instalado() and transcrever == "1":
        estado.gravacoes.marcar_transcricao(id_, "fila")
        estado.fila_voz.put(id_)
        g = estado.gravacoes.obter(id_) or g
    return _com_progresso(g)


@app.get("/api/gravacoes/{id_}")
def gravacoes_obter(id_: int) -> dict:
    g = estado.gravacoes.obter(id_)
    if not g:
        raise HTTPException(status_code=404, detail="gravação não encontrada")
    return _com_progresso(g)


@app.post("/api/gravacoes/{id_}/transcrever")
def gravacoes_transcrever(id_: int) -> dict:
    """Poe a gravacao na fila do Whisper. Volta na hora; a tela acompanha pelo estado."""
    g = estado.gravacoes.obter(id_)
    if not g:
        raise HTTPException(status_code=404, detail="gravação não encontrada")
    if not g["existe"]:
        raise HTTPException(status_code=404, detail="o áudio não está mais no disco")
    if not estado.transcritor.instalado():
        raise HTTPException(status_code=503, detail="o modelo de voz não está baixado nesta máquina")
    if g["transcricao_estado"] in ("fila", "transcrevendo"):
        return _com_progresso(g)
    estado.gravacoes.marcar_transcricao(id_, "fila")
    estado.fila_voz.put(id_)
    return _com_progresso(estado.gravacoes.obter(id_) or {})


@app.post("/api/gravacoes/{id_}/resumo")
def gravacoes_resumo(id_: int) -> dict:
    """
    O resumo da gravacao, escrito pelo modelo local sobre a transcricao.

    Uma reuniao de uma hora nao cabe na janela do modelo: o texto e cortado
    em blocos, cada bloco ganha um resumo parcial e os parciais viram um so.
    Demora alguns minutos em CPU; a tela avisa.
    """
    g = estado.gravacoes.obter(id_)
    if not g:
        raise HTTPException(status_code=404, detail="gravação não encontrada")
    if not g.get("trechos"):
        raise HTTPException(status_code=400, detail="a gravação ainda não foi transcrita")
    disponivel, motivo = check_ollama(estado.client.model)
    if not disponivel:
        raise HTTPException(status_code=503, detail=motivo)
    texto = estado.gravacoes.texto_da_transcricao(id_)
    try:
        resumo = _resumir_em_blocos(texto)
    except OllamaError as exc:
        raise HTTPException(status_code=503, detail=str(exc)) from exc
    estado.gravacoes.guardar_resumo(id_, resumo)
    if g.get("servico_id"):
        estado.servicos.trilha(int(g["servico_id"]), "Resumo da gravação: " + g["titulo"], "Assistente")
    return _com_progresso(estado.gravacoes.obter(id_) or {})


def _resumir_em_blocos(texto: str, tamanho: int = 9000) -> str:
    if len(texto) <= tamanho:
        return _limpar_sugestao(estado.client.ask(gravacoes_mod.INSTRUCAO_RESUMO, texto))
    linhas = texto.split("\n")
    blocos, atual = [], ""
    for linha in linhas:
        if len(atual) + len(linha) + 1 > tamanho and atual:
            blocos.append(atual)
            atual = ""
        atual += linha + "\n"
    if atual:
        blocos.append(atual)
    parciais = [
        f"Parte {i + 1} de {len(blocos)}:\n" + _limpar_sugestao(estado.client.ask(gravacoes_mod.INSTRUCAO_RESUMO, bloco))
        for i, bloco in enumerate(blocos)
    ]
    return _limpar_sugestao(estado.client.ask(gravacoes_mod.INSTRUCAO_JUNTAR, "\n\n".join(parciais)))


def _mover_audio_da_gravacao(id_: int, servico_id: int | None) -> None:
    """
    Ligar a gravacao a um servico leva o audio para a pasta do servico no
    Acervo; desligar traz de volta para a pasta das gravacoes.
    """
    pasta = estado.servicos.pasta_de(servico_id, criar=True) if servico_id else None
    try:
        estado.gravacoes.mover_para(id_, pasta)
    except OSError as exc:
        raise HTTPException(status_code=500, detail=f"não consegui mover o áudio: {exc}") from exc


@app.post("/api/gravacoes/{id_}")
def gravacoes_atualizar(id_: int, payload: dict) -> dict:
    antes = estado.gravacoes.obter(id_)
    if not antes or not estado.gravacoes.atualizar(id_, payload):
        raise HTTPException(status_code=404, detail="gravação não encontrada")
    depois = estado.gravacoes.obter(id_) or {}
    if (antes.get("servico_id") or None) != (depois.get("servico_id") or None):
        _mover_audio_da_gravacao(id_, depois.get("servico_id"))
        if depois.get("servico_id"):
            estado.servicos.trilha(int(depois["servico_id"]), "Gravação ligada: " + depois["titulo"])
        depois = estado.gravacoes.obter(id_) or depois
    return depois


@app.post("/api/gravacoes/{id_}/anotacoes")
def gravacoes_anotar(id_: int, payload: dict) -> dict:
    try:
        return {"anotacoes": estado.gravacoes.anotar(id_, str(payload.get("texto", "")), _quem_sou())}
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc


@app.post("/api/gravacoes/{id_}/anotacoes/{indice}")
def gravacoes_anotacao_editar(id_: int, indice: int, payload: dict) -> dict:
    try:
        return {"anotacoes": estado.gravacoes.anotacao_editar(id_, indice, str(payload.get("texto", "")), _quem_sou())}
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc


@app.delete("/api/gravacoes/{id_}/anotacoes/{indice}")
def gravacoes_anotacao_remover(id_: int, indice: int) -> dict:
    try:
        return {"anotacoes": estado.gravacoes.anotacao_remover(id_, indice)}
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc


@app.delete("/api/gravacoes/{id_}")
def gravacoes_apagar(id_: int) -> dict:
    g = estado.gravacoes.obter(id_)
    if not g:
        raise HTTPException(status_code=404, detail="gravação não encontrada")
    caminho = estado.gravacoes.caminho(id_)
    entrada = estado.lixeira.apagar_linha("gravacao", id_, g["titulo"], "Gravações · " + g["tipo_rotulo"] + " · " + g["duracao_texto"],
                                          arquivos=[str(caminho)] if caminho else [])
    return _foi_para_lixeira(entrada, "apagada", id_)


def _nome_para_baixar(id_: int, caminho: Path) -> str:
    g = estado.gravacoes.obter(id_) or {}
    titulo = re.sub(r'[\\/:*?"<>|]+', "-", str(g.get("titulo") or caminho.stem)).strip() or caminho.stem
    return titulo + caminho.suffix.lower()


@app.get("/api/gravacoes/{id_}/audio")
def gravacoes_audio(id_: int, baixar: int = 0) -> FileResponse:
    caminho = estado.gravacoes.caminho(id_)
    if not caminho:
        raise HTTPException(status_code=404, detail="o áudio não está mais no disco")
    tipo = gravacoes_mod.MEDIA_TYPES.get(caminho.suffix.lower(), "application/octet-stream")
    if baixar:
        return FileResponse(caminho, media_type=tipo, filename=_nome_para_baixar(id_, caminho))
    return FileResponse(caminho, media_type=tipo)


@app.post("/api/gravacoes/{id_}/audio/salvar")
def gravacoes_audio_salvar(id_: int, payload: dict) -> dict:
    """Baixar o audio na janela do programa: o "Salvar como" do Windows escolhe onde, e aqui o arquivo e copiado."""
    caminho = estado.gravacoes.caminho(id_)
    if not caminho:
        raise HTTPException(status_code=404, detail="o áudio não está mais no disco")
    destino = Path(str(payload.get("caminho", "")))
    if not destino.name:
        raise HTTPException(status_code=400, detail="escolha onde salvar")
    if not destino.suffix:
        destino = destino.with_suffix(caminho.suffix)
    try:
        shutil.copy2(caminho, destino)
    except OSError as exc:
        raise HTTPException(status_code=500, detail=f"não consegui salvar: {exc.strerror or exc}") from exc
    return {"nome": destino.name, "caminho": str(destino)}


@app.get("/api/gravacoes/{id_}/audio/nome")
def gravacoes_audio_nome(id_: int) -> dict:
    caminho = estado.gravacoes.caminho(id_)
    if not caminho:
        raise HTTPException(status_code=404, detail="o áudio não está mais no disco")
    return {"nome": _nome_para_baixar(id_, caminho)}


@app.post("/api/gravacoes/{id_}/corrigir")
def gravacoes_corrigir(id_: int, payload: dict) -> dict:
    try:
        trocados = estado.gravacoes.corrigir(id_, str(payload.get("de", "")), str(payload.get("para", "")))
    except ValueError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc
    g = estado.gravacoes.obter(id_) or {}
    g["trocados"] = trocados
    return _com_progresso(g)


def _docx_da_gravacao(id_: int) -> tuple[str, bytes]:
    try:
        titulo, html = estado.gravacoes.html_para_exportar(id_)
    except ValueError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc
    return titulo, documento.para_docx(documento.ler_html(html), titulo)


@app.get("/api/gravacoes/{id_}/transcricao.docx")
def gravacoes_docx(id_: int):
    """A transcricao (com o resumo e as notas) em Word, para continuar fora do PAULUS."""
    from fastapi.responses import Response

    titulo, dados = _docx_da_gravacao(id_)
    return Response(dados, media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    headers=_anexo(_arquivo(titulo) + " - transcricao.docx"))


@app.post("/api/gravacoes/{id_}/exportar")
def gravacoes_exportar(id_: int) -> dict:
    """Grava o .docx em data/exportacoes e devolve o caminho: e o que vai anexo no e-mail."""
    titulo, dados = _docx_da_gravacao(id_)
    EXPORTACOES_DIR.mkdir(parents=True, exist_ok=True)
    nome = _arquivo(titulo) + " - transcricao.docx"
    caminho = EXPORTACOES_DIR / nome
    caminho.write_bytes(dados)
    return {"path": str(caminho), "nome": nome, "mb": round(len(dados) / (1024 * 1024), 2)}


@app.post("/api/gravacoes/{id_}/marcadores")
def gravacoes_marcar(id_: int, payload: dict) -> dict:
    try:
        return {"marcadores": estado.gravacoes.marcar(id_, int(payload.get("t") or 0), str(payload.get("texto", "")))}
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc


@app.delete("/api/gravacoes/{id_}/marcadores/{indice}")
def gravacoes_tirar_marcador(id_: int, indice: int) -> dict:
    try:
        return {"marcadores": estado.gravacoes.tirar_marcador(id_, indice)}
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc


# -------------------------------------------------------------- lixeira


def _foi_para_lixeira(entrada: dict | None, chave: str, id_) -> dict:
    """A resposta de um DELETE: o que saiu e o numero na lixeira, para o Desfazer do aviso."""
    if not entrada:
        raise HTTPException(status_code=404, detail="não encontrado")
    return {chave: id_, "lixeira": entrada["id"], "titulo": entrada["titulo"],
            "aviso": entrada["tipo_rotulo"] + " “" + entrada["titulo"] + "” foi para a lixeira · fica " + str(lixeira_mod.DIAS) + " dias"}


@app.get("/api/lixeira")
def lixeira_listar() -> dict:
    estado.lixeira.esvaziar_vencidos()
    itens = estado.lixeira.listar()
    return {"itens": itens, "total": len(itens), "dias": lixeira_mod.DIAS, "pasta": str(LIXEIRA_DIR)}


@app.post("/api/lixeira/esvaziar")
def lixeira_esvaziar() -> dict:
    return {"apagados": estado.lixeira.esvaziar()}


@app.post("/api/lixeira/{id_}/restaurar")
def lixeira_restaurar(id_: int) -> dict:
    try:
        e = estado.lixeira.restaurar(id_)
    except ValueError as exc:
        raise HTTPException(status_code=409, detail=str(exc)) from exc
    if e["tipo"] == "conversa":
        estado.trabalhos.recarregar_um(e["alvo_id"])
    return {"restaurado": e["alvo_id"], "tipo": e["tipo"], "titulo": e["titulo"],
            "aviso": e["tipo_rotulo"] + " “" + e["titulo"] + "” voltou"}


@app.delete("/api/lixeira/{id_}")
def lixeira_tirar(id_: int) -> dict:
    if not estado.lixeira.tirar(id_):
        raise HTTPException(status_code=404, detail="essa entrada não está na lixeira")
    return {"apagado_de_vez": id_}


# ------------------------------------------------------------- conexoes


@app.get("/api/conexoes")
def conexoes_ver() -> dict:
    dados = estado.conexoes.para_tela()
    dados["janela_embutida"] = _tem_janela()
    return dados


def _tem_janela() -> bool:
    """A janela embutida so existe quando o programa roda como desktop."""
    try:
        import webview

        return bool(webview.windows)
    except Exception:
        return False


@app.post("/api/conexoes/abrir")
def conexoes_abrir(payload: dict | None = None) -> dict:
    """
    Abre a janela do serviço, presa a um endereço só.

    Fora do modo desktop nao ha janela embutida - e a tela diz isso em vez de
    abrir o navegador comum fingindo que e a mesma coisa, porque a sessao nao
    seria a mesma.
    """
    endereco = str((payload or {}).get("endereco", "")) or conexoes.WHATSAPP
    if not endereco.startswith(conexoes.WHATSAPP):
        raise HTTPException(status_code=400, detail="esta janela só abre o WhatsApp Web")

    if not _tem_janela():
        return {
            "abriu": False,
            "endereco": endereco,
            "motivo": (
                "A janela embutida existe quando o PAULUS roda como programa "
                "(python src/desktop.py). No navegador, a sessão seria a do próprio "
                "navegador, não a minha — então prefiro dizer isso a fingir que é igual."
            ),
        }

    import webview

    webview.create_window(
        "WhatsApp Web · PAULUS", endereco,
        width=1100, height=800, min_size=(720, 560),
    )
    return {"abriu": True, "endereco": endereco}


@app.post("/api/conexoes/mensagem")
def conexoes_mensagem(payload: MensagemWhats) -> dict:
    """
    Prepara a conversa com o texto pronto - e para aí.

    Apertar enviar continua sendo seu. O programa não clica na tela de um site
    de terceiro: quebraria a cada mudança de layout, e o preço de errar é a
    conta do escritório ser derrubada.
    """
    endereco = conexoes.link_de_conversa(payload.telefone, payload.texto)
    if not endereco:
        raise HTTPException(status_code=400, detail="informe um telefone com DDD")

    anexo = ""
    if payload.anexo:
        alvo = Path(payload.anexo)
        if not alvo.exists():
            raise HTTPException(status_code=400, detail="não achei esse arquivo")
        anexo = str(alvo)

    estado.conexoes.anotar(payload.telefone, payload.nome, anexo)
    return {
        "endereco": endereco,
        "anexo": Path(anexo).name if anexo else "",
        "pasta_do_anexo": str(Path(anexo).parent) if anexo else "",
        "aviso": (
            "Abro a conversa com o texto escrito. Você confere e aperta enviar."
            + (" O anexo você arrasta da pasta que eu abri." if anexo else "")
        ),
    }


@app.post("/api/conexoes/sessao/apagar")
def conexoes_apagar_sessao() -> dict:
    apagou = estado.conexoes.apagar_sessao()
    return {
        "apagou": apagou,
        "aviso": "A sessão foi apagada desta máquina. O WhatsApp vai pedir o código de novo."
        if apagou else "Não havia sessão guardada.",
    }


@app.get("/api/conexoes/contatos")
def conexoes_contatos() -> dict:
    """Quem tem telefone no cadastro - para não digitar número na mão."""
    achados = []
    for f in estado.cadastros.listar():
        if (f.get("telefone") or "").strip():
            achados.append({
                "id": f["id"], "nome": f["nome"], "telefone": f["telefone"],
                "numero": conexoes.numero_whatsapp(f["telefone"]),
            })
    return {"contatos": achados}

# ------------------------------------------------------------------ leis


class ImportarLei(BaseModel):
    caminho: str = ""
    codigo: str = ""


@app.get("/api/leis")
def leis_listar() -> dict:
    """O que está no disco, e o que falta."""
    return {
        "codigos": estado.leis.instalados(),
        "contagem": estado.leis.contagem(),
        "porque": (
            "Com o texto no disco, citar um artigo é consulta a um índice. Sem ele, eu "
            "não cito: pedir o número a um modelo de 3 bilhões de parâmetros devolve "
            "artigo plausível e errado dentro de um contrato."
        ),
        "como_baixar": (
            "Abra a página do código no Planalto, salve como “Página da Web, completa” "
            "e escolha o arquivo aqui. É o texto compilado oficial, com as alterações "
            "já incorporadas."
        ),
    }


@app.post("/api/leis/importar")
def leis_importar(payload: ImportarLei) -> dict:
    """
    Lê o texto oficial e guarda artigo por artigo.

    O que não é artigo do código fica de fora e é contado: o decreto que aprova
    a consolidação tem artigos próprios, e as disposições finais citam artigos
    de outras leis. Guardar essas citações como artigos deste código seria
    inventar artigo que não existe.
    """
    alvo = Path(payload.caminho)
    if not alvo.exists():
        raise HTTPException(status_code=400, detail="não achei esse arquivo")
    if alvo.suffix.lower() not in (".html", ".htm"):
        raise HTTPException(status_code=400, detail="salve a página do Planalto como HTML")

    try:
        resultado = estado.leis.importar(alvo, payload.codigo)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc

    return {**resultado, "codigos": estado.leis.instalados(),
            "contagem": estado.leis.contagem()}


@app.post("/api/leis/importar-pasta")
def leis_importar_pasta(payload: dict) -> dict:
    """Importa de uma vez tudo o que a pasta tiver e o programa reconhecer."""
    pasta = Path(str(payload.get("pasta", "")))
    if not pasta.is_dir():
        raise HTTPException(status_code=400, detail="essa pasta não existe")

    feitos, ignorados = [], []
    for arquivo in sorted(list(pasta.glob("*.html")) + list(pasta.glob("*.htm"))):
        codigo = leis.reconhecer(arquivo)
        if not codigo:
            ignorados.append({"arquivo": arquivo.name, "motivo": "não reconheci de qual código é"})
            continue
        try:
            feitos.append(estado.leis.importar(arquivo, codigo))
        except ValueError as exc:
            ignorados.append({"arquivo": arquivo.name, "motivo": str(exc)})

    return {"importados": feitos, "ignorados": ignorados,
            "codigos": estado.leis.instalados(), "contagem": estado.leis.contagem()}


@app.delete("/api/leis/{codigo}")
def leis_apagar(codigo: str) -> dict:
    estado.leis.apagar(codigo)
    return {"codigos": estado.leis.instalados(), "contagem": estado.leis.contagem()}


@app.get("/api/leis/procurar")
def leis_procurar(termo: str = "", codigo: str = "", limite: int = 20) -> dict:
    """
    Busca por número de artigo ou por palavra.

    Quando não há código instalado, a resposta diz isso em vez de devolver
    vazio: lista vazia parece "não existe artigo sobre isso", e o certo é
    "ainda não tenho o texto da lei aqui".
    """
    if not estado.leis.contagem()["codigos"]:
        return {"achados": [], "sem_codigo": True,
                "aviso": "Nenhum código instalado ainda. Importe o texto oficial primeiro."}
    return {"achados": estado.leis.procurar(termo, codigo, limite), "sem_codigo": False}


@app.get("/api/leis/artigo")
def leis_artigo(codigo: str, numero: str) -> dict:
    a = estado.leis.artigo(codigo, numero)
    if not a:
        raise HTTPException(
            status_code=404,
            detail=f"não achei o art. {numero} no {leis.CODIGOS.get(codigo, {}).get('nome', codigo)}",
        )
    return {"artigo": a, "vizinhos": estado.leis.vizinhos(codigo, a["ordem"])}

def main() -> None:
    import argparse

    import uvicorn

    parser = argparse.ArgumentParser(description="PAULUS Legal - interface web (local)")
    parser.add_argument("--contracts", type=Path, default=CONTRACTS_DIR, help="pasta com os contratos")
    parser.add_argument("--model", default=DEFAULT_MODEL, help=f"modelo Ollama (padrao: {DEFAULT_MODEL})")
    parser.add_argument("--port", type=int, default=8000, help="porta (padrao: 8000)")
    args = parser.parse_args()

    estado.pasta = args.contracts
    estado.porta = args.port
    estado.client = LlamaClient(model=args.model)

    # host fixo em 127.0.0.1: o servidor nao deve ficar exposto na rede local,
    # os documentos sao de cliente.
    uvicorn.run(app, host="127.0.0.1", port=args.port, log_level="warning")


if __name__ == "__main__":
    main()
