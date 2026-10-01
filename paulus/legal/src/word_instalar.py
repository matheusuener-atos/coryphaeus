"""
Instalar o suplemento no Word deste computador (W1, montagem A).

O que a W0 mediu e decidiu (docs/PROGRESSO-WORD.md):

- o painel vem de https://localhost:<porta fixa>: a porta é escolhida uma vez
  e guardada, porque o manifesto leva o endereço;
- o certificado é de uma autoridade desta instalação **restrita a localhost e
  127.0.0.1** (name constraints), cuja chave é descartada logo depois de
  emitir o do servidor: copiada, a pasta não emite certificado para outro
  nome. Ele entra em CurrentUser\\Root sem administrador; o Windows pergunta,
  e a pessoa clica em Sim;
- confiado ou não, **quem diz é o repositório lido depois**, nunca o retorno
  do comando: na W0 o Import-Certificate respondeu erro com o certificado
  instalado;
- o Word acha o suplemento pelo registro de desenvolvedor do usuário
  (HKCU\\...\\WEF\\Developer), sem administrador; o catálogo em pasta
  compartilhada fica como segunda via, com o passo a passo na tela;
- "instalado" só depois que o painel abriu num Word e avisou (POST
  /api/word/carregou).
"""

from __future__ import annotations

import ipaddress
import json
import socket
import ssl
import subprocess
import sys
import threading
import time
import uuid
from datetime import datetime, timedelta, timezone
from pathlib import Path
from string import Template
from xml.sax.saxutils import escape

MODELO_DO_MANIFESTO = Path(__file__).resolve().parent.parent / "word" / "manifesto.xml"
FAIXA_DE_PORTAS = range(46300, 46400)
DIAS_DO_CERTIFICADO = 825
RENOVAR_COM_DIAS = 30
CHAVE_DO_REGISTRO = r"Software\Microsoft\Office\16.0\WEF\Developer"
ESPACO_DOS_IDS = uuid.UUID("2f0c9e57-1d0b-4c43-9f0e-5e6b1a7f4d21")
NOME_NO_WORD = "PAVLVS"


class FalhaNaInstalacao(Exception):
    pass


# ------------------------------------------------------------ certificado

def gerar_certificado(pasta: Path, dias: int = DIAS_DO_CERTIFICADO) -> dict:
    """A autoridade só-localhost, o certificado do servidor e a chave dele."""
    from cryptography import x509
    from cryptography.hazmat.primitives import hashes, serialization
    from cryptography.hazmat.primitives.asymmetric import ec
    from cryptography.x509.oid import ExtendedKeyUsageOID, NameOID

    pasta.mkdir(parents=True, exist_ok=True)
    agora = datetime.now(timezone.utc)
    fim = agora + timedelta(days=dias)
    dns, ips = ["localhost"], [ipaddress.ip_address("127.0.0.1")]

    chave_ac = ec.generate_private_key(ec.SECP256R1())
    nome_ac = x509.Name([
        x509.NameAttribute(NameOID.COMMON_NAME, "PAULUS - Word nesta instalacao (so localhost)"),
        x509.NameAttribute(NameOID.ORGANIZATION_NAME, "PAULUS"),
    ])
    ac = (x509.CertificateBuilder()
          .subject_name(nome_ac).issuer_name(nome_ac).public_key(chave_ac.public_key())
          .serial_number(x509.random_serial_number())
          .not_valid_before(agora - timedelta(minutes=5)).not_valid_after(fim)
          .add_extension(x509.BasicConstraints(ca=True, path_length=0), critical=True)
          .add_extension(x509.NameConstraints(
              permitted_subtrees=[x509.DNSName(n) for n in dns]
              + [x509.IPAddress(ipaddress.ip_network(f"{ip}/32")) for ip in ips],
              excluded_subtrees=None), critical=True)
          .add_extension(x509.KeyUsage(digital_signature=False, content_commitment=False, key_encipherment=False,
                                       data_encipherment=False, key_agreement=False, key_cert_sign=True,
                                       crl_sign=True, encipher_only=False, decipher_only=False), critical=True)
          .add_extension(x509.SubjectKeyIdentifier.from_public_key(chave_ac.public_key()), critical=False)
          .sign(chave_ac, hashes.SHA256()))
    chave_srv = ec.generate_private_key(ec.SECP256R1())
    srv = (x509.CertificateBuilder()
           .subject_name(x509.Name([x509.NameAttribute(NameOID.COMMON_NAME, "localhost")]))
           .issuer_name(nome_ac).public_key(chave_srv.public_key())
           .serial_number(x509.random_serial_number())
           .not_valid_before(agora - timedelta(minutes=5)).not_valid_after(fim)
           .add_extension(x509.SubjectAlternativeName([x509.DNSName(n) for n in dns] + [x509.IPAddress(ip) for ip in ips]),
                          critical=False)
           .add_extension(x509.BasicConstraints(ca=False, path_length=None), critical=True)
           .add_extension(x509.ExtendedKeyUsage([ExtendedKeyUsageOID.SERVER_AUTH]), critical=False)
           .add_extension(x509.AuthorityKeyIdentifier.from_issuer_public_key(chave_ac.public_key()), critical=False)
           .sign(chave_ac, hashes.SHA256()))
    del chave_ac  # nunca gravada

    der = ac.public_bytes(serialization.Encoding.DER)
    (pasta / "autoridade.cer").write_bytes(der)
    (pasta / "servidor.pem").write_bytes(srv.public_bytes(serialization.Encoding.PEM)
                                         + ac.public_bytes(serialization.Encoding.PEM))
    (pasta / "servidor.key").write_bytes(chave_srv.private_bytes(
        serialization.Encoding.PEM, serialization.PrivateFormat.PKCS8, serialization.NoEncryption()))
    info = {"impressao": ac.fingerprint(hashes.SHA1()).hex().upper(), "vence": fim.isoformat(timespec="seconds")}
    (pasta / "certificado.json").write_text(json.dumps(info), encoding="utf-8")
    return info


def confiado(der: bytes) -> bool:
    """O certificado está entre as raízes confiáveis deste usuário? Lido do repositório."""
    if sys.platform != "win32":
        return False
    try:
        return any(cert == der for cert, _cod, _uso in ssl.enum_certificates("ROOT"))
    except OSError:
        return False


def pedir_confianca(arquivo: Path) -> int:
    """certutil com a janela do Windows ("Aviso de Segurança"); devolve o código de saída."""
    r = subprocess.run(["certutil", "-user", "-addstore", "Root", str(arquivo)], capture_output=True,
                       creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0), timeout=600)
    return r.returncode


# ------------------------------------------------------------ registro do Word

def registrar_no_word(manifesto: Path) -> None:
    import winreg

    with winreg.CreateKey(winreg.HKEY_CURRENT_USER, CHAVE_DO_REGISTRO) as k:
        winreg.SetValueEx(k, str(manifesto), 0, winreg.REG_SZ, str(manifesto))


def tirar_do_word(manifesto: Path) -> None:
    import winreg

    try:
        with winreg.OpenKey(winreg.HKEY_CURRENT_USER, CHAVE_DO_REGISTRO, 0, winreg.KEY_SET_VALUE) as k:
            winreg.DeleteValue(k, str(manifesto))
    except OSError:
        pass


def registrado_no_word(manifesto: Path) -> bool:
    if sys.platform != "win32":
        return False
    import winreg

    try:
        with winreg.OpenKey(winreg.HKEY_CURRENT_USER, CHAVE_DO_REGISTRO) as k:
            valor, _ = winreg.QueryValueEx(k, str(manifesto))
            return Path(valor) == manifesto
    except OSError:
        return False


def word_no_computador() -> bool:
    """Há Word instalado? Pelo registro do Windows (App Paths), sem abrir o Word."""
    if sys.platform != "win32":
        return False
    import winreg

    caminho = r"Software\Microsoft\Windows\CurrentVersion\App Paths\Winword.exe"
    for raiz in (winreg.HKEY_CURRENT_USER, winreg.HKEY_LOCAL_MACHINE):
        for vista in (winreg.KEY_WOW64_64KEY, winreg.KEY_WOW64_32KEY):
            try:
                with winreg.OpenKey(raiz, caminho, 0, winreg.KEY_READ | vista):
                    return True
            except OSError:
                continue
    return False


# ------------------------------------------------------------ o documento leva o PAVLVS

# No Word 2021 o registro de desenvolvedor só faz o Word conhecer o
# suplemento: a aba PAVLVS aparece no documento que traz a referência a ele
# (medido em 01/10/2026: Word em branco sem a aba; o mesmo Word, com um
# documento que traz o PAVLVS, com a aba e o painel; "Meus Suplementos" diz
# "Sem Suplementos"). Por isso o PAULUS põe a referência nos .docx que ele
# gera - o Word grava essas partes junto quando a pessoa salva.
_NS_WE = "http://schemas.microsoft.com/office/webextensions/webextension/2010/11"
_REL_TASKPANES = "http://schemas.microsoft.com/office/2011/relationships/webextensiontaskpanes"


def com_pavlvs(dados: bytes, *, id_: str, versao: str, abrir_painel: bool) -> bytes:
    """
    O .docx com a referência ao PAVLVS (as partes webextension). Sem tocar no
    texto: só acrescenta partes. Documento que já traz o PAVLVS volta igual;
    documento que já tem outro suplemento também volta igual (não se mexe no
    que não é nosso).
    """
    import io
    import zipfile

    try:
        zin = zipfile.ZipFile(io.BytesIO(dados))
    except zipfile.BadZipFile:
        return dados
    with zin:
        nomes = set(zin.namelist())
        if any(n.startswith("word/webextensions/") for n in nomes) or "[Content_Types].xml" not in nomes \
                or "_rels/.rels" not in nomes:
            return dados
        propriedade = ('<we:property name="Office.AutoShowTaskpaneWithDocument" value="true"/>' if abrir_painel else "")
        we = (f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
              f'<we:webextension xmlns:we="{_NS_WE}" id="{{{str(uuid.uuid4()).upper()}}}">'
              f'<we:reference id="{escape(id_)}" version="{escape(versao_do_manifesto(versao))}" store="developer" storeType="Registry"/>'
              f'<we:alternateReferences/><we:properties>{propriedade}</we:properties><we:bindings/>'
              f'<we:snapshot xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"/></we:webextension>')
        visivel = "1" if abrir_painel else "0"
        tp = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
              '<wetp:taskpanes xmlns:wetp="http://schemas.microsoft.com/office/webextensions/taskpanes/2010/11">'
              f'<wetp:taskpane dockstate="right" visibility="{visivel}" width="380" row="4">'
              '<wetp:webextensionref xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" r:id="rId1"/>'
              '</wetp:taskpane></wetp:taskpanes>')
        tp_rels = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
                   '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                   '<Relationship Id="rId1" Type="http://schemas.microsoft.com/office/2011/relationships/webextension" '
                   'Target="webextension1.xml"/></Relationships>')
        saida = io.BytesIO()
        with zipfile.ZipFile(saida, "w", zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                parte = zin.read(item.filename)
                if item.filename == "[Content_Types].xml":
                    parte = parte.replace(b"</Types>", (
                        b'<Override PartName="/word/webextensions/taskpanes.xml" ContentType="application/vnd.ms-office.webextensiontaskpanes+xml"/>'
                        b'<Override PartName="/word/webextensions/webextension1.xml" ContentType="application/vnd.ms-office.webextension+xml"/>'
                        b"</Types>"))
                elif item.filename == "_rels/.rels":
                    parte = parte.replace(b"</Relationships>", (
                        f'<Relationship Id="rIdPavlvs" Type="{_REL_TASKPANES}" Target="word/webextensions/taskpanes.xml"/>'
                        "</Relationships>").encode())
                zout.writestr(item, parte)
            zout.writestr("word/webextensions/taskpanes.xml", tp)
            zout.writestr("word/webextensions/_rels/taskpanes.xml.rels", tp_rels)
            zout.writestr("word/webextensions/webextension1.xml", we)
    return saida.getvalue()


def documento_base(paragrafos: list[tuple[str, str]] | None = None, *, modelo: bool = False) -> bytes:
    """
    Um .docx (ou .dotx, com `modelo`) do jeito que o Word cria: modo de
    compatibilidade 15 (sem a faixa "Modo de Compatibilidade" do modelo do
    python-docx) e Calibri 11. `paragrafos`: [(estilo, texto)].
    """
    import io
    import zipfile

    import docx
    from docx.oxml.ns import qn
    from docx.shared import Pt

    d = docx.Document()
    normal = d.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    for estilo, texto in paragrafos or []:
        if estilo == "titulo":
            d.add_heading(texto, level=1)
        else:
            d.add_paragraph(texto)
    configuracoes = d.settings.element
    compat = configuracoes.find(qn("w:compat"))
    if compat is None:
        compat = configuracoes.makeelement(qn("w:compat"), {})
        configuracoes.append(compat)
    for velho in compat.findall(qn("w:compatSetting")):
        if velho.get(qn("w:name")) == "compatibilityMode":
            compat.remove(velho)
    item = compat.makeelement(qn("w:compatSetting"), {qn("w:name"): "compatibilityMode",
                                                      qn("w:uri"): "http://schemas.microsoft.com/office/word",
                                                      qn("w:val"): "15"})
    compat.append(item)
    b = io.BytesIO()
    d.save(b)
    if not modelo:
        return b.getvalue()
    saida = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(b.getvalue())) as zin, zipfile.ZipFile(saida, "w", zipfile.ZIP_DEFLATED) as zout:
        for parte in zin.infolist():
            dados = zin.read(parte.filename)
            if parte.filename == "[Content_Types].xml":
                dados = dados.replace(b"wordprocessingml.document.main+xml", b"wordprocessingml.template.main+xml")
            zout.writestr(parte, dados)
    return saida.getvalue()


def word_aberto_desde() -> float | None:
    """Desde quando há um Word aberto (o mais antigo), em segundos epoch; None sem Word."""
    try:
        import psutil
    except ImportError:
        return None
    inicios = []
    for p in psutil.process_iter(["name", "create_time"]):
        if (p.info.get("name") or "").lower() == "winword.exe" and p.info.get("create_time"):
            inicios.append(p.info["create_time"])
    return min(inicios) if inicios else None


def tem_pavlvs(dados: bytes, id_: str) -> bool:
    import io
    import zipfile

    try:
        with zipfile.ZipFile(io.BytesIO(dados)) as z:
            return any(n.startswith("word/webextensions/webextension") and id_.encode() in z.read(n)
                       for n in z.namelist())
    except zipfile.BadZipFile:
        return False


# ------------------------------------------------------------ manifesto

def versao_do_manifesto(versao: str) -> str:
    partes = [p for p in str(versao).split(".") if p.isdigit()][:4]
    return ".".join((partes + ["0", "0", "0", "0"])[:4])


def montar_manifesto(*, id_: str, base: str, versao: str) -> str:
    modelo = Template(MODELO_DO_MANIFESTO.read_text(encoding="utf-8"))
    return modelo.substitute(ID=id_, BASE=escape(base.rstrip("/")), VERSAO=versao_do_manifesto(versao),
                             NOME=NOME_NO_WORD)


# ------------------------------------------------------------ a instalação

class Instalacao:
    def __init__(self, estado, pasta_dados: Path) -> None:
        self.estado = estado
        self.pasta = Path(pasta_dados) / "word"
        self.pasta_cert = self.pasta / "certificado"
        self._servidor = None
        self._fio = None
        self._trava = threading.Lock()

    # --- o que está guardado

    def _prefs(self) -> dict:
        return dict(self.estado.prefs.dados.get("word") or {})

    def _id(self, montagem: str) -> str:
        base = self._prefs().get("id") or ""
        if not base:
            base = uuid.uuid4().hex
            self.estado.prefs.atualizar({"word": {"id": base}})
        return str(uuid.uuid5(ESPACO_DOS_IDS, f"{base}:{montagem}"))

    def porta(self) -> int:
        return int(self._prefs().get("porta") or 0)

    def porta_aberta(self) -> int:
        """A porta do servidor HTTPS no ar, ou 0. O porteiro pergunta a cada
        pedido: sem servidor não há pedido nessa porta, e não se lê nada."""
        servidor = self._servidor
        return servidor.config.port if servidor is not None and getattr(servidor, "started", False) else 0

    def manifesto_local(self) -> Path:
        return self.pasta / "manifesto-local.xml"

    def certificado(self) -> dict:
        try:
            info = json.loads((self.pasta_cert / "certificado.json").read_text(encoding="utf-8"))
            der = (self.pasta_cert / "autoridade.cer").read_bytes()
        except (OSError, ValueError):
            return {"existe": False, "confiado": False}
        vence = info.get("vence", "")
        try:
            dias = (datetime.fromisoformat(vence) - datetime.now(timezone.utc)).days
        except ValueError:
            dias = 0
        return {"existe": True, "confiado": confiado(der), "impressao": info.get("impressao", ""),
                "vence": vence[:10], "dias": dias}

    def anotar_carregamento(self, registro: dict) -> None:
        if registro.get("origem") != "local":
            return
        self.pasta.mkdir(parents=True, exist_ok=True)
        (self.pasta / "carregou.json").write_text(json.dumps(registro, ensure_ascii=False), encoding="utf-8")

    def ultimo_carregamento(self) -> dict | None:
        try:
            return json.loads((self.pasta / "carregou.json").read_text(encoding="utf-8"))
        except (OSError, ValueError):
            return None

    def endereco_de_fora(self) -> str:
        try:
            situacao = self.estado.acesso_de_fora.situacao()
        except Exception:  # noqa: BLE001
            return ""
        nome = str(situacao.get("hostname") or "")
        return f"https://{nome}" if nome and situacao.get("ligado") else ""

    def situacao(self) -> dict:
        porta = self.porta()
        cert = self.certificado()
        manifesto = self.manifesto_local()
        registrado = manifesto.exists() and registrado_no_word(manifesto)
        instalado_em = self._prefs().get("instalado_em") or ""
        carregou = self.ultimo_carregamento()
        # "Carregado" vale só se o painel abriu depois da última instalação.
        carregado = bool(carregou and instalado_em and carregou.get("quando", "") >= instalado_em)
        return {"porta": porta, "endereco": f"https://localhost:{porta}/word/painel.html" if porta else "",
                "servidor": self._servidor is not None and getattr(self._servidor, "started", False),
                "certificado": cert, "registrado": registrado, "instalado_em": instalado_em,
                "carregado": carregado, "carregou": carregou, "endereco_de_fora": self.endereco_de_fora(),
                "word_no_computador": word_no_computador(), "instalado": self.instalado(),
                "atalhos": self._atalhos_sem_falhar()}

    def _atalhos_sem_falhar(self) -> dict:
        try:
            return self.atalhos()
        except Exception:  # noqa: BLE001 - a tela mostra o resto
            return {"disponivel": False, "area_de_trabalho": False, "menu_iniciar": False, "botao_direito": False}

    # --- instalar

    def _escolher_porta(self) -> int:
        atual = self.porta()
        if atual:
            return atual
        for p in FAIXA_DE_PORTAS:
            with socket.socket(socket.AF_INET, socket.SOCK_STREAM) as s:
                try:
                    s.bind(("127.0.0.1", p))
                except OSError:
                    continue
            self.estado.prefs.atualizar({"word": {"porta": p}})
            return p
        raise FalhaNaInstalacao("nenhuma porta livre entre 46300 e 46399")

    def manifesto(self, montagem: str) -> str:
        from versao import VERSAO

        if montagem == "fora":
            base = self.endereco_de_fora()
            if not base:
                raise FalhaNaInstalacao("o acesso de fora não está ligado e conectado: sem ele não há endereço para o Word de outro computador")
        else:
            porta = self.porta()
            if not porta:
                raise FalhaNaInstalacao("instale no Word deste computador primeiro: a porta ainda não foi escolhida")
            base = f"https://localhost:{porta}"
        return montar_manifesto(id_=self._id(montagem), base=base, versao=VERSAO)

    def instalar(self, *, pedir=pedir_confianca, registrar=registrar_no_word,
                 atalhos=lambda inst: inst.criar_atalhos()) -> list[dict]:
        passos: list[dict] = []
        with self._trava:
            cert = self.certificado()
            if not cert["existe"] or cert.get("dias", 0) < RENOVAR_COM_DIAS:
                info = gerar_certificado(self.pasta_cert)
                passos.append({"passo": "certificado", "resultado": f"gerado, vale até {info['vence'][:10]}"})
                cert = self.certificado()
            else:
                passos.append({"passo": "certificado", "resultado": f"o mesmo, vale até {cert['vence']}"})
            if not cert["confiado"]:
                codigo = pedir(self.pasta_cert / "autoridade.cer")
                cert = self.certificado()
                if not cert["confiado"]:
                    raise FalhaNaInstalacao("o Windows não confiou no certificado (a janela de aviso foi recusada ou fechada"
                                            f"; código {codigo}). Sem ele o Word não abre o painel.")
                passos.append({"passo": "confiança", "resultado": "o Windows confiou (conferido no repositório)"})
            else:
                passos.append({"passo": "confiança", "resultado": "já confiado"})
            porta = self._escolher_porta()
            passos.append({"passo": "porta", "resultado": f"https://localhost:{porta}"})
            self.pasta.mkdir(parents=True, exist_ok=True)
            self.manifesto_local().write_text(self.manifesto("local"), encoding="utf-8")
            passos.append({"passo": "manifesto", "resultado": str(self.manifesto_local())})
            registrar(self.manifesto_local())
            passos.append({"passo": "registro", "resultado": "o Word deste usuário acha o PAVLVS ao abrir"})
            self.estado.prefs.atualizar({"word": {"instalado_em": datetime.now().strftime("%Y-%m-%d %H:%M:%S")}})
        try:
            feitos = atalhos(self) if atalhos else []
        except Exception as exc:  # noqa: BLE001 - sem atalho, o resto funciona
            feitos = []
            passos.append({"passo": "atalhos", "resultado": f"não criados: {exc}"})
        else:
            if feitos:
                passos.append({"passo": "atalhos", "resultado": ", ".join(feitos)})
        aberto = self.aplicar()
        passos.append({"passo": "servidor", "resultado": "no ar" if aberto else "não abriu: a porta está ocupada"})
        return passos

    # --- o documento leva o PAVLVS

    def instalado(self) -> bool:
        prefs = self._prefs()
        return bool(prefs.get("ligado") and prefs.get("instalado_em") and self.porta())

    def no_documento(self, dados: bytes, *, abrir_painel: bool = True) -> bytes:
        """O .docx que o PAULUS gera, com o PAVLVS (se ligado e instalado)."""
        if not self.instalado():
            return dados
        from versao import VERSAO

        return com_pavlvs(dados, id_=self._id("local"), versao=VERSAO, abrir_painel=abrir_painel)

    def precisa_fechar_o_word(self, *, desde=None) -> bool:
        """
        O Word lê o registro só ao abrir: um Word aberto antes da instalação
        não conhece o PAVLVS, e o documento abriria com "suplemento não
        disponível". Aí é preciso esperar ele fechar.
        """
        inicio = (desde or word_aberto_desde)()
        instalado_em = self._prefs().get("instalado_em") or ""
        if inicio is None or not instalado_em:
            return False
        try:
            marco = datetime.strptime(instalado_em, "%Y-%m-%d %H:%M:%S").timestamp()
        except ValueError:
            return False
        return inicio < marco

    def comece_aqui(self) -> Path:
        """O documento que abre o Word com o PAVLVS depois de ativar."""
        texto = documento_base([
            ("titulo", "PAVLVS no Word"),
            ("", "Este documento abriu o Word com o PAVLVS: a aba PAVLVS lá em cima e o painel ao lado."),
            ("", "Para começar, clique em Conectar ao PAULUS no painel e confira o código na janela do PAULUS."),
            ("", "Para escrever um documento novo com o PAVLVS, use o atalho Word com PAVLVS da área de trabalho. "
                 "Para abrir um arquivo seu com ele, clique com o botão direito no arquivo .docx e escolha Abrir no Word "
                 "com o PAVLVS (no Windows 11, em Mostrar mais opções). Dali em diante, o duplo clique já abre com o PAVLVS."),
            ("", "No Editor do PAULUS, o botão Abrir no Word também abre o documento com o PAVLVS, e todo documento Word que "
                 "o PAULUS gera já vem com ele. O PAVLVS continua no documento depois que você salva."),
        ])
        destino = self.pasta / "PAVLVS - comece aqui.docx"
        self.pasta.mkdir(parents=True, exist_ok=True)
        dados = self.no_documento(texto)
        try:
            destino.write_bytes(dados)
        except OSError:
            pass  # aberto no Word: abre o que já está lá
        return destino

    def modelo_novo(self) -> Path:
        """O modelo do atalho "Word com PAVLVS": o Word abre um documento novo com o PAVLVS."""
        destino = self.pasta / "PAVLVS.dotx"
        self.pasta.mkdir(parents=True, exist_ok=True)
        try:
            destino.write_bytes(self.no_documento(documento_base(modelo=True)))
        except OSError:
            pass  # em uso: o que já está lá serve
        return destino

    def por_no_arquivo(self, caminho: Path) -> str:
        """
        O botão direito "Abrir no Word com o PAVLVS": põe a referência no
        arquivo .docx que a pessoa escolheu (o texto não muda) e devolve
        "posto", "ja_tinha" ou "preso" (aberto em outro programa: abre assim).
        Grava numa cópia ao lado e troca de uma vez: arquivo nunca fica pela
        metade.
        """
        import os

        if caminho.suffix.lower() != ".docx" or not caminho.is_file():
            raise FalhaNaInstalacao("só arquivo .docx")
        dados = caminho.read_bytes()
        novo = self.no_documento(dados)
        if novo == dados:
            return "ja_tinha"
        temporario = caminho.with_name(f".{caminho.name}.pavlvs")
        try:
            temporario.write_bytes(novo)
            os.replace(temporario, caminho)
        except OSError:
            temporario.unlink(missing_ok=True)
            return "preso"
        return "posto"

    # --- os atalhos

    def atalhos(self) -> dict:
        from acesso import energia

        import word_atalhos as wa

        exe = energia.exe_do_programa()
        mesa = wa.lnk_em(wa.pasta_do_windows("Desktop"))
        menu = wa.lnk_em(wa.pasta_do_windows("Programs"))
        return {"disponivel": bool(exe), "area_de_trabalho": bool(mesa and mesa.exists()),
                "menu_iniciar": bool(menu and menu.exists()), "botao_direito": wa.botao_direito_ligado(exe)}

    def criar_atalhos(self, *, area_de_trabalho: bool = True, menu_iniciar: bool = False,
                      botao_direito: bool = True) -> list[str]:
        """Os caminhos do Windows até o Word com o PAVLVS (src/word_atalhos.py)."""
        from acesso import energia

        import word_atalhos as wa

        exe = energia.exe_do_programa()
        if not exe or sys.platform != "win32":
            return []
        feitos = []
        icone = wa.icone_do_word() or f"{exe},0"
        for querer, pasta in ((area_de_trabalho, "Desktop"), (menu_iniciar, "Programs")):
            lnk = wa.lnk_em(wa.pasta_do_windows(pasta)) if querer else None
            if lnk:
                wa.criar_atalho(lnk, Path(exe), "--word", icone, "Abre um documento novo no Word, com o PAVLVS")
                feitos.append("área de trabalho" if pasta == "Desktop" else "menu Iniciar")
        if botao_direito:
            wa.ligar_botao_direito(Path(exe))
            feitos.append("botão direito dos .docx")
        return feitos

    def tirar_atalhos(self) -> None:
        import word_atalhos as wa

        for pasta in ("Desktop", "Programs"):
            lnk = wa.lnk_em(wa.pasta_do_windows(pasta))
            if lnk and lnk.exists():
                try:
                    lnk.unlink()
                except OSError:
                    pass
        if sys.platform == "win32":
            wa.desligar_botao_direito()

    def desinstalar(self, *, tirar=tirar_do_word) -> None:
        tirar(self.manifesto_local())
        try:
            self.tirar_atalhos()
        except Exception:  # noqa: BLE001 - o atalho que sobrar não abre nada: o PAULUS recusa
            pass
        self.estado.prefs.atualizar({"word": {"instalado_em": ""}})

    # --- o servidor HTTPS

    def aplicar(self) -> bool:
        """Liga o servidor HTTPS se o Word está ligado e instalado; senão, fecha."""
        prefs = self._prefs()
        cert_ok = (self.pasta_cert / "servidor.pem").exists() and (self.pasta_cert / "servidor.key").exists()
        if prefs.get("ligado") and self.porta() and cert_ok:
            return self.abrir()
        self.fechar()
        return False

    def abrir(self) -> bool:
        app = getattr(self.estado, "app_para_word", None)
        porta = self.porta()
        if app is None or not porta:
            return False
        servidor = self._servidor
        if servidor is not None and getattr(servidor, "started", False) and servidor.config.port == porta:
            return True
        self.fechar()
        with socket.socket(socket.AF_INET, socket.SOCK_STREAM) as s:
            try:
                s.bind(("127.0.0.1", porta))
            except OSError:
                return False
        import uvicorn

        # lifespan desligado, como a porta do túnel: o arranque já aconteceu.
        servidor = uvicorn.Server(uvicorn.Config(app, host="127.0.0.1", port=porta, log_level="warning",
                                                 lifespan="off", ssl_certfile=str(self.pasta_cert / "servidor.pem"),
                                                 ssl_keyfile=str(self.pasta_cert / "servidor.key")))
        self._fio = threading.Thread(target=servidor.run, name="porta-do-word", daemon=True)
        self._fio.start()
        fim = time.time() + 5
        while time.time() < fim and not servidor.started:
            time.sleep(0.05)
        if not servidor.started:
            servidor.should_exit = True
            return False
        self._servidor = servidor
        return True

    def fechar(self) -> None:
        if self._servidor is not None:
            self._servidor.should_exit = True
            fio = self._fio
            if fio is not None and fio is not threading.current_thread():
                fio.join(timeout=8)
        self._servidor = None
        self._fio = None
