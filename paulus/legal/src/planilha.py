"""
PAULUS - Planilha.

Grade com formulas, no formato que um escritorio brasileiro digita: ponto e
virgula separa argumento, virgula e decimal, e as funcoes tem nome em portugues
- =SE(C4="pago";0;B4*0,1) e o que a pessoa escreve no Excel dela.

Escrever um interpretador de formula da trabalho e por isso a tentacao e usar
`eval`. Nao da: formula vem de arquivo que veio de fora - uma planilha que o
cliente mandou - e `eval` num texto de terceiro executa o que estiver la. Aqui
ha um analisador de verdade, que so conhece numeros, textos, referencias e as
funcoes desta lista. O que ele nao conhece vira erro na celula, nunca acao.

Referencia circular tambem e erro visivel, nao travamento: A1 que soma A1
devolve #CIRCULAR e a planilha continua respondendo.
"""

from __future__ import annotations

import json
import re
import unicodedata
from dataclasses import dataclass, field, replace
from datetime import date, datetime
from pathlib import Path

LINHAS_PADRAO = 40
COLUNAS_PADRAO = 8
MAX_LINHAS = 2000
# CZ. A grade cresce com a rolagem ate aqui; o teto existe porque a tela
# desenha tudo que tem, e uma tabela de escritorio nao passa de algumas
# dezenas de colunas.
MAX_COLUNAS = 104

# Largura e altura sao em pixels da tela, e so as que a pessoa mudou ficam
# guardadas: as outras seguem o padrao.
LARGURA_MIN, LARGURA_MAX = 24, 900
ALTURA_MIN, ALTURA_MAX = 18, 400

FORMATOS = {
    "": "Texto",
    "numero": "Número",
    "moeda": "R$",
    "porcento": "%",
    "data": "Data",
}

ERRO_CIRCULAR = "#CIRCULAR"
ERRO_NOME = "#NOME?"
ERRO_VALOR = "#VALOR!"
ERRO_DIV = "#DIV/0!"
ERRO_REF = "#REF!"

RE_CELULA = re.compile(r"^([A-Z]{1,2})(\d{1,4})$")


# ------------------------------------------------------------- a grade


@dataclass
class Celula:
    valor: str = ""            # o que a pessoa digitou, inclusive a formula
    formato: str = ""
    negrito: bool = False
    italico: bool = False
    # Borda existe para separar um quadro do resto da folha - total de
    # honorarios, cabecalho de tabela - e por isso vai junto para o XLSX: uma
    # planilha que perde a moldura ao ser aberta no Excel nao foi exportada
    # inteira.
    borda: bool = False
    # Quantas colunas esta celula ocupa. 1 e o normal; maior que 1 e uma
    # celula juntada com as da direita - um titulo atravessando o quadro.
    juntar: int = 1
    # Nota presa na celula, como o comentario do Excel: aparece ao passar o
    # mouse e vai junto para o XLSX.
    comentario: str = ""

    @property
    def e_formula(self) -> bool:
        return self.valor.startswith("=")

    @property
    def vazia(self) -> bool:
        return (not self.valor and not self.formato and not self.negrito
                and not self.italico and not self.borda and self.juntar <= 1
                and not self.comentario)

    def to_dict(self) -> dict:
        saida = {
            "valor": self.valor, "formato": self.formato,
            "negrito": self.negrito, "italico": self.italico,
            "borda": self.borda, "juntar": self.juntar,
        }
        if self.comentario:
            saida["comentario"] = self.comentario
        return saida


@dataclass
class Aba:
    nome: str = "Página 1"
    celulas: dict[str, Celula] = field(default_factory=dict)
    linhas: int = LINHAS_PADRAO
    colunas: int = COLUNAS_PADRAO
    congelar_cabecalho: bool = False
    # Só o que foi mudado: {"B": 180} e {"3": 40}. Coluna por letra e linha
    # pelo numero que aparece na tela, que e como a pessoa pensa nelas.
    larguras: dict[str, int] = field(default_factory=dict)
    alturas: dict[str, int] = field(default_factory=dict)
    colunas_ocultas: list[str] = field(default_factory=list)
    linhas_ocultas: list[int] = field(default_factory=list)

    def obter(self, ref: str) -> Celula | None:
        return self.celulas.get(ref.upper())

    def gravar(self, ref: str, dados: dict) -> Celula:
        ref = ref.upper()
        celula = self.celulas.get(ref) or Celula()
        if "valor" in dados:
            celula.valor = str(dados["valor"] if dados["valor"] is not None else "")
        for campo in ("formato", "negrito", "italico", "borda"):
            if campo in dados:
                setattr(celula, campo, str(dados[campo] or "") if campo == "formato" else bool(dados[campo]))
        if "juntar" in dados:
            celula.juntar = max(1, min(int(dados["juntar"] or 1), MAX_COLUNAS))
        if "comentario" in dados:
            celula.comentario = str(dados["comentario"] or "").strip()[:2000]

        if celula.vazia:
            self.celulas.pop(ref, None)
            return celula

        self.celulas[ref] = celula
        linha, coluna = _posicao(ref)
        self.linhas = max(self.linhas, min(linha + 4, MAX_LINHAS))
        self.colunas = max(self.colunas, min(coluna + 2, MAX_COLUNAS))
        return celula

    def to_dict(self) -> dict:
        return {
            "nome": self.nome,
            "linhas": self.linhas,
            "colunas": self.colunas,
            "congelar_cabecalho": self.congelar_cabecalho,
            "larguras": dict(self.larguras),
            "alturas": dict(self.alturas),
            "colunas_ocultas": list(self.colunas_ocultas),
            "linhas_ocultas": list(self.linhas_ocultas),
            "celulas": {k: c.to_dict() for k, c in self.celulas.items()},
        }

    # ------------------------------------------------ a forma da grade

    def medir(self, larguras: dict | None = None, alturas: dict | None = None) -> None:
        """
        Muda largura de coluna e altura de linha. `None` volta ao padrao.

        Medida fora do razoavel e presa no limite em vez de recusada: quem
        arrasta a borda ate o fim da tela quer a coluna larga, nao um erro.
        """
        for letra, px in (larguras or {}).items():
            letra = str(letra).upper().strip()
            if not re.fullmatch(r"[A-Z]{1,2}", letra):
                continue
            if px is None:
                self.larguras.pop(letra, None)
            else:
                self.larguras[letra] = max(LARGURA_MIN, min(int(px), LARGURA_MAX))
        for linha, px in (alturas or {}).items():
            chave = str(linha).strip()
            if not chave.isdigit() or not 1 <= int(chave) <= MAX_LINHAS:
                continue
            if px is None:
                self.alturas.pop(chave, None)
            else:
                self.alturas[chave] = max(ALTURA_MIN, min(int(px), ALTURA_MAX))

    def ocultar(self, colunas=(), linhas=(), ocultar: bool = True) -> None:
        """Esconde ou mostra colunas e linhas. Esconder nao apaga nada."""
        cols = {str(c).upper().strip() for c in colunas if re.fullmatch(r"[A-Za-z]{1,2}", str(c).strip())}
        lins = {int(l) for l in linhas if str(l).strip().isdigit() and 1 <= int(l) <= MAX_LINHAS}
        if ocultar:
            self.colunas_ocultas = sorted(set(self.colunas_ocultas) | cols, key=indice_da_coluna)
            self.linhas_ocultas = sorted(set(self.linhas_ocultas) | lins)
        else:
            self.colunas_ocultas = [c for c in self.colunas_ocultas if c not in cols]
            self.linhas_ocultas = [l for l in self.linhas_ocultas if l not in lins]


def letra_da_coluna(indice: int) -> str:
    """0 vira A, 25 vira Z, 26 vira AA."""
    nome = ""
    indice += 1
    while indice:
        indice, resto = divmod(indice - 1, 26)
        nome = chr(65 + resto) + nome
    return nome


def indice_da_coluna(letras: str) -> int:
    total = 0
    for c in letras.upper():
        total = total * 26 + (ord(c) - 64)
    return total - 1


def _posicao(ref: str) -> tuple[int, int]:
    achado = RE_CELULA.match(ref.upper())
    if not achado:
        return 0, 0
    return int(achado.group(2)) - 1, indice_da_coluna(achado.group(1))


# ------------------------------------------------------ numeros e texto


def numero(texto) -> float | None:
    """
    Le numero como um brasileiro escreve.

    "12.000,00" e doze mil, nao doze. Ler isso errado num campo de honorarios
    troca o valor por mil vezes menos, e ninguem confere planilha somando de
    cabeca.
    """
    if isinstance(texto, (int, float)) and not isinstance(texto, bool):
        return float(texto)
    bruto = str(texto or "").strip().replace("R$", "").replace("%", "").strip()
    if not bruto:
        return None

    negativo = bruto.startswith("-") or (bruto.startswith("(") and bruto.endswith(")"))
    bruto = bruto.strip("()-").strip()

    if "," in bruto:
        # Virgula e o decimal; ponto, se houver, e separador de milhar.
        limpo = bruto.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"\d{1,3}(\.\d{3})+", bruto):
        limpo = bruto.replace(".", "")
    else:
        limpo = bruto

    try:
        valor = float(limpo)
    except ValueError:
        return None
    return -valor if negativo else valor


def formatar(valor, formato: str) -> str:
    """O valor como ele aparece na celula."""
    if valor is None or valor == "":
        return ""
    if isinstance(valor, str):
        return valor

    if formato == "porcento":
        return _br(valor * 100 if abs(valor) <= 1 else valor, 2).rstrip("0").rstrip(",") + "%"
    if formato == "moeda":
        return "R$ " + _br(valor, 2)
    if formato == "data":
        try:
            return (date(1899, 12, 30).toordinal() + int(valor)) and \
                date.fromordinal(date(1899, 12, 30).toordinal() + int(valor)).strftime("%d/%m/%Y")
        except (ValueError, OverflowError):
            return _br(valor, 0)
    if formato == "numero":
        return _br(valor, 2)

    if float(valor).is_integer():
        return _br(valor, 0)
    return _br(valor, 2)


def _br(valor: float, casas: int) -> str:
    texto = f"{valor:,.{casas}f}"
    return texto.replace(",", "\x00").replace(".", ",").replace("\x00", ".")


def _sem_acento(texto: str) -> str:
    normal = unicodedata.normalize("NFD", texto or "")
    return "".join(c for c in normal if unicodedata.category(c) != "Mn").upper()


# --------------------------------------------------------- as formulas


class ErroFormula(Exception):
    def __init__(self, codigo: str) -> None:
        super().__init__(codigo)
        self.codigo = codigo


TOKEN = re.compile(r"""
    (?P<numero>\d+(?:,\d+)?|\,\d+)
  | (?P<texto>"[^"]*")
  | (?P<refperdida>\#REF!)
  | (?P<faixa>\$?[A-Za-z]{1,2}\$?\d{1,4}\s*:\s*\$?[A-Za-z]{1,2}\$?\d{1,4})
  | (?P<celula>\$?[A-Za-z]{1,2}\$?\d{1,4})
  | (?P<funcao>[A-Za-zÀ-ÿ][A-Za-zÀ-ÿ0-9._]*)
  | (?P<op><=|>=|<>|[-+*/^()=<>;%&])
  | (?P<espaco>\s+)
""", re.VERBOSE)


def _tokenizar(formula: str) -> list[tuple[str, str]]:
    saida, pos = [], 0
    while pos < len(formula):
        achado = TOKEN.match(formula, pos)
        if not achado:
            raise ErroFormula(ERRO_VALOR)
        pos = achado.end()
        tipo = achado.lastgroup
        if tipo != "espaco":
            saida.append((tipo, achado.group()))
    return saida


class _Analisador:
    """
    Analisador descendente recursivo.

    Escrito a mao de proposito: a alternativa seria `eval`, e formula chega de
    planilha que veio de fora. Aqui o pior que uma formula estranha faz e
    devolver #NOME? numa celula.
    """

    def __init__(self, tokens: list[tuple[str, str]], resolver) -> None:
        self.tokens = tokens
        self.i = 0
        self.resolver = resolver          # (ref) -> valor da celula

    def _olhar(self) -> tuple[str, str] | None:
        return self.tokens[self.i] if self.i < len(self.tokens) else None

    def _comer(self, texto: str | None = None) -> tuple[str, str]:
        atual = self._olhar()
        if atual is None or (texto and atual[1].upper() != texto.upper()):
            raise ErroFormula(ERRO_VALOR)
        self.i += 1
        return atual

    # -------------------------------------------------------- gramatica

    def expressao(self):
        """Comparacoes: o nivel mais frouxo."""
        esquerda = self.soma()
        while (atual := self._olhar()) and atual[1] in ("=", "<>", "<", ">", "<=", ">="):
            self._comer()
            direita = self.soma()
            esquerda = _comparar(atual[1], esquerda, direita)
        return esquerda

    def soma(self):
        valor = self.produto()
        while (atual := self._olhar()) and atual[1] in ("+", "-", "&"):
            self._comer()
            outro = self.produto()
            if atual[1] == "&":
                valor = _como_texto(valor) + _como_texto(outro)
            else:
                a, b = _como_numero(valor), _como_numero(outro)
                valor = a + b if atual[1] == "+" else a - b
        return valor

    def produto(self):
        valor = self.potencia()
        while (atual := self._olhar()) and atual[1] in ("*", "/"):
            self._comer()
            outro = _como_numero(self.potencia())
            if atual[1] == "*":
                valor = _como_numero(valor) * outro
            else:
                if outro == 0:
                    raise ErroFormula(ERRO_DIV)
                valor = _como_numero(valor) / outro
        return valor

    def potencia(self):
        valor = self.unario()
        if (atual := self._olhar()) and atual[1] == "^":
            self._comer()
            return _como_numero(valor) ** _como_numero(self.potencia())
        return valor

    def unario(self):
        atual = self._olhar()
        if atual and atual[1] == "-":
            self._comer()
            return -_como_numero(self.unario())
        if atual and atual[1] == "+":
            self._comer()
            return self.unario()
        return self.posfixo()

    def posfixo(self):
        valor = self.primario()
        if (atual := self._olhar()) and atual[1] == "%":
            self._comer()
            return _como_numero(valor) / 100
        return valor

    def primario(self):
        atual = self._olhar()
        if atual is None:
            raise ErroFormula(ERRO_VALOR)
        tipo, texto = atual

        if tipo == "numero":
            self._comer()
            return float(texto.replace(",", "."))
        if tipo == "texto":
            self._comer()
            return texto[1:-1]
        if tipo == "refperdida":
            # A celula que a formula usava foi excluida junto com a linha ou
            # a coluna dela.
            raise ErroFormula(ERRO_REF)
        if tipo == "faixa":
            self._comer()
            return _Faixa(texto.replace("$", "").replace(" ", ""))
        if tipo == "celula":
            self._comer()
            return self.resolver(texto.replace("$", "").upper())
        if texto == "(":
            self._comer()
            valor = self.expressao()
            self._comer(")")
            return valor
        if tipo == "funcao":
            return self.funcao()
        raise ErroFormula(ERRO_VALOR)

    def funcao(self):
        _, nome = self._comer()
        chave = _sem_acento(nome).replace(".", "")
        if not (self._olhar() and self._olhar()[1] == "("):
            # VERDADEIRO e FALSO sao constantes, nao funcoes.
            if chave == "VERDADEIRO":
                return True
            if chave == "FALSO":
                return False
            raise ErroFormula(ERRO_NOME)

        self._comer("(")
        argumentos = []
        if self._olhar() and self._olhar()[1] != ")":
            argumentos.append(self.expressao())
            while self._olhar() and self._olhar()[1] == ";":
                self._comer(";")
                argumentos.append(self.expressao())
        self._comer(")")

        funcao = FUNCOES.get(chave)
        if funcao is None:
            raise ErroFormula(ERRO_NOME)
        return funcao(argumentos)


@dataclass
class _Faixa:
    """A1:B5 ainda sem resolver - quem resolve e a planilha."""

    texto: str
    valores: list = field(default_factory=list)


def _como_numero(valor) -> float:
    if isinstance(valor, _Faixa):
        valor = valor.valores[0] if valor.valores else 0
    if isinstance(valor, bool):
        return 1.0 if valor else 0.0
    if isinstance(valor, (int, float)):
        return float(valor)
    if valor in (None, ""):
        return 0.0
    convertido = numero(valor)
    if convertido is None:
        raise ErroFormula(ERRO_VALOR)
    return convertido


def _como_texto(valor) -> str:
    if isinstance(valor, bool):
        return "VERDADEIRO" if valor else "FALSO"
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    return "" if valor is None else str(valor)


def _comparar(operador: str, a, b) -> bool:
    if isinstance(a, str) or isinstance(b, str):
        x, y = _como_texto(a).strip().lower(), _como_texto(b).strip().lower()
    else:
        x, y = _como_numero(a), _como_numero(b)
    return {
        "=": x == y, "<>": x != y,
        "<": x < y, ">": x > y, "<=": x <= y, ">=": x >= y,
    }[operador]


def _achatar(argumentos) -> list:
    saida = []
    for a in argumentos:
        if isinstance(a, _Faixa):
            saida.extend(a.valores)
        else:
            saida.append(a)
    return saida


def _numeros(argumentos) -> list[float]:
    achados = []
    for v in _achatar(argumentos):
        if isinstance(v, bool) or v in (None, ""):
            continue
        if isinstance(v, (int, float)):
            achados.append(float(v))
            continue
        convertido = numero(v)
        if convertido is not None:
            achados.append(convertido)
    return achados


def _criterio(valores: list, criterio) -> list[bool]:
    """
    O criterio de SOMASE e CONT.SE: ">100", "pago", "<>0".

    Texto compara sem diferenciar maiuscula, que e como o Excel faz e como a
    pessoa espera - "Pago" e "pago" sao o mesmo na planilha do escritorio.
    """
    bruto = _como_texto(criterio).strip()
    achado = re.match(r"^(<=|>=|<>|<|>|=)\s*(.*)$", bruto)
    operador, alvo = (achado.group(1), achado.group(2)) if achado else ("=", bruto)

    # ">1000" compara numero com numero.
    #
    # O criterio chega sempre como texto, e sem esta conversao a comparacao
    # caia no ramo de texto: letra a letra, "500" > "1000" e verdadeiro, porque
    # "5" vem depois de "1". Com isso =SOMASE(B1:B4;">1000") somava a coluna
    # inteira - 13.700 onde o certo era 13.200 - e =SOMASE(B1:B4;"<1000")
    # devolvia zero. A celula mostrava um numero plausivel, que e a pior forma
    # de errar valor.
    numerico = numero(alvo)
    if numerico is None:
        saida = []
        for v in valores:
            try:
                saida.append(_comparar(operador, v, alvo))
            except ErroFormula:
                saida.append(False)
        return saida

    saida = []
    for v in valores:
        n = v if isinstance(v, (int, float)) and not isinstance(v, bool) else numero(v)
        if n is None:
            # Celula de texto ou vazia nao e maior nem menor que 1000; so
            # "diferente de" e verdadeiro sobre ela.
            saida.append(operador == "<>")
            continue
        try:
            saida.append(_comparar(operador, float(n), numerico))
        except ErroFormula:
            saida.append(False)
    return saida


def _se(args):
    if len(args) < 2:
        raise ErroFormula(ERRO_VALOR)
    condicao = args[0]
    verdade = bool(condicao) if isinstance(condicao, bool) else _como_numero(condicao) != 0
    if verdade:
        return args[1]
    return args[2] if len(args) > 2 else False


def _somase(args):
    if len(args) < 2:
        raise ErroFormula(ERRO_VALOR)
    faixa = args[0].valores if isinstance(args[0], _Faixa) else [args[0]]
    soma_faixa = args[2].valores if len(args) > 2 and isinstance(args[2], _Faixa) else faixa
    marcas = _criterio(faixa, args[1])
    total = 0.0
    for i, marca in enumerate(marcas):
        if marca and i < len(soma_faixa):
            total += _numeros([soma_faixa[i]])[0] if _numeros([soma_faixa[i]]) else 0.0
    return total


def _contse(args):
    if len(args) < 2:
        raise ErroFormula(ERRO_VALOR)
    faixa = args[0].valores if isinstance(args[0], _Faixa) else [args[0]]
    return float(sum(1 for m in _criterio(faixa, args[1]) if m))


def _media(args):
    valores = _numeros(args)
    if not valores:
        raise ErroFormula(ERRO_DIV)
    return sum(valores) / len(valores)


def _se_erro(args):
    """SE.ERRO ja recebe o valor pronto; o erro vira excecao antes de chegar."""
    if not args:
        raise ErroFormula(ERRO_VALOR)
    return args[0]


def _hoje(args):
    hoje = date.today()
    return float(hoje.toordinal() - date(1899, 12, 30).toordinal())


FUNCOES = {
    "SOMA": lambda a: sum(_numeros(a)),
    "MEDIA": _media,
    "MAXIMO": lambda a: max(_numeros(a)) if _numeros(a) else 0.0,
    "MINIMO": lambda a: min(_numeros(a)) if _numeros(a) else 0.0,
    "CONT": lambda a: float(len([v for v in _achatar(a) if v not in (None, "")])),
    "CONTNUM": lambda a: float(len(_numeros(a))),
    "CONTSE": _contse,
    "SOMASE": _somase,
    "SE": _se,
    "SEERRO": _se_erro,
    "ARRED": lambda a: round(_como_numero(a[0]), int(_como_numero(a[1])) if len(a) > 1 else 0),
    "ABS": lambda a: abs(_como_numero(a[0])),
    "INT": lambda a: float(int(_como_numero(a[0]))),
    "HOJE": _hoje,
    "CONCATENAR": lambda a: "".join(_como_texto(v) for v in _achatar(a)),
    "MAIUSCULA": lambda a: _como_texto(a[0]).upper(),
    "MINUSCULA": lambda a: _como_texto(a[0]).lower(),
    "E": lambda a: all(bool(v) if isinstance(v, bool) else _como_numero(v) != 0 for v in _achatar(a)),
    "OU": lambda a: any(bool(v) if isinstance(v, bool) else _como_numero(v) != 0 for v in _achatar(a)),
    "NAO": lambda a: not (bool(a[0]) if isinstance(a[0], bool) else _como_numero(a[0]) != 0),
}

# O que a tela mostra no menu "Fórmulas", com a explicacao em portugues.
AJUDA_FUNCOES = [
    ("SOMA", "=SOMA(B1:B5)", "Soma os números da faixa."),
    ("MÉDIA", "=MÉDIA(B1:B5)", "Média dos números da faixa."),
    ("SE", '=SE(C4="pago";0;B4*0,1)', "Um valor se a condição for verdadeira, outro se não."),
    ("SOMASE", '=SOMASE(C1:C5;"em aberto";B1:B5)', "Soma só as linhas que atendem ao critério."),
    ("CONT.SE", '=CONT.SE(C1:C5;"em aberto")', "Conta quantas linhas atendem ao critério."),
    ("MÁXIMO", "=MÁXIMO(B1:B5)", "O maior número da faixa."),
    ("MÍNIMO", "=MÍNIMO(B1:B5)", "O menor número da faixa."),
    ("ARRED", "=ARRED(B4*0,1;2)", "Arredonda com o número de casas que você pedir."),
    ("HOJE", "=HOJE()", "A data de hoje."),
    ("CONCATENAR", '=CONCATENAR(A1;" - ";B1)', "Junta textos."),
]


# ------------------------------------------------------------ calcular


def calcular_aba(aba: Aba) -> dict:
    """
    Resolve a aba inteira e devolve o valor de cada celula.

    Formula que depende de formula e resolvida por recursao com memoria; a
    pilha de visitados pega referencia circular e devolve #CIRCULAR na celula,
    em vez de estourar a pilha do programa.
    """
    prontos: dict[str, object] = {}
    visitando: set[str] = set()

    def valor_de(ref: str):
        ref = ref.upper()
        if ref in prontos:
            return prontos[ref]
        if ref in visitando:
            raise ErroFormula(ERRO_CIRCULAR)
        if not RE_CELULA.match(ref):
            raise ErroFormula(ERRO_REF)

        celula = aba.celulas.get(ref)
        if celula is None or not celula.valor:
            prontos[ref] = ""
            return ""

        if not celula.e_formula:
            convertido = numero(celula.valor)
            resultado = convertido if convertido is not None else celula.valor
            prontos[ref] = resultado
            return resultado

        visitando.add(ref)
        try:
            resultado = _avaliar(celula.valor[1:], valor_de)
        except ErroFormula as exc:
            resultado = exc.codigo
        except RecursionError:
            resultado = ERRO_CIRCULAR
        except Exception:
            resultado = ERRO_VALOR
        finally:
            visitando.discard(ref)

        prontos[ref] = resultado
        return resultado

    saida = {}
    for ref in list(aba.celulas):
        bruto = valor_de(ref)
        celula = aba.celulas[ref]
        # Celula digitada e sem formato aparece exatamente como foi digitada.
        # Quem escreve "12.000,00" nao espera ver "12.000" de volta - a conta
        # usa o numero, mas a tela devolve o que a pessoa pôs ali.
        if not celula.e_formula and not celula.formato:
            texto = celula.valor
        else:
            texto = bruto if isinstance(bruto, str) else formatar(bruto, celula.formato)

        saida[ref] = {
            "bruto": bruto if not isinstance(bruto, float) else round(bruto, 10),
            "texto": texto,
            "erro": isinstance(bruto, str) and bruto.startswith("#"),
            "formula": celula.valor if celula.e_formula else "",
        }
    return saida


def _avaliar(formula: str, valor_de):
    def resolver(ref: str):
        return valor_de(ref)

    def faixa_resolvida(alvo: _Faixa) -> _Faixa:
        inicio, fim = alvo.texto.split(":")
        l1, c1 = _posicao(inicio)
        l2, c2 = _posicao(fim)
        if (l1, c1) > (l2, c2):
            l1, l2, c1, c2 = l2, l1, c2, c1
        valores = []
        for linha in range(min(l1, l2), max(l1, l2) + 1):
            for coluna in range(min(c1, c2), max(c1, c2) + 1):
                valores.append(valor_de(f"{letra_da_coluna(coluna)}{linha + 1}"))
        alvo.valores = valores
        return alvo

    tokens = _tokenizar(formula)
    analisador = _Analisador(tokens, resolver)
    # As faixas so sabem seus valores depois de resolvidas; o analisador as
    # devolve cruas e este envelope preenche na hora de usar.
    original = analisador.primario

    def primario():
        valor = original()
        return faixa_resolvida(valor) if isinstance(valor, _Faixa) else valor

    analisador.primario = primario
    resultado = analisador.expressao()
    if analisador.i < len(analisador.tokens):
        raise ErroFormula(ERRO_VALOR)
    if isinstance(resultado, _Faixa):
        return resultado.valores[0] if resultado.valores else ""
    return resultado


def resumo(aba: Aba, calculado: dict) -> dict:
    """Os numeros que a barra de status e o painel lateral mostram."""
    numeros = [
        v["bruto"] for v in calculado.values()
        if isinstance(v["bruto"], (int, float)) and not isinstance(v["bruto"], bool)
    ]
    return {
        "celulas": len(aba.celulas),
        "formulas": sum(1 for c in aba.celulas.values() if c.e_formula),
        "erros": sum(1 for v in calculado.values() if v["erro"]),
        "soma": round(sum(numeros), 2) if numeros else 0,
        "media": round(sum(numeros) / len(numeros), 2) if numeros else 0,
        "maximo": round(max(numeros), 2) if numeros else 0,
        "minimo": round(min(numeros), 2) if numeros else 0,
    }


# ------------------------------------------------------ trabalhar em faixa


def cantos(faixa: str) -> tuple[int, int, int, int]:
    """(linha1, coluna1, linha2, coluna2) de "B9:D2", ja em ordem e contando de 0."""
    inicio, _, fim = str(faixa or "").upper().partition(":")
    l1, c1 = _posicao(inicio)
    l2, c2 = _posicao(fim or inicio)
    return min(l1, l2), min(c1, c2), max(l1, l2), max(c1, c2)


def refs_da_faixa(faixa: str) -> list[str]:
    """As celulas de "B2:D9", em ordem de leitura. Uma celula so vira ela mesma."""
    if not str(faixa or "").strip():
        return []
    l1, c1, l2, c2 = cantos(faixa)
    return [f"{letra_da_coluna(c)}{l + 1}"
            for l in range(l1, l2 + 1) for c in range(c1, c2 + 1)]


def resumo_selecao(aba: Aba, calculado: dict, refs: list[str]) -> dict:
    """
    O que a seleção tem dentro.

    Soma e média respondem a pergunta comum; mínimo, máximo e quantas estão
    vazias respondem a segunda: "esse total está certo?". Uma célula vazia no
    meio de uma coluna de parcelas muda a média e não muda a soma, e é o tipo
    de coisa que não se vê olhando.
    """
    numeros, textos, vazias, erros = [], 0, 0, 0
    for ref in refs:
        pronta = calculado.get(ref.upper()) or {}
        if pronta.get("erro"):
            erros += 1
            continue
        valor = pronta.get("bruto")
        if isinstance(valor, (int, float)) and not isinstance(valor, bool):
            numeros.append(valor)
        elif str(pronta.get("texto", "")).strip():
            textos += 1
        else:
            vazias += 1

    return {
        "quantas": len(refs),
        "com_numero": len(numeros),
        "com_texto": textos,
        "vazias": vazias,
        "erros": erros,
        "soma": round(sum(numeros), 2) if numeros else 0,
        "media": round(sum(numeros) / len(numeros), 2) if numeros else 0,
        "minimo": round(min(numeros), 2) if numeros else 0,
        "maximo": round(max(numeros), 2) if numeros else 0,
    }


def serie_da_faixa(aba: Aba, calculado: dict, faixa: str) -> dict:
    """
    A seleção virando uma série para desenhar: rótulos e valores.

    A regra é a que a pessoa já usa sem pensar: a última coluna com número é o
    valor, e a primeira coluna de texto à esquerda dela são os rótulos. Pedir
    para escolher coluna de valor e coluna de rótulo antes de ver o desenho
    seria um formulário no lugar de um gráfico.

    O que não dá para desenhar devolve `pode=False` com o motivo, e não um
    gráfico vazio: gráfico de nada é pior que gráfico nenhum, porque parece um
    resultado.
    """
    l1, c1, l2, c2 = cantos(faixa)

    # Quantos números tem cada coluna da seleção.
    numericas, textuais = {}, {}
    for coluna in range(c1, c2 + 1):
        numeros = textos = 0
        for linha in range(l1, l2 + 1):
            pronta = calculado.get(f"{letra_da_coluna(coluna)}{linha + 1}") or {}
            valor = pronta.get("bruto")
            if isinstance(valor, (int, float)) and not isinstance(valor, bool):
                numeros += 1
            elif str(pronta.get("texto", "")).strip():
                textos += 1
        numericas[coluna] = numeros
        textuais[coluna] = textos

    com_numero = [c for c in range(c1, c2 + 1) if numericas[c]]
    if not com_numero:
        return {"pode": False, "porque": "não há número nenhum nesta seleção"}

    valores_em = com_numero[-1]
    rotulos_em = next((c for c in range(c1, valores_em) if textuais[c] > numericas[c]), None)

    pontos = []
    for linha in range(l1, l2 + 1):
        pronta = calculado.get(f"{letra_da_coluna(valores_em)}{linha + 1}") or {}
        valor = pronta.get("bruto")
        if not isinstance(valor, (int, float)) or isinstance(valor, bool):
            continue
        if rotulos_em is None:
            rotulo = f"{letra_da_coluna(valores_em)}{linha + 1}"
        else:
            rotulo = str((calculado.get(f"{letra_da_coluna(rotulos_em)}{linha + 1}")
                          or {}).get("texto", "")).strip()
        pontos.append({
            "rotulo": rotulo or f"linha {linha + 1}",
            "valor": round(float(valor), 2),
            "texto": pronta.get("texto", ""),
        })

    if len(pontos) < 2:
        return {"pode": False,
                "porque": "um valor só não vira gráfico — selecione a coluna inteira"}

    return {
        "pode": True,
        "pontos": pontos,
        "coluna_valores": letra_da_coluna(valores_em),
        "coluna_rotulos": letra_da_coluna(rotulos_em) if rotulos_em is not None else "",
        "maior": max(p["valor"] for p in pontos),
        "menor": min(p["valor"] for p in pontos),
        "soma": round(sum(p["valor"] for p in pontos), 2),
    }


def pode_mesclar(aba: Aba, faixa: str) -> tuple[bool, str]:
    """
    Se dá para juntar as células desta faixa numa só.

    Só junta na horizontal, e só se o que vai ficar coberto estiver vazio. O
    Excel junta e joga fora o que estava debaixo, avisando numa caixa que todo
    mundo clica em OK sem ler. Aqui o que tem conteúdo não é coberto: a pessoa
    esvazia antes, se for isso mesmo que ela quer.
    """
    l1, c1, l2, c2 = cantos(faixa)
    # A altura primeiro: uma selecao em pe tem uma coluna so, e cairia no aviso
    # de "duas celulas lado a lado" - que nao e o problema dela.
    if l2 != l1:
        return False, "por enquanto só junto células da mesma linha"
    if c2 == c1:
        return False, "selecione pelo menos duas células lado a lado"

    ocupadas = [f"{letra_da_coluna(c)}{l1 + 1}" for c in range(c1 + 1, c2 + 1)
                if (aba.celulas.get(f"{letra_da_coluna(c)}{l1 + 1}") or Celula()).valor]
    if ocupadas:
        return False, ("as células " + ", ".join(ocupadas[:4]) +
                       " têm conteúdo — esvazie antes, para nada sumir sem você ver")
    return True, ""


def mesclar(aba: Aba, faixa: str) -> dict:
    """Junta as celulas da faixa numa so, a partir da primeira."""
    pode, porque = pode_mesclar(aba, faixa)
    if not pode:
        raise ValueError(porque)

    l1, c1, _, c2 = cantos(faixa)
    ref = f"{letra_da_coluna(c1)}{l1 + 1}"
    aba.gravar(ref, {"juntar": c2 - c1 + 1})
    for c in range(c1 + 1, c2 + 1):
        aba.celulas.pop(f"{letra_da_coluna(c)}{l1 + 1}", None)
    return {"ref": ref, "colunas": c2 - c1 + 1}


def separar(aba: Aba, faixa: str) -> dict:
    """Desfaz a juncao das celulas da faixa."""
    soltas = 0
    for ref in refs_da_faixa(faixa):
        celula = aba.celulas.get(ref)
        if celula and celula.juntar > 1:
            celula.juntar = 1
            soltas += 1
    return {"soltas": soltas}


# Referencia de celula dentro de uma formula. O olhar para tras evita cortar
# "AB12" no meio, e o olhar para frente evita confundir nome de funcao.
RE_REF_FORMULA = re.compile(r"(?<![A-Z0-9_$])([A-Z]{1,2})(\d{1,4})(?![0-9(])")


def _trocar_linha(valor: str, de: int, para: int) -> str:
    """
    Referencias a linha `de` passam a apontar para a linha `para`.

    Numa tabela de parcelas, =B7*0,1 na linha 7 tem que virar =B9*0,1 quando a
    linha 7 for para a 9 - senao ordenar troca silenciosamente a conta de cada
    linha, que e o pior estrago que uma planilha pode sofrer sem avisar.

    O que aponta para OUTRA linha fica onde estava: e referencia a um total, a
    uma taxa, a uma celula fora da tabela.
    """
    if de == para or not str(valor).startswith("="):
        return valor

    saida = []
    # Fora das aspas apenas: ="A1" e texto, nao referencia.
    for pedaco in re.split(r'("(?:[^"]|"")*")', valor):
        if pedaco.startswith('"'):
            saida.append(pedaco)
            continue
        saida.append(RE_REF_FORMULA.sub(
            lambda m: f"{m.group(1)}{para}" if int(m.group(2)) == de else m.group(0),
            pedaco))
    return "".join(saida)


def _chave_de_ordem(pronta: dict):
    """Numero antes de texto, e vazio sempre por ultimo, nos dois sentidos."""
    valor = pronta.get("bruto")
    if isinstance(valor, (int, float)) and not isinstance(valor, bool):
        return (1, float(valor), "")
    texto = str(pronta.get("texto", "")).strip()
    if not texto:
        return (9, 0.0, "")
    return (2, 0.0, _sem_acento(texto).lower())


def ordenar(aba: Aba, calculado: dict, faixa: str, coluna: str,
            crescente: bool = True, com_cabecalho: bool = False) -> dict:
    """
    Ordena as linhas de uma faixa por uma coluna.

    A linha inteira anda junto. Ordenar so a coluna escolhida embaralha a
    tabela - o valor da linha 7 passa a valer para o cliente da linha 3 - e
    esse estrago nao aparece olhando: cada celula continua com um numero
    plausivel. Por isso a faixa e a tabela inteira, e nao uma coluna.

    Ordena pelo valor CALCULADO: uma coluna de formulas se ordena pelo
    resultado, que e o que a pessoa ve.
    """
    l1, c1, l2, c2 = cantos(faixa)
    alvo = indice_da_coluna(str(coluna or "").upper().strip() or letra_da_coluna(c1))
    if not (c1 <= alvo <= c2):
        raise ValueError("a coluna de ordenação está fora da faixa")

    primeira = l1 + 1 if com_cabecalho else l1
    if primeira >= l2:
        return {"linhas": 0, "formulas": 0}

    # Cada linha vira uma copia antes de qualquer escrita: ordenar no lugar
    # sobrescreveria celulas que ainda nao foram lidas.
    linhas = []
    for linha in range(primeira, l2 + 1):
        celulas = {}
        for coluna_i in range(c1, c2 + 1):
            ref = f"{letra_da_coluna(coluna_i)}{linha + 1}"
            existente = aba.celulas.get(ref)
            if existente:
                celulas[coluna_i] = replace(existente)
        chave = _chave_de_ordem(calculado.get(f"{letra_da_coluna(alvo)}{linha + 1}") or {})
        linhas.append((chave, linha, celulas))

    # Vazio por ultimo nos dois sentidos: inverter a ordem nao pode jogar as
    # linhas em branco para o topo da tabela.
    vazias = [x for x in linhas if x[0][0] == 9]
    cheias = [x for x in linhas if x[0][0] != 9]
    cheias.sort(key=lambda x: x[0], reverse=not crescente)
    ordenadas = cheias + vazias

    ajustadas = 0
    for posicao, (_, de_linha, celulas) in enumerate(ordenadas):
        para_linha = primeira + posicao
        for coluna_i in range(c1, c2 + 1):
            ref = f"{letra_da_coluna(coluna_i)}{para_linha + 1}"
            celula = celulas.get(coluna_i)
            if celula is None:
                aba.celulas.pop(ref, None)
                continue
            novo = _trocar_linha(celula.valor, de_linha + 1, para_linha + 1)
            if novo != celula.valor:
                ajustadas += 1
            celula.valor = novo
            aba.celulas[ref] = celula

    return {"linhas": len(ordenadas), "formulas": ajustadas}


def filtrar(aba: Aba, calculado: dict, faixa: str, coluna: str, criterio: str,
            com_cabecalho: bool = True) -> dict:
    """
    Quais linhas da faixa NAO casam com o criterio.

    Devolve as linhas a esconder, e nao uma planilha nova: filtro e jeito de
    olhar, nao alteracao do documento. Um filtro gravado esconderia linhas de
    quem abrisse o arquivo depois sem saber que havia filtro - e uma tabela de
    parcelas com linhas faltando e um erro que ninguem percebe.

    O criterio e o mesmo de SOMASE: "pago", ">1000", "<>0".
    """
    l1, c1, l2, c2 = cantos(faixa)
    alvo = indice_da_coluna(str(coluna or "").upper().strip() or letra_da_coluna(c1))
    if not (c1 <= alvo <= c2):
        raise ValueError("a coluna do filtro está fora da faixa")

    primeira = l1 + 1 if com_cabecalho else l1
    valores, numeros = [], []
    for linha in range(primeira, l2 + 1):
        pronta = calculado.get(f"{letra_da_coluna(alvo)}{linha + 1}") or {}
        bruto = pronta.get("bruto")
        valores.append(bruto if bruto is not None else pronta.get("texto", ""))
        numeros.append(linha)

    casa = _criterio(valores, criterio)
    esconder = [linha + 1 for linha, ok in zip(numeros, casa) if not ok]
    return {
        "esconder": esconder,
        "mostrando": len(numeros) - len(esconder),
        "de": len(numeros),
        "coluna": letra_da_coluna(alvo),
    }


# ---------------------------------------- inserir e excluir linha e coluna

# Referencia com ou sem $, sozinha ou em faixa. O olhar para tras evita pegar
# o meio de um nome; o olhar para frente evita nome de funcao ("LOG10(").
RE_REF_COMPLETA = re.compile(
    r"(?<![A-Za-z0-9_$.])(\$?)([A-Za-z]{1,2})(\$?)(\d{1,4})(?![0-9A-Za-z_(])"
    r"(?:(\s*:\s*)(\$?)([A-Za-z]{1,2})(\$?)(\d{1,4})(?![0-9A-Za-z_(]))?")


def _mover(x: int, em: int, quantas: int, inserir: bool) -> int | None:
    """Onde a posicao `x` vai parar. None e posicao que foi excluida."""
    if inserir:
        return x + quantas if x >= em else x
    if x < em:
        return x
    if x >= em + quantas:
        return x - quantas
    return None


def _mover_faixa(a: int, b: int, em: int, quantas: int, inserir: bool) -> tuple[int, int] | None:
    """
    Uma faixa encolhe ou estica com a estrutura, como no Excel.

    =SOMA(B2:B10) com uma linha inserida na 5 vira =SOMA(B2:B11): a linha nova
    entra na conta. Excluir as linhas 9 a 12 encolhe para B2:B8. So quando a
    faixa inteira some a formula vira #REF!.
    """
    if inserir:
        return _mover(a, em, quantas, True), _mover(b, em, quantas, True)
    fim = em + quantas - 1
    novo_a = a if a < em else (em if a <= fim else a - quantas)
    novo_b = b if b < em else (em - 1 if b <= fim else b - quantas)
    if novo_a > novo_b:
        return None
    return novo_a, novo_b


def ajustar_formula(valor: str, eixo: str, em: int, quantas: int, inserir: bool) -> str:
    """
    As referencias de uma formula depois de inserir ou excluir linhas/colunas.

    Aqui o $ nao prende: a celula mudou de lugar, e a formula tem que seguir a
    celula - e o que o Excel faz. `em` conta de 1 para linhas e de 0 para
    colunas. Referencia a uma celula excluida vira #REF!, visivel na hora, em
    vez de passar a apontar para a vizinha sem ninguem ver.
    """
    if not str(valor).startswith("="):
        return valor

    def trocar(m: re.Match) -> str:
        d1, col1, d2, lin1, sep, d3, col2, d4, lin2 = m.groups()
        if sep is None:
            if eixo == "linhas":
                nova = _mover(int(lin1), em, quantas, inserir)
                return "#REF!" if nova is None else f"{d1}{col1.upper()}{d2}{nova}"
            nova = _mover(indice_da_coluna(col1), em, quantas, inserir)
            return "#REF!" if nova is None else f"{d1}{letra_da_coluna(nova)}{d2}{lin1}"

        if eixo == "linhas":
            faixa = _mover_faixa(int(lin1), int(lin2), em, quantas, inserir)
            if faixa is None:
                return "#REF!"
            return f"{d1}{col1.upper()}{d2}{faixa[0]}:{d3}{col2.upper()}{d4}{faixa[1]}"
        faixa = _mover_faixa(indice_da_coluna(col1), indice_da_coluna(col2), em, quantas, inserir)
        if faixa is None:
            return "#REF!"
        return (f"{d1}{letra_da_coluna(faixa[0])}{d2}{lin1}:"
                f"{d3}{letra_da_coluna(faixa[1])}{d4}{lin2}")

    saida = []
    for pedaco in re.split(r'("(?:[^"]|"")*")', valor):
        saida.append(pedaco if pedaco.startswith('"') else RE_REF_COMPLETA.sub(trocar, pedaco))
    return "".join(saida)


def reestruturar(aba: Aba, eixo: str, acao: str, em: int, quantas: int = 1) -> dict:
    """
    Insere ou exclui linhas (eixo "linhas", `em` contando de 1) ou colunas
    (eixo "colunas", `em` contando de 0).

    Tudo anda junto: as celulas, as formulas de toda a aba, as larguras, as
    alturas, o que estava oculto e as celulas juntadas. Inserir que empurraria
    conteudo para fora da grade e recusado, em vez de cortar o fim da tabela.
    """
    if eixo not in ("linhas", "colunas") or acao not in ("inserir", "excluir"):
        raise ValueError("operação desconhecida")
    inserir = acao == "inserir"
    limite = MAX_LINHAS if eixo == "linhas" else MAX_COLUNAS
    base = 1 if eixo == "linhas" else 0
    quantas = int(quantas or 1)
    em = int(em)
    if quantas < 1 or quantas > limite:
        raise ValueError("quantidade inválida")
    if not base <= em < limite + base:
        raise ValueError("posição fora da grade")

    def posicao(ref: str) -> tuple[int, int]:
        linha, coluna = _posicao(ref)
        return linha + 1, coluna

    if inserir:
        maior = max((posicao(r)[0 if eixo == "linhas" else 1] for r in aba.celulas), default=base - 1)
        if maior >= em and maior + quantas >= limite + base:
            raise ValueError("não cabe: o conteúdo passaria do fim da grade")

    novas: dict[str, Celula] = {}
    removidas = 0
    for ref, celula in aba.celulas.items():
        linha, coluna = posicao(ref)
        if eixo == "linhas":
            nova = _mover(linha, em, quantas, inserir)
            if nova is None:
                removidas += 1
                continue
            destino = f"{letra_da_coluna(coluna)}{nova}"
        else:
            if celula.juntar > 1:
                faixa = _mover_faixa(coluna, coluna + celula.juntar - 1, em, quantas, inserir)
                if faixa is None or _mover(coluna, em, quantas, inserir) is None:
                    removidas += 1
                    continue
                celula.juntar = faixa[1] - faixa[0] + 1
            nova = _mover(coluna, em, quantas, inserir)
            if nova is None:
                removidas += 1
                continue
            destino = f"{letra_da_coluna(nova)}{linha}"
        novas[destino] = celula

    formulas = 0
    for celula in novas.values():
        if celula.e_formula:
            novo = ajustar_formula(celula.valor, eixo, em, quantas, inserir)
            if novo != celula.valor:
                formulas += 1
                celula.valor = novo
    aba.celulas = novas

    if eixo == "linhas":
        aba.alturas = {str(n): px for k, px in aba.alturas.items()
                       if (n := _mover(int(k), em, quantas, inserir)) is not None}
        aba.linhas_ocultas = sorted(n for l in aba.linhas_ocultas
                                    if (n := _mover(int(l), em, quantas, inserir)) is not None)
        aba.linhas = max(1, min(aba.linhas + (quantas if inserir else -quantas), MAX_LINHAS))
    else:
        aba.larguras = {letra_da_coluna(n): px for k, px in aba.larguras.items()
                        if (n := _mover(indice_da_coluna(k), em, quantas, inserir)) is not None}
        aba.colunas_ocultas = [letra_da_coluna(n) for c in aba.colunas_ocultas
                               if (n := _mover(indice_da_coluna(c), em, quantas, inserir)) is not None]
        aba.colunas = max(1, min(aba.colunas + (quantas if inserir else -quantas), MAX_COLUNAS))

    return {"eixo": eixo, "acao": acao, "em": em, "quantas": quantas,
            "formulas": formulas, "removidas": removidas}


def gravar_lote(aba: Aba, itens: list) -> int:
    """
    Varias celulas de uma vez: colar, preencher arrastando a alca, limpar.

    Uma chamada por celula num colar de 200 celulas seriam 200 versoes do
    documento no Historico - e 200 idas ao disco.
    """
    feitas = 0
    for item in itens:
        if not isinstance(item, dict):
            continue
        ref = str(item.get("ref", "")).upper()
        dados = item.get("dados")
        if not RE_CELULA.match(ref) or not isinstance(dados, dict):
            continue
        linha, coluna = _posicao(ref)
        if linha >= MAX_LINHAS or coluna >= MAX_COLUNAS:
            continue
        permitidos = {k: v for k, v in dados.items()
                      if k in ("valor", "formato", "negrito", "italico", "borda", "comentario")}
        if permitidos:
            aba.gravar(ref, permitidos)
            feitas += 1
    return feitas


# --------------------------------------------------------- entrar e sair


def de_csv(texto: str, nome: str = "Importado") -> Aba:
    """
    Le CSV, adivinhando o separador.

    Planilha exportada no Brasil quase sempre usa ponto e virgula, porque a
    virgula ja e o decimal. Adivinhar errado joga a linha inteira numa celula.
    """
    import csv
    import io

    amostra = texto[:4000]
    try:
        dialeto = csv.Sniffer().sniff(amostra, delimiters=";,\t")
        separador = dialeto.delimiter
    except csv.Error:
        separador = ";" if amostra.count(";") >= amostra.count(",") else ","

    aba = Aba(nome=nome)
    leitor = csv.reader(io.StringIO(texto), delimiter=separador)
    for l, linha in enumerate(leitor):
        if l >= MAX_LINHAS:
            break
        for c, valor in enumerate(linha):
            if c >= MAX_COLUNAS:
                break
            if str(valor).strip():
                aba.gravar(f"{letra_da_coluna(c)}{l + 1}", {"valor": str(valor).strip()})
    return aba


def para_csv(aba: Aba, calculado: dict) -> str:
    """Sai com os valores calculados, nao com as formulas."""
    import csv
    import io

    saida = io.StringIO()
    escritor = csv.writer(saida, delimiter=";", lineterminator="\n")
    for linha in range(aba.linhas):
        fila = []
        for coluna in range(aba.colunas):
            ref = f"{letra_da_coluna(coluna)}{linha + 1}"
            fila.append((calculado.get(ref) or {}).get("texto", ""))
        while fila and not fila[-1]:
            fila.pop()
        escritor.writerow(fila)
    return saida.getvalue()


def de_xlsx(dados: bytes) -> list[Aba]:
    import io

    from openpyxl import load_workbook

    livro = load_workbook(io.BytesIO(dados), data_only=False)
    abas = []
    for folha in livro.worksheets[:6]:
        aba = Aba(nome=str(folha.title)[:24])
        for linha in folha.iter_rows(max_row=min(folha.max_row or 1, MAX_LINHAS),
                                     max_col=min(folha.max_column or 1, MAX_COLUNAS)):
            for celula in linha:
                if celula.value is None:
                    continue
                valor = celula.value
                if isinstance(valor, datetime):
                    valor = valor.strftime("%d/%m/%Y")
                aba.gravar(celula.coordinate, {"valor": str(valor)})
        abas.append(aba)
    return abas or [Aba()]


def para_xlsx(abas: list[Aba], calculados: list[dict]) -> bytes:
    """
    Exporta com formula, nao so com o resultado.

    Quem abre no Excel espera continuar calculando. Mandar so o numero
    entregaria uma foto da planilha em vez da planilha.
    """
    import io

    from openpyxl import Workbook

    livro = Workbook()
    livro.remove(livro.active)

    for aba, calculado in zip(abas, calculados):
        folha = livro.create_sheet(title=aba.nome[:31] or "Página")
        for ref, celula in aba.celulas.items():
            if celula.e_formula:
                folha[ref] = celula.valor
            else:
                convertido = numero(celula.valor)
                folha[ref] = convertido if convertido is not None else celula.valor
            if celula.formato == "moeda":
                folha[ref].number_format = 'R$ #,##0.00'
            elif celula.formato == "porcento":
                folha[ref].number_format = "0.0%"
            elif celula.formato == "numero":
                folha[ref].number_format = "#,##0.00"
            if celula.negrito or celula.italico:
                from openpyxl.styles import Font

                folha[ref].font = Font(bold=celula.negrito, italic=celula.italico)
            if celula.borda:
                from openpyxl.styles import Border, Side

                fio = Side(style="thin")
                folha[ref].border = Border(left=fio, right=fio, top=fio, bottom=fio)
        for ref, celula in aba.celulas.items():
            if celula.juntar > 1:
                linha, coluna = _posicao(ref)
                fim = f"{letra_da_coluna(coluna + celula.juntar - 1)}{linha + 1}"
                folha.merge_cells(f"{ref}:{fim}")
        if aba.congelar_cabecalho:
            folha.freeze_panes = "A2"
        # Largura em pixels vira a unidade do Excel (mais ou menos 7 px por
        # caractere) e altura vira pontos: a coluna que a pessoa alargou para
        # caber o nome do cliente chega larga do outro lado.
        for letra, px in aba.larguras.items():
            folha.column_dimensions[letra].width = round(px / 7, 2)
        for linha, px in aba.alturas.items():
            folha.row_dimensions[int(linha)].height = round(px * 0.75, 2)
        for letra in aba.colunas_ocultas:
            folha.column_dimensions[letra].hidden = True
        for linha in aba.linhas_ocultas:
            folha.row_dimensions[int(linha)].hidden = True
        for ref, celula in aba.celulas.items():
            if celula.comentario:
                from openpyxl.comments import Comment

                folha[ref].comment = Comment(celula.comentario, "PAULUS")

    saida = io.BytesIO()
    livro.save(saida)
    return saida.getvalue()


def de_dict(bruto: dict) -> list[Aba]:
    abas = []
    for item in bruto.get("abas", []):
        aba = Aba(
            nome=item.get("nome", "Página"),
            linhas=int(item.get("linhas", LINHAS_PADRAO)),
            colunas=int(item.get("colunas", COLUNAS_PADRAO)),
            congelar_cabecalho=bool(item.get("congelar_cabecalho")),
        )
        # Planilha gravada antes de existir largura e altura simplesmente nao
        # tem esses campos - e abre com tudo no padrao.
        if isinstance(item.get("larguras"), dict) or isinstance(item.get("alturas"), dict):
            try:
                aba.medir(item.get("larguras") or {}, item.get("alturas") or {})
            except (TypeError, ValueError):
                pass
        aba.ocultar(item.get("colunas_ocultas") or [], item.get("linhas_ocultas") or [])
        for ref, dados in (item.get("celulas") or {}).items():
            if isinstance(dados, dict):
                aba.celulas[ref.upper()] = Celula(
                    valor=str(dados.get("valor", "")),
                    formato=str(dados.get("formato", "")),
                    negrito=bool(dados.get("negrito")),
                    italico=bool(dados.get("italico")),
                    borda=bool(dados.get("borda")),
                    juntar=max(1, int(dados.get("juntar") or 1)),
                    comentario=str(dados.get("comentario") or ""),
                )
        abas.append(aba)
    return abas or [Aba()]


def para_json(abas: list[Aba]) -> str:
    return json.dumps({"versao": 1, "abas": [a.to_dict() for a in abas]}, ensure_ascii=False)
