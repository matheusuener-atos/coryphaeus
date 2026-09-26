"""
A base de demonstração do PAULUS: um escritório fictício, pronto para mostrar.

Cria data/demo/ (ou a pasta que vier no argumento) com:

- seis documentos fictícios do escritório (contrato de transporte, aditivo,
  locação, procuração, notificação e contrato de honorários), em DOCX, na
  pasta do Acervo da demonstração;
- cadastros, compromissos, tarefas, lançamentos do Financeiro e serviços,
  com datas contadas a partir de HOJE - a demonstração não envelhece;
- o nome da advogada e do escritório nas preferências.

Tudo é inventado: nomes, CPFs, CNPJs e valores. Os números de documento
têm dígito verificador válido para a conferência do editor não reclamar.

Nada aqui toca em data/ de verdade. Para abrir o programa na demonstração:

    tools\\demo\\abrir_demo.bat

Para refazer do zero (apaga só a pasta da demonstração):

    venv\\Scripts\\python.exe tools\\demo\\criar_demo.py --refazer
"""

from __future__ import annotations

import shutil
import sys
from datetime import date, timedelta
from pathlib import Path

RAIZ = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(RAIZ / "src"))

import agenda as agenda_mod  # noqa: E402
import base as base_mod  # noqa: E402
import cadastros as cadastros_mod  # noqa: E402
import financeiro as financeiro_mod  # noqa: E402
import servicos as servicos_mod  # noqa: E402
import tarefas as tarefas_mod  # noqa: E402
from config import Preferencias  # noqa: E402

PADRAO = RAIZ / "data" / "demo"

ESCRITORIO = "Moura & Campos Advocacia"
ADVOGADA = "Helena Moura"
OAB = "PA 12.345"


# ------------------------------------------------------------ documentos


def _dv_cpf(base9: str) -> str:
    d = [int(c) for c in base9]
    for peso0 in (10, 11):
        s = sum(v * p for v, p in zip(d, range(peso0, 1, -1)))
        r = (s * 10) % 11
        d.append(0 if r == 10 else r)
    t = "".join(map(str, d))
    return f"{t[:3]}.{t[3:6]}.{t[6:9]}-{t[9:]}"


def _dv_cnpj(base12: str) -> str:
    d = [int(c) for c in base12]
    for pesos in ([5, 4, 3, 2, 9, 8, 7, 6, 5, 4, 3, 2], [6, 5, 4, 3, 2, 9, 8, 7, 6, 5, 4, 3, 2]):
        s = sum(v * p for v, p in zip(d, pesos))
        r = s % 11
        d.append(0 if r < 2 else 11 - r)
    t = "".join(map(str, d))
    return f"{t[:2]}.{t[2:5]}.{t[5:8]}/{t[8:12]}-{t[12:]}"


CNPJ_COOPERATIVA = _dv_cnpj("274315620001")
CNPJ_TRANSPORTES = _dv_cnpj("183904570001")
CNPJ_CLINICA = _dv_cnpj("315528410001")
CPF_JOAO = _dv_cpf("418326095")
CPF_LOCADOR = _dv_cpf("205771348")

DOCUMENTOS = {
    "Cooperativa Vale Verde/Contrato de transporte - Cooperativa Vale Verde x Rio Fresco.docx": [
        ("t", "CONTRATO DE PRESTAÇÃO DE SERVIÇOS DE TRANSPORTE"),
        ("p", f"CONTRATANTE: COOPERATIVA AGRÍCOLA VALE VERDE, inscrita no CNPJ sob o nº {CNPJ_COOPERATIVA}, "
              "com sede na Rodovia PA-279, km 12, zona rural de Tucumã/PA."),
        ("p", f"CONTRATADA: TRANSPORTES RIO FRESCO LTDA., inscrita no CNPJ sob o nº {CNPJ_TRANSPORTES}, "
              "com sede na Avenida Tapajós, 880, Santarém/PA."),
        ("h", "Cláusula 1ª - Do objeto"),
        ("p", "A CONTRATADA transportará a produção de cacau e castanha da CONTRATANTE entre os armazéns de "
              "Tucumã e o porto de Santarém, em até 8 (oito) viagens por mês."),
        ("h", "Cláusula 2ª - Do preço"),
        ("p", "Pelos serviços, a CONTRATANTE pagará o valor mensal de R$ 18.500,00 (dezoito mil e quinhentos "
              "reais), até o dia 10 (dez) de cada mês, mediante nota fiscal."),
        ("h", "Cláusula 3ª - Do atraso"),
        ("p", "O atraso no pagamento sujeita a CONTRATANTE à multa de 2% (dois por cento) sobre o valor devido, "
              "mais juros de 1% (um por cento) ao mês."),
        ("h", "Cláusula 4ª - Da vigência"),
        ("p", "O contrato vigora por 12 (doze) meses, de 1º de março de 2026 a 28 de fevereiro de 2027, e se "
              "renova automaticamente por igual período, salvo aviso de não renovação enviado por escrito com "
              "antecedência mínima de 60 (sessenta) dias."),
        ("h", "Cláusula 5ª - Da rescisão"),
        ("p", "Qualquer das partes pode rescindir o contrato sem motivo, com aviso prévio de 30 (trinta) dias. "
              "Quem rescindir pagará multa de 10% (dez por cento) sobre o valor das parcelas restantes."),
        ("h", "Cláusula 6ª - Do foro"),
        ("p", "Fica eleito o foro da Comarca de Santarém/PA para dirimir as questões deste contrato."),
        ("p", "Tucumã/PA, 20 de fevereiro de 2026."),
    ],
    "Cooperativa Vale Verde/Aditivo n. 1 - contrato de transporte.docx": [
        ("t", "PRIMEIRO TERMO ADITIVO AO CONTRATO DE PRESTAÇÃO DE SERVIÇOS DE TRANSPORTE"),
        ("p", "As partes do contrato de transporte firmado em 20 de fevereiro de 2026, COOPERATIVA AGRÍCOLA VALE "
              "VERDE e TRANSPORTES RIO FRESCO LTDA., resolvem aditá-lo nos termos abaixo."),
        ("h", "Cláusula 1ª - Do novo valor"),
        ("p", "A partir de 1º de setembro de 2026, o valor mensal passa a ser de R$ 19.980,00 (dezenove mil, "
              "novecentos e oitenta reais), em razão da inclusão da rota Altamira."),
        ("h", "Cláusula 2ª - Da nova rota"),
        ("p", "Fica incluída a rota Tucumã-Altamira, com até 2 (duas) viagens por mês, além das 8 (oito) já previstas."),
        ("p", "Permanecem inalteradas as demais cláusulas. Tucumã/PA, 25 de agosto de 2026."),
    ],
    "Cooperativa Vale Verde/Procuração - Cooperativa Vale Verde.docx": [
        ("t", "PROCURAÇÃO AD JUDICIA ET EXTRA"),
        ("p", f"OUTORGANTE: COOPERATIVA AGRÍCOLA VALE VERDE, CNPJ {CNPJ_COOPERATIVA}, neste ato representada "
              "por seu presidente."),
        ("p", f"OUTORGADA: {ADVOGADA.upper()}, advogada, inscrita na OAB/{OAB}, integrante de {ESCRITORIO}."),
        ("p", "PODERES: os da cláusula ad judicia et extra, para o foro em geral, e os especiais para negociar, "
              "transigir, receber e dar quitação, firmar acordos e notificar extrajudicialmente, em especial nas "
              "questões do contrato de transporte com a Transportes Rio Fresco Ltda."),
        ("p", "Não estão incluídos os poderes para renunciar a créditos da OUTORGANTE nem para confessar."),
        ("p", "Tucumã/PA, 3 de março de 2026."),
    ],
    "Cooperativa Vale Verde/Notificação - atraso nas entregas (Rio Fresco).docx": [
        ("t", "NOTIFICAÇÃO EXTRAJUDICIAL"),
        ("p", "NOTIFICANTE: COOPERATIVA AGRÍCOLA VALE VERDE. NOTIFICADA: TRANSPORTES RIO FRESCO LTDA."),
        ("p", "A NOTIFICANTE comunica que, em agosto de 2026, 3 (três) entregas previstas no contrato de transporte "
              "chegaram ao porto de Santarém com mais de 48 horas de atraso."),
        ("p", "Fica a NOTIFICADA constituída em mora e concedido o prazo de 10 (dez) dias, contados do recebimento "
              "desta, para regularizar o cronograma, sob pena de aplicação da multa da Cláusula 5ª e rescisão."),
        ("p", f"Tucumã/PA, 15 de setembro de 2026. {ADVOGADA}, OAB/{OAB}."),
    ],
    "Clínica Bem Viver/Contrato de locação - Clínica Bem Viver.docx": [
        ("t", "CONTRATO DE LOCAÇÃO NÃO RESIDENCIAL"),
        ("p", f"LOCADOR: MARCOS ANTÔNIO LEMOS, CPF {CPF_LOCADOR}."),
        ("p", f"LOCATÁRIA: CLÍNICA BEM VIVER LTDA., CNPJ {CNPJ_CLINICA}."),
        ("h", "Cláusula 1ª - Do imóvel"),
        ("p", "Sala comercial nº 402 do Edifício Centro Médico Xingu, Rua 7 de Setembro, 1210, São Félix do Xingu/PA."),
        ("h", "Cláusula 2ª - Do aluguel"),
        ("p", "O aluguel mensal é de R$ 6.200,00 (seis mil e duzentos reais), com vencimento no dia 5 (cinco) de "
              "cada mês, reajustado anualmente pela variação do IGP-M."),
        ("h", "Cláusula 3ª - Do prazo"),
        ("p", "A locação é de 36 (trinta e seis) meses, de 1º de julho de 2024 a 30 de junho de 2027."),
        ("h", "Cláusula 4ª - Da garantia"),
        ("p", "Como garantia, a LOCATÁRIA entrega caução em dinheiro equivalente a 3 (três) aluguéis."),
        ("h", "Cláusula 5ª - Da rescisão antecipada"),
        ("p", "Se a LOCATÁRIA devolver o imóvel antes do prazo, pagará multa de 3 (três) aluguéis, proporcional ao "
              "tempo que faltar."),
    ],
    "João Batista Ferreira/Contrato de honorários - João Batista Ferreira.docx": [
        ("t", "CONTRATO DE HONORÁRIOS ADVOCATÍCIOS"),
        ("p", f"CONTRATANTE: JOÃO BATISTA FERREIRA, CPF {CPF_JOAO}."),
        ("p", f"CONTRATADA: {ESCRITORIO}, por sua sócia {ADVOGADA}, OAB/{OAB}."),
        ("h", "Cláusula 1ª - Do objeto"),
        ("p", "Ajuizamento e acompanhamento de reclamação trabalhista contra o antigo empregador do CONTRATANTE, "
              "em primeira e segunda instâncias."),
        ("h", "Cláusula 2ª - Dos honorários"),
        ("p", "Honorários fixos de R$ 4.500,00 (quatro mil e quinhentos reais), em 3 (três) parcelas de R$ 1.500,00, "
              "com vencimentos em 10 de outubro, 10 de novembro e 10 de dezembro de 2026, mais honorários de êxito "
              "de 20% (vinte por cento) sobre o proveito econômico obtido."),
        ("p", "São Félix do Xingu/PA, 22 de setembro de 2026."),
    ],
}


def escrever_documentos(pasta: Path) -> int:
    from docx import Document

    for relativo, blocos in DOCUMENTOS.items():
        destino = pasta / relativo
        destino.parent.mkdir(parents=True, exist_ok=True)
        doc = Document()
        for tipo, texto in blocos:
            if tipo == "t":
                doc.add_heading(texto, level=1)
            elif tipo == "h":
                doc.add_heading(texto, level=2)
            else:
                doc.add_paragraph(texto)
        doc.save(str(destino))
    return len(DOCUMENTOS)


# ---------------------------------------------------------------- banco


def semear(dados: Path, pasta_acervo: Path, hoje: date) -> dict:
    def dia(n: int) -> str:
        return (hoje + timedelta(days=n)).isoformat()

    b = base_mod.Base(dados / "paulus.db")
    cad = cadastros_mod.Cadastros(b, dados / "cadastros_ignorados.json")
    ag = agenda_mod.Agenda(b)
    tar = tarefas_mod.Tarefas(b)
    fin = financeiro_mod.Financeiro(b)
    serv = servicos_mod.Servicos(b, lambda: ADVOGADA.split()[0], lambda: pasta_acervo)

    coop = cad.salvar({"tipo": "cliente", "nome": "Cooperativa Agrícola Vale Verde", "documento": CNPJ_COOPERATIVA,
                       "telefone": "(94) 99812-4410", "email": "diretoria@valeverde.coop.exemplo"})
    clinica = cad.salvar({"tipo": "cliente", "nome": "Clínica Bem Viver Ltda.", "documento": CNPJ_CLINICA,
                          "telefone": "(94) 3435-2201", "email": "administrativo@bemviver.exemplo"})
    joao = cad.salvar({"tipo": "cliente", "nome": "João Batista Ferreira", "documento": CPF_JOAO,
                       "telefone": "(94) 99160-3378", "email": "joao.ferreira@exemplo.com"})
    cad.salvar({"tipo": "colaborador", "nome": "Lívia Santos", "observacao": "estagiária", "vinculo": "estagio"})
    cad.salvar({"tipo": "despesa", "nome": "Aluguel do escritório", "honorario": "2.800,00", "dia_vencimento": 10})

    ag.salvar({"titulo": "Ligação com a Clínica Bem Viver - reajuste do aluguel", "data": dia(0), "hora": "16:00",
               "duracao": 30, "onde": "telefone", "cadastro_id": clinica})
    ag.salvar({"titulo": "Reunião com a Cooperativa Vale Verde - renovação do transporte", "data": dia(1),
               "hora": "10:00", "duracao": 60, "onde": "escritorio", "cadastro_id": coop, "avisar_min": 30})
    ag.salvar({"titulo": "Audiência de conciliação - João Batista Ferreira", "data": dia(2), "hora": "14:30",
               "duracao": 90, "onde": "online", "cadastro_id": joao, "avisar_min": 60})
    ag.salvar({"titulo": "Visita à sala 402 com o locador", "data": dia(6), "hora": "09:00", "duracao": 60,
               "cadastro_id": clinica})

    tar.salvar({"titulo": "Enviar a minuta do aditivo à Cooperativa", "prazo": dia(-2), "importante": True,
                "cadastro_id": coop})
    tar.salvar({"titulo": "Conferir o reajuste IGP-M da locação da Clínica", "prazo": dia(0), "cadastro_id": clinica})
    tar.salvar({"titulo": "Acompanhar o prazo da notificação à Rio Fresco", "prazo": dia(3), "importante": True,
                "cadastro_id": coop})
    tar.salvar({"titulo": "Organizar a pasta do João Batista", "cadastro_id": joao})

    mes = hoje.replace(day=1)
    fin.salvar({"tipo": "recebimento", "descricao": "Honorários mensais - Cooperativa Vale Verde", "valor": "7.500,00",
                "categoria": "honorarios", "cadastro_id": coop, "vencimento": mes.replace(day=5).isoformat(),
                "liquidado_em": mes.replace(day=5).isoformat()})
    fin.salvar({"tipo": "recebimento", "descricao": "Honorários - Clínica Bem Viver", "valor": "2.400,00",
                "categoria": "honorarios", "cadastro_id": clinica, "vencimento": dia(-6)})
    fin.salvar({"tipo": "recebimento", "descricao": "Honorários - João Batista (parcela 1 de 3)", "valor": "1.500,00",
                "categoria": "honorarios", "cadastro_id": joao, "vencimento": dia(14)})
    fin.salvar({"tipo": "despesa", "descricao": "Aluguel do escritório", "valor": "2.800,00", "categoria": "aluguel",
                "vencimento": mes.replace(day=10).isoformat(), "liquidado_em": mes.replace(day=10).isoformat()})
    fin.salvar({"tipo": "despesa", "descricao": "Contabilidade", "valor": "650,00", "categoria": "outros",
                "vencimento": dia(4)})

    s1 = serv.salvar({"nome": "Renovação do contrato de transporte", "cadastro_id": coop,
                      "descricao": "Aditivo com a rota Altamira e decisão sobre a renovação automática."})
    serv.etapa_adicionar(s1, "Enviar a minuta do aditivo", dia(-2))
    serv.etapa_adicionar(s1, "Decidir se avisa a não renovação", dia(20))
    s2 = serv.salvar({"nome": "Reclamação trabalhista - João Batista", "cadastro_id": joao,
                      "descricao": "Ação contra o antigo empregador; audiência de conciliação marcada."})
    serv.etapa_adicionar(s2, "Audiência de conciliação", dia(2))

    return {"cadastros": 5, "compromissos": 4, "tarefas": 4, "lancamentos": 5, "servicos": 2}


def main() -> int:
    argumentos = [a for a in sys.argv[1:] if not a.startswith("--")]
    dados = Path(argumentos[0]).resolve() if argumentos else PADRAO
    if dados.exists() and any(dados.iterdir()):
        if "--refazer" not in sys.argv:
            print(f"{dados} já existe. Use --refazer para apagar e criar de novo.")
            return 1
        # Só apaga o que é claramente da demonstração.
        if dados.name != "demo" or dados == (RAIZ / "data").resolve():
            print("por segurança, --refazer só apaga uma pasta chamada 'demo'")
            return 1
        shutil.rmtree(dados)

    acervo = dados / "test_contracts"
    acervo.mkdir(parents=True, exist_ok=True)
    n = escrever_documentos(acervo)
    contas = semear(dados, acervo, date.today())
    Preferencias(dados / "preferencias.json").atualizar({
        "pessoa": {"nome": ADVOGADA, "oab": OAB, "email": "helena@mouracampos.exemplo"},
        "escritorio": {"nome": ESCRITORIO},
    })
    print(f"demonstração criada em {dados}")
    print(f"  {n} documentos no Acervo")
    print("  " + ", ".join(f"{v} {k}" for k, v in contas.items()))
    print("abra com tools\\demo\\abrir_demo.bat")
    return 0


if __name__ == "__main__":
    sys.exit(main())
